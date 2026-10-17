"""Several runs at once: one combined results table (isoharness 0.2.0).

Each run is corrected on its own by ea_irms, with its own standards and its own choices;
this module only puts the runs' Results rows one after another, with a Run column first,
describes each run on one row, and lists the samples measured in more than one run. It
computes no correction. The means and differences on the Repeats sheet are arithmetic on
the engine's values. From isoharness 0.3.0 one sample sheet may go with the batch: its
columns follow the Results columns, and each run's information follows its row on Runs.
Each run may come with a configuration of its own (finding W-11): the endpoints take one
configuration per file, and the page names each run's configuration in config_name. The
samples on Repeats are paired by name as the sample sheet pairs names, with the spaces taken
out when nothing fits as written (finding W-12).
"""
from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ea_irms.cells import put
from ea_irms.metadata import RUN_FIELDS, norm, output_columns, output_header, output_values, summary_text
from ea_irms.pipeline import Config
from ea_irms.workbook import VERSION

from .serialize import csv_cells, results_header, results_values

ARIAL = "Arial"
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_H1 = Font(name=ARIAL, size=11, bold=True)
F_HDR = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")
F_SM = Font(name=ARIAL, size=9)
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F4E79")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
FILL_RERUN = PatternFill("solid", fgColor="F8CBAD")
FILL_CHECK = PatternFill("solid", fgColor="FCE4D6")

# The Results sheet's widths and number formats (ea_irms.workbook, Results), by column name.
WIDTHS = {"Row": 6, "Identifier": 18, "Type": 10, "Mass (mg)": 10, "%C": 8, "1σ(%C)": 8, "%N": 8, "1σ(%N)": 8,
          "C:N molar": 10, "1σ(C:N molar)": 10, "C:N mass": 10, "1σ(C:N mass)": 10, "d13C (permil)": 12,
          "1σ(d13C)": 9, "d15N (permil)": 12, "1σ(d15N)": 9, "Flags": 54, "Rerun": 9, "Why": 46,
          "Suggested mass (mg)": 24, "Elemental drift factor N": 10, "Elemental drift factor C": 10,
          "%C as measured": 9, "%N as measured": 9}
FORMATS = {"Mass (mg)": "0.000", "1σ(d13C)": "0.000", "1σ(d15N)": "0.000",
           "Elemental drift factor N": "0.0000", "Elemental drift factor C": "0.0000"}
TWO = ("%C", "1σ(%C)", "%N", "1σ(%N)", "C:N molar", "1σ(C:N molar)", "C:N mass", "1σ(C:N mass)",
       "d13C (permil)", "d15N (permil)", "%C as measured", "%N as measured")
# Results columns 5 to 16, the reported values and their 1σ, are green as on the workbook's sheet.
GREEN = ("%C", "1σ(%C)", "%N", "1σ(%N)", "C:N molar", "1σ(C:N molar)", "C:N mass", "1σ(C:N mass)",
         "d13C (permil)", "1σ(d13C)", "d15N (permil)", "1σ(d15N)")


@dataclass
class BatchRun:
    """One corrected run: the file it came from, its positions and fits, and its configuration."""
    file: str
    positions: list
    fits: dict
    config: Config
    run_date: str | None = None
    label: str = ""
    index: int = 0
    extra: dict = field(default_factory=dict)
    run_info: dict = field(default_factory=dict)       # Operator, Instrument, Project, Notes (0.3.0)
    config_notes: list = field(default_factory=list)   # what the configuration set for another run (0.4.1)
    config_name: str | None = None                     # the name of its configuration, as the page gives it (0.5.0)


def label_runs(runs: list[BatchRun]) -> list[BatchRun]:
    """Number the runs from 1 and name each by its file name without the extension; a name that
    two files share is told apart by its number, as in 20260915_raw (2)."""
    seen: dict[str, int] = {}
    for i, r in enumerate(runs, start=1):
        r.index = i
        base = Path(r.file).stem or f"run {i}"
        seen[base.casefold()] = seen.get(base.casefold(), 0) + 1
        r.label = base if seen[base.casefold()] == 1 else f"{base} ({seen[base.casefold()]})"
    return runs


def _both(n: bool, c: bool, yes: str = "N and C", no: str = "no") -> str:
    return yes if n and c else "N only" if n else "C only" if c else no


def run_summary(r: BatchRun) -> dict:
    """The run's row on the Runs sheet."""
    F, P = r.fits, r.positions
    roles = [p.role for p in P]
    counts = F.get("rerun", {}).get("counts", {})
    drift = _both(F["drift_n"]["in_force"], F["drift_c"]["in_force"])
    models = {F[f"drift_{el}"]["model"] for el in ("n", "c") if F[f"drift_{el}"]["in_force"]}
    if models:
        drift += " (" + " and ".join("against position" if m == "Row" else "against the background"
                                     for m in sorted(models, key=lambda m: m != "Row")) + ")"
    return {
        "Run": r.label, "File": r.file, "Run date": r.run_date or None, "Positions": len(P),
        "Samples": roles.count("Sample"), "Anchor analyses": roles.count("Anchor"),
        "Check analyses": roles.count("Check"),
        "Blank subtracted": _both(F["blank_n"]["in_force"], F["blank_c"]["in_force"]),
        "Linearity applied": _both(F["lin_n"]["in_force"], F["lin_c"]["in_force"]),
        "Drift applied": drift,
        "Elemental drift corrected": _both(F["elem_n"]["in_force"], F["elem_c"]["in_force"]),
        "Calibration slope d15N": F["cal_n"].get("slope"), "Calibration slope d13C": F["cal_c"].get("slope"),
        "Samples to rerun": sum(counts.get(k, 0) for k in ("N", "C", "N and C")),
        "Samples to check": counts.get("Check", 0),
        "Problems of the whole run": F.get("rerun", {}).get("run_level_text", ""),
    }


AS_WRITTEN, SPACES_OUT = "as written", "spaces ignored"


def _repeat_groups(runs: list[BatchRun]) -> list[tuple[tuple, list]]:
    """The names found in two or more runs, each as (key, analyses), where each analysis is (run, position,
    name, how): how its name matched the group's name, AS_WRITTEN or SPACES_OUT.

    Names are compared as the sample sheet compares them (ea_irms.metadata, finding W-12): as written first,
    without regard to case, spaces at either end, or runs of spaces (norm); and a name that no name in
    another run fits that way is compared again with every space taken out, so that 366 - H+S and 366-H+S
    are one sample. Within one key without spaces, a name already found in two runs as written keeps its own
    group, and a name of one run joins the first of them; with none, the names of one run each that differ
    only in their spaces form one group. A name whose runs no other name reaches stays out."""
    written: dict[tuple, list] = {}           # (role, name as written) -> [(run, position, name)]
    for r in runs:
        for p in r.positions:
            if p.role == "Sample":
                name = (p.identifier or "").strip()
            elif p.role == "Check" and p.std_index is not None:
                name = str(r.config.standards[p.std_index][0]).strip()
            else:
                continue
            key = (p.role, norm(name))
            if not key[1]:
                continue
            written.setdefault(key, []).append((r, p, name))
    runs_of = {k: {r.index for r, _, _ in v} for k, v in written.items()}
    nospace: dict[tuple, list] = {}           # (role, name without spaces) -> keys, in order of appearance
    for key in written:
        nospace.setdefault((key[0], key[1].replace(" ", "")), []).append(key)
    out = []
    for keys in nospace.values():
        paired = [k for k in keys if len(runs_of[k]) >= 2]
        alone = [k for k in keys if len(runs_of[k]) < 2
                 and any(runs_of[o] - runs_of[k] for o in keys if o != k)]
        for k in paired:
            out.append((k, [(r, p, n, AS_WRITTEN) for r, p, n in written[k]]))
        if alone and paired:
            out[len(out) - len(paired)][1].extend((r, p, n, SPACES_OUT) for k in alone for r, p, n in written[k])
        elif alone:
            head = alone[0]
            out.append((head, [(r, p, n, AS_WRITTEN if k == head else SPACES_OUT)
                               for k in alone for r, p, n in written[k]]))
    return sorted(out, key=lambda g: (g[0][0] != "Sample", g[0][1]))


def repeats(runs: list[BatchRun]) -> list[dict]:
    """Samples and check standards measured in more than one run, one entry per analysis.

    A sample is known by its identifier and a check standard by its row of the standards table, since each
    run may write a number after its name; the names are compared as the sample sheet compares them, as
    written and, when nothing fits that way, with the spaces taken out (_repeat_groups). Only names found in
    two or more runs are listed, with every analysis of them, including repeats within one run, and each
    entry says how its name matched the group's (Name matched). The mean is over the analyses that have the
    value."""
    out = []
    for _, group in _repeat_groups(runs):
        items = [(r, p, name) for r, p, name, _ in group]
        how = {(r.index, p.row): h for r, p, _, h in group}
        if len({r.index for r, _, _ in items}) < 2:
            continue
        stats = {}
        for attr in ("d15n_final", "d13c_final"):
            vals = [getattr(p, attr) for _, p, _ in items if getattr(p, attr) is not None]
            mean = sum(vals) / len(vals) if vals else None
            sd = math.sqrt(sum((v - mean) ** 2 for v in vals) / (len(vals) - 1)) if len(vals) >= 2 else None
            stats[attr] = (mean, sd, (max(vals) - min(vals)) if len(vals) >= 2 else None, len(vals))
        for r, p, name in sorted(items, key=lambda t: (t[0].index, t[1].row)):
            def minus(attr):
                v, (mean, _, _, n) = getattr(p, attr), stats[attr]
                return v - mean if v is not None and mean is not None and n >= 2 else None
            out.append({
                "Name": items[0][2], "Type": p.role, "Runs": len({x.index for x, _, _ in items}),
                "Analyses": len(items), "Run": r.label, "Row": p.row, "Identifier": p.identifier,
                "d15N (permil)": p.d15n_final, "1σ(d15N)": p.u_d15n, "d15N minus mean": minus("d15n_final"),
                "d13C (permil)": p.d13c_final, "1σ(d13C)": p.u_d13c, "d13C minus mean": minus("d13c_final"),
                "Mean d15N": stats["d15n_final"][0], "SD d15N": stats["d15n_final"][1],
                "Mean d13C": stats["d13c_final"][0], "SD d13C": stats["d13c_final"][1],
                "%N": p.pct_n_reported, "%C": p.pct_c_reported, "C:N molar": p.cn_molar_reported,
                "Rerun": p.rerun, "Flags": "; ".join(p.flags), "Name matched": how[(r.index, p.row)],
            })
    return out


def sheet_columns(runs: list[BatchRun], sheet):
    """The sample sheet's columns the combined table writes: those with a value for a sample of any run."""
    return output_columns(sheet, [r.positions for r in runs])


def combined_rows(runs: list[BatchRun], sheet=None):
    """Every position of every run, in run order: the Run label, then the run's Results row, then the
    sample sheet's columns when one was given."""
    cols = sheet_columns(runs, sheet)
    for r in runs:
        for p in r.positions:
            yield [r.label] + results_values(p) + output_values(p, sheet, cols)


def combined_csv(runs: list[BatchRun], sheet=None) -> str:
    """The combined results table as CSV: Run, then the Results columns at full precision."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(csv_cells(["Run"] + results_header("1sigma") + output_header(sheet, sheet_columns(runs, sheet))))
    for row in combined_rows(runs, sheet):
        w.writerow(csv_cells(row))
    return buf.getvalue()


def _table(ws, top: int, header: list[str], rows, widths: dict, formats: dict | None = None):
    for j, h in enumerate(header, start=1):
        c = put(ws, top, j, h)      # sample-sheet column names are text (ea_irms 1.4.10)
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 12)
    n = 0
    for i, row in enumerate(rows, start=top + 1):
        n += 1
        for j, v in enumerate(row, start=1):
            c = put(ws, i, j, v)     # every value here is data, never a formula
            c.font = F_SM
            fmt = (formats or {}).get(header[j - 1])
            if fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = fmt
    return n


def combined_xlsx(runs: list[BatchRun], config_name: str | None = None, sheet=None) -> bytes:
    """The combined results as an Excel file: README, Results, Runs, and Repeats. A run that carries its
    configuration's name (BatchRun.config_name, from config_names, isoharness 0.5.0) gets it in a Configuration
    column of Runs."""
    wb = openpyxl.Workbook()
    rd = wb.active
    rd.title = "README"
    rd.sheet_view.showGridLines = False
    rd.column_dimensions["A"].width = 3
    rd.column_dimensions["B"].width = 110
    reps = repeats(runs)
    names = sorted({e["Name"] for e in reps}, key=str.casefold)
    spaced = {e["Name"] for e in reps if e["Name matched"] == SPACES_OUT}
    lines = [
        ("title", f"Combined results of {len(runs)} run{'s' if len(runs) != 1 else ''}"),
        ("note", f"Written by isoharness with ea_irms {VERSION}. Each run was corrected on its own, with the "
                 "standards in that run and the choices made for it."),
        ("h1", "Sheets"),
        ("text", "Results: every position of every run, in the order the runs were given, with the run's name "
                 "in column A. The other columns are those of the run's Results sheet in its workbook, each value "
                 "with its 1σ (one standard deviation), then Flags, the rerun mark with its reason and a suggested "
                 "mass, the elemental drift factors, and %C and %N as measured. The values are numbers at full "
                 "precision; each run's workbook holds the formulas behind them."),
        ("text", "Runs: one row per run, with its file, run date, counts of positions, which corrections were "
                 "in force, the calibration slopes, the samples marked for a rerun or a check, and the problems "
                 "the run's Rerun sheet lists for the whole run; then the run information typed on the page "
                 "(Operator, Instrument, Project, Notes), with a sample sheet how many samples matched it, and "
                 "Configuration notes when the configuration set positions for another run or named none."),
        ("text", "Repeats: samples, and check standards, measured in more than one of these runs, one row per "
                 "analysis. A sample is matched by its identifier, and a check standard by its row of the "
                 "standards table, as the sample sheet matches names: as written first, ignoring case, spaces "
                 "at either end, and runs of spaces, and when no name in another run fits that way, with every "
                 "space taken out, so that 366 - H+S and 366-H+S are one sample. When a name was matched the "
                 "second way, a last column, Name matched, says how each analysis's name matched the name in "
                 "column A. Every analysis of a matched name is listed, "
                 "including repeats within one run. 'minus mean' is the analysis's value less the mean of that "
                 "name's analyses with a value, and SD is their standard deviation (n − 1); these are arithmetic "
                 "on the values in Results, and each analysis keeps its own 1σ, which leaves out differences "
                 "between runs."),
        ("h1", "These runs"),
    ]
    for r in runs:
        s = run_summary(r)
        lines.append(("text", f"{r.index}. {r.label}: {r.file}"
                              + (f", run of {r.run_date}" if r.run_date else "") + f", {s['Positions']} positions, "
                              f"{s['Samples']} samples."))
    # W-11: the page names each run's configuration when a run has one of its own, as in "lab.xlsx for runs
    # 1 and 3; own.json for run 2"
    lines.append(("text", "Configuration: " + (config_name or "the package defaults") + ". Each run's choices "
                          "were written into its configuration; the Runs sheet shows which corrections were in "
                          "force" + (", and in its Configuration column which configuration each run used"
                                     if any(r.config_name for r in runs) else "") + "."))
    if sheet is not None:
        cols = sheet_columns(runs, sheet)
        lines.append(("text", f"Sample sheet: {sheet.name}, ID column {sheet.id_column}, {len(sheet.rows)} rows. "
                              "Results carries, after its own columns, the ID each sample matched (Sample ID "
                              "(sheet)), how it matched (exact or prefix, with spaces ignored when only that fits, or not in "
                              "sheet), and "
                              + (", ".join(h for _, h in cols) if cols else "no other column")
                              + ". A sample matches a row when its name equals the ID or begins with it followed "
                                "by a character that is neither a letter nor a digit; the longest such ID wins. "
                                "Columns with no value for any sample of these runs are left out."))
    listed = [n + (" (spaces ignored)" if n in spaced else "") for n in names]
    lines.append(("text", f"Repeats: {len(names)} name{'s' if len(names) != 1 else ''} found in more than one run"
                          + (": " + ", ".join(listed[:12]) + (", and others" if len(names) > 12 else "") + "."
                             if names else ".")
                          + (f" {len(spaced)} of them matched only with the spaces taken out."
                             if spaced else "")))
    row = 1
    for kind, text in lines:
        c = rd.cell(row, 2, text)
        c.font = {"title": F_TITLE, "h1": F_H1, "note": F_NOTE}.get(kind, F_SM)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 2 if kind in ("title", "note") else 1

    # Results
    rs = wb.create_sheet("Results")
    mhead = output_header(sheet, sheet_columns(runs, sheet))
    header = ["Run"] + results_header("1σ") + mhead
    widths = dict(WIDTHS, Run=max([8] + [min(40, len(r.label) + 2) for r in runs]))
    widths.update({h: 16 if h == "Sample ID (sheet)" else 11 if h == "Matched by" else 14 for h in mhead})
    formats = dict({h: "0.00" for h in TWO}, **FORMATS)
    n = _table(rs, 1, header, combined_rows(runs, sheet), widths, formats)
    for j, h in enumerate(header, start=1):
        if h in GREEN:
            for i in range(2, n + 2):
                rs.cell(i, j).fill = FILL_OUT
    rs.freeze_panes = "D2"
    last = get_column_letter(len(header))
    rerun_col = get_column_letter(header.index("Rerun") + 1)
    if n:
        rs.auto_filter.ref = f"A1:{last}{n + 1}"
        rs.conditional_formatting.add(f"A2:{last}{n + 1}",
                                      FormulaRule(formula=[f'${rerun_col}2="Check"'], fill=FILL_CHECK))
        rs.conditional_formatting.add(f"A2:{last}{n + 1}",
                                      FormulaRule(formula=[f'AND(${rerun_col}2<>"",${rerun_col}2<>"Check")'],
                                                  fill=FILL_RERUN))

    # Runs, with the run information and the sheet's matches after the 0.2.0 columns (0.3.0)
    ru = wb.create_sheet("Runs")
    sums = [run_summary(r) for r in runs]
    fields = [f for f in RUN_FIELDS if any(r.run_info.get(f) for r in runs)]
    for r, srow in zip(runs, sums):
        for f in fields:
            srow[f] = r.run_info.get(f) or None
        if sheet is not None:
            srow["Samples matched"] = summary_text(r.fits.get("metadata", {}).get("match", {}))
    if any(r.config_name for r in runs):            # 0.5.0 (finding W-11): each run's own configuration
        for r, srow in zip(runs, sums):
            srow["Configuration"] = r.config_name or "the package defaults"
    if any(r.config_notes for r in runs):           # 0.4.1: only when a configuration said something
        for r, srow in zip(runs, sums):
            srow["Configuration notes"] = "; ".join(r.config_notes) or None
    rh = list(sums[0].keys()) if sums else ["Run"]
    rwid = {"Run": widths["Run"], "File": 30, "Run date": 11, "Positions": 9, "Samples": 9,
            "Anchor analyses": 9, "Check analyses": 9, "Blank subtracted": 10, "Linearity applied": 10,
            "Drift applied": 26, "Elemental drift corrected": 11, "Calibration slope d15N": 11,
            "Calibration slope d13C": 11, "Samples to rerun": 9, "Samples to check": 9,
            "Problems of the whole run": 60, "Operator": 16, "Instrument": 16, "Project": 18, "Notes": 40,
            "Samples matched": 40, "Configuration": 30, "Configuration notes": 60}
    _table(ru, 1, rh, ([s[h] for h in rh] for s in sums), rwid,
           {"Calibration slope d15N": "0.0000", "Calibration slope d13C": "0.0000"})
    ru.freeze_panes = "B2"

    # Repeats
    rp = wb.create_sheet("Repeats")
    ph = ["Name", "Type", "Runs", "Analyses", "Run", "Row", "Identifier", "d15N (permil)", "1σ(d15N)",
          "d15N minus mean", "d13C (permil)", "1σ(d13C)", "d13C minus mean", "Mean d15N", "SD d15N",
          "Mean d13C", "SD d13C", "%N", "%C", "C:N molar", "Rerun", "Flags"]
    if spaced:                  # W-12: only when a name was matched with its spaces taken out
        ph.append("Name matched")
    pwid = {"Name": 18, "Type": 8, "Runs": 6, "Analyses": 8, "Run": widths["Run"], "Row": 6, "Identifier": 18,
            "Rerun": 9, "Flags": 54, "Name matched": 14}
    pfmt = {h: "0.00" for h in ("d15N (permil)", "d13C (permil)", "d15N minus mean", "d13C minus mean",
                                "Mean d15N", "Mean d13C", "%N", "%C", "C:N molar")}
    pfmt.update({h: "0.000" for h in ("1σ(d15N)", "1σ(d13C)", "SD d15N", "SD d13C")})
    if reps:
        k = _table(rp, 1, ph, ([e[h] for h in ph] for e in reps), pwid, pfmt)
        rp.freeze_panes = "B2"
        rp.auto_filter.ref = f"A1:{get_column_letter(len(ph))}{k + 1}"
    else:
        rp.cell(1, 1, "No sample or check standard was measured in more than one of these runs.").font = F_NOTE
        rp.column_dimensions["A"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
