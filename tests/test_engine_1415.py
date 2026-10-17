"""ea_irms 1.4.15: the engine's own results for what 1.4.15 adds.

1. The size series sets the floor of the size flags. Walking up from the smallest linearity standard while three or
   more larger ones remain, a standard further from the line through the larger ones (after the blank step) than
   floor_offset_permil, or floor_sd_multiple times that line's residual SD where that is larger, moves the floor to
   the next standard. min_area_n and _c set the floor outright, and a run without a size series takes
   no_series_min_area_n and _c. A peak below the floor and not below the smallest standard is "N below series
   floor", one below a lab minimum "N below the minimum peak" (and C), both reasons for a rerun.
2. Blank capsules that gave no peak put the blank below detection.
3. The 1σ of the final d15N and d13C holds the Kragten terms of the blank, linearity, and drift corrections in force
   (kragten_corrections) and the drift monitors' SD in place of the repeatability where they scatter more
   (monitor_repeatability). Every acetanilide analysis is held against the repeatability the samples' 1σ carries
   (ace_scatter_check): where they scatter more, every sample with that final value is flagged, last, and marked for
   a check unless a rerun reason marks it, and the whole run lists it as a problem.
4. A drift monitor whose standards row gives an accepted d15N or d13C is flagged when further from it than
   check_offset_permil, with an item in the problems of the whole run.
5. The corrections fitted but not applied, each with how far applying it would move the samples' final values.

Each test runs on the 2026-09-15 test run or on a copy of it whose parsed positions are edited before process, as
test_review_fixes.py's R-09 test does: standards renamed, peak areas and delta values set to illustrative values.
The Kragten terms and the shifts of the corrections not applied are checked against the four steps written out
here, the 1σ of 1.4.14 against a numerical perturbation of the line through the anchors, and the F test against
scipy where it is installed and against closed forms where it is not.
"""
import io
import json
import math
import statistics
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from ea_irms.labconfig import RANGES, SETTINGS, LabConfigError, config_to_xlsx, read_config_json, read_lab_config
from ea_irms.parse import parse_export
from ea_irms.pipeline import (DEFAULT_STANDARDS, Config, _xl, config_from_dict, config_to_dict, f_p_right, process,
                              t_p_two_sided)

DATA = Path(__file__).resolve().parent / "data"
RUN = DATA / "20260915_raw.csv"
ACE = {"n": "acetanilide d15N scatter above 1σ", "c": "acetanilide d13C scatter above 1σ"}
KEYS = {"n": {"raw": "d15n", "area": "area_n", "acc": "acc15", "u_acc": "u_acc15", "drift": "d15n_drift",
              "final": "d15n_final", "u": "u_d15n", "name": "d15N"},
        "c": {"raw": "d13c", "area": "area_c", "acc": "acc13", "u_acc": "u_acc13", "drift": "d13c_drift",
              "final": "d13c_final", "u": "u_d13c", "name": "d13C"}}
FITS = {"blank": "blank", "linearity": "lin", "drift": "drift"}       # each correction's block in the fits
ACE_OFF = Config(ace_scatter_check="Off")
# the drift monitors given an accepted d13C and d15N (illustrative, near the mean of what they read, -26.61 and
# -5.24 permil), as tests/test_workbook_1415.py gives them
ACCEPTED = [tuple(s[:2]) + (-26.6, -5.2) + tuple(s[4:]) if s[0] == "Ace Drift" else s for s in DEFAULT_STANDARDS]


def run(config=None, edit=None):
    """The test run processed, after ``edit`` has changed its parsed positions, given as {row: Position}."""
    positions = parse_export(RUN)
    if edit is not None:
        edit({p.row: p for p in positions})
    return process(positions, config or Config())


def with_roles(**roles):
    """The default standards table with the rows of some roles given another (for example Linearity="Exclude")."""
    return [(s[0], roles.get(s[1], s[1])) + tuple(s[2:]) for s in DEFAULT_STANDARDS]


def _line(pairs):
    """Least squares through the pairs with both values: slope, intercept, the slope's SE, the residual SD."""
    pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
    xs, ys = [x for x, _ in pairs], [y for _, y in pairs]
    slope, intercept = statistics.linear_regression(xs, ys)
    n, mx = len(pairs), statistics.fmean(xs)
    sd = math.sqrt(sum((y - intercept - slope * x) ** 2 for x, y in pairs) / (n - 2)) if n > 2 else None
    se = sd / math.sqrt(sum((x - mx) ** 2 for x in xs)) if sd is not None else None
    return slope, intercept, se, sd


# ------------------------------------------------------------------------------------- 1. the floor of the size flags
SERIES_ROWS = (6, 7, 8, 9, 31, 35)          # the four Ace Lin capsules, then two samples made standards
AREAS = (5.0, 10.0, 20.0, 30.0, 40.0, 50.0)


def on_line(area):
    """An illustrative size effect: -5 permil at 30 Vs, 0.02 permil per Vs."""
    return -5.0 + 0.02 * (area - 30.0)


def series(offsets, areas=AREAS):
    """An edit that gives the test run a nitrogen size series on on_line, standard k at areas[k] Vs and offsets[k]
    permil off the line. Rows of SERIES_ROWS not needed are left as they were, an Ace Lin capsule among them
    excluded."""
    def edit(by):
        for k, row in enumerate(SERIES_ROWS):
            p = by[row]
            if k < len(areas):
                p.identifier, p.area_n, p.d15n = f"Ace Lin{k + 1}", areas[k], on_line(areas[k]) + offsets[k]
            elif p.identifier.startswith("Ace Lin"):
                p.identifier = "Ace Bypass"
    return edit


def walk_by_hand(pts, offset=0.3, multiple=3.0):
    """The walk up a series written out: pts [(area, delta)] smallest first; returns the floor and each step tested,
    (area, residual, limit, moved)."""
    steps, floor = [], pts[0][0]
    for i in range(len(pts) - 3):
        rest = pts[i + 1:]
        slope, intercept, _, sd = _line(rest)
        limit = offset if multiple is None else max(offset, multiple * sd)
        resid = pts[i][1] - (intercept + slope * pts[i][0])
        steps.append((pts[i][0], resid, limit, abs(resid) > limit))
        if abs(resid) <= limit:
            break
        floor = rest[0][0]
    return floor, steps


@pytest.mark.parametrize("offsets,areas,multiple,floor,moved", [
    # on the line: the smallest standard is the floor
    ((0, 0, 0, 0, 0, 0), AREAS, 3.0, 5.0, [False]),
    # the smallest 0.5 permil off a line whose own SD is 0: the offset, 0.3 permil, is the limit
    ((0.5, 0, 0, 0, 0, 0), AREAS, 3.0, 10.0, [True, False]),
    # the larger standards 0.15 permil about their line, so 3 x its SD, 0.49, is the limit and the smallest, 0.375 off
    # the line, stays the floor; with the offset alone the limit is 0.3 and it does not
    ((0.45, 0, 0.15, -0.15, 0.15, -0.15), AREAS, 3.0, 5.0, [False]),
    ((0.45, 0, 0.15, -0.15, 0.15, -0.15), AREAS, None, 10.0, [True, False]),
    # two standards off: the floor moves twice, and the walk stops at the first standard on its line
    ((2.0, 0.6, 0, 0, 0, 0), AREAS, 3.0, 20.0, [True, True, False]),
    # five standards: after two steps only two remain above 20 Vs, fewer than the three a step needs, so the walk ends
    ((4.0, 1.0, 0, 0, 0), AREAS[:5], 3.0, 20.0, [True, True]),
])
def test_the_floor_walk_on_a_constructed_series(offsets, areas, multiple, floor, moved):
    positions, F = run(Config(blank_correction="Off", floor_sd_multiple=multiple), series(offsets, areas))
    fl = F["floor_n"]
    pts = sorted((p.area_n, p.d15n_blank) for p in positions if p.role == "Linearity")
    assert pts == [(a, on_line(a) + o) for a, o in zip(areas, offsets)]
    by_hand, steps = walk_by_hand(pts, 0.3, multiple)
    assert (fl["smallest"], fl["floor"], fl["moved"]) == (5.0, floor, floor != 5.0) and by_hand == floor
    assert [s["moved"] for s in fl["steps"]] == [m for *_, m in steps] == moved
    for s, (area, resid, limit, _) in zip(fl["steps"], steps):
        assert s["area"] == area and s["n_line"] == sum(1 for a in areas if a > area)
        assert s["resid"] == pytest.approx(resid, abs=1e-9) and s["limit"] == pytest.approx(limit, abs=1e-9)
    L = F["size_limits"]
    assert (L["min_n"], L["min_n_source"]) == (floor, "series floor" if floor != 5.0 else "smallest standard")


def test_the_test_runs_series_and_its_flags():
    """The test run's four standards allow one step. Its smallest N2 standard, 8.889 Vs, reads 1.02 permil off the
    line through the three larger ones, more than 3 x that line's SD (0.80), so the floor is 12.058 Vs; the carbon
    series' smallest is 0.014 permil off its line and stays the floor."""
    positions, F = run(ACE_OFF)
    assert not F["blank_n"]["in_force"] and not F["blank_c"]["in_force"]   # the deltas after the blank are as read
    by = {p.row: p for p in positions}
    for el, floor, moved in (("n", 12.058, True), ("c", 37.895, False)):
        k = KEYS[el]
        pts = sorted((getattr(p, k["area"]), getattr(p, k["raw"])) for p in positions if p.role == "Linearity")
        by_floor, steps = walk_by_hand(pts)
        fl = F[f"floor_{el}"]
        assert len(fl["steps"]) == len(steps) == 1 and fl["floor"] == by_floor == floor and fl["moved"] == moved
        assert fl["steps"][0]["resid"] == pytest.approx(steps[0][1], abs=1e-9)
        assert fl["steps"][0]["limit"] == pytest.approx(steps[0][2], abs=1e-9)
    assert round(F["floor_n"]["steps"][0]["resid"], 2) == 1.02 and round(F["floor_n"]["steps"][0]["limit"], 2) == 0.80
    assert F["size_limits"]["min_n_source"] == "series floor" and F["size_limits"]["min_c_source"] == "smallest standard"
    # 21 (10.11 Vs) lies between the smallest standard and the floor, 40 (2.809 Vs) below the smallest, 32 (12.229
    # Vs) above the floor
    assert by[21].flags[0] == "N below series floor" and by[21].rerun == "N"
    assert by[40].flags[0] == "N below linearity range" and not any("series floor" in f for f in by[40].flags)
    assert not any(f.startswith("N below") for f in by[32].flags)

    # the smallest carbon standard moved 0.5 permil off its line raises the carbon floor to the next, 54.07 Vs
    def c_off(by):
        by[6].d13c += 0.5
    positions, F = run(ACE_OFF, c_off)
    by = {p.row: p for p in positions}
    assert (F["size_limits"]["min_c"], F["size_limits"]["min_c_source"]) == (54.07, "series floor")
    assert by[21].flags[:2] == ["N below series floor", "C below series floor"] and by[21].rerun == "N and C"
    assert by[21].rerun_why == "N below series floor; C below series floor; ran before the background settled"
    assert "C below linearity range" in by[23].flags                       # 18.131 Vs, below the smallest


def test_fewer_than_four_standards_or_no_offset_leave_the_smallest_as_the_floor():
    def three(by):
        by[6].identifier = "Ace Bypass"                                    # the series is 7, 8 and 9
    positions, F = run(ACE_OFF, three)
    assert F["floor_n"] == {"smallest": 12.058, "floor": 12.058, "steps": [], "moved": False}
    assert F["size_limits"]["min_n_source"] == "smallest standard"
    assert next(p for p in positions if p.row == 21).flags[0] == "N below linearity range"
    # floor_offset_permil empty: no walk, and the smallest standard is the floor, as before 1.4.15
    positions, F = run(Config(floor_offset_permil=None, ace_scatter_check="Off"))
    assert F["floor_n"]["steps"] == [] and F["size_limits"]["min_n"] == 8.889
    p21 = next(p for p in positions if p.row == 21)
    assert not any(f.startswith("N below") for f in p21.flags) and p21.rerun == "Check"


def test_a_minimum_given_sets_the_floor_outright():
    positions, F = run(Config(min_area_n=15.0, ace_scatter_check="Off"))
    L = F["size_limits"]
    assert (L["min_n"], L["min_n_source"], L["min_c_source"]) == (15.0, "configuration", "smallest standard")
    assert F["floor_n"]["moved"]                                          # the walk is still reported, and not used
    below = sorted(p.row for p in positions if "N below linearity range" in p.flags)
    assert below == sorted(p.row for p in positions if p.area_n is not None and p.area_n < 15.0)
    assert {21, 32, 39} <= set(below) and not any("series floor" in f for p in positions for f in p.flags)
    # a minimum of 0 is a value: no N peak is below it
    positions, F = run(Config(min_area_n=0.0))
    assert F["size_limits"]["min_n_source"] == "configuration"
    assert not any(f.startswith("N below") for p in positions for f in p.flags)


def test_a_run_without_a_size_series_takes_the_lab_minimum():
    no_series = with_roles(Linearity="Exclude")
    positions, F = run(Config(standards=no_series, ace_scatter_check="Off"))
    assert F["floor_n"] == F["floor_c"] == {"smallest": None, "floor": None, "steps": [], "moved": False}
    L = F["size_limits"]
    assert (L["min_n"], L["min_c"], L["min_n_source"], L["min_c_source"]) == (0, 0, None, None)
    assert not any(" linearity range" in f or "minimum peak" in f for p in positions for f in p.flags)
    positions, F = run(Config(standards=no_series, no_series_min_area_n=20.0, no_series_min_area_c=60.0,
                              ace_scatter_check="Off"))
    L = F["size_limits"]
    assert (L["min_n"], L["min_c"], L["min_n_source"], L["min_c_source"]) == (20.0, 60.0, "lab minimum", "lab minimum")
    for p in positions:
        n_low = p.area_n is not None and p.area_n < 20.0
        c_low = p.area_c is not None and p.area_c < 60.0
        assert ("N below the minimum peak" in p.flags) == n_low and ("C below the minimum peak" in p.flags) == c_low
        if p.role == "Sample" and (n_low or c_low):
            assert p.rerun == ("N and C" if n_low and c_low else "N" if n_low else "C"), p.row
    by = {p.row: p for p in positions}
    assert (by[21].rerun, by[21].rerun_why) == (
        "N and C", "N below the minimum peak; C below the minimum peak; ran before the background settled")
    assert (by[39].rerun, by[27].rerun, by[53].rerun) == ("N", "C", "N")   # 13.845 Vs of N2; 57.484 of CO2; no CO2
    # with a size series the lab minimum is not used
    _, F = run(Config(no_series_min_area_n=20.0, no_series_min_area_c=60.0))
    assert F["size_limits"]["min_n_source"] == "series floor" and F["size_limits"]["min_n"] == 12.058


# ------------------------------------------------------------------------------------- 2. the blank below detection
def test_blank_capsules_without_a_peak_put_the_blank_below_detection():
    _, F = run()
    # three blank capsules, 4, 5 and 56: two gave an N2 peak, none a CO2 peak
    assert [F["blank_n"][k] for k in ("capsules", "n", "below_detection")] == [3, 2, False]
    assert [F["blank_c"][k] for k in ("capsules", "n", "below_detection")] == [3, 0, True]
    assert not F["blank_c"]["in_force"] and F["blank_c"]["area"] == 0.0
    assert not any(u["correction"] == "blank" and u["el"] == "c" for u in F["unapplied"])

    def no_n2(by):
        for row in (4, 56):
            by[row].area_n = by[row].ampl_n = by[row].d15n = None
    _, F = run(Config(blank_correction="On"), no_n2)
    assert [F["blank_n"][k] for k in ("capsules", "n", "below_detection", "in_force")] == [3, 0, True, False]
    assert "blank" not in F["unc_n"]["corrections"]
    # with no blank capsule there is nothing below detection
    _, F = run(Config(standards=with_roles(Blank="Exclude")))
    assert all([F[f"blank_{el}"][k] for k in ("capsules", "n", "below_detection")] == [0, 0, False] for el in "nc")


# ------------------------------------------------------------------------------------- 3. the 1σ of the final values
def stages(positions, F, config, el, move=None, apply=()):
    """The values of element el after each of the engine's four steps, written out here: {step: {row: value}}.

    Blank: the mean area and delta of the blank capsules with a peak, applied to every larger peak but the blanks'.
    Linearity: the slope fitted on the size series after the blank, applied about the series' mean area. Drift: the
    slope fitted on the drift monitors in use after linearity (from the first drift position, without the positions
    left out), applied about their mean position (Row) or reference background (Background; a position without one
    is left as it is). Calibration: accepted on measured through the anchors. The run's own choice of what is in
    force holds; ``move`` is (correction, amount) added to the blank's delta or to that slope, every later step
    refitted; ``apply`` names corrections not in force to apply as fitted."""
    k = KEYS[el]
    kind, amount = move or (None, 0.0)

    def on(name):
        return F[f"{FITS[name]}_{el}"]["in_force"] or name in apply
    d = {p.row: getattr(p, k["raw"]) for p in positions}
    caps = [p for p in positions if p.role == "Blank" and getattr(p, k["area"]) is not None]
    if on("blank") and caps:
        ab = statistics.fmean(getattr(p, k["area"]) for p in caps)
        db = statistics.fmean(getattr(p, k["raw"]) for p in caps) + (amount if kind == "blank" else 0.0)
        for p in positions:
            a = getattr(p, k["area"])
            if d[p.row] is not None and p.role != "Blank" and a > ab:
                d[p.row] = (a * d[p.row] - ab * db) / (a - ab)
    out = {"blank": dict(d)}
    series = [p for p in positions if p.role == "Linearity" and getattr(p, k["area"]) is not None]
    pivot = statistics.fmean(getattr(p, k["area"]) for p in series)
    b = _line([(getattr(p, k["area"]), d[p.row]) for p in series])[0] if on("linearity") else 0.0
    b += amount if kind == "linearity" else 0.0
    for p in positions:
        if d[p.row] is not None:
            d[p.row] -= b * (getattr(p, k["area"]) - pivot)
    out["linearity"] = dict(d)
    out["monitors"] = monitor_points(positions, F, config, el, d)
    x_pivot = statistics.fmean(x for x, _ in out["monitors"])
    s = _line(out["monitors"])[0] if on("drift") else 0.0
    s += amount if kind == "drift" else 0.0
    for p in positions:
        x = drift_x(F, el, p)
        if d[p.row] is not None and x is not None:
            d[p.row] -= s * (x - x_pivot)
    m, c, _, _ = _line([(d[p.row], getattr(p, k["acc"])) for p in positions if p.role == "Anchor"])
    out["final"] = {p.row: (None if d[p.row] is None else m * d[p.row] + c) for p in positions}
    return out


def drift_x(F, el, p):
    return float(p.row) if F[f"drift_{el}"]["model"] == "Row" else p.ref_bgd_n


def monitor_points(positions, F, config, el, values):
    """(x, value) of the drift monitors in the fit: from the first drift position, without the positions left out."""
    first = config.first_drift_position or 0
    left_out = set(getattr(config, f"exclude_drift_{el}"))
    pts = [(drift_x(F, el, p), values[p.row]) for p in positions if p.role == "Drift" and p.row >= first
           and p.row not in left_out]
    return [(x, y) for x, y in pts if x is not None and y is not None]


def chain(positions, F, config, el, move=None, apply=()):
    """The final values {row: value} of stages."""
    return stages(positions, F, config, el, move, apply)["final"]


def moved_by(positions, F, config, el):
    """What each correction in force is moved by for its term, found here: the blank's delta by the SD of its
    capsules' deltas, the linearity and drift slopes by their standard errors."""
    k = KEYS[el]
    s = stages(positions, F, config, el)
    out = {}
    caps = [getattr(p, k["raw"]) for p in positions if p.role == "Blank" and getattr(p, k["area"]) is not None]
    if F[f"blank_{el}"]["in_force"] and len(caps) >= 2:
        out["blank"] = statistics.stdev(caps)
    if F[f"lin_{el}"]["in_force"]:
        out["linearity"] = _line([(getattr(p, k["area"]), s["blank"][p.row]) for p in positions
                                  if p.role == "Linearity"])[2]
    if F[f"drift_{el}"]["in_force"]:
        out["drift"] = _line(s["monitors"])[2]
    return out


def carbon_blank(by):
    """Each blank capsule given a small CO2 peak (illustrative values), which no run in the repository has."""
    for row, (area, d13c) in zip((4, 5, 56), ((1.10, -27.9), (1.25, -26.8), (0.95, -28.4))):
        by[row].area_c, by[row].ampl_c, by[row].d13c = area, 60.0, d13c


KRAGTEN_CASES = {
    "defaults": (Config(), None, {("n", "drift"), ("c", "linearity"), ("c", "drift")}),
    "all in force": (Config(blank_correction="On", linearity_correction="On", drift_correction="On"), None,
                     {("n", "blank"), ("n", "linearity"), ("n", "drift"), ("c", "linearity"), ("c", "drift")}),
    "carbon blank": (Config(blank_correction="On", linearity_correction="On"), carbon_blank,
                     {("n", "blank"), ("n", "linearity"), ("n", "drift"), ("c", "blank"), ("c", "linearity"),
                      ("c", "drift")}),
    "Background model": (Config(drift_model="Background", drift_correction="On", linearity_correction="On"), None,
                         {("n", "linearity"), ("n", "drift"), ("c", "linearity"), ("c", "drift")}),
    "first position, left out": (Config(first_drift_position=20, exclude_drift_n=[38], exclude_drift_c=[42],
                                        drift_correction="On"), None,
                                 {("n", "drift"), ("c", "linearity"), ("c", "drift")}),
}


@pytest.mark.parametrize("name", list(KRAGTEN_CASES))
def test_the_kragten_terms_are_a_perturbation_of_the_chain(name):
    config, edit, terms_in_force = KRAGTEN_CASES[name]
    positions, F = run(config, edit)
    off, _ = run(replace(config, kragten_corrections="Off"), edit)
    seen = set()
    for el in ("n", "c"):
        k = KEYS[el]
        base = chain(positions, F, config, el)
        for p in positions:                                   # the chain written here gives the engine's values
            v = getattr(p, k["final"])
            assert (v is None) == (base[p.row] is None) and (v is None or abs(v - base[p.row]) < 1e-9), (el, p.row)
        moves = moved_by(positions, F, config, el)
        assert set(F[f"unc_{el}"]["corrections"]) == set(moves), (el, moves)
        seen |= {(el, c) for c in moves}
        for correction, amount in moves.items():
            moved = chain(positions, F, config, el, (correction, amount))
            terms = []
            for p in positions:
                got = getattr(p, f"u_terms_{el}")[correction]
                want = None if base[p.row] is None else abs(moved[p.row] - base[p.row])
                assert (got is None) == (want is None), (el, correction, p.row)
                if want is not None:
                    assert abs(got - want) < 1e-9, (el, correction, p.row)
                    terms.append(want)
            assert F[f"unc_{el}"]["corrections"][correction] == pytest.approx(max(terms), abs=1e-12)
        # the terms add to the 1σ without them in quadrature
        for p, q in zip(positions, off):
            if getattr(p, k["u"]) is None:
                assert getattr(q, k["u"]) is None and getattr(q, f"u_terms_{el}") == {}
                continue
            extra = sum(t * t for t in getattr(p, f"u_terms_{el}").values() if t is not None)
            assert getattr(p, k["u"]) ** 2 == pytest.approx(getattr(q, k["u"]) ** 2 + extra, rel=1e-12), (el, p.row)
    assert seen == terms_in_force


def calibration_sigma(positions, el):
    """1σ of each position's final value from the calibration alone, the 1σ of 1.4.14: a Kragten perturbation of the
    line through the anchors (each material's accepted value by its 1σ, the mean of its analyses by their SD over
    the square root of their number) and the anchors' pooled SD times the slope, combined in quadrature."""
    k = KEYS[el]
    anchors = [p for p in positions if p.role == "Anchor" and getattr(p, k["drift"]) is not None]
    xs, ys = [getattr(p, k["drift"]) for p in anchors], [getattr(p, k["acc"]) for p in anchors]
    groups = {}
    for p in anchors:
        groups.setdefault(p.std_index, []).append(getattr(p, k["drift"]))
    pooled = math.sqrt(sum((len(g) - 1) * statistics.variance(g) for g in groups.values())
                       / sum(len(g) - 1 for g in groups.values()))
    slope, _ = statistics.linear_regression(xs, ys)
    eps = 1e-6

    def at(x, idx, dx, dy):
        b, a = statistics.linear_regression([v + (dx if p.std_index == idx else 0.0) for v, p in zip(xs, anchors)],
                                            [v + (dy if p.std_index == idx else 0.0) for v, p in zip(ys, anchors)])
        return a + b * x
    out = {}
    for s in positions:
        x = getattr(s, k["drift"])
        if x is None:
            continue
        terms = [slope * pooled]
        for idx, g in groups.items():
            u_acc = next(getattr(p, k["u_acc"]) for p in anchors if p.std_index == idx) or 0.0
            terms.append((at(x, idx, 0.0, eps) - at(x, idx, 0.0, -eps)) / (2 * eps) * u_acc)
            u_mean = statistics.stdev(g) / math.sqrt(len(g))
            terms.append((at(x, idx, eps, 0.0) - at(x, idx, -eps, 0.0)) / (2 * eps) * u_mean)
        out[s.row] = math.sqrt(sum(t * t for t in terms))
    return out


def test_kragten_off_gives_the_1sigma_of_1414():
    positions, F = run(Config(kragten_corrections="Off", monitor_repeatability="Off"))
    on, G = run()
    for el in ("n", "c"):
        k = KEYS[el]
        assert F[f"unc_{el}"]["corrections"] == {} and all(getattr(p, f"u_terms_{el}") == {} for p in positions)
        assert not G[f"unc_{el}"]["monitors"]["used"]          # with the defaults the monitors' SD is not in use
        sigma = calibration_sigma(positions, el)
        assert len(sigma) > 40
        for p, q in zip(positions, on):
            if p.row not in sigma:
                assert getattr(p, k["u"]) is None and getattr(q, k["u"]) is None
                continue
            assert abs(getattr(p, k["u"]) - sigma[p.row]) < 1e-9, (el, p.row)
            # the defaults: the terms of the corrections in force on top
            extra = sum(t * t for t in getattr(q, f"u_terms_{el}").values() if t is not None)
            assert extra > 0 and getattr(q, k["u"]) ** 2 == pytest.approx(getattr(p, k["u"]) ** 2 + extra, rel=1e-12)


def monitors_kept(positions, F, config, el):
    """The final values of the drift monitors that give their SD: from the first drift position, without the
    positions left out of the fit or the outliers."""
    first = config.first_drift_position or 0
    left_out = set(getattr(config, f"exclude_drift_{el}")) | set(F[f"drift_{el}"]["outliers"])
    return [getattr(p, KEYS[el]["final"]) for p in positions if p.role == "Drift" and p.row >= first
            and p.row not in left_out and getattr(p, KEYS[el]["final"]) is not None]


def repeatability(positions, config, el):
    """The repeatability the samples' 1σ carries on the calibrated scale, found here, and the anchors' pooled degrees
    of freedom: the configuration's repeatability, or else the anchors' pooled SD, times the calibration slope."""
    k = KEYS[el]
    groups = {}
    for p in positions:
        if p.role == "Anchor" and getattr(p, k["drift"]) is not None:
            groups.setdefault(p.std_index, []).append(getattr(p, k["drift"]))
    df = sum(len(g) - 1 for g in groups.values())
    pooled = math.sqrt(sum((len(g) - 1) * statistics.variance(g) for g in groups.values()) / df)
    u = getattr(config, f"u_sample_{el}")
    slope = _line([(getattr(p, k["drift"]), getattr(p, k["acc"])) for p in positions if p.role == "Anchor"])[0]
    return abs((pooled if u is None else u) * slope), df


@pytest.mark.parametrize("switch,u_sample_c,used_c", [
    ("Off", None, False),
    ("Auto", None, False),             # d13C: the monitors' SD 0.044 permil against 0.021, larger, but p 0.083
    ("On", None, True),                # On: larger is enough
    ("Auto", 0.01, True),              # against a repeatability of 0.01 permil the F test finds them larger
    ("Off", 0.01, False),
])
def test_the_drift_monitors_scatter_in_place_of_the_repeatability(switch, u_sample_c, used_c):
    config = Config(monitor_repeatability=switch, u_sample_c=u_sample_c, kragten_corrections="Off")
    positions, F = run(config)
    base, _ = run(replace(config, monitor_repeatability="Off"))
    for el, used in (("n", False), ("c", used_c)):
        mon = F[f"unc_{el}"]["monitors"]
        values = monitors_kept(positions, F, config, el)
        sd = statistics.stdev(values)
        rep, df = repeatability(positions, config, el)
        assert (mon["n"], mon["df"], mon["switch"]) == (13, 4, switch)
        assert mon["sd"] == pytest.approx(sd, rel=1e-12) and mon["repeatability"] == pytest.approx(rep, rel=1e-12)
        assert mon["p"] == pytest.approx(f_p_right((sd / rep) ** 2, len(values) - 1, df), rel=1e-12)
        assert mon["larger"] == (sd > rep) and mon["used"] == used
        assert el == "c" or sd < rep            # d15N: 0.140 permil against 0.173, never in use
        for p, q in zip(positions, base):
            u, u0 = getattr(p, KEYS[el]["u"]), getattr(q, KEYS[el]["u"])
            if u is None:
                assert u0 is None
                continue
            assert u ** 2 == pytest.approx(u0 ** 2 + (sd * sd - rep * rep if used else 0.0), rel=1e-12), (el, p.row)
    p_c = F["unc_c"]["monitors"]["p"]                # the F test against 0.01 permil, or against the anchors' 0.021
    assert (p_c < 0.05) if u_sample_c else (0.05 < p_c < 0.1)


def test_the_monitors_left_out_are_not_counted():
    def outlier(by):
        by[30].d13c += 1.0                       # an illustrative monitor 1 permil off the others
    config = Config(first_drift_position=20, exclude_drift_c=[42], monitor_repeatability="On")
    positions, F = run(config, outlier)
    assert F["drift_c"]["outliers"] == [30] and F["drift_n"]["outliers"] == []
    values = monitors_kept(positions, F, config, "c")
    # thirteen monitors less the three before position 20 (10, 14 and 18), 42 left out, and 30, an outlier
    assert len(values) == F["unc_c"]["monitors"]["n"] == 8 and F["unc_n"]["monitors"]["n"] == 10
    assert F["unc_c"]["monitors"]["sd"] == pytest.approx(statistics.stdev(values), rel=1e-12)


def test_the_acetanilide_check_fires_on_the_test_run():
    """The size series ran before the background settled and reads up to 2.7 permil from the drift monitors in
    d15N, so the seventeen acetanilide analyses scatter more than the repeatability the samples carry."""
    off, _ = run(ACE_OFF)
    positions, F = run()
    a = F["ace_n"]
    assert F["settling"]["settled_from"] == 27 and F["settling"]["linearity_before"] == 4
    monitors = statistics.fmean(p.d15n_final for p in positions if p.role == "Drift")
    assert 2.7 < max(abs(p.d15n_final - monitors) for p in positions if p.role == "Linearity") < 2.8
    finals = [p.d15n_final for p in positions if p.role in ("Linearity", "Drift")]
    sd = statistics.stdev(finals)
    rep, df = repeatability(positions, Config(), "n")
    p_value = f_p_right((sd / rep) ** 2, len(finals) - 1, df)
    assert (a["switch"], a["n"], a["n_series"], a["n_monitors"], a["df"]) == ("On", 17, 4, 13, 4)
    assert a["sd"] == pytest.approx(sd, rel=1e-12) and a["basis"] == pytest.approx(rep, rel=1e-12)
    assert a["p"] == pytest.approx(p_value, rel=1e-12) and p_value < 0.05 and a["over"]
    # d13C: 0.043 permil against 0.021, p 0.086, not at the significance level
    assert not F["ace_c"]["over"] and F["ace_c"]["p"] > 0.05
    # the flag, last, on every sample with a final d15N and nowhere else; a check unless a rerun reason marks it
    assert sum(1 for q in positions if q.role == "Sample" and q.d15n_final is not None) == 26
    for p, q in zip(off, positions):
        has = q.role == "Sample" and q.d15n_final is not None
        assert q.flags == p.flags + ([ACE["n"]] if has else []), q.row
        assert q.rerun == ((p.rerun or "Check") if has else p.rerun), q.row
        if has:
            assert q.rerun_why == "; ".join(filter(None, [p.rerun_why, ACE["n"]])), q.row
            assert q.rerun_mass == p.rerun_mass
    assert F["rerun"]["counts"] == {"N": 2, "C": 1, "N and C": 6, "Check": 17}
    item = ("acetanilide d15N, size series and drift monitors, SD 0.759 permil against the samples' repeatability "
            "0.173 (p 0.006): the samples' d15N marked for a check")
    assert item == (f"acetanilide d15N, size series and drift monitors, SD {_xl(sd, 3)} permil against the samples' "
                    f"repeatability {_xl(rep, 3)} (p {_xl(p_value, 3)}): the samples' d15N marked for a check")
    assert F["rerun"]["run_level"] == [item]


def test_the_acetanilide_check_off_or_within_its_limit():
    positions, F = run(ACE_OFF)
    assert F["ace_n"]["switch"] == "Off" and F["ace_n"]["p"] is None and not F["ace_n"]["over"]
    assert not any(f in ACE.values() for p in positions for f in p.flags) and F["rerun"]["run_level"] == []
    # with the monitors' SD in use for d13C the basis is theirs, with their degrees of freedom, and the seventeen
    # analyses scatter no more
    positions, F = run(Config(monitor_repeatability="On"))
    a, mon = F["ace_c"], F["unc_c"]["monitors"]
    assert mon["used"] and a["basis"] == mon["sd"] and a["df"] == mon["n"] - 1 == 12
    assert a["p"] == pytest.approx(f_p_right((a["sd"] / a["basis"]) ** 2, 16, 12), rel=1e-12) and not a["over"]
    assert not any(ACE["c"] in p.flags for p in positions)


def test_both_elements_and_the_order_of_the_problems():
    # against a d13C repeatability of 0.01 permil, with the monitors' SD not in its place, d13C fires too
    positions, F = run(Config(u_sample_c=0.01, monitor_repeatability="Off"))
    assert F["ace_n"]["over"] and F["ace_c"]["over"]
    for p in positions:
        if p.role == "Sample":
            tail = [ACE[el] for el in "nc" if getattr(p, KEYS[el]["final"]) is not None]
            assert p.flags[-len(tail):] == tail, p.row
    assert next(p for p in positions if p.row == 53).flags[-1] == ACE["n"]       # no CO2 peak, no d13C
    assert [t.split(",")[0] for t in F["rerun"]["run_level"]] == ["acetanilide d15N", "acetanilide d13C"]
    # after the check standards and the drift monitors off their accepted values, before the anchors outside k σ
    data = json.loads((DATA / "config_check_standard.json").read_text())
    data["standards"] = [r[:2] + [-26.6, -5.2] + r[4:] if r[0] == "Ace Drift" else r for r in data["standards"]]
    data.update(check_offset_permil=0.05, coverage_k=0.5)
    _, F = run(config_from_dict(data))
    items = F["rerun"]["run_level"]
    at = [next(k for k, t in enumerate(items) if t.startswith(head))
          for head in ("check standard off by", "drift monitor off", "acetanilide d15N,", "anchors outside 0.5σ")]
    assert at == [0, 1, 2, 3], items


# ------------------------------------------------------------------------------------- 4. monitors off an accepted value
def test_a_drift_monitor_off_its_accepted_value():
    positions, F = run(Config(standards=ACCEPTED, check_offset_permil=0.15))
    flagged = {}
    for p in positions:
        mine = [f for f in p.flags if f.startswith("monitor ")]
        if p.role != "Drift":
            assert not mine
            continue
        want = [f"monitor {name} off by {_xl(got - acc, 2, sign=True)}"
                for got, acc, name in ((p.d15n_final, -5.2, "d15N"), (p.d13c_final, -26.6, "d13C"))
                if abs(got - acc) > 0.15]
        assert mine == want and p.rerun is None, p.row
        if mine:
            flagged[p.row] = mine
    assert flagged == {10: ["monitor d15N off by +0.20"], 30: ["monitor d15N off by -0.18"],
                       38: ["monitor d15N off by -0.31"]}
    assert F["rerun"]["run_level"][0] == ("drift monitor off its accepted value by more than 0.15 permil at "
                                          "positions 10, 30, 38")
    _, F = run(Config(standards=ACCEPTED, check_offset_permil=0.25))
    assert F["rerun"]["run_level"][0] == "drift monitor off its accepted value by more than 0.25 permil at position 38"
    # d13C too, at a tighter offset: 14 reads 0.08 permil under -26.6, and 0.03 over -5.2
    positions, _ = run(Config(standards=ACCEPTED, check_offset_permil=0.07))
    assert [f for f in next(p for p in positions if p.row == 14).flags if f.startswith("monitor ")] == [
        "monitor d13C off by -0.08"]
    # a standards table without accepted values for the monitors holds them against none
    positions, F = run(Config(check_offset_permil=0.0))
    assert not any(f.startswith("monitor ") for p in positions for f in p.flags)
    assert not any(t.startswith("drift monitor") for t in F["rerun"]["run_level"])


# ------------------------------------------------------------------------------------- 5. corrections not applied
def shift_range(positions, F, config, el, correction):
    """How far applying a correction not in force moves the samples' final values: (lowest, highest, how many)."""
    base, moved = chain(positions, F, config, el), chain(positions, F, config, el, apply=(correction,))
    shifts = [moved[p.row] - base[p.row] for p in positions if p.role == "Sample" and base[p.row] is not None]
    return min(shifts), max(shifts), len(shifts)


def move_text(lo, hi, name):
    return f"applied, it would move the samples' {name} by {_xl(lo, 2, sign=True)} to {_xl(hi, 2, sign=True)} permil"


def test_corrections_fitted_but_not_applied():
    config = Config()
    positions, F = run(config)
    lines = F["rerun"]["unapplied"]
    assert [(u["el"], u["correction"], u["mode"]) for u in F["unapplied"]] == [("n", "blank", "Auto"),
                                                                              ("n", "linearity", "Auto")]
    assert F["rerun"]["unapplied_text"] == "; ".join(lines) and len(lines) == 2
    # the blank of the two capsules that gave an N2 peak, 4 and 56, which Auto leaves out (fewer than three)
    caps = [p for p in positions if p.role == "Blank" and p.area_n is not None]
    area, d, sd = (statistics.fmean(p.area_n for p in caps), statistics.fmean(p.d15n for p in caps),
                   statistics.stdev(p.d15n for p in caps))
    lo, hi, n = shift_range(positions, F, config, "n", "blank")
    assert n == 26 and lines[0] == (f"d15N blank of 2 capsules ({_xl(area, 2)} Vs, {_xl(d, 2, sign=True)} permil, "
                                    f"SD {_xl(sd, 2)}) not applied (Auto); " + move_text(lo, hi, "d15N"))
    assert lines[0] == ("d15N blank of 2 capsules (1.87 Vs, +4.50 permil, SD 15.95) not applied (Auto); applied, "
                        "it would move the samples' d15N by -10.44 to +0.56 permil")
    # the nitrogen linearity slope, fitted on the series as read (the blank is not in force) and not significant at
    # 0.05
    slope, _, se, _ = _line([(p.area_n, p.d15n) for p in positions if p.role == "Linearity"])
    lo, hi, _ = shift_range(positions, F, config, "n", "linearity")
    assert lines[1] == (f"d15N linearity fitted, not applied (Auto): slope {_xl(slope, 4)} per Vs, SE {_xl(se, 4)}, "
                        f"p {_xl(t_p_two_sided(slope / se, 2), 3)}; " + move_text(lo, hi, "d15N"))
    assert lines[1] == ("d15N linearity fitted, not applied (Auto): slope -0.0570 per Vs, SE 0.0149, p 0.062; "
                        "applied, it would move the samples' d15N by -2.65 to +1.01 permil")
    # the drift correction off: each drift slope, fitted on the thirteen monitors by position, after the others
    config = Config(drift_correction="Off")
    positions, F = run(config)
    assert [(u["el"], u["correction"], u["mode"]) for u in F["unapplied"]] == [
        ("n", "blank", "Auto"), ("n", "linearity", "Auto"), ("n", "drift", "Off"), ("c", "drift", "Off")]
    for line, el in zip(F["rerun"]["unapplied"][2:], "nc"):
        k = KEYS[el]
        s = stages(positions, F, config, el)
        slope, _, se, _ = _line(s["monitors"])
        lo, hi, n = shift_range(positions, F, config, el, "drift")
        assert line == (f"{k['name']} drift (Row model) fitted, not applied (Off): slope {_xl(slope, 4)} per position, "
                        f"SE {_xl(se, 4)}, p {_xl(t_p_two_sided(slope / se, 11), 3)}; " + move_text(lo, hi, k["name"]))
        assert n == (26 if el == "n" else 25)                                # 53 has no CO2 peak
    assert F["rerun"]["unapplied"][3].startswith("d13C drift (Row model) fitted, not applied (Off): slope 0.0026 per "
                                                 "position")


def test_the_unapplied_lines_in_their_other_forms():
    # the Background model reads per mV, and a carbon switch of its own names its mode
    _, F = run(Config(drift_model="Background", drift_correction="Off", linearity_correction_c="Off"))
    lines = F["rerun"]["unapplied"]
    assert lines[0].startswith("d15N blank of 2 capsules (1.87 Vs, +4.50 permil, SD 15.95) not applied (Auto); ")
    assert [ln.split(":")[0] for ln in lines[1:]] == [
        "d15N linearity fitted, not applied (Auto)", "d15N drift (Background model) fitted, not applied (Off)",
        "d13C linearity fitted, not applied (Off)", "d13C drift (Background model) fitted, not applied (Off)"]
    assert " per mV, SE " in lines[2] and " per mV, SE " in lines[4] and " per Vs, SE " in lines[3]

    # a blank not applied while the linearity correction is: its shift has the slope refitted on the series so
    # corrected, as the chain written here refits it
    config = Config(blank_correction="Off", linearity_correction="On")
    positions, F = run(config)
    lo, hi, _ = shift_range(positions, F, config, "n", "blank")
    assert F["rerun"]["unapplied"][0] == ("d15N blank of 2 capsules (1.87 Vs, +4.50 permil, SD 15.95) not applied "
                                          "(Off); " + move_text(lo, hi, "d15N"))
    after = stages(positions, F, config, "n", apply=("blank",))["blank"]
    refit = _line([(p.area_n, after[p.row]) for p in positions if p.role == "Linearity"])[0]
    assert abs(refit - F["lin_n"]["applied"]) > 0.01                     # the refit moves the slope here

    # a blank of one capsule has no SD
    def one_capsule(by):
        by[4].area_n = by[4].ampl_n = by[4].d15n = None
    positions, F = run(Config(), one_capsule)
    p56 = next(p for p in positions if p.row == 56)
    assert F["rerun"]["unapplied"][0].startswith(
        f"d15N blank of 1 capsule ({_xl(p56.area_n, 2)} Vs, {_xl(p56.d15n, 2, sign=True)} permil) not applied (Auto); "
        "applied, it would move the samples' d15N by ")
    # no sample to move
    _, F = run(Config(standards=list(DEFAULT_STANDARDS) + [("Random", "Exclude")]))
    assert F["rerun"]["unapplied"][0].endswith("not applied (Auto); no sample has a d15N to move")
    # every correction fitted is in force
    _, F = run(Config(blank_correction="On", linearity_correction="On", drift_correction="On"))
    assert F["unapplied"] == [] and F["rerun"]["unapplied"] == [] and F["rerun"]["unapplied_text"] == "none"


# ------------------------------------------------------------------------------------- the F test
# scipy.stats.f.sf(f, df1, df2) as scipy 1.18.1 gives it. scipy is not among the requirements, and CI allows no
# skipped test, so these are the reference, and where scipy is installed the test also asks it again.
F_SF = {
    (1.0, 1, 1): 0.5000000000000001,
    (19.2, 16, 4): 0.005601008470794665,
    (4.0, 2, 30): 0.028844623295696835,
    (2.5, 12, 4): 0.19498661292400293,
    (0.5, 3, 10): 0.6906222455335577,
    (3.0, 5, 2): 0.26868270504761954,
    (100.0, 1, 1): 0.06345103486110713,
    (0.0, 4, 4): 1.0,
    (1e-06, 1, 1000): 0.9992023150184528,
    (7.3, 30, 100): 1.823651129856346e-14,
    (1.7, 2.5, 7.5): 0.2477806838580318,
}


@pytest.mark.parametrize("f,df1,df2", list(F_SF))
def test_f_p_right_is_the_f_distributions_right_tail(f, df1, df2):
    assert f_p_right(f, df1, df2) == pytest.approx(F_SF[(f, df1, df2)], rel=1e-9, abs=1e-15)
    try:
        from scipy import stats
    except ImportError:
        return
    assert F_SF[(f, df1, df2)] == pytest.approx(stats.f.sf(f, df1, df2), rel=1e-12, abs=1e-16)


def test_f_p_right_against_closed_forms():
    for f in (0.0, 0.3, 1.0, 2.0, 9.5, 40.0):
        for d in (1, 2, 3, 4, 7, 12, 30):
            # two degrees of freedom above: (d2 / (d2 + 2f))^(d2/2); two below: 1 - (d1 f / (d1 f + 2))^(d1/2)
            assert f_p_right(f, 2, d) == pytest.approx((d / (d + 2 * f)) ** (d / 2), rel=1e-10, abs=1e-15)
            assert f_p_right(f, d, 2) == pytest.approx(1 - (d * f / (d * f + 2)) ** (d / 2), rel=1e-10, abs=1e-15)
        # one degree of freedom above: the two-sided t test of the square root
        for d in (1, 4, 11):
            assert f_p_right(f, 1, d) == pytest.approx(t_p_two_sided(math.sqrt(f), d), rel=1e-10, abs=1e-15)
    for d in (1, 3, 8, 25):
        assert f_p_right(1.0, d, d) == pytest.approx(0.5, abs=1e-12)          # F(d, d) has its median at 1
    for bad in ((None, 2, 2), (1.0, None, 2), (1.0, 2, None), (1.0, 0, 2), (1.0, 2, 0.5), (-0.1, 2, 2)):
        assert f_p_right(*bad) is None


# ------------------------------------------------------------------------------------- the settings
NEW = {"floor_offset_permil": 0.5, "floor_sd_multiple": None, "no_series_min_area_n": 12.5, "no_series_min_area_c": 40.0,
       "kragten_corrections": "Off", "monitor_repeatability": "On", "ace_scatter_check": "Off"}


def _settings_row(ws, key):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 4).value == key)


def _xlsx(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_new_settings_round_trip():
    assert [getattr(Config(), k) for k in NEW] == [0.3, 3.0, None, None, "On", "Auto", "On"]
    kinds = {key: kind for key, _, kind, _ in SETTINGS if key}
    assert {k: kinds[k] for k in NEW} == {
        "floor_offset_permil": "number_or_empty", "floor_sd_multiple": "number_or_empty",
        "no_series_min_area_n": "number_or_empty", "no_series_min_area_c": "number_or_empty",
        "kragten_corrections": "onoff", "monitor_repeatability": "switch", "ace_scatter_check": "onoff"}
    cfg = Config(**NEW)
    d = config_to_dict(cfg)
    assert {k: d[k] for k in NEW} == NEW
    assert config_from_dict(json.loads(json.dumps(d))) == cfg
    assert read_config_json(json.loads(json.dumps(d))) == (cfg, [])
    assert read_lab_config(config_to_xlsx(cfg)) == (cfg, [])
    ws = openpyxl.load_workbook(io.BytesIO(config_to_xlsx(cfg)))["Settings"]
    assert ws.cell(_settings_row(ws, "floor_sd_multiple"), 2).value is None           # empty: the offset alone
    # the switches in any case, and two of them take On or Off only
    back = config_from_dict({"kragten_corrections": "off", "monitor_repeatability": "AUTO", "ace_scatter_check": " on "})
    assert (back.kragten_corrections, back.monitor_repeatability, back.ace_scatter_check) == ("Off", "Auto", "On")
    assert config_from_dict({"kragten_corrections": False}).kragten_corrections == "Off"
    for key, bad in (("kragten_corrections", "Auto"), ("ace_scatter_check", "Auto"), ("monitor_repeatability", "yes")):
        with pytest.raises(ValueError):
            Config(**{key: bad})
        with pytest.raises(LabConfigError, match=key):
            read_config_json({key: bad})
    # a lab configuration written before 1.4.15 takes the defaults, and says so
    wb = openpyxl.load_workbook(io.BytesIO(config_to_xlsx(cfg)))
    st = wb["Settings"]
    for r in range(st.max_row, 0, -1):
        if st.cell(r, 4).value in NEW:
            st.delete_rows(r)
    back, notes = read_lab_config(_xlsx(wb))
    assert [getattr(back, k) for k in NEW] == [getattr(Config(), k) for k in NEW]
    assert any(t.startswith("not in the file, so the package defaults were used: ") and all(k in t for k in NEW)
               for t in notes)


@pytest.mark.parametrize("key,bad,good", [
    ("floor_offset_permil", 0.0, 0.05),            # a limit of 0 would move the floor at any residual
    ("floor_offset_permil", -0.3, 2.0),
    ("floor_sd_multiple", -1.0, 0.0),
    ("no_series_min_area_n", -0.5, 0.0),
    ("no_series_min_area_c", -2.0, 0.0),
])
def test_the_new_numbers_are_checked(key, bad, good):
    assert key in RANGES
    with pytest.raises(LabConfigError, match=key):
        read_config_json({key: bad})
    assert getattr(read_config_json({key: good})[0], key) == good
    wb = openpyxl.load_workbook(io.BytesIO(config_to_xlsx(Config())))
    row = _settings_row(wb["Settings"], key)
    wb["Settings"].cell(row, 2).value = bad
    with pytest.raises(LabConfigError) as err:
        read_lab_config(_xlsx(wb))
    assert err.value.problems == [f"Settings B{row} ({key}) is {bad:g}; it needs a number "
                                  f"{'above' if key == 'floor_offset_permil' else 'of at least'} 0"]
    wb["Settings"].cell(row, 2).value = good
    assert getattr(read_lab_config(_xlsx(wb))[0], key) == good
