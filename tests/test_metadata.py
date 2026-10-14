"""ea_irms 1.4.8: the sample sheet and the run information (ea_irms.metadata).

A sample sheet is read from .xlsx or .csv: the sheet named Samples or the first, the header
found below a title line, the ID column chosen or guessed. Each Sample position matches a
row when its identifier equals the row's ID or begins with it followed by a character that
is neither a letter nor a digit, the longest such ID winning (Tony, 2026-10-13); when no ID
fits as written, the spaces are taken out of both and they are compared again (Tony, the same
day, after PD1-H missed the coral sheet's PD 1). Replicates are matched samples of one run with
the same identifier or the same Replicate of.
"""
import datetime as dt
import io
import math

import openpyxl
import pytest

from ea_irms.metadata import (CRITICAL_RANGE, MATCH_HEADER, NOT_MATCHED, SAMPLE_COLUMNS, SampleSheetError,
                              _range_quantile, clean_run_info, critical_range_factor, match_positions,
                              match_summary, norm, output_columns, output_header, output_values,
                              read_sample_sheet, sample_names, summary_text, template_xlsx)


class P:
    """A position as the matching sees it."""

    def __init__(self, row, identifier, role="Sample"):
        self.row, self.identifier, self.role = row, identifier, role


def csv_sheet(text, name="samples.csv", id_column=None):
    return read_sample_sheet(text.encode("utf-8"), name, id_column)


def xlsx_bytes(sheets):
    """An Excel file from {sheet name: rows}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ reading
def test_an_excel_sheet_with_a_title_line_and_dates():
    data = xlsx_bytes({
        "Notes": [["not this sheet"]],
        "Samples": [["Bocas corals, 2025"], [],
                    ["Coral ID", "Species", "site", "Collected", "Depth", None, "Notes", "Notes"],
                    ["CI 1-1", "Porites astreoides", "CI", dt.datetime(2025, 6, 14), 3.5, None, "a", "b"],
                    [61, "Siderastrea siderea", "Sids", dt.datetime(2025, 6, 15, 10, 30), 12, "x", None, None],
                    [None, "a row without an ID", None, None, None, None, None, None]],
    })
    sh = read_sample_sheet(data, "Bocas.xlsx")
    assert sh.sheet_name == "Samples" and sh.header_row == 3
    assert sh.columns == ["Coral ID", "Species", "site", "Collected", "Depth", "Column F", "Notes", "Notes (2)"]
    assert sh.id_column == "Coral ID"                       # a column with the word ID
    assert sh.rows[0][3] == "2025-06-14" and sh.rows[1][3] == "2025-06-15 10:30"
    assert sh.rows[0][4] == 3.5 and sh.sample_id(1) == "61"
    assert sh.notes == ["the header was read from row 3", "1 row without a value in Coral ID was left out"]


def test_a_text_sheet_with_semicolons_and_a_chosen_column():
    text = "Lab code;Sample name;Depth;Code\n001;CI 3-2;12.5;3-2\n002;CI 3-3;-4;7\n"
    sh = read_sample_sheet(text.encode("cp1252"), "lab.csv", id_column="sample NAME")
    assert sh.id_column == "Sample name"
    assert sh.rows[0] == ["001", "CI 3-2", 12.5, "3-2"]       # 001 and 3-2 stay text, 12.5 is a number
    assert sh.rows[1][2] == -4 and sh.rows[1][3] == 7
    assert read_sample_sheet(text.encode(), "lab.csv").id_column == "Sample name"     # guessed: a sample's name
    with pytest.raises(SampleSheetError, match="has no column named 'Mass'; its columns are Lab code, Sample name"):
        read_sample_sheet(text.encode(), "lab.csv", id_column="Mass")


def test_sheets_that_are_refused():
    with pytest.raises(SampleSheetError, match="not a sample sheet ea_irms reads"):
        read_sample_sheet(b"%PDF", "notes.pdf")
    with pytest.raises(SampleSheetError, match="is empty"):
        read_sample_sheet(b"", "empty.csv")
    with pytest.raises(SampleSheetError, match="has no row with a value in its ID column"):
        csv_sheet("Sample ID,Project\n,P1\n")
    with pytest.raises(SampleSheetError, match="could not be opened as an Excel workbook"):
        read_sample_sheet(b"not a zip", "broken.xlsx")


def test_duplicate_ids_keep_the_first_row():
    sh = csv_sheet("Sample ID,Project\nR1,first\nr1 ,second\nR2,third\n")
    assert sh.duplicates == {"r1": 2}
    assert sh.notes == ["the ID R1 is in 2 rows; the first is used"]
    assert sh.match("R1") == (0, "exact")


# ------------------------------------------------------------------ matching
SHEET = csv_sheet("Sample ID,Project\nCI 3-2,Bocas\nCI 3,Bocas\nHP-\nPD 1,Bocas\n61,Sids\nR1,lab\n")


@pytest.mark.parametrize("name, sid, how", [
    ("CI 3-2", "CI 3-2", "exact"),
    ("  ci   3-2 ", "CI 3-2", "exact"),           # case, spaces at either end, and runs of spaces
    ("CI 3-2-S", "CI 3-2", "prefix"),
    ("CI 3-2 H+S", "CI 3-2", "prefix"),
    ("CI 3-2_rep", "CI 3-2", "prefix"),
    ("CI 3-2.1", "CI 3-2", "prefix"),
    ("CI 3-20", "CI 3", "prefix"),                # not CI 3-2: the 0 continues the number
    ("CI 3-2-H", "CI 3-2", "prefix"),             # the longest ID that fits wins over CI 3
    ("HP-36", "HP-", "prefix"),                   # an ID that ends in a separator
    ("61-H+S", "61", "prefix"),                   # a number in the sheet
    ("PD1-H", "PD 1", "prefix, spaces ignored"),  # the sheet writes PD 1; no ID fits as written
    ("PD1", "PD 1", "exact, spaces ignored"),
    ("pd 1", "PD 1", "exact"),
    ("PD 1-S", "PD 1", "prefix"),
    ("CI3-2-H", "CI 3-2", "prefix, spaces ignored"),
    ("R 1", "R1", "exact, spaces ignored"),
    ("PD10", None, None),                         # without spaces the 0 still continues the number
    ("R 10", None, None),
    ("R10", None, None),
    ("CI", None, None),
    ("", None, None),
])
def test_matching(name, sid, how):
    i, got = SHEET.match(name)
    assert got == how and (SHEET.sample_id(i) if i is not None else None) == sid


def test_only_samples_are_matched_and_replicates_are_found():
    sh = csv_sheet("Sample ID,Project,Replicate of\nCI 3-2,Bocas,\nR1,lab,\nR2 a,lab,R2\nR2 b,lab,r2\nUSGS 64,x,\n")
    ps = [P(1, "USGS 64", "Anchor"), P(20, "CI 3-2-H"), P(21, "CI 3-2-S"), P(22, "R1"), P(23, "Random"),
          P(24, "r1"), P(25, "R2 a"), P(26, "R2 b"), P(27, "R2 c"), P(28, "Random"), P(29, "CI 3-2-H")]
    groups = match_positions(ps, sh)
    by = {p.row: p for p in ps}
    assert by[1].sheet_index is None and by[1].matched_by is None           # a standard is not matched
    assert [(p.row, p.sample_id, p.matched_by) for p in ps if p.sheet_index is not None] == [
        (20, "CI 3-2", "prefix"), (21, "CI 3-2", "prefix"), (22, "R1", "exact"), (24, "R1", "exact"),
        (25, "R2 a", "exact"), (26, "R2 b", "exact"), (29, "CI 3-2", "prefix")]
    assert groups == [
        {"group": 1, "key": "ci 3-2-h", "name": "CI 3-2-H", "rows": [20, 29], "by": "name"},
        {"group": 2, "key": "r1", "name": "R1", "rows": [22, 24], "by": "name"},
        {"group": 3, "key": "r2", "name": "R2", "rows": [25, 26], "by": "Replicate of"}]
    # fractions of one colony matched by a shorter ID, and unmatched names, are not replicates
    assert by[21].replicate_group is None and by[23].replicate_group is None and by[27].replicate_group is None
    m = match_summary(ps, sh)
    assert (m["samples"], m["matched"], m["exact"], m["prefix"]) == (10, 7, 4, 3)
    assert m["unmatched"] == [{"row": 23, "identifier": "Random"}, {"row": 27, "identifier": "R2 c"},
                              {"row": 28, "identifier": "Random"}]
    assert summary_text(m) == "7 of 10 samples matched (4 by the exact name, 3 by an ID the name begins with)"
    assert m["spaces"] == 0
    assert match_positions(ps, None) == [] and all(p.sheet_index is None for p in ps)
    assert summary_text(match_summary(ps, None)) == "no sample sheet"


def test_a_match_with_spaces_ignored_is_counted():
    ps = [P(1, "PD1-H"), P(2, "PD 1-S"), P(3, "CI 3-2"), P(4, "PD 2")]
    match_positions(ps, SHEET)
    assert [p.matched_by for p in ps] == ["prefix, spaces ignored", "prefix", "exact", None]
    m = match_summary(ps, SHEET)
    assert (m["matched"], m["exact"], m["prefix"], m["spaces"]) == (3, 1, 2, 1)
    assert summary_text(m) == ("3 of 4 samples matched (1 by the exact name, 2 by an ID the name begins with; "
                               "1 of the 3 only with spaces ignored)")


def test_results_columns():
    sh = csv_sheet("Sample ID,Project,Site,Type,Rerun,1σ(x),Depth\nA,P,,core,yes,1,\nB,,,,,,7\n")
    ps = [P(1, "A"), P(2, "zzz"), P(3, "Ace Drift", "Drift")]
    match_positions(ps, sh)
    cols = output_columns(sh, [ps])                  # Site and Depth hold nothing for a matched sample
    assert [h for _, h in cols] == ["Project", "Type (sheet)", "Rerun (sheet)", "1σ(x) (sheet)"]
    assert output_header(sh, cols) == list(MATCH_HEADER) + [h for _, h in cols]
    assert output_values(ps[0], sh, cols) == ["A", "exact", "P", "core", "yes", 1]
    assert output_values(ps[1], sh, cols) == [None, NOT_MATCHED, None, None, None, None]
    assert output_values(ps[2], sh, cols) == [None] * 6
    assert output_columns(None, [ps]) == [] and output_header(None, []) == [] and output_values(ps[0], None, []) == []


# ------------------------------------------------------------------ the template
def test_template_lists_each_name_once_and_reads_back():
    names = sample_names([["CI 3-2-H", "ci 3-2-h ", "R1"], ["R1", "CI  3-2-S", ""]])
    assert names == ["CI 3-2-H", "R1", "CI 3-2-S"]
    wb = openpyxl.load_workbook(io.BytesIO(template_xlsx(names)))
    assert wb.sheetnames == ["Samples", "How to fill"]
    ws = wb["Samples"]
    assert [c.value for c in ws[1]] == list(SAMPLE_COLUMNS)
    assert [ws.cell(r, 1).value for r in range(2, 5)] == names and ws.cell(5, 1).value is None
    ws["B2"], ws["G3"], ws["I4"] = "Bocas", 12.5, "CI 3-2-H"
    buf = io.BytesIO()
    wb.save(buf)
    sh = read_sample_sheet(buf.getvalue(), "filled.xlsx")
    assert sh.id_column == "Sample ID" and len(sh.rows) == 3
    ps = [P(5, "CI 3-2-H"), P(6, "CI 3-2-S")]
    assert [g["rows"] for g in match_positions(ps, sh)] == [[5, 6]]       # CI 3-2-S repeats CI 3-2-H
    assert [h for _, h in output_columns(sh, [ps])] == ["Project", "Replicate of"]


# ------------------------------------------------------------------ run information
def test_run_information():
    assert clean_run_info({"operator": "  XT  Wang ", "NOTES": " two  spaces ", "Project": "", "Instrument": None}) == {
        "Operator": "XT Wang", "Notes": "two  spaces"}
    assert list(clean_run_info({"Notes": "n", "Operator": "o"})) == ["Operator", "Notes"]
    assert clean_run_info(None) == {} and clean_run_info({}) == {}
    assert len(clean_run_info({"Notes": "x" * 900})["Notes"]) == 500
    with pytest.raises(ValueError, match=r"fields ea_irms does not read \(Analyst\)"):
        clean_run_info({"Analyst": "XT"})
    with pytest.raises(ValueError, match="must be text"):
        clean_run_info({"Operator": ["XT"]})
    with pytest.raises(ValueError, match="must be a JSON object"):
        clean_run_info(["XT"])


# ------------------------------------------------------------------ the critical range of ISO 5725-6
def test_critical_range_factors():
    assert critical_range_factor(1) is None and critical_range_factor(0) is None
    z975 = 1.959963984540054
    assert abs(critical_range_factor(2) - z975 * math.sqrt(2)) < 1e-6        # r = 2.8 s, unrounded
    for n in (2, 3, 4, 10, 20):
        assert abs(CRITICAL_RANGE[n] - _range_quantile(n)) <= 1.5e-6, n
    assert [round(CRITICAL_RANGE[n], 1) for n in range(2, 11)] == [2.8, 3.3, 3.6, 3.9, 4.0, 4.2, 4.3, 4.4, 4.5]
    assert all(CRITICAL_RANGE[n] < CRITICAL_RANGE[n + 1] for n in range(2, 20))
    assert critical_range_factor(25) > CRITICAL_RANGE[20]
    assert norm(" A  b ") == "a b"
