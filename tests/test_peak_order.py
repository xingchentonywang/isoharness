"""ea_irms 1.4.11: a position's peaks are read in Peak Nr order, whatever order the export lists them in.

An export can list each position's peaks from the last to the first. Read in file order, the engine took
such an export's undesignated reference pulses for the sample peaks; from 1.4.11 it puts each position's
peaks in Peak Nr order first. These tests write the 2026-09-15 run with each position's records reversed,
which lists every position from Peak Nr 6 down to 1 with the TCD rows last, and shuffled, and hold the
engine, the service, and the workbook to the run as exported. The workbook test has LibreOffice compute
the formulas and skips without it.
"""
import csv
import json
import math
import random
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ea_irms.parse import parse_export            # noqa: E402
from ea_irms.pipeline import Config, process      # noqa: E402
from ea_irms.workbook import build_workbook       # noqa: E402
from isoharness.app import app                    # noqa: E402

DATA = Path(__file__).resolve().parent / "data"
RUN = DATA / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def _rows():
    with open(RUN, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], rows[1:]


def _write(dst, reorder):
    """Write the run with each position's records in the order ``reorder`` puts them, positions as exported."""
    head, body = _rows()
    i = head.index("Row")
    by_pos = {}
    for r in body:
        by_pos.setdefault(r[i], []).append(r)
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        for recs in by_pos.values():
            w.writerows(reorder(list(recs)))
    return dst


def _peak_numbers(path):
    """Each position's Peak Nr values in file order."""
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    i, k = rows[0].index("Row"), rows[0].index("Peak Nr")
    out = {}
    for r in rows[1:]:
        if r[k].strip():
            out.setdefault(r[i], []).append(float(r[k]))
    return out


def _same(a, b):
    if isinstance(a, float) and isinstance(b, float):
        return a == b or (math.isnan(a) and math.isnan(b))
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(_same(a[k], b[k]) for k in a)
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
        return len(a) == len(b) and all(_same(x, y) for x, y in zip(a, b))
    return a == b


@pytest.fixture(scope="module")
def reversed_run(tmp_path_factory):
    return _write(tmp_path_factory.mktemp("order") / "20260915_raw.csv", lambda recs: recs[::-1])


def test_the_reversed_copy_lists_every_position_from_its_last_peak(reversed_run):
    """The copy is the case the fix is for: in every position the Peak Nr values fall, as in the export
    that showed the problem, and the same records are all there."""
    original, flipped = _peak_numbers(RUN), _peak_numbers(reversed_run)
    assert original.keys() == flipped.keys() and len(original) == 56
    for pos, nr in flipped.items():
        assert nr == sorted(nr, reverse=True) and nr[0] > nr[-1], pos
        assert sorted(nr) == sorted(original[pos]), pos


def test_reversed_peaks_parse_as_exported(reversed_run):
    want, got = parse_export(RUN), parse_export(reversed_run)
    assert [p.row for p in got] == [p.row for p in want]
    for a, b in zip(got, want):
        assert a == b, (a.row, {k for k in vars(a) if getattr(a, k) != getattr(b, k)})
    assert sum(p.d15n is not None for p in want) >= 50 and sum(p.d13c is not None for p in want) >= 40


@pytest.mark.parametrize("seed", [1, 2, 3, 4, 5])
def test_shuffled_peaks_parse_as_exported(tmp_path, seed):
    rng = random.Random(seed)

    def shuffle(recs):
        rng.shuffle(recs)
        return recs
    got = parse_export(_write(tmp_path / f"shuffled_{seed}.csv", shuffle))
    assert got == parse_export(RUN)


@pytest.mark.parametrize("cfg", [Config(), Config(blank_correction="On", linearity_correction="On",
                                                   drift_correction="On", drift_model="Background")])
def test_reversed_peaks_process_as_exported(reversed_run, cfg):
    (pa, fa), (pb, fb) = process(parse_export(reversed_run), cfg), process(parse_export(RUN), cfg)
    assert len(pa) == len(pb)
    for a, b in zip(pa, pb):
        assert _same(vars(a), vars(b)), a.row
    assert _same(fa, fb)


def test_the_service_reads_reversed_peaks_as_exported(reversed_run):
    client = TestClient(app)
    got, want = (client.post("/api/process", files={"file": (RUN.name, path.read_bytes())}).json()
                 for path in (reversed_run, RUN))
    for key in ("positions", "fits", "port", "config_notes", "metadata"):
        assert _same(got[key], want[key]), key


@pytest.mark.skipif(not SOFFICE, reason="needs LibreOffice to compute the formulas")
def test_the_workbook_of_reversed_peaks_agrees_with_the_engine(reversed_run, tmp_path):
    """The workbook keeps the records in the export's order on Raw and finds each peak by its Peak Nr, so
    built from the reversed copy it gives the engine's numbers for the run as exported."""
    out = tmp_path / "reversed.xlsx"
    build_workbook(reversed_run, out, Config())
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path / "calc"),
                    str(out)], check=True, capture_output=True, timeout=600)
    wb = openpyxl.load_workbook(tmp_path / "calc" / "reversed.xlsx", data_only=True)
    raw, wk, ft, rs = wb["Raw"], wb["Work"], wb["Fits"], wb["Results"]
    head, _ = _rows()
    with open(reversed_run, newline="", encoding="utf-8-sig") as fh:
        body = list(csv.reader(fh))[1:]
    k = head.index("Peak Nr")
    for n, r in enumerate(body, start=2):                       # Raw is the export as given
        cell = raw.cell(n, 4).value
        assert (cell in (None, "")) if not r[k].strip() else cell == float(r[k]), n
    positions, fits = process(parse_export(RUN), Config())
    columns = {"area_n": "L", "ampl_n": "M", "bgd_n": "N", "d15n": "O", "area_c": "P", "ampl_c": "Q",
               "d13c": "R", "d15n_final": "AA", "d13c_final": "AF", "pct_n_measured": "AI",
               "pct_c_measured": "AJ", "cn_molar": "AK", "ref_bgd_n": "BI", "ref_pulse_d15n": "BT",
               "ref_pulse_d13c": "BW"}
    compared = 0
    for attr, col in columns.items():
        for p in positions:
            mine, theirs = getattr(p, attr), wk[f"{col}{5 + p.row}"].value
            theirs = theirs if isinstance(theirs, (int, float)) else None
            assert (mine is None) == (theirs is None), (p.row, attr, mine, theirs)
            if mine is not None:
                assert abs(mine - theirs) < 1e-9 * max(1.0, abs(mine)), (p.row, attr, mine, theirs)
                compared += 1
    assert compared > 500
    for cell, (key, sub) in {"B56": ("cal_n", "slope"), "B63": ("cal_c", "slope"),
                             "B201": ("ref_gas_n", "n"), "B208": ("ref_gas_c", "n")}.items():
        assert abs(ft[cell].value - fits[key][sub]) < 1e-9 * max(1.0, abs(fits[key][sub])), cell
    flags = [c for c in range(1, rs.max_column + 1) if rs.cell(4, c).value == "Flags"][0]
    for p in positions:
        text = rs.cell(4 + p.row, flags).value or ""
        assert set(p.flags) == {t.strip() for t in text.split(";") if t.strip()}, p.row
    assert json.dumps(sorted(p.row for p in positions)) == json.dumps(list(range(1, 57)))
