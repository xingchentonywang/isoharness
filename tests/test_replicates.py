"""ea_irms 1.4.8: replicates of a sample in one run, from the sample sheet, in the engine and the workbook.

Tony chose on 2026-10-13 to mark replicates that differ by more than the repeatability limit,
2.8 s for two analyses and the critical range of ISO 5725-6 for more, once a sample sheet says
which analyses hold one sample. The engine compares each group's range with f(n) x s, s the
repeatability of one analysis in use times the calibration slope, and marks every member for a
rerun of that element. The test run names its samples Random; here twelve of them are renamed so
that three groups form (by the same name, by Replicate of, and with one analysis's CO2 peak moved
0.3 permil), and two fractions of one colony and a near miss (CI 3-20) match without being
replicates. The workbook must agree with the engine after LibreOffice has computed it.
"""
import csv
import io
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from ea_irms.metadata import (NOT_MATCHED, critical_range_factor, output_columns, output_values,
                              read_sample_sheet, summary_text)
from ea_irms.parse import parse_export
from ea_irms.pipeline import Config, config_from_dict, process
from ea_irms.workbook import build_workbook
from isoharness.serialize import results_csv

DATA = Path(__file__).resolve().parent / "data"
RUN = DATA / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NAMES = {21: "R1", 24: "R1", 49: "r1 ", 31: "R2 a", 32: "R2 b", 33: "R2 c", 35: "R3", 48: "R3",
         36: "CI 3-2-H", 37: "CI 3-2-S", 39: "CI 3-20", 51: "Unknown 9"}
SHEET = ("Sample ID,Project,Site,Species or material,Fraction,Collection date,Depth,Notes,Replicate of,Type,Colony\n"
         "R1,Alpha,,acetanilide,,2025-06-14,10,first,,lab,\n"
         "R1,Beta,,ignored duplicate,,,,,,,\n"
         "R2 a,Alpha,,acetanilide,,,12.5,,R2,lab,\n"
         "R2 b,Alpha,,acetanilide,,,12.5,,R2,lab,\n"
         "R2 c,Alpha,,acetanilide,,,12.5,,r2,lab,\n"
         "R3,Gamma,,acetanilide,,,,,,,\n"
         "CI 3-2,Bocas,,Porites,,2025-07-01,3,,,coral,CI 3-2\n"
         "CI 3,Bocas,,Porites,,,,,,coral,CI 3\n"
         ",orphan row without an ID,,,,,,,,,\n")
INFO = {"Operator": "XT Wang", "Instrument": "Delta V Plus", "Notes": "test run, renamed"}


def variant(folder, names=NAMES, shift_c=((48, 0.3),)):
    """The test run with some samples renamed and the sample CO2 peak of others moved (permil)."""
    rows = list(csv.reader(io.StringIO(RUN.read_text(encoding="utf-8"))))
    h = rows[0]
    col = {k: h.index(k) for k in ("Row", "Identifier 1", "Ampl  44", "Is Ref _", "d 13C/12C", "Peak Nr")}
    for r in rows[1:]:
        if r and r[col["Row"]].strip().isdigit() and int(r[col["Row"]]) in names:
            r[col["Identifier 1"]] = names[int(r[col["Row"]])]
    for row, delta in shift_c:
        # the sample CO2 peak is the mass-44 peak before the first reference one (ea_irms.parse)
        recs = [r for r in rows[1:] if r and r[col["Row"]].strip() == str(row) and r[col["Peak Nr"]].strip()
                and len(r) > col["Ampl  44"] and r[col["Ampl  44"]].strip()]
        first_ref = next(i for i, r in enumerate(recs) if r[col["Is Ref _"]].strip() == "1")
        s = recs[first_ref - 1]
        s[col["d 13C/12C"]] = repr(float(s[col["d 13C/12C"]]) + delta)
    out = Path(folder) / "replicates.csv"
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    out.write_text(buf.getvalue(), encoding="utf-8")
    return out


@pytest.fixture(scope="module")
def export(tmp_path_factory):
    return variant(tmp_path_factory.mktemp("replicates"))


def sheet():
    return read_sample_sheet(SHEET.encode(), "samples.csv")


def test_groups_and_limits(export):
    positions, F = process(parse_export(export), Config(), samples=sheet())
    by = {p.row: p for p in positions}
    R = F["replicates"]
    s_n = F["unc_n"]["u_sample"] * abs(F["cal_n"]["slope"])
    s_c = F["unc_c"]["u_sample"] * abs(F["cal_c"]["slope"])
    assert R["on"] and R["s"] == {"n": s_n, "c": s_c}
    assert [(g["name"], g["rows"], g["by"]) for g in R["groups"]] == [
        ("R1", [21, 24, 49], "name"), ("R2", [31, 32, 33], "Replicate of"), ("R3", [35, 48], "name")]
    for g in R["groups"]:
        for el, attr, s in (("n", "d15n_final", s_n), ("c", "d13c_final", s_c)):
            vals = [getattr(by[r], attr) for r in g["rows"]]
            e = g[el]
            assert e["n"] == len(vals) and e["range"] == max(vals) - min(vals)
            assert e["factor"] == critical_range_factor(len(vals)) and e["limit"] == e["factor"] * s
            assert e["over"] == (e["range"] > e["limit"])
    over = {(g["name"], el) for g in R["groups"] for el in ("n", "c") if g[el]["over"]}
    assert over == {("R1", "c"), ("R2", "n"), ("R3", "c")}
    # the 0.3 permil moved into position 48 is what puts R3 over: unmoved, its d13C range is 0.015
    r3 = R["groups"][2]
    assert 0.25 < r3["c"]["range"] < 0.35 and r3["n"]["range"] < r3["n"]["limit"]
    # fractions of one colony match the colony's row and are not replicates; CI 3-20 is not CI 3-2
    assert (by[36].sample_id, by[36].matched_by, by[36].replicate_group) == ("CI 3-2", "prefix", None)
    assert (by[37].sample_id, by[39].sample_id) == ("CI 3-2", "CI 3")
    assert by[51].sheet_index is None and by[20].sheet_index is None
    m = F["metadata"]["match"]
    assert summary_text(m) == "11 of 26 samples matched (8 by the exact name, 3 by an ID the name begins with)"


def test_marks_and_reasons(export):
    plain, _ = process(parse_export(export), Config())
    positions, F = process(parse_export(export), Config(), samples=sheet())
    before = {p.row: p.rerun for p in plain}
    by = {p.row: p for p in positions}
    rows = "21, 24, 49"
    g1 = F["replicates"]["groups"][0]["c"]
    assert by[24].rerun == "C" and by[24].rerun_why == (
        f"ran before the background settled; replicates at positions {rows} differ by "
        f"{g1['range']:.2f} in d13C, above {g1['limit']:.2f}")
    assert before[24] == "Check"                                   # a check before, a rerun of C now
    g2 = F["replicates"]["groups"][1]["n"]
    for r in (31, 32, 33):
        assert by[r].rerun_why.endswith(f"replicates at positions 31, 32, 33 differ by {g2['range']:.2f} in d15N, "
                                        f"above {g2['limit']:.2f}")
    assert by[33].rerun == "N"                                     # 33's peaks above the range mark nothing (1.4.7)
    assert {by[35].rerun, by[48].rerun} == {"C"}
    changed = {r for r in before if before[r] != by[r].rerun}
    assert changed == {21, 24, 31, 32, 33, 35, 48, 49}
    # Off: the groups are still described, and nothing is marked for them
    off, F_off = process(parse_export(export), config_from_dict({"rerun_replicates": "Off"}), samples=sheet())
    assert not F_off["replicates"]["on"] and not any(g[el]["over"] for g in F_off["replicates"]["groups"]
                                                     for el in ("n", "c"))
    assert {p.row: p.rerun for p in off} == before
    assert config_from_dict({"rerun_replicates": False}).rerun_replicates == "Off"
    # no sheet, no groups
    _, F_none = process(parse_export(export), Config())
    assert F_none["replicates"]["groups"] == [] and F_none["metadata"]["sheet"] is None


def test_the_repeatability_in_use_sets_the_limit(export):
    positions, F = process(parse_export(export), config_from_dict({"u_sample_n": 0.05, "u_sample_c": 0.5}),
                           samples=sheet())
    assert F["replicates"]["s"]["n"] == 0.05 * abs(F["cal_n"]["slope"])
    over = {(g["name"], el) for g in F["replicates"]["groups"] for el in ("n", "c") if g[el]["over"]}
    assert over == {("R1", "n"), ("R2", "n")}       # both groups of three over the tighter d15N limit, no d13C


def test_results_table_carries_the_sheet(export):
    sh = sheet()
    positions, _ = process(parse_export(export), Config(), samples=sh)
    text = results_csv(positions, Config(), sh)
    head = next(csv.reader(io.StringIO(text)))
    assert head[24:] == ["Sample ID (sheet)", "Matched by", "Project", "Species or material", "Collection date",
                         "Depth", "Notes", "Replicate of", "Type (sheet)", "Colony"]
    rows = {r["Row"]: r for r in csv.DictReader(io.StringIO(text))}
    assert [rows["21"][h] for h in ("Sample ID (sheet)", "Matched by", "Project", "Collection date", "Depth")] == [
        "R1", "exact", "Alpha", "2025-06-14", "10"]
    assert rows["20"]["Matched by"] == NOT_MATCHED and rows["1"]["Matched by"] == ""
    assert rows["37"]["Colony"] == "CI 3-2" and rows["32"]["Depth"] == "12.5"
    plain = results_csv(process(parse_export(export), Config())[0], Config())
    assert [ln.split(",")[:3] for ln in plain.splitlines()[1:]] == [ln.split(",")[:3] for ln in text.splitlines()[1:]]


CASES = ("replicates", "off", "nosheet")


@pytest.fixture(scope="module")
def computed(export, tmp_path_factory):
    if not SOFFICE:
        pytest.skip("needs LibreOffice")
    folder = tmp_path_factory.mktemp("replicates_book")
    cases = {"replicates": (Config(), sheet(), INFO),
             "off": (config_from_dict({"rerun_replicates": "Off"}), sheet(), {}),
             "nosheet": (Config(), None, {})}
    files = []
    for name, (cfg, sh, info) in cases.items():
        out = folder / f"{name}.xlsx"
        build_workbook(export, out, cfg, samples=sh, run_info=info, file_name="renamed run.csv")
        files.append(out)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + [str(f) for f in files], check=True, capture_output=True, timeout=900)
    return folder, cases


def _same(mine, theirs):
    if mine is None or mine == "":
        return theirs in (None, "")
    if isinstance(mine, (int, float)) and not isinstance(mine, bool):
        return isinstance(theirs, (int, float)) and abs(mine - theirs) <= 1e-9 * max(1.0, abs(mine))
    return mine == theirs


@pytest.mark.parametrize("name", CASES)
def test_workbook_agrees_with_engine(computed, export, name):
    folder, cases = computed
    cfg, sh, info = cases[name]
    positions, F = process(parse_export(export), cfg, samples=sh)
    wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
    assert wb.sheetnames == ["README", "Raw", "Setup", "Work", "Fits", "Results", "Rerun", "Replicates", "QC",
                             "Metadata"]
    rs, rp, md, st = wb["Results"], wb["Replicates"], wb["Metadata"], wb["Setup"]
    head = [rs.cell(4, c).value for c in range(1, rs.max_column + 1)]
    cols = output_columns(sh, [positions])
    assert len(head) == 24 + (2 + len(cols) if sh else 0)
    for p in positions:
        i = 4 + p.row
        for h, attr in (("Rerun", "rerun"), ("Why", "rerun_why")):
            assert (getattr(p, attr) or "") == (rs.cell(i, head.index(h) + 1).value or ""), (name, p.row, h)
        for j, v in enumerate(output_values(p, sh, cols)):
            assert _same(v, rs.cell(i, 25 + j).value), (name, p.row, head[24 + j])
    assert st["B71"].value == cfg.rerun_replicates and st["A7"].value == "Operator"
    assert (st["B7"].value or None) == (info.get("Operator") or None)
    R = F["replicates"]
    assert _same(R["s"]["n"], rp["B4"].value) and _same(R["s"]["c"], rp["B5"].value)
    top = 9
    for g in R["groups"]:
        assert [rp.cell(top + k, 3).value for k in range(len(g["rows"]))] == g["rows"]
        for el, cols_ in (("n", "GHIJ"), ("c", "KLMN")):
            e = g[el]
            got = [rp[f"{c}{top}"].value for c in cols_]
            assert got[0] == e["n"] and _same(e["range"], got[1]) and _same(e["limit"], got[2]), (name, g["name"], el)
            assert (got[3] == "yes") == e["over"], (name, g["name"], el)
        assert rp.cell(top, 15).value == ", ".join(str(r) for r in g["rows"])
        top += len(g["rows"])
    if not R["groups"]:
        assert "No sample sheet was given" in rp.cell(9, 1).value
    # Metadata: the run information, then the sheet and this run's samples
    labels = {md.cell(r, 1).value: md.cell(r, 2).value for r in range(5, 12)}
    assert labels["File"] == "renamed run.csv" and labels["Run"] == "replicates"
    for k in ("Operator", "Instrument", "Project", "Notes"):
        assert (labels[k] or None) == (info.get(k) or None), (name, k)
    if sh is None:
        assert md["A14"].value == "No sample sheet was given."
    else:
        assert md["B19"].value == summary_text(F["metadata"]["match"]) and md["B20"].value == len(R["groups"])
        table = [[md.cell(r, c).value for c in range(1, 6)] for r in range(24, 50) if md.cell(r, 1).value]
        assert [r[0] for r in table] == [p.row for p in positions if p.role == "Sample"]
        assert [r[3] for r in table][:2] == [NOT_MATCHED, "exact"]
    errors = [x.coordinate for s in (rs, rp, md, wb["Rerun"], wb["QC"], wb["Fits"]) for row in s.iter_rows() for x in row
              if isinstance(x.value, str) and x.value.startswith("#") and x.value[1:4].isupper()]
    assert not errors, errors[:5]


def test_another_run_pasted_into_raw_keeps_no_stale_group(computed):
    """The groups belong to the run the sheet was matched to: Work keeps one only while the identifier
    at that position is the one matched."""
    folder, _ = computed
    wb = openpyxl.load_workbook(folder / "replicates.xlsx")
    wk = wb["Work"]
    col = next(c for c in range(1, wk.max_column + 1) if wk.cell(5, c).value == "rep group")
    assert wk.cell(5 + 21, col).value == '=IF(TRIM($B26)="R1",1,"")'
    assert wk.cell(5 + 49, col).value == '=IF(TRIM($B54)="r1",1,"")'
    assert wk.cell(5 + 20, col).value is None


def test_command_line_takes_a_sample_sheet(export, tmp_path, capsys):
    from ea_irms.cli import main
    path = tmp_path / "samples.csv"
    path.write_text(SHEET)
    out = tmp_path / "run.xlsx"
    assert main([str(export), "-o", str(out), "--samples", str(path), "--no-recalc"]) == 0
    assert ("sample sheet samples.csv: 11 of 26 samples matched (8 by the exact name, 3 by an ID the name begins "
            "with); 3 groups of replicates, 3 over the critical range") in capsys.readouterr().out
    assert openpyxl.load_workbook(out).sheetnames[-2:] == ["QC", "Metadata"]
    assert main([str(export), "-o", str(out), "--samples", str(path), "--id-column", "Mass", "--no-recalc"]) == 2
    assert "has no column named 'Mass'" in capsys.readouterr().err
