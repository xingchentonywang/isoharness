"""ea_irms 1.4.15: the workbook against the engine in the cases the reference runs do not show.

1.4.15 has the size series set the floor below which a peak is marked for a rerun (Fits rows 411 to 422 and the
walk up the series from row 486), a run without a size series take the lab minimum, blank capsules that gave no
peak put the blank below detection, the 1σ of the final d15N and d13C hold the corrections in force (a Kragten
term for each, Fits rows 424 to 451) and the drift monitors' scatter where it is larger, every acetanilide
analysis held against the samples' 1σ, a drift monitor held against an accepted value the standards table gives
it, and the corrections fitted but not applied listed on the Rerun sheet. The 19 reference runs carry none of the
Background drift model, accepted values for the drift monitors, a run without a size series with a lab minimum, a
blank in force, or a series longer than 12, so copies of the 2026-09-15 test run show them, with the three
switches, the floor's settings, a floor set outright (with a minimum of 0, under which a peak above the range was
left unflagged before 1.4.15), and a first drift position. Each workbook is computed by LibreOffice, all in one
call (skipped without it), and compared with the engine by scripts/reference_check.py's check_values, which holds
every value, flag, mark, and fit the reference check holds, 1.4.15's included.
"""
import csv
import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from ea_irms.parse import parse_export                                                  # noqa: E402
from ea_irms.pipeline import DEFAULT_STANDARDS, Config, process                          # noqa: E402
from ea_irms.workbook import RERUN_FIRST_ROW, RERUN_UNAPPLIED_ROW, build_workbook        # noqa: E402
from isoharness.serialize import clean, position_record                                  # noqa: E402
from port_cases import carbon_blank_variant                                              # noqa: E402
from reference_check import check_values                                                 # noqa: E402

RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
# the drift monitors given an accepted d13C and d15N (illustrative, near the mean of what they read, -26.61 and
# -5.24 permil), so that with a check offset of 0.15 permil three of the thirteen lie off it in d15N
ACCEPTED = [tuple(s[:2]) + (-26.6, -5.2) + tuple(s[4:]) if s[0] == "Ace Drift" else s for s in DEFAULT_STANDARDS]


def _rows():
    with open(RUN, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], rows[1:]


def renamed(dst, names):
    """The test run with the Identifier 1 of some positions changed."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    out = []
    for r in body:
        if r and r[i["Row"]].strip().isdigit() and int(r[i["Row"]]) in names:
            r = list(r)
            r[i["Identifier 1"]] = names[int(r[i["Row"]])]
        out.append(r)
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        w.writerows(out)
    return dst


def cases(folder):
    """{name: (export, configuration)}: copies of the test run, each showing a case."""
    folder = Path(folder)
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    ids = {int(r[i["Row"]]): r[i["Identifier 1"]].strip() for r in body if r and r[i["Row"]].strip().isdigit()}
    series = sorted(r for r, name in ids.items() if name.startswith("Ace Lin"))
    samples = sorted(r for r, name in ids.items() if name == "Random")
    assert len(series) == 4 and len(samples) >= 10
    # no size series: the four Ace Lin capsules renamed to the text of a row with the role Exclude
    no_series = renamed(folder / "no_series.csv", {r: "Ace Bypass" for r in series})
    # a size series of 14: ten samples renamed to linearity standards
    many = renamed(folder / "many.csv", {r: f"Ace Lin x{k}" for k, r in enumerate(samples[:10])})
    carbon_blank = carbon_blank_variant(RUN, folder / "carbon_blank.csv", Config().standards)
    return {
        "background": (RUN, Config(drift_model="Background", drift_correction="On")),
        "accepted": (RUN, Config(standards=ACCEPTED, check_offset_permil=0.15)),
        "no_series": (no_series, Config(no_series_min_area_n=20.0, no_series_min_area_c=60.0)),
        "no_series_no_limit": (no_series, Config()),
        "switches": (RUN, Config(kragten_corrections="Off", monitor_repeatability="On", ace_scatter_check="Off")),
        "floor": (RUN, Config(floor_offset_permil=0.01, floor_sd_multiple=None, linearity_correction="On")),
        "floor_off": (RUN, Config(floor_offset_permil=None, monitor_repeatability="Off")),
        "limits": (RUN, Config(min_area_n=15.0, min_area_c=0.0)),
        "carbon_blank": (carbon_blank, Config(blank_correction="On")),
        "first_drift": (RUN, Config(first_drift_position=20, exclude_drift_n=[40], drift_correction="On")),
        "many": (many, Config(floor_sd_multiple=None)),
    }


@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    if not SOFFICE:
        pytest.skip("needs LibreOffice to compute the formulas")
    folder = tmp_path_factory.mktemp("w1415")
    runs = cases(folder)
    books = []
    for name, (export, cfg) in runs.items():
        build_workbook(export, folder / f"{name}.xlsx", cfg)
        books.append(str(folder / f"{name}.xlsx"))
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + books, check=True, capture_output=True, timeout=1800)
    out = {}
    for name, (export, cfg) in runs.items():
        positions, F = process(parse_export(export), cfg)
        wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
        out[name] = (wb, positions, F)
    return out


@pytest.mark.parametrize("name", ["background", "accepted", "no_series", "no_series_no_limit", "switches", "floor",
                                  "floor_off", "limits", "carbon_blank", "first_drift", "many"])
def test_the_workbook_agrees_with_the_engine(computed, name):
    wb, positions, F = computed[name]
    body = json.loads(json.dumps({"positions": [position_record(p) for p in positions], "fits": clean(F)}))
    n, bad, errors = check_values(body, wb)
    assert not errors, errors[:5]
    assert not bad, (name, len(bad), bad[:8])
    assert n > 3000, n


def test_the_cases_show_what_they_are_for(computed):
    def of(name):
        return computed[name][1], computed[name][2]
    # the Background model, with the drift term of 1σ in force
    P, F = of("background")
    assert F["drift_n"]["model"] == F["drift_c"]["model"] == "Background" and "drift" in F["unc_n"]["corrections"]
    # drift monitors off their accepted value: flagged, and a problem of the whole run
    P, F = of("accepted")
    off = [p.row for p in P if any(f.startswith("monitor d15N off by ") for f in p.flags)]
    assert off == [10, 30, 38] and not any(f.startswith("monitor d13C") for p in P for f in p.flags)
    assert ("drift monitor off its accepted value by more than 0.15 permil at positions 10, 30, 38"
            in F["rerun"]["run_level"])
    # a run without a size series takes the lab minimum, or no limit
    P, F = of("no_series")
    assert F["size_limits"]["min_n_source"] == "lab minimum" and F["size_limits"]["min_n"] == 20.0
    assert any("N below the minimum peak" in p.flags for p in P if p.role == "Sample")
    P, F = of("no_series_no_limit")
    assert F["size_limits"]["min_n_source"] is None and F["size_limits"]["min_n"] == 0
    # the switches: no Kragten term, the monitors' SD in use wherever larger, no acetanilide check
    P, F = of("switches")
    assert F["unc_n"]["corrections"] == {} == F["unc_c"]["corrections"]
    assert F["unc_c"]["monitors"]["used"] and F["ace_n"]["p"] is None and not F["ace_n"]["over"]
    # the floor at 0.01 permil with no SD multiple, which raises the carbon floor too, and switched off
    P, F = of("floor")
    assert F["size_limits"]["min_c_source"] == "series floor" and F["lin_n"]["in_force"]
    assert any("C below series floor" in p.flags for p in P)
    P, F = of("floor_off")
    assert F["floor_n"]["steps"] == [] and F["size_limits"]["min_n_source"] == "smallest standard"
    # a floor set outright, and a minimum of 0, under which a peak above the range is still flagged
    P, F = of("limits")
    assert F["size_limits"]["min_n_source"] == "configuration" and F["size_limits"]["min_c"] == 0.0
    assert any("C above linearity range" in p.flags for p in P)
    # a carbon blank in force, with its Kragten term
    P, F = of("carbon_blank")
    assert F["blank_c"]["in_force"] and "blank" in F["unc_c"]["corrections"]
    # a series of 14 standards: the walk's table grows to hold them; without the SD multiple the nitrogen walk
    # moves the floor at each of its 11 steps, and the carbon walk stops at its eighth
    wb, P, F = computed["many"]
    assert [s["moved"] for s in F["floor_n"]["steps"]] == [True] * 11
    assert [s["moved"] for s in F["floor_c"]["steps"]] == [True] * 7 + [False]
    ft = wb["Fits"]
    top = next(r for r in range(486, ft.max_row + 1)
               if str(ft.cell(r, 1).value).startswith("Size series floor, carbon"))
    assert top == 488 + 14 + 2 and ft.cell(487, 1).value == "Rank by area" and ft.cell(488 + 13, 1).value == 14
    assert ft["B412"].value == 14


def test_the_layout(computed):
    wb, positions, F = computed["accepted"]
    st, rr = wb["Setup"], wb["Rerun"]
    assert [st[f"B{r}"].value for r in range(78, 86)] == [0.3, 3, None, None, None, "On", "Auto", "On"]
    assert st["A77"].value == "Size series floor (1.4.15)"
    assert st["A82"].value == "1σ of the final d15N and d13C (1.4.15)"
    assert rr["A7"].value == "Corrections fitted but not applied" and RERUN_UNAPPLIED_ROW == 8
    assert rr.cell(RERUN_FIRST_ROW - 1, 1).value == "Row"
    assert [rr.cell(RERUN_UNAPPLIED_ROW + k, 1).value for k in range(len(F["rerun"]["unapplied"]))] == \
        F["rerun"]["unapplied"]
