"""ea_irms 1.4.13: two checks of the whole run, and four workbook items.

The elemental drift correction against the anchors. The anchors carry certified %N and %C and the correction
does not use them, so they test it: their RMS deviation from certified as measured and as corrected, and a
Note or a Caution when the correction adds more than half the gate Auto applies it at (Setup B36), or all of
it. In PanamaSids2 the correction moved the anchors' %N from 0.65% RMS to 6.5%.

The anchors read as one material (finding W-13). Analyses of one material agree in their C/N peak-area ratio
and their raw delta values; in a real run every acetanilide capsule was labeled USGS 64 Lin, Drift, or Bypass
and was read as a USGS 64 anchor.

Each test writes a copy of the 2026-09-15 run that shows its case: three drift monitors beside the anchors
renamed USGS 64 Drift, and the drift monitor between the two anchor groups weighed about 10% light. The
workbook tests have LibreOffice compute the formulas and skip without it; the port test needs node.

The workbook items: the charts' overlay and axis settings, the README's text on C:N without a mass, on the
first-order propagation of 1σ and on a sample with neither peak, and the Replicates sheet when no group formed
(finding R-34).
"""
import csv
import json
import math
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from ea_irms.metadata import read_sample_sheet                                                   # noqa: E402
from ea_irms.parse import parse_export                                                           # noqa: E402
from ea_irms.pipeline import (ANCHOR_RATIO_LIMIT, ANCHOR_SPREAD_PERMIL, DEFAULT_STANDARDS, Config,  # noqa: E402
                              _xl, config_to_dict, process)
from ea_irms.workbook import NSTD, build_workbook                                                # noqa: E402
from isoharness.serialize import port_input                                                      # noqa: E402

RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NODE = shutil.which("node")
MIXED_TEXT = ("anchors labeled USGS 64 do not read as one material: C/N peak-area ratio 1.19 to 4.81, 4.02 times; "
              "raw d15N -5.95 to 1.27 permil; raw d13C -23.69 to -9.45 permil")
# Fits rows of 1.4.13: the anchors as one material from 369 (a row per line of the standards table from 373), the
# elemental drift correction against the anchors from 393
AM_TOP, AM_LEVEL, AM_TEXT = 373, 389, 390
EC_BASE = {"n": 394, "c": 399}
EC_TEXT = 404
AM_COLS = ["name", "n", "n_ratio", "ratio_min", "ratio_max", "ratio_x", "d15n_min", "d15n_max", "d15n_spread",
           "d13c_min", "d13c_max", "d13c_spread", "level"]


def _rows(path=RUN):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], rows[1:]


def _write(dst, head, body):
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        w.writerows(body)
    return dst


def _positions_of(body, i, test):
    out = []
    for r in body:
        if test(r[i["Identifier 1"]].strip()) and int(r[i["Row"]]) not in out:
            out.append(int(r[i["Row"]]))
    return out


def mixed_variant(dst):
    """The test run with the three drift monitors nearest the anchors, Ace Drift1 to Ace Drift3 at positions 10,
    14 and 18, renamed USGS 64 Drift1 to USGS 64 Drift3, which the match text USGS 64 reads as anchors."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    anchors = _positions_of(body, i, lambda s: s in ("USGS 64", "USGS 65"))
    monitors = _positions_of(body, i, lambda s: s.startswith("Ace Drift"))
    near = sorted(monitors, key=lambda r: (min(abs(r - a) for a in anchors), r))[:3]
    assert sorted(near) == [10, 14, 18]
    out = []
    for r in body:
        if int(r[i["Row"]]) in near:
            r = list(r)
            r[i["Identifier 1"]] = r[i["Identifier 1"]].replace("Ace Drift", "USGS 64 Drift")
        out.append(r)
    return _write(dst, head, out)


def light_variant(dst, factor=0.9):
    """The test run with the drift monitor between the two anchor groups, Ace Drift2 at position 14 between the
    USGS 64 analyses at 11 to 13 and the USGS 65 analyses at 15 to 17, weighed about 10% light (0.428 mg entered
    as 0.385), as a capsule weighed or typed wrongly would be."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    last64 = max(_positions_of(body, i, lambda s: s == "USGS 64"))
    first65 = min(_positions_of(body, i, lambda s: s == "USGS 65"))
    between = [r for r in _positions_of(body, i, lambda s: s.startswith("Ace Drift")) if last64 < r < first65]
    assert between == [14]
    out = []
    for r in body:
        if int(r[i["Row"]]) == 14 and r[i["Amount"]].strip():
            r = list(r)
            r[i["Amount"]] = f"{float(r[i['Amount']]) * factor:.3f}"
        out.append(r)
    return _write(dst, head, out)


# The anchors without a certified %N or %C, so the elemental correction (On) has no anchor to be tested against
UNCERTIFIED = Config(elemental_correction="On", standards=[
    tuple(s[:4]) + (None, None) + tuple(s[6:]) if s[1] == "Anchor" else s for s in DEFAULT_STANDARDS])


def _rms(values):
    return math.sqrt(sum(v * v for v in values) / len(values))


def _by_hand(positions, el):
    """The check of the elemental correction worked out here from the values as measured: each drift monitor's
    measured over certified value is its factor, interpolated in position between monitors and held beyond the
    first and the last; each anchor's corrected value is its measured value over the factor at its position."""
    meas, cert = f"pct_{el}_measured", f"pct_{el}"
    mons = sorted((p.row, getattr(p, meas) / getattr(p, cert)) for p in positions
                  if p.role == "Drift" and getattr(p, meas) is not None and getattr(p, cert))

    def factor(row):
        if row <= mons[0][0]:
            return mons[0][1]
        if row >= mons[-1][0]:
            return mons[-1][1]
        for (x0, f0), (x1, f1) in zip(mons, mons[1:]):
            if x0 <= row <= x1:
                return f0 + (row - x0) / (x1 - x0) * (f1 - f0)
    anchors = [p for p in positions if p.role == "Anchor" and getattr(p, meas) is not None]
    as_measured = _rms([getattr(p, meas) / getattr(p, cert) - 1 for p in anchors])
    corrected = _rms([getattr(p, meas) / factor(p.row) / getattr(p, cert) - 1 for p in anchors])
    return len(anchors), as_measured, corrected


def _level(added, gate_pct=5.0):
    return "Caution" if added > gate_pct / 100 else "Note" if added > gate_pct / 200 else "OK"


# ------------------------------------------------------------------------------------------ the engine
def test_the_test_run_passes_both_checks():
    positions, F = process(parse_export(RUN), Config())
    assert F["elem_n"]["in_force"] and not F["elem_c"]["in_force"]
    n = F["elem_check"]["n"]
    assert F["elem_check"]["c"] is None                     # the carbon correction is not in force
    assert (n["n"], n["level"], n["gate_pct"]) == (6, "OK", 5.0)
    assert round(n["rms_measured"] * 100, 2) == 1.15 and round(n["rms_reported"] * 100, 2) == 1.62
    assert n["added"] == n["rms_reported"] - n["rms_measured"]
    k, as_measured, corrected = _by_hand(positions, "n")
    assert k == 6 and abs(n["rms_measured"] - as_measured) < 1e-12 and abs(n["rms_reported"] - corrected) < 1e-12
    assert n["rms_reported"] == F["unc_elem"]["n"]["anchor_rms"]       # the trueness term reads the same anchors
    am = F["anchor_material"]
    assert am["level"] == "OK" and am["limits"] == {"ratio": ANCHOR_RATIO_LIMIT, "spread": ANCHOR_SPREAD_PERMIL}
    assert [(m["index"], m["name"], m["n"], m["positions"], m["level"]) for m in am["materials"]] == [
        (0, "USGS 64", 3, [11, 12, 13], "OK"), (1, "USGS 65", 3, [15, 16, 17], "OK")]
    for m in am["materials"]:
        by = {p.row: p for p in positions}
        ratios = [by[r].area_c / by[r].area_n for r in m["positions"]]
        assert (m["ratio_min"], m["ratio_max"]) == (min(ratios), max(ratios)) and m["ratio_x"] < 1.02
        assert m["d15n_spread"] < 0.4 and m["d13c_spread"] < 0.04 and m["over"] == []
    # neither check is a problem of the whole run. 1.4.15: the one problem the test run has is the acetanilide's
    # d15N scatter, the check 1.4.15 adds (test_engine_1415.py)
    assert [t.split(",")[0] for t in F["rerun"]["run_level"]] == ["acetanilide d15N"]


def test_anchors_of_two_materials_under_one_label(tmp_path):
    positions, F = process(parse_export(mixed_variant(tmp_path / "mixed.csv")), Config())
    assert [p.role for p in positions if p.row in (10, 14, 18)] == ["Anchor"] * 3
    am = F["anchor_material"]
    m64, m65 = am["materials"]
    assert am["level"] == "Caution" and (m64["level"], m65["level"]) == ("Caution", "OK")
    assert m64["n"] == m64["n_ratio"] == 6 and m64["positions"] == [10, 11, 12, 13, 14, 18]
    by = {p.row: p for p in positions}
    ratios = [by[r].area_c / by[r].area_n for r in m64["positions"]]
    assert m64["ratio_x"] == max(ratios) / min(ratios) and 4.0 < m64["ratio_x"] < 4.1    # glycine 1.19, acetanilide 4.8
    assert m64["d15n_spread"] == max(by[r].d15n for r in m64["positions"]) - min(by[r].d15n for r in m64["positions"])
    assert m64["over"] == ["C/N peak-area ratio 1.19 to 4.81, 4.02 times", "raw d15N -5.95 to 1.27 permil",
                           "raw d13C -23.69 to -9.45 permil"]
    items = F["rerun"]["run_level"]
    assert MIXED_TEXT in items and MIXED_TEXT in F["rerun"]["run_level_text"]
    # in the order of the workbook's concatenation: after the calibration line, before the anchors outside k sigma
    at = items.index(MIXED_TEXT)
    assert all(k < at for k, t in enumerate(items) if "calibration slope" in t or "residual SD" in t)
    assert all(k > at for k, t in enumerate(items) if t.startswith(("check standard", "anchors outside",
                                                                     "elemental drift", "N2 ref", "CO2 ref")))


def test_one_analysis_or_one_value_is_not_judged(tmp_path):
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    out = []
    for r in body:                          # the second and third USGS 65 analyses renamed, so one is left
        if r[i["Identifier 1"]].strip() == "USGS 65" and r[i["Row"]] in ("16", "17"):
            r = list(r)
            r[i["Identifier 1"]] = "USGS 65 left out"
        out.append(r)
    cfg = Config(standards=list(DEFAULT_STANDARDS) + [
        ("USGS 65 left out", "Exclude", None, None, None, None, "left out for the test", None, None, None, None)])
    _, F = process(parse_export(_write(tmp_path / "one65.csv", head, out)), cfg)
    m64, m65 = F["anchor_material"]["materials"]
    assert m64["level"] == "OK" and m65["n"] == 1 and m65["level"] is None
    assert m65["ratio_min"] == m65["ratio_max"] and m65["ratio_x"] is None and m65["d15n_spread"] is None
    assert F["anchor_material"]["level"] == "OK"
    # an element in force with no anchor to test it against has a check without a level
    _, F = process(parse_export(RUN), UNCERTIFIED)
    assert [F["elem_check"][el]["n"] for el in "nc"] == [0, 0]
    assert [F["elem_check"][el]["level"] for el in "nc"] == [None, None]
    _, F = process(parse_export(RUN), Config(elemental_correction="Off"))
    assert F["elem_check"] == {"n": None, "c": None}


def test_a_light_drift_monitor_beside_the_anchors(tmp_path):
    # 1.4.15: the acetanilide's whole-run check, which fires on the test run, is off, so the problems of the whole
    # run are this check's alone
    positions, F = process(parse_export(light_variant(tmp_path / "light.csv")),
                           Config(elemental_correction="On", ace_scatter_check="Off"))
    texts = []
    for el, name in (("n", "%N"), ("c", "%C")):
        c = F["elem_check"][el]
        k, as_measured, corrected = _by_hand(positions, el)
        assert c["n"] == k == 6
        assert abs(c["rms_measured"] - as_measured) < 1e-12 and abs(c["rms_reported"] - corrected) < 1e-12
        assert c["level"] == _level(corrected - as_measured)
        texts.append(f"elemental drift correction moves the anchors' {name} {_xl((corrected - as_measured) * 100, 2)} "
                     f"points further from certified (RMS {_xl(as_measured * 100, 2)}% as measured, "
                     f"{_xl(corrected * 100, 2)}% corrected)")
    # By hand: the monitor's %N reads about 1/0.9, 11%, high, so its factor rises from 0.995 to 1.101, and the
    # anchors on both sides are divided by up to 1.07. %N moves 3.34 points (a Note, between 2.5 and 5), %C 5.35
    # (a Caution, above 5).
    assert [F["elem_check"][el]["level"] for el in "nc"] == ["Note", "Caution"]
    assert texts == ["elemental drift correction moves the anchors' %N 3.34 points further from certified "
                     "(RMS 1.17% as measured, 4.51% corrected)",
                     "elemental drift correction moves the anchors' %C 5.35 points further from certified "
                     "(RMS 3.17% as measured, 8.52% corrected)"]
    assert F["rerun"]["run_level"] == texts
    assert F["anchor_material"]["level"] == "OK"
    # a wider gate makes both OK; left out of the correction, the monitor no longer moves the anchors
    _, wide = process(parse_export(tmp_path / "light.csv"), Config(elemental_correction="On",
                                                                   elemental_threshold_pct=20.0))
    assert [wide["elem_check"][el]["level"] for el in "nc"] == ["OK", "OK"]
    _, left = process(parse_export(tmp_path / "light.csv"), Config(elemental_correction="On", exclude_elemental=[14]))
    assert [left["elem_check"][el]["level"] for el in "nc"] == ["OK", "OK"]


def test_the_port_block_carries_the_anchor_check(tmp_path):
    export = mixed_variant(tmp_path / "mixed.csv")
    positions, F = process(parse_export(export), Config())
    block = port_input(positions, F, Config(), export.name, export.stat().st_size)
    am = block["anchorMaterial"]
    assert am["level"] == "Caution" and am["limits"] == {"ratio": 1.5, "spread": 5.0}
    m = am["materials"][0]
    assert set(m) == {"name", "si", "n", "ratio", "ratioX", "d15n", "d13c", "positions", "over", "level"}
    assert (m["name"], m["si"], m["n"], m["level"]) == ("USGS 64", 0, 6, "Caution")
    src = F["anchor_material"]["materials"][0]
    assert m["ratio"] == [src["ratio_min"], src["ratio_max"]] and m["ratioX"] == src["ratio_x"]
    assert m["d15n"] == [src["d15n_min"], src["d15n_max"]] and m["d13c"] == [src["d13c_min"], src["d13c_max"]]
    assert m["positions"] == [10, 11, 12, 13, 14, 18] and len(m["over"]) == 3
    assert block["settings"]["refSdLimit"] == 0.1
    off = replace(Config(), ref_gas_sd_limit=None)
    assert port_input(positions, F, off, export.name, 1)["settings"]["refSdLimit"] is None


# ------------------------------------------------------------------------------------------ the port
@pytest.mark.skipif(not NODE, reason="needs node for the preview port")
@pytest.mark.parametrize("variant", ["light", "mixed"])
def test_the_port_agrees_on_the_elemental_check(tmp_path, variant):
    export = (light_variant if variant == "light" else mixed_variant)(tmp_path / f"{variant}.csv")
    (tmp_path / "cfg.json").write_text(json.dumps(config_to_dict(Config())))
    out = tmp_path / "port"
    gen = subprocess.run([sys.executable, str(ROOT / "scripts" / "port_cases.py"), str(export), "--out", str(out),
                          "--config", str(tmp_path / "cfg.json"), "--random", "40", "--per-variant", "2"],
                         capture_output=True, text=True, cwd=ROOT)
    assert gen.returncode == 0, gen.stderr
    run = subprocess.run([NODE, str(ROOT / "port" / "test_port.js"), str(out / "port_input.json"),
                          str(out / "cases.json")], capture_output=True, text=True, cwd=ROOT)
    assert run.returncode == 0, run.stdout[-2000:] + run.stderr[-2000:]
    assert "value mismatches 0, flag-list mismatches 0, fit mismatches 0" in run.stdout, run.stdout[-500:]
    cases = json.loads((out / "cases.json").read_text())
    levels = {(k, (c["fits"][k] or [None])[0]) for c in cases for k in ("elemCheckN", "elemCheckC")}
    if variant == "light":            # the comparison reaches every level, and the correction switched off
        assert {("elemCheckN", "Note"), ("elemCheckC", "Caution"), ("elemCheckN", "OK"),
                ("elemCheckN", None)} <= levels
    block = json.loads((out / "port_input.json").read_text())
    assert block["anchorMaterial"]["level"] == ("Caution" if variant == "mixed" else "OK")


# ------------------------------------------------------------------------------------------ the workbook
@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    """The workbooks of the cases, built by ea_irms and computed by LibreOffice in one call."""
    if not SOFFICE:
        pytest.skip("needs LibreOffice to compute the formulas")
    folder = tmp_path_factory.mktemp("checks1413")
    cases = {"default": (RUN, Config()),
             "mixed": (mixed_variant(folder / "mixed.csv"), Config()),
             "light": (light_variant(folder / "light.csv"), Config(elemental_correction="On")),
             "uncertified": (RUN, UNCERTIFIED)}
    books = []
    for name, (export, cfg) in cases.items():
        build_workbook(export, folder / f"{name}.xlsx", cfg)
        books.append(str(folder / f"{name}.xlsx"))
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + books, check=True, capture_output=True, timeout=1200)
    out = {}
    for name, (export, cfg) in cases.items():
        wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
        out[name] = (wb, cfg, *process(parse_export(export), cfg))
    return out


def _same(mine, theirs):
    if mine is None:
        return theirs in (None, "")
    if isinstance(mine, str):
        return mine == theirs
    return isinstance(theirs, (int, float)) and abs(mine - theirs) <= 1e-9 * max(1.0, abs(mine))


def _verdict_elem(F, el, name):
    c = F["elem_check"][el]
    if c is None:
        return f"not judged: the elemental drift correction is not in force for {name}"
    if c["level"] is None:
        return f"not judged: fewer than two anchors with a {name}"
    if c["level"] == "OK":
        return "ok"
    return (f"CHECK, {c['level']}: the correction moves the anchors {_xl(c['added'] * 100, 2)} points further from "
            f"certified {name}; check the drift monitors beside them and Setup B37, or set Setup B35 to Off")


@pytest.mark.parametrize("name", ["default", "mixed", "light", "uncertified"])
def test_the_workbook_agrees_with_the_engine(computed, name):
    wb, cfg, positions, F = computed[name]
    ft, qc = wb["Fits"], wb["QC"]
    # the anchors as one material, one row per line of the standards table
    am = F["anchor_material"]
    mats = {m["index"]: m for m in am["materials"]}
    for k in range(NSTD):
        m = mats.get(k)
        for j, key in enumerate(AM_COLS, start=1):
            got = ft.cell(AM_TOP + k, j).value
            if m is not None:
                want = m[key]
            elif key == "name":
                want = cfg.standards[k][0] if k < len(cfg.standards) else None
            else:
                want = 0 if key in ("n", "n_ratio") else None
            assert _same(want, got), (name, AM_TOP + k, key, want, got)
    assert (ft["B370"].value, ft["B371"].value) == (ANCHOR_RATIO_LIMIT, ANCHOR_SPREAD_PERMIL)
    assert _same(am["level"], ft[f"B{AM_LEVEL}"].value)
    pieces = [t for t in F["rerun"]["run_level"] if t.startswith("anchors labeled ")]
    assert _same("".join(t + "; " for t in pieces) or None, ft[f"B{AM_TEXT}"].value), name
    # the elemental drift correction against the anchors
    for el in ("n", "c"):
        c = F["elem_check"][el]
        for off, key in enumerate(("n", "rms_measured", "rms_reported", "added", "level")):
            want = None if c is None else c[key]
            assert _same(want, ft.cell(EC_BASE[el] + off, 2).value), (name, el, key, want)
    pieces = [t for t in F["rerun"]["run_level"] if t.startswith("elemental drift correction moves")]
    assert _same("".join(t + "; " for t in pieces) or None, ft[f"B{EC_TEXT}"].value), name
    # QC: the levels, the RMS deviations in percent, and the verdicts
    assert _same(am["level"], qc["B177"].value)
    assert qc["D177"].value == ({"OK": "ok", None: "not judged: no anchor material has two analyses"}.get(am["level"])
                                or qc["D177"].value)
    assert (am["level"] == "Caution") == str(qc["D177"].value).startswith("CHECK: the anchors of a row")
    for el, r, pct in (("n", 178, "%N"), ("c", 180, "%C")):
        c = F["elem_check"][el]
        for row, key in ((r, "rms_measured"), (r + 1, "rms_reported")):
            want = None if c is None or c[key] is None else c[key] * 100
            assert _same(want, qc[f"B{row}"].value), (name, el, key)
        assert qc[f"D{r + 1}"].value == _verdict_elem(F, el, pct), (name, el)
    # the problems of the whole run, in the engine's order
    assert ft["B356"].value == F["rerun"]["run_level_text"]
    assert wb["Rerun"]["A5"].value == F["rerun"]["run_level_text"]
    errors = [x.coordinate for s in (ft, qc, wb["Work"]) for row in s.iter_rows() for x in row
              if isinstance(x.value, str) and x.value.startswith("#") and x.value[1:4].isupper()
              and not (s.title == "Work" and x.value == "#N/A")]
    assert not errors, errors[:5]


def test_the_cases_cover_each_outcome(computed):
    levels = {name: (F["anchor_material"]["level"], *[(F["elem_check"][el] or {}).get("level", "off") for el in "nc"])
              for name, (_, _, _, F) in computed.items()}
    assert levels == {"default": ("OK", "OK", "off"), "mixed": ("Caution", "OK", "OK"),
                      "light": ("OK", "Note", "Caution"), "uncertified": ("OK", None, None)}


# ------------------------------------------------------------------------------------------ the workbook items
@pytest.fixture(scope="module")
def plain(tmp_path_factory):
    """The test run's workbook as openpyxl writes it, formulas without values."""
    folder = tmp_path_factory.mktemp("plain1413")
    return build_workbook(RUN, folder / "plain.xlsx", Config())


def test_charts_keep_their_text_off_the_plot(plain):
    with zipfile.ZipFile(plain) as z:
        charts = [z.read(n).decode() for n in sorted(z.namelist()) if re.match(r"xl/charts/chart\d+\.xml$", n)]
    assert len(charts) == 9
    tag = r"(?:c:)?"
    for x in charts:
        title = re.search(rf"<{tag}chart><{tag}title>.*?</{tag}title>", x, re.S).group(0)
        legend = re.search(rf"<{tag}legend>.*?</{tag}legend>", x, re.S).group(0)
        assert re.findall(r'overlay val="(\w+)"', title) == ["0"]
        assert re.findall(r'overlay val="(\w+)"', legend) == ["0"]
        axes = re.findall(rf"<{tag}valAx>.*?</{tag}valAx>", x, re.S)
        assert len(axes) == 2
        x_axis, y_axis = axes
        assert re.findall(r'axPos val="(\w+)"', x_axis) == ["b"] and re.findall(r'axPos val="(\w+)"', y_axis) == ["l"]
        for ax in axes:
            assert re.findall(r'tickLblPos val="(\w+)"', ax) == ["low"]
            assert re.findall(r'overlay val="(\w+)"', ax) == ["0"]        # the axis title


def test_readme_qc_and_replicates_text(plain, tmp_path):
    wb = openpyxl.load_workbook(plain)
    readme = " ".join(str(c.value) for c in wb["README"]["B"] if c.value)
    assert "no %C, %N, or C:N" not in readme
    assert "Without a mass there is no %C or %N for that row; C:N, which comes from the two peaks" in readme
    assert "as a Kragten spreadsheet would" not in readme and "to first order" in readme
    assert "A sample with neither peak is marked for a rerun of N and C." in readme
    assert "QC row 177" in readme and "Fits rows 393 to 404" in readme
    qc = " ".join(str(c.value) for row in wb["QC"].iter_rows(max_col=1) for c in row if c.value)
    assert "as a Kragten spreadsheet would" not in qc and "to first order" in qc
    assert "for a rerun of N and C when it has neither peak" in qc
    assert "C:N comes from the two peaks and needs none" in wb["Results"]["A2"].value
    rp = wb["Replicates"]["A9"].value
    assert rp.startswith("No sample sheet was given, so no replicate group formed") and "Dup" in rp
    sheet = read_sample_sheet(b"Sample ID,Replicate of\nXYZ 1,\nXYZ 1 Dup,XYZ 1\n", "samples.csv")
    out = build_workbook(RUN, tmp_path / "sheet.xlsx", Config(), samples=sheet)
    rp = openpyxl.load_workbook(out)["Replicates"]["A9"].value
    assert rp.startswith("No replicate group formed in this run") and "Replicate of column names them" in rp
