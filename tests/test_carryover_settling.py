"""Carryover flag and settled position (1.4.0), in the engine and in the workbook.

The scenarios are the 2026-09-15 run from the reference template's Raw sheet,
edited where a scenario needs it: a mid-run spike on the reference-peak
background, a run that is flat from the start, one that never settles, an
oversized sample followed by a raised background, a spike on a reading that
an oversized sample exposes, and two low readings in the middle of the decline. Before 1.4.0 the spike moved the settled position
from 27 to 44 and the flat run came out as never settled (57 of 56 positions).
The engine is checked against the expected values, and the workbook, computed by
LibreOffice, against the engine.
"""
import csv
import os
import random
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ea_irms.parse import parse_export            # noqa: E402
from ea_irms.pipeline import Config, process      # noqa: E402
from ea_irms.workbook import build_workbook       # noqa: E402

TEMPLATE = Path(os.environ.get(
    "EA_TEMPLATE",
    Path(__file__).resolve().parent / "data" / "EA-IRMS_Correction_Template_v1.xlsx"))
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
if not TEMPLATE.exists():
    pytest.skip("needs the reference template (EA_TEMPLATE)", allow_module_level=True)

RAW = [list(r) for r in openpyxl.load_workbook(TEMPLATE)["Raw"].iter_rows(max_col=16, values_only=True)
       if r[0] is not None]
COL = {h: k for k, h in enumerate(RAW[0])}
POSITIONS = sorted({r[0] for r in RAW[1:]})


def _n_rows(rows, pos):
    return [r for r in rows[1:] if r[0] == pos and r[COL["Ampl  28"]] is not None]


def _edit(ref_bgd=None, sample_bgd=None, sample_area_n=None):
    """The Raw rows with the designated N2 reference peak's or the sample peak's values replaced."""
    rows = [list(r) for r in RAW]
    for pos, v in (ref_bgd or {}).items():
        for r in _n_rows(rows, pos):
            if r[COL["Is Ref _"]] == 1:
                r[COL["BGD 28"]] = v
    for pos, v in (sample_bgd or {}).items():
        max(_n_rows(rows, pos), key=lambda r: r[COL["Peak Nr"]])[COL["BGD 28"]] = v
    for pos, v in (sample_area_n or {}).items():
        max(_n_rows(rows, pos), key=lambda r: r[COL["Peak Nr"]])[COL["Area 28"]] = v
    return rows


def _flat(seed=1):
    rnd = random.Random(seed)
    return {p: round(23.1 + rnd.gauss(0, 0.1), 1) for p in POSITIONS}


SPIKE = {40: 32.3, 41: 37.0, 42: 37.0, 43: 25.7}          # the PanamaSids3 shape
NEVER = {p: 40.0 if k < 5 else (30.0 if k % 10 == 0 else 35.0) for k, p in enumerate(POSITIONS)}
SCENARIOS = {
    "measured":   {},
    "spike":      {"ref_bgd": SPIKE},
    "flat":       {"ref_bgd": _flat()},
    "flat_spike": {"ref_bgd": {**_flat(), **SPIKE}},
    "never":      {"ref_bgd": NEVER},
    "carry":      {"sample_area_n": {36: 200.0}, "sample_bgd": {37: 33.5, 38: 27.0}},
    "skip":       {"sample_area_n": {27: 200.0}, "ref_bgd": {28: 30.0}},
    "dip":        {"ref_bgd": {20: 23.2, 21: 23.2}},      # two low readings are not three
}
# settled position, flagged positions at the default 2 mV
EXPECTED = {"measured": (27, []), "spike": (27, []), "flat": (1, []), "flat_spike": (1, []),
            "never": (POSITIONS[-1] + 1, []), "carry": (27, [37, 38]), "skip": (27, []),
            "dip": (27, [])}


@pytest.fixture(scope="module")
def exports(tmp_path_factory):
    folder = tmp_path_factory.mktemp("scenarios")
    out = {}
    for name, edits in SCENARIOS.items():
        path = folder / f"{name}.csv"
        with open(path, "w", newline="") as fh:
            w = csv.writer(fh)
            for r in _edit(**edits):
                w.writerow(["" if v is None else v for v in r])
        out[name] = path
    return out


@pytest.mark.parametrize("name", list(SCENARIOS))
def test_engine_settles_and_flags_as_designed(exports, name):
    positions, fits = process(parse_export(exports[name]), Config())
    settled, flagged = EXPECTED[name]
    assert fits["settling"]["settled_from"] == settled
    assert fits["carryover"]["flagged"] == flagged
    for p in positions:
        assert ("possible carryover from an oversized sample" in p.flags) == (p.row in flagged)


def test_flat_run_counts_as_settled_because_its_scale_is_noise(exports):
    _, fits = process(parse_export(exports["flat"]), Config())
    assert fits["settling"]["excess"] < 1.0
    assert fits["settling"]["anchors_before"] == 0


def test_rise_is_measured_from_the_last_unexposed_background(exports):
    positions, fits = process(parse_export(exports["carry"]), Config())
    by_row = {p.row: p for p in positions}
    assert by_row[36].oversized and by_row[37].carry_exposed and by_row[38].carry_exposed
    assert abs(by_row[37].carry_rise - (33.5 - by_row[36].bgd_n)) < 1e-12
    assert abs(by_row[38].carry_rise - (27.0 - by_row[36].bgd_n)) < 1e-12
    assert fits["carryover"]["followed"] == 2      # 23 before 24, and 36 before 37


def test_threshold_decides_and_empty_switches_the_flag_off(exports):
    _, strict = process(parse_export(exports["carry"]), Config(carryover_mv=3.0))
    assert strict["carryover"]["flagged"] == [37]
    positions, off = process(parse_export(exports["carry"]), Config(carryover_mv=None))
    assert off["carryover"]["flagged"] == []
    assert not any("carryover" in f for p in positions for f in p.flags)


def test_an_exposed_spike_does_not_break_the_streak(exports):
    """Position 27 is oversized, so 28 and 29 are exposed; a spike at 28 is skipped
    and the streak that started at 27 stands. Without the skip it would restart at 29."""
    _, fits = process(parse_export(exports["skip"]), Config())
    assert fits["settling"]["settled_from"] == 27


@pytest.mark.skipif(not SOFFICE, reason="needs LibreOffice to compute the workbook")
def test_workbook_agrees_with_engine(exports, tmp_path):
    configs = {name: Config() for name in SCENARIOS}
    configs["carry_strict"] = Config(carryover_mv=3.0)
    sources = dict(exports, carry_strict=exports["carry"])
    files = []
    for name, cfg in configs.items():
        out = tmp_path / f"{name}.xlsx"
        build_workbook(sources[name], out, cfg)
        files.append(out)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path / "calc")]
                   + [str(f) for f in files], check=True, capture_output=True, timeout=600)
    for name, cfg in configs.items():
        positions, fits = process(parse_export(sources[name]), cfg)
        wb = openpyxl.load_workbook(tmp_path / "calc" / f"{name}.xlsx", data_only=True)
        wk, ft, rs = wb["Work"], wb["Fits"], wb["Results"]
        header = {wk.cell(5, c).value: c for c in range(1, wk.max_column + 1)}
        feeds = {c for h, c in header.items() if str(h or "").startswith("ch ")}
        errors = [(s.title, c.coordinate) for s in wb for row in s.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper()
                  and not (s.title == "Work" and c.column in feeds and c.value == "#N/A")]
        assert not errors, (name, errors[:5])
        assert ft["B86"].value == fits["settling"]["settled_from"], name
        assert abs(ft["B85"].value - fits["settling"]["excess"]) < 1e-9, name
        assert ft["B229"].value == fits["carryover"]["oversized"], name
        assert ft["B230"].value == fits["carryover"]["followed"], name
        assert ft["B231"].value == len(fits["carryover"]["flagged"]), name
        mx = fits["carryover"]["max_rise"]
        assert (ft["B232"].value in (None, "")) if mx is None else abs(ft["B232"].value - mx) < 1e-9, name
        for p in positions:
            k = 5 + p.row
            for attr, col in (("carry_rise", header["carry rise"]),
                              ("settling_fraction", header["settling frac"])):
                mine, theirs = getattr(p, attr), wk.cell(k, col).value
                theirs = theirs if isinstance(theirs, (int, float)) else None
                assert (mine is None) == (theirs is None), (name, p.row, attr, mine, theirs)
                if mine is not None:
                    assert abs(mine - theirs) < 1e-9 * max(1.0, abs(mine)), (name, p.row, attr)
            text = rs[f"K{4 + p.row}"].value or ""
            assert set(p.flags) == {t.strip() for t in text.split(";") if t.strip()}, (name, p.row)
