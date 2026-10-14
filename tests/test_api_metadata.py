"""isoharness 0.3.0: the sample sheet and the run information through the API.

The sheet goes with the request that needs it, as the export does, and is read for that
request only: /api/process, /api/results.csv, /api/workbook and the two batch endpoints take
it in a sample_sheet field (with id_column to choose its ID column); the run information goes
in run_info, or run_infos for a batch. /api/samplesheet/read matches the samples the page
lists, and /api/samplesheet/template.xlsx writes a sheet to fill in.
"""
import csv
import io
import json

import openpyxl
from fastapi.testclient import TestClient

from ea_irms.workbook import VERSION
from isoharness.app import XLSX, app, limiter

from test_replicates import SHEET, variant

client = TestClient(app)


def files_for(export, sheet=SHEET, name="samples.csv"):
    out = {"file": (export.name, export.read_bytes())}
    if sheet is not None:
        out["sample_sheet"] = (name, sheet.encode() if isinstance(sheet, str) else sheet)
    return out


def test_process_matches_the_samples_and_compares_replicates(tmp_path):
    export = variant(tmp_path)
    r = client.post("/api/process", files=files_for(export),
                    data={"run_info": json.dumps({"operator": " XT  Wang ", "Notes": "renamed"})})
    assert r.status_code == 200
    body = r.json()
    meta = body["metadata"]
    assert meta["run_info"] == {"Operator": "XT Wang", "Notes": "renamed"}
    assert meta["sheet"]["id_column"] == "Sample ID" and meta["sheet"]["rows"] == 8
    assert meta["match"]["matched"] == 11 and len(meta["samples"]) == 26
    assert meta["columns"][:3] == ["Sample ID (sheet)", "Matched by", "Project"]
    assert {s["row"]: s["group"] for s in meta["samples"] if s["group"]} == {21: 1, 24: 1, 49: 1, 31: 2, 32: 2,
                                                                             33: 2, 35: 3, 48: 3}
    marks = {p["row"]: p["rerun"] for p in body["positions"] if p["rerun"]}
    assert marks[35] == marks[48] == "C" and marks[31] == "N"
    assert [g["rows"] for g in body["fits"]["replicates"]["groups"]] == [[21, 24, 49], [31, 32, 33], [35, 48]]
    plain = client.post("/api/process", files=files_for(export, None)).json()
    assert plain["metadata"]["sheet"] is None and plain["metadata"]["columns"] == []
    assert plain["metadata"]["run_info"] == {} and plain["fits"]["replicates"]["groups"] == []


def test_results_table_and_workbook_carry_the_sheet(tmp_path):
    export = variant(tmp_path)
    r = client.post("/api/results.csv", files=files_for(export))
    assert r.status_code == 200
    head = next(csv.reader(io.StringIO(r.text)))
    assert head[24:26] == ["Sample ID (sheet)", "Matched by"] and "Colony" in head
    r = client.post("/api/workbook", files=files_for(export),
                    data={"values": "false", "run_info": json.dumps({"Instrument": "Delta V Plus", "Operator": "XT"})})
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    md = wb["Metadata"]
    assert (md["B6"].value, md["B9"].value, wb["Setup"]["B7"].value) == ("replicates.csv", "Delta V Plus", "XT")
    assert md["B16"].value == "Sample ID" and md["C23"].value == "Sample ID (sheet)"
    assert [c.value for c in wb["Results"][4]][24:26] == ["Sample ID (sheet)", "Matched by"]


def test_choosing_the_id_column(tmp_path):
    export = variant(tmp_path)
    r = client.post("/api/process", files=files_for(export), data={"id_column": "colony"})
    assert r.status_code == 200
    meta = r.json()["metadata"]
    assert meta["sheet"]["id_column"] == "Colony"
    # CI 3-2-H and CI 3-2-S match the colony CI 3-2; CI 3-20 matches CI 3
    assert {s["row"]: s["sample_id"] for s in meta["samples"] if s["sample_id"]} == {36: "CI 3-2", 37: "CI 3-2",
                                                                                     39: "CI 3"}


def test_refusals(tmp_path):
    export = variant(tmp_path)
    r = client.post("/api/process", files=files_for(export, b"%PDF", "notes.pdf"))
    assert r.status_code == 415 and "not a sample sheet ea_irms reads" in r.json()["detail"]
    r = client.post("/api/process", files=files_for(export), data={"id_column": "Mass"})
    assert r.status_code == 422 and r.json()["detail"].startswith(
        "the sample sheet could not be read: samples.csv has no column named 'Mass'")
    r = client.post("/api/process", files=files_for(export), data={"run_info": "{not json"})
    assert r.status_code == 422 and "run_info could not be read as JSON" in r.json()["detail"]
    r = client.post("/api/process", files=files_for(export), data={"run_info": json.dumps({"Analyst": "XT"})})
    assert r.status_code == 422 and "fields ea_irms does not read (Analyst)" in r.json()["detail"]
    r = client.post("/api/workbook", files=files_for(export, "Sample ID\n"), data={"values": "false"})
    assert r.status_code == 422 and "has no row with a value in its ID column" in r.json()["detail"]
    r = client.post("/api/samplesheet/read", files={"file": ("samples.csv", SHEET.encode())}, data={"runs": "[1]"})
    assert r.status_code == 422 and r.json()["detail"] == "run 1 of runs must be an object with a samples list"
    r = client.post("/api/samplesheet/read", files={"file": ("samples.csv", SHEET.encode())},
                    data={"runs": json.dumps([{"samples": [["20", "R1"]]}])})
    assert r.status_code == 422 and r.json()["detail"] == "each sample of run 1 must be [position, identifier]"
    r = client.post("/api/samplesheet/template.xlsx", json={"names": "R1"})
    assert r.status_code == 422
    r = client.post("/api/samplesheet/template.xlsx", content=b"not json", headers={"Content-Type": "application/json"})
    assert r.status_code == 422


def test_the_page_reads_a_sheet_and_matches_its_runs():
    runs = [{"name": "a.csv", "samples": [[21, "R1"], [24, "R1"], [36, "CI 3-2-H"], [39, "CI 3-20"], [51, "x"]]},
            {"name": "b.csv", "samples": [[5, "R2 a"], [6, "R2 b"], [7, "R2 c"]]}]
    r = client.post("/api/samplesheet/read", files={"file": ("samples.csv", SHEET.encode())},
                    data={"runs": json.dumps(runs)})
    assert r.status_code == 200
    body = r.json()
    assert body["ea_irms"] == VERSION and body["sheet"]["replicate_column"] == "Replicate of"
    assert body["sheet"]["notes"] == ["1 row without a value in Sample ID was left out",
                                      "the ID R1 is in 2 rows; the first is used"]
    a, b = body["runs"]
    assert [(s["row"], s["sample_id"], s["matched_by"], s["group"]) for s in a["samples"]] == [
        (21, "R1", "exact", 1), (24, "R1", "exact", 1), (36, "CI 3-2", "prefix", None), (39, "CI 3", "prefix", None),
        (51, None, None, None)]
    assert a["match"]["unmatched"] == [{"row": 51, "identifier": "x"}]
    assert a["groups"] == [{"group": 1, "name": "R1", "rows": [21, 24], "by": "name"}]
    assert b["groups"] == [{"group": 1, "name": "R2", "rows": [5, 6, 7], "by": "Replicate of"}]
    assert body["columns"] == ["Sample ID (sheet)", "Matched by", "Project", "Species or material", "Collection date",
                               "Depth", "Notes", "Replicate of", "Type (sheet)", "Colony"]
    r = client.post("/api/samplesheet/read", files={"file": ("samples.csv", SHEET.encode())})
    assert r.status_code == 200 and r.json()["runs"] == []


def test_template_from_the_page():
    r = client.post("/api/samplesheet/template.xlsx", json={"names": ["CI 3-2-H", "ci 3-2-h", "R1", 61]})
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == f'attachment; filename="sample_sheet_ea_irms_v{VERSION}.xlsx"'
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Samples"]
    assert [ws.cell(i, 1).value for i in range(1, 6)] == ["Sample ID", "CI 3-2-H", "R1", "61", None]


def test_batch_with_a_sheet_and_run_information(tmp_path):
    export = variant(tmp_path)
    a, b = tmp_path / "run_a.csv", tmp_path / "run_b.csv"
    a.write_bytes(export.read_bytes())
    b.write_bytes(export.read_bytes())
    files = [("files", (f.name, f.read_bytes())) for f in (a, b)]
    plain = client.post("/api/batch/results.csv", files=files).text.splitlines()
    sheet = [("sample_sheet", ("samples.csv", SHEET.encode()))]
    r = client.post("/api/batch/results.csv", files=files + sheet, data={"run_infos": json.dumps({"Operator": "XT"})})
    assert r.status_code == 200
    lines = r.text.splitlines()
    assert len(lines) == len(plain) == 113
    one = client.post("/api/results.csv", files=files_for(a)).text.splitlines()
    assert lines[0] == "Run," + one[0] and lines[1:57] == ["run_a," + ln for ln in one[1:]]
    assert lines[0].endswith(",Sample ID (sheet),Matched by,Project,Species or material,Collection date,Depth,Notes,"
                             "Replicate of,Type (sheet),Colony")
    # the sheet's columns follow the run's own; the replicates change the marks of their members only
    changed = [int(row[1]) for old, row in zip(csv.reader(plain[1:57]), csv.reader(lines[1:57])) if row[:25] != old]
    assert changed == [21, 24, 31, 32, 33, 35, 48, 49]
    r = client.post("/api/batch/results.xlsx", files=files + sheet,
                    data={"run_infos": json.dumps([{"Operator": "XT", "Project": "P1"}, {"Notes": "second"}])})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    runs = list(wb["Runs"].iter_rows(values_only=True))
    assert list(runs[0][-4:]) == ["Operator", "Project", "Notes", "Samples matched"]
    assert list(runs[1][-4:]) == ["XT", "P1", None, "11 of 26 samples matched (8 by the exact name, 3 by an ID "
                                                    "the name begins with)"]
    assert list(runs[2][-4:-1]) == [None, None, "second"]
    head = [c.value for c in wb["Results"][1]]
    assert head[25:27] == ["Sample ID (sheet)", "Matched by"]
    readme = " ".join(str(c.value) for c in wb["README"]["B"] if c.value)
    assert "Sample sheet: samples.csv, ID column Sample ID, 8 rows." in readme
    r = client.post("/api/batch/results.csv", files=files, data={"run_infos": json.dumps([{}])})
    assert r.status_code == 422 and "one entry for each of the 2 files" in r.json()["detail"]
    r = client.post("/api/batch/results.csv", files=files, data={"run_infos": json.dumps([{}, {"Analyst": "x"}])})
    assert r.status_code == 422 and r.json()["detail"].startswith("the run information for run_b.csv: ")


def test_the_new_requests_count_as_uploads(monkeypatch):
    monkeypatch.setenv("ISOHARNESS_READS_PER_HOUR", "1")
    limiter.reset()
    try:
        assert client.post("/api/samplesheet/template.xlsx", json={"names": []}).status_code == 200
        r = client.post("/api/samplesheet/read", files={"file": ("samples.csv", SHEET.encode())})
        assert r.status_code == 429
    finally:
        limiter.reset()
