"""isoharness 0.2.0: several runs at once, one combined results table.

Two runs are made from the 2026-09-15 test run. In both, each sample is named by its
position (S20, S21, ...), since the test run calls every sample Random; in the second,
each sample's N2 peak reads 0.5 permil higher. Every correction is then the same in the
two runs, so each sample's two final d15N values differ by 0.5 times the calibration
slope and its d13C values are equal, which is what the Repeats sheet must show.

Finding W-12: the samples measured in more than one run are paired as the sample sheet
pairs names, as written and then with the spaces taken out, on the server and on the page.
"""
import csv
import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from ea_irms.parse import parse_export
from ea_irms.pipeline import Config, config_to_dict, process
from ea_irms.workbook import VERSION
from isoharness.app import XLSX, app
from isoharness.batch import BatchRun, combined_csv, label_runs, repeats
from isoharness.serialize import csv_cells, results_values

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pagelogic import NODE, renamed, run_page  # noqa: E402

RUN = Path(__file__).resolve().parent / "data" / "20260915_raw.csv"
client = TestClient(app)
SHIFT = 0.5


def variant(shift=0.0, name="S{}"):
    """The test run with each sample named S<position>, or by the pattern name, and its N2 peak's d15N moved
    by shift."""
    text = RUN.read_text(encoding="utf-8")
    rows = list(csv.reader(io.StringIO(text)))
    h = rows[0]
    col = {k: h.index(k) for k in ("Row", "Identifier 1", "Ampl  28", "Is Ref _", "d 15N/14N")}
    positions, _ = process(parse_export(RUN), Config())
    samples = {str(p.row) for p in positions if p.role == "Sample"}
    last28 = {}
    for i, r in enumerate(rows[1:], start=1):
        if len(r) > col["Ampl  28"] and r[col["Ampl  28"]].strip():
            last28[r[col["Row"]].strip()] = i
    for i, r in enumerate(rows[1:], start=1):
        row = r[col["Row"]].strip() if r else ""
        if row in samples:
            r[col["Identifier 1"]] = name.format(row)
            if shift and last28.get(row) == i and r[col["Is Ref _"]].strip() != "1":
                r[col["d 15N/14N"]] = repr(float(r[col["d 15N/14N"]]) + shift)
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


@pytest.fixture(scope="module")
def pair(tmp_path_factory):
    d = tmp_path_factory.mktemp("batch")
    a, b = d / "run_a.csv", d / "run_b.csv"
    a.write_bytes(variant())
    b.write_bytes(variant(SHIFT))
    return a, b


def send(path, files, **form):
    return client.post(path, files=[("files", (Path(f).name, Path(f).read_bytes())) for f in files], data=form)


def test_combined_results_hold_every_position_of_every_run(pair):
    r = send("/api/batch/results.xlsx", pair, run_dates=json.dumps(["2026-09-15", None]))
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == f'attachment; filename="combined_2_runs_results_ea_irms_v{VERSION}.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["README", "Results", "Runs", "Repeats"]
    rs = wb["Results"]
    header = [c.value for c in rs[1]]
    assert header[:4] == ["Run", "Row", "Identifier", "Type"] and "1σ(d15N)" in header and len(header) == 25
    got = [list(row) for row in rs.iter_rows(min_row=2, values_only=True)]
    want = []
    for f in pair:
        positions, _ = process(parse_export(f), Config())
        want += [[f.stem] + [v if v != "" else None for v in results_values(p)] for p in positions]
    assert len(got) == len(want) == 112
    for g, w in zip(got, want):
        for x, y in zip(g, w):
            same = x == y or (x in (None, "") and y in (None, ""))
            assert same or abs(x - y) <= 1e-12 * max(1, abs(y)), (g[:3], x, y)
    runs = list(wb["Runs"].iter_rows(values_only=True))
    assert [r[:4] for r in runs[1:]] == [("run_a", "run_a.csv", "2026-09-15", 56), ("run_b", "run_b.csv", None, 56)]


def test_repeats_show_each_sample_across_the_runs(pair):
    wb = openpyxl.load_workbook(io.BytesIO(send("/api/batch/results.xlsx", pair).content))
    rows = list(wb["Repeats"].iter_rows(values_only=True))
    head, body = rows[0], rows[1:]
    ix = {h: i for i, h in enumerate(head)}
    positions, fits = process(parse_export(pair[0]), Config())
    samples = [p for p in positions if p.role == "Sample"]
    assert len(body) == 2 * len(samples)
    half = SHIFT * fits["cal_n"]["slope"] / 2
    for row in body:
        assert row[ix["Runs"]] == 2 and row[ix["Analyses"]] == 2 and row[ix["Name"]] == row[ix["Identifier"]]
        if row[ix["d15N (permil)"]] is None:
            continue
        sign = -1 if row[ix["Run"]] == "run_a" else 1
        assert abs(row[ix["d15N minus mean"]] - sign * half) < 1e-9
        if row[ix["d13C minus mean"]] is not None:
            assert abs(row[ix["d13C minus mean"]]) < 1e-9
        assert abs(row[ix["SD d15N"]] - SHIFT * fits["cal_n"]["slope"] / 2 ** 0.5) < 1e-9


def test_the_combined_csv_is_the_results_table_with_a_run_column(pair):
    r = send("/api/batch/results.csv", pair)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    lines = r.text.splitlines()
    assert lines[0].startswith("Run,Row,Identifier,Type,Mass (mg),%C,1sigma(%C)") and lines[0].isascii()
    assert len(lines) == 1 + 112
    one = client.post("/api/results.csv", files={"file": (pair[0].name, pair[0].read_bytes())}).text.splitlines()
    assert lines[1:57] == ["run_a," + line for line in one[1:]]


def test_each_run_takes_its_own_configuration(pair):
    off = config_to_dict(Config())
    on = dict(off, linearity_correction="On")
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([off, on]), config_name="lab.xlsx")
    runs = list(openpyxl.load_workbook(io.BytesIO(r.content))["Runs"].iter_rows(values_only=True))
    col = runs[0].index("Linearity applied")
    assert runs[1][col] == "C only" and runs[2][col] == "N and C"
    same = send("/api/batch/results.xlsx", pair, configs=json.dumps(on))
    runs = list(openpyxl.load_workbook(io.BytesIO(same.content))["Runs"].iter_rows(values_only=True))
    assert runs[1][col] == runs[2][col] == "N and C"
    readme = [c.value for c in openpyxl.load_workbook(io.BytesIO(r.content))["README"]["B"] if c.value]
    assert any("lab.xlsx" in t for t in readme)


def test_runs_name_each_runs_configuration_when_given(pair):
    """isoharness 0.5.0 (finding W-11): config_names gives the Runs sheet a Configuration column; without it the
    sheet is as before."""
    off = config_to_dict(Config())
    on = dict(off, linearity_correction="On")
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([off, on]),
             config_name="lab.xlsx for run 1; own.json for run 2", config_names=json.dumps(["lab.xlsx", "own.json"]))
    runs = list(openpyxl.load_workbook(io.BytesIO(r.content))["Runs"].iter_rows(values_only=True))
    col = runs[0].index("Configuration")
    assert [runs[1][col], runs[2][col]] == ["lab.xlsx", "own.json"]
    readme = [c.value for c in openpyxl.load_workbook(io.BytesIO(r.content))["README"]["B"] if c.value]
    assert any("Configuration column" in t for t in readme)
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([off, on]), config_names=json.dumps([None, "own.json"]))
    runs = list(openpyxl.load_workbook(io.BytesIO(r.content))["Runs"].iter_rows(values_only=True))
    assert runs[1][runs[0].index("Configuration")] == "the package defaults"
    plain = send("/api/batch/results.xlsx", pair, configs=json.dumps([off, on]))
    assert "Configuration" not in next(openpyxl.load_workbook(io.BytesIO(plain.content))["Runs"].iter_rows(values_only=True))
    bad = send("/api/batch/results.xlsx", pair, configs=json.dumps([off, on]), config_names=json.dumps(["one"]))
    assert bad.status_code == 422 and "config_names" in bad.text


def test_two_files_of_one_name_are_told_apart():
    runs = [BatchRun(file="20260915_raw.csv", positions=[], fits={}, config=Config()) for _ in range(3)]
    assert [r.label for r in label_runs(runs)] == ["20260915_raw", "20260915_raw (2)", "20260915_raw (3)"]
    assert [r.index for r in runs] == [1, 2, 3]


def test_a_sample_in_one_run_only_is_not_a_repeat(pair):
    positions, fits = process(parse_export(pair[0]), Config())
    runs = label_runs([BatchRun(file="a.csv", positions=positions, fits=fits, config=Config())])
    assert repeats(runs) == []
    text = combined_csv(runs)
    assert next(csv.reader(io.StringIO(text.splitlines()[1]))) == [str(v) for v in
                                                                   csv_cells(["a"] + results_values(positions[0]))]
    wb = openpyxl.load_workbook(io.BytesIO(send("/api/batch/results.xlsx", [RUN]).content))
    assert "No sample or check standard" in wb["Repeats"]["A1"].value


def test_refusals(pair, monkeypatch):
    r = client.post("/api/batch/results.xlsx", data={"configs": "{}"})
    assert r.status_code == 422
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([{}]))
    assert r.status_code == 422 and "one entry for each of the 2 files" in r.json()["detail"]
    r = send("/api/batch/results.xlsx", pair, configs="{not json")
    assert r.status_code == 422 and "configs could not be read" in r.json()["detail"]
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([{}, {"significance_level": "abc"}]))
    assert r.status_code == 422 and "run_b.csv" in r.json()["detail"]
    r = send("/api/batch/results.xlsx", pair, configs=json.dumps([{}, ["not", "an", "object"]]))
    assert r.status_code == 422 and r.json()["detail"] == "the configuration for run_b.csv is not a JSON object"
    r = client.post("/api/batch/results.xlsx", files=[("files", ("notes.pdf", b"%PDF"))])
    assert r.status_code == 415
    monkeypatch.setenv("ISOHARNESS_MAX_BATCH_RUNS", "1")
    r = send("/api/batch/results.xlsx", pair)
    assert r.status_code == 422 and "at most 1 runs" in r.json()["detail"]
    monkeypatch.setenv("ISOHARNESS_MAX_UPLOAD_MB", "0.03")
    monkeypatch.setenv("ISOHARNESS_MAX_BATCH_RUNS", "20")
    r = send("/api/batch/results.csv", pair)
    assert r.status_code == 413 and "together" in r.json()["detail"]


def test_batch_requests_count_as_uploads(pair, monkeypatch):
    from isoharness.app import limiter
    monkeypatch.setenv("ISOHARNESS_READS_PER_HOUR", "1")
    limiter.reset()
    try:
        assert send("/api/batch/results.csv", pair).status_code == 200
        assert send("/api/batch/results.xlsx", pair).status_code == 429
    finally:
        limiter.reset()


# ---------------------------------------------------------------- finding W-12: names that differ in their spaces
def _runs(*changes):
    """A batch of copies of the test run, the identifiers at some positions of each changed ({row: name})."""
    runs = []
    for i, changed in enumerate(changes, start=1):
        positions, fits = process(parse_export(RUN), Config())
        for p in positions:
            if p.row in changed:
                p.identifier = changed[p.row]
        runs.append(BatchRun(file=f"run{i}.csv", positions=positions, fits=fits, config=Config()))
    return label_runs(runs)


def _pairs(runs):
    """The Repeats entries of the changed names, the test run's other samples being all Random, as
    {Name: [(run, row, identifier, how its name matched)]}."""
    out = {}
    for e in repeats(runs):
        if e["Name"] != "Random":
            out.setdefault(e["Name"], []).append((e["Run"], e["Row"], e["Identifier"], e["Name matched"]))
    return out


def test_repeats_pair_names_that_differ_only_in_their_spaces():
    # colony 366 written two ways in two runs (finding W-12)
    assert _pairs(_runs({20: "366 - H+S"}, {20: "366-H+S"})) == {
        "366 - H+S": [("run1", 20, "366 - H+S", "as written"), ("run2", 20, "366-H+S", "spaces ignored")]}


def test_repeats_compare_names_as_written_first():
    # case and runs of spaces are ignored as written
    assert _pairs(_runs({21: "CI 3-2  H"}, {21: "ci 3-2 h"})) == {
        "CI 3-2  H": [("run1", 21, "CI 3-2  H", "as written"), ("run2", 21, "ci 3-2 h", "as written")]}
    # a name of one run joins, with its spaces taken out, the name that two other runs share as written
    assert _pairs(_runs({23: "PD 1"}, {23: "PD 1"}, {23: "PD1"})) == {
        "PD 1": [("run1", 23, "PD 1", "as written"), ("run2", 23, "PD 1", "as written"),
                 ("run3", 23, "PD1", "spaces ignored")]}
    # two names each found in two runs as written stay two samples, though they differ only in a space
    assert _pairs(_runs({24: "A B"}, {24: "A B", 25: "AB"}, {25: "AB"})) == {
        "A B": [("run1", 24, "A B", "as written"), ("run2", 24, "A B", "as written")],
        "AB": [("run2", 25, "AB", "as written"), ("run3", 25, "AB", "as written")]}


def test_repeats_pair_nothing_within_one_run_or_beyond_the_spaces():
    # two spellings in one run are not measured in two runs, and Dup is more than a space
    assert _pairs(_runs({20: "X Y", 21: "XY"}, {})) == {}
    assert _pairs(_runs({20: "366-H+S"}, {20: "366-H+S Dup"})) == {}


def test_the_repeats_sheet_says_which_names_matched_with_the_spaces_taken_out(tmp_path, pair):
    a, b = tmp_path / "run_a.csv", tmp_path / "run_b.csv"
    a.write_bytes(variant(name="S {}"))
    b.write_bytes(variant(SHIFT, name="S{}"))
    wb = openpyxl.load_workbook(io.BytesIO(send("/api/batch/results.xlsx", [a, b]).content))
    rows = list(wb["Repeats"].iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    assert rows[0][-1] == "Name matched" and len(rows) - 1 == 52
    for r in rows[1:]:
        first = r[ix["Run"]] == "run_a"
        assert r[ix["Name"]] == f"S {r[ix['Row']]}" and r[ix["Identifier"]] == (f"S {r[ix['Row']]}" if first else f"S{r[ix['Row']]}")
        assert r[ix["Name matched"]] == ("as written" if first else "spaces ignored") and r[ix["Runs"]] == r[ix["Analyses"]] == 2
    readme = [c.value for c in wb["README"]["B"] if c.value]
    assert any(t.startswith("Repeats: 26 names found in more than one run: S 20 (spaces ignored), S 21 (spaces ignored),")
               and t.endswith(", and others. 26 of them matched only with the spaces taken out.") for t in readme), readme
    assert any("366 - H+S and 366-H+S are one sample" in t for t in readme)
    # names that match as written leave the sheet as it was, without the column
    same = openpyxl.load_workbook(io.BytesIO(send("/api/batch/results.xlsx", pair).content))
    head = next(same["Repeats"].iter_rows(values_only=True))
    assert "Name matched" not in head and head[-1] == "Flags"


@pytest.mark.skipif(not NODE, reason="needs node for the page's logic")
def test_the_page_pairs_names_as_the_combined_results_do(tmp_path):
    changes = [{20: "366 - H+S", 21: "CI 3-2  H", 23: "PD 1", 24: "A B", 27: "X Y", 28: "XY"},
               {20: "366-H+S", 21: "ci 3-2 h", 23: "PD 1", 24: "A B", 25: "AB", 31: "366-H+S Dup"},
               {23: "PD1", 25: "AB", 27: "XY"}]
    files = [renamed(tmp_path, names, f"run{i}.csv") for i, names in enumerate(changes, start=1)]
    bodies = [client.post("/api/process", files={"file": (f.name, f.read_bytes())}).json() for f in files]
    page = run_page(bodies, "return L.repeats().map((g) => ({ name: g.name, runs: [...g.runs].sort(), count: g.count, alt: g.alt }));")
    got = {g["name"]: (g["runs"], g["count"], sorted(g["alt"])) for g in page}
    runs = []
    for f in files:
        positions, fits = process(parse_export(f), Config())
        runs.append(BatchRun(file=f.name, positions=positions, fits=fits, config=Config()))
    want = {}
    for e in repeats(label_runs(runs)):
        if e["Type"] != "Sample":
            continue
        g = want.setdefault(e["Name"], [set(), e["Analyses"], set()])
        g[0].add(int(e["Run"][3:]) - 1)
        if e["Name matched"] == "spaces ignored":
            g[2].add(e["Identifier"])
    want = {k: (sorted(v[0]), v[1], sorted(v[2])) for k, v in want.items()}
    assert got == want
    assert got["366 - H+S"] == ([0, 1], 2, ["366-H+S"]) and got["PD 1"] == ([0, 1, 2], 3, ["PD1"])
    assert got["XY"] == ([0, 2], 3, ["X Y"]) and got["A B"] == ([0, 1], 2, []) and got["AB"] == ([1, 2], 2, [])
    assert got["CI 3-2  H"] == ([0, 1], 2, []) and "366-H+S Dup" not in got
