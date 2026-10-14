"""The generator and the engine must agree on every number, new columns included.

Builds workbooks from the 2026-09-15 run under three configurations, has
LibreOffice compute the formulas, and compares each per-position column, the
fitted statistics and the flags against the Python pipeline. Skips when
LibreOffice or the reference template is not available.
"""
import csv
import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ea_irms.parse import parse_export            # noqa: E402
from ea_irms.pipeline import Config, process      # noqa: E402
from ea_irms.workbook import build_workbook       # noqa: E402

TEMPLATE = Path(os.environ.get(
    "EA_TEMPLATE",
    Path(__file__).resolve().parent / "data" / "EA-IRMS_Correction_Template_v1.xlsx"))
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
if not TEMPLATE.exists() or not SOFFICE:
    pytest.skip("needs the reference template (EA_TEMPLATE) and LibreOffice", allow_module_level=True)

COLUMNS = {"d15n_final": "AA", "d13c_final": "AF", "pct_n_measured": "AI", "pct_c_measured": "AJ",
           "cn_molar": "AK", "cn_mass": "AL", "ref_bgd_n": "BI", "settling_fraction": "BJ",
           "ref_pulse_d15n": "BT", "ref_pulse_d13c": "BW", "t_drift_n": "CB", "t_drift_c": "CC",
           "u_d15n": "CL", "U_d15n": "CM", "u_d13c": "CN", "U_d13c": "CO"}
FITS = {"B20": ("lin_n", "applied"), "B95": ("lin_n", "fit", "p"), "B96": ("lin_n", "in_force"),
        "B28": ("lin_c", "applied"), "B102": ("lin_c", "fit", "p"), "B103": ("lin_c", "in_force"),
        "B39": ("drift_n", "applied"), "B114": ("drift_n", "fit", "p"), "B115": ("drift_n", "in_force"),
        "B51": ("drift_c", "applied"), "B127": ("drift_c", "fit", "p"), "B128": ("drift_c", "in_force"),
        "B56": ("cal_n", "slope"), "B63": ("cal_c", "slope"),
        "B135": ("unc_n", "pooled_sd"), "B155": ("unc_c", "pooled_sd"),
        "B138": ("cal_n", "acc_min"), "B159": ("cal_c", "acc_max"),
        "B201": ("ref_gas_n", "n"), "B205": ("ref_gas_n", "fit", "slope"), "B206": ("ref_gas_n", "fit", "p"),
        "B208": ("ref_gas_c", "n"), "B213": ("ref_gas_c", "fit", "p"),
        "B217": ("checks", "n"), "B97": ("size_limits", "max_n"),
        "B85": ("settling", "excess"), "B86": ("settling", "settled_from"),
        "B229": ("carryover", "oversized"), "B230": ("carryover", "followed")}
CONFIGS = {
    "auto": Config(),
    "on_excluded": Config(blank_correction="On", linearity_correction="On", drift_correction="On",
                          exclude_drift_n=[40], exclude_drift_c=[40, 47], u_sample_c=0.03),
    "lin_split": Config(linearity_correction="On", linearity_correction_c="Off"),   # 1.4.1, Setup B18
}


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    folder = tmp_path_factory.mktemp("gen")
    raw = openpyxl.load_workbook(TEMPLATE)["Raw"]
    export = folder / "raw.csv"
    with open(export, "w", newline="") as fh:
        w = csv.writer(fh)
        for record in raw.iter_rows(max_col=16, values_only=True):
            if record[0] is not None:
                w.writerow(["" if v is None else v for v in record])
    files = []
    for name, cfg in CONFIGS.items():
        out = folder / f"{name}.xlsx"
        build_workbook(export, out, cfg)
        files.append(out)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(folder / "calc")] + [str(f) for f in files],
                   check=True, capture_output=True, timeout=600)
    return export, folder / "calc"


def _dig(fits, path):
    node = fits
    for key in path:
        node = node.get(key) if isinstance(node, dict) else None
        if node is None:
            return None
    return node


@pytest.mark.parametrize("name", list(CONFIGS))
def test_workbook_agrees_with_engine(built, name):
    export, calc = built
    positions, fits = process(parse_export(export), CONFIGS[name])
    wb = openpyxl.load_workbook(calc / f"{name}.xlsx", data_only=True)
    feeds = {c for c in range(1, wb["Work"].max_column + 1)
             if str(wb["Work"].cell(5, c).value or "").startswith("ch ")}   # NA() by design
    errors = [c.value for s in wb for row in s.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper()
              and not (s.title == "Work" and c.column in feeds and c.value == "#N/A")
              and not (s.title == "Raw" and str(s.cell(1, c.column).value or "").startswith("chart ")
                       and c.value == "#N/A")]   # 1.4.10: the pulses chart feeds, NA() by design
    assert not errors, errors[:5]
    wk, ft, rs = wb["Work"], wb["Fits"], wb["Results"]
    for attr, col in COLUMNS.items():
        for p in positions:
            mine = getattr(p, attr, None)
            theirs = wk[f"{col}{5 + p.row}"].value
            theirs = theirs if isinstance(theirs, (int, float)) else None
            assert (mine is None) == (theirs is None), f"{name} row {p.row} {attr}: {mine} vs {theirs}"
            if mine is not None:
                assert abs(mine - theirs) < 1e-9 * max(1.0, abs(mine)), f"{name} row {p.row} {attr}"
    for cell, path in FITS.items():
        mine = _dig(fits, path)
        theirs = ft[cell].value
        if isinstance(mine, bool):
            mine = "Yes" if mine else "No"
        if isinstance(mine, float):
            assert isinstance(theirs, (int, float)) and abs(mine - theirs) < 1e-9 * max(1.0, abs(mine)), \
                f"{name} {cell} {path}: {mine} vs {theirs}"
        else:
            assert mine == theirs, f"{name} {cell} {path}: {mine} vs {theirs}"
    flags_col = [c for c in range(1, rs.max_column + 1) if rs.cell(4, c).value == "Flags"][0]
    for p in positions:
        text = rs.cell(4 + p.row, flags_col).value or ""
        theirs = {t.strip() for t in text.split(";") if t.strip()}
        assert set(p.flags) == theirs, f"{name} row {p.row}: {set(p.flags) ^ theirs}"
