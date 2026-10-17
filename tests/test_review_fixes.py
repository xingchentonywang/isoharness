"""isoharness 0.4.1 and ea_irms 1.4.10: the fixes for the High and Medium findings of the 2026-10-14 review.

Each test names the finding it holds in place (R-01 to R-21; the review's workbook, sheet Findings). The
workbook tests have LibreOffice compute the formulas and skip without it, and the preview-port tests run
node and skip without it. The page's own fixes (R-03, R-14, R-15, R-17, R-18, and R-19) are checked in a
browser by scripts/e2e_web.py.
"""
import csv
import io
import json
import os
import shutil
import statistics
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from ea_irms.labconfig import LabConfigError, config_to_xlsx, read_config_json, read_lab_config  # noqa: E402
from ea_irms.metadata import read_sample_sheet                                                     # noqa: E402
from ea_irms.parse import ExportFormatError, parse_export                                          # noqa: E402
from ea_irms.pipeline import (DEFAULT_STANDARDS, MONITORS, Config, _xl, config_from_dict,          # noqa: E402
                              config_to_dict, process, run_key, scope_to_run)
from ea_irms.workbook import MAX_POSITION, MAX_RECORDS, VERSION, bake_values, build_workbook      # noqa: E402
from ea_irms.xlsxguard import MAX_READ_ROWS, UploadTooLarge, check_xlsx                            # noqa: E402
from isoharness.app import app, limiter                                                           # noqa: E402
from isoharness.serialize import results_csv                                                      # noqa: E402

DATA = ROOT / "tests" / "data"
RUN = DATA / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NODE = shutil.which("node")
client = TestClient(app)


def rows_of(path=RUN):
    return list(csv.reader(io.StringIO(Path(path).read_text(encoding="utf-8-sig"))))


def write_rows(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)
    return path


def renamed(folder, names, path_name="renamed.csv"):
    """The test run with the identifiers of some positions changed."""
    rows = rows_of()
    h = rows[0]
    row, ident = h.index("Row"), h.index("Identifier 1")
    for r in rows[1:]:
        if r and r[row].strip().isdigit() and int(r[row]) in names:
            r[ident] = names[int(r[row])]
    return write_rows(rows, Path(folder) / path_name)


def renumbered(folder, offset=100):
    """The test run with every position number moved up by ``offset``."""
    rows = rows_of()
    row = rows[0].index("Row")
    for r in rows[1:]:
        if r and r[row].strip().isdigit():
            r[row] = str(int(r[row]) + offset)
    return write_rows(rows, Path(folder) / f"renumbered_{offset + 1}.csv")


def post(path, name=RUN.name, content=None, **form):
    content = RUN.read_bytes() if content is None else content
    return client.post(path, files={"file": (name, content)}, data=form)


# ---------------------------------------------------------------- R-01, one anchor material
def test_r01_one_anchor_material_calibrates_nothing():
    """With USGS 65 out of the standards table, every anchor carries 1.76 permil in d15N. Before 1.4.10 every
    sample then read 1.76 with a small 1σ and no flag; now nothing is calibrated and the Rerun sheet says why."""
    std = [r for r in DEFAULT_STANDARDS if r[0] != "USGS 65"]
    positions, fits = process(parse_export(RUN), Config(standards=std))
    for el, iso in (("n", "d15N"), ("c", "d13C")):
        cal = fits[f"cal_{el}"]
        assert cal["calibrated"] is False and cal["values"] == 1 and cal["analyses"] == 3
        assert cal["slope"] is None
        assert (f"no {iso} calibrated: every anchor carries the same accepted {iso}, and the scale needs two "
                f"materials with different values") in fits["rerun"]["run_level_text"]
    samples = [p for p in positions if p.role == "Sample"]
    assert samples and all(p.d15n_final is None and p.u_d15n is None for p in samples)
    assert all(p.d13c_final is None and p.u_d13c is None for p in samples)


def test_r01_two_materials_still_calibrate():
    positions, fits = process(parse_export(RUN), Config())
    assert fits["cal_n"]["calibrated"] and fits["cal_n"]["values"] == 2 and fits["cal_n"]["analyses"] == 6
    assert "no d15N calibrated" not in fits["rerun"]["run_level_text"]


# ---------------------------------------------------------------- R-02, the workbook's size
def test_r02_positions_above_120_and_the_limits(tmp_path):
    """A run numbered 101 to 156 fills Work and Results to position 156; past the limits the workbook is
    refused with the reason instead of being cut short."""
    export = renumbered(tmp_path)
    out = build_workbook(export, tmp_path / "renumbered.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb["Work"].cell(5 + 156, 1).value == 156
    assert wb["Results"].cell(4 + 156, 1).value == '=IF(Work!$B161="","",Work!$A161)'

    rows = rows_of()
    row = rows[0].index("Row")
    for r in rows[1:]:
        if r and r[row].strip() == "56":
            r[row] = str(MAX_POSITION + 1)
    with pytest.raises(ValueError, match=f"positions run to {MAX_POSITION + 1}"):
        build_workbook(write_rows(rows, tmp_path / "too_far.csv"), tmp_path / "too_far.xlsx")

    rows = rows_of()
    extra = [list(r) for r in rows[1:] if r and r[row].strip() == "56"]
    copies = MAX_RECORDS // len(extra) + 1
    for k in range(copies):
        for r in extra:
            q = list(r)
            q[row] = str(57 + k % 300)
            rows.append(q)
    with pytest.raises(ValueError, match=f"holds {MAX_RECORDS}"):
        build_workbook(write_rows(rows, tmp_path / "too_many.csv"), tmp_path / "too_many.xlsx")


def test_r02_workbook_refusal_is_a_422(tmp_path):
    rows = rows_of()
    row = rows[0].index("Row")
    for r in rows[1:]:
        if r and r[row].strip() == "56":
            r[row] = str(MAX_POSITION + 1)
    data = write_rows(rows, tmp_path / "far.csv").read_bytes()
    r = post("/api/workbook", name="far.csv", content=data, values="false")
    assert r.status_code == 422 and f"positions 1 to {MAX_POSITION}" in r.json()["detail"]


# ---------------------------------------------------------------- R-04, uploads that unpack large
def fake_xlsx(size_mb, parts=3):
    """An .xlsx-shaped zip whose sheet unpacks to size_mb of spaces."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        for i in range(parts - 2):
            z.writestr(f"xl/part{i}.xml", "<x/>")
        z.writestr("xl/worksheets/sheet1.xml", b" " * int(size_mb * 1024 * 1024))
    return buf.getvalue()


def test_r04_an_xlsx_that_unpacks_large_is_refused_before_it_is_opened(tmp_path):
    blob = fake_xlsx(30)
    assert len(blob) < 200_000
    with pytest.raises(UploadTooLarge, match="unpacks to 30 MB, more than the 25 MB"):
        check_xlsx(blob, what="big.xlsx")
    path = tmp_path / "big.xlsx"
    path.write_bytes(blob)
    with pytest.raises(ExportFormatError, match="unpacks to 30 MB"):
        parse_export(path)
    r = post("/api/process", name="big.xlsx", content=blob)
    assert r.status_code == 422 and "unpacks to 30 MB" in r.json()["detail"]
    with pytest.raises(LabConfigError, match="unpacks to 6 MB, more than the 5 MB"):
        read_lab_config(fake_xlsx(6))
    with pytest.raises(UploadTooLarge, match="holds 600 files"):
        check_xlsx(fake_xlsx(0.001, parts=600))
    check_xlsx(b"not a zip at all")          # left to the reader, which says what is wrong in its own words


def test_r04_rows_past_the_limit_are_refused(tmp_path):
    rows = rows_of()
    body = [r for r in rows[1:] if r]
    while len(body) <= MAX_READ_ROWS:
        body += body
    path = write_rows([rows[0]] + body, tmp_path / "long.csv")
    with pytest.raises(ExportFormatError, match=f"more than {MAX_READ_ROWS} rows"):
        parse_export(path)


# ---------------------------------------------------------------- R-05, the visitor's address
def test_r05_a_made_up_forwarded_for_does_not_reset_the_limit(monkeypatch):
    """Each proxy adds the address it received the request from at the right-hand end, so on
    isoharness.org (two proxies) the visitor is second from the right, whatever the visitor wrote."""
    monkeypatch.setenv("ISOHARNESS_READS_PER_HOUR", "2")
    monkeypatch.setenv("ISOHARNESS_PROXY_HOPS", "2")
    limiter.reset()
    try:
        def send(chain):
            return client.post("/api/config/read", headers={"X-Forwarded-For": chain}).status_code
        assert send("10.0.0.1, 203.0.113.7, 198.51.100.2") == 422      # no file: refused, and counted
        assert send("10.0.0.2, 203.0.113.7, 198.51.100.2") == 422
        assert send("10.0.0.3, 203.0.113.7, 198.51.100.2") == 429      # the same visitor, a new made-up entry
        assert send("203.0.113.8, 198.51.100.2") == 422                 # another visitor
        seen = client.get("/api/client", headers={"X-Forwarded-For": "1.1.1.1, 203.0.113.7, 198.51.100.2"}).json()
        assert seen == {"address": "203.0.113.7", "forwarded_for_entries": 3, "proxy_hops": 2}
    finally:
        limiter.reset()


def test_r05_visitors_behind_one_proxy_are_counted_apart(monkeypatch):
    """Every visitor of isoharness.org reaches the service through the same proxies, so the address nearest
    the service is the same for all of them and must not carry a limit of its own."""
    monkeypatch.setenv("ISOHARNESS_READS_PER_HOUR", "1")
    monkeypatch.setenv("ISOHARNESS_PROXY_HOPS", "2")
    limiter.reset()
    try:
        codes = [client.post("/api/config/read", headers={"X-Forwarded-For": f"10.0.1.{i}, 198.51.100.9"}).status_code
                 for i in range(40)]
        assert codes == [422] * 40
    finally:
        limiter.reset()


# ---------------------------------------------------------------- R-06, text that reads as a formula
def test_r06_uploaded_text_stays_text(tmp_path):
    export = renamed(tmp_path, {21: "=1+1", 24: '=HYPERLINK("http://example.invalid","x")', 25: "@SUM(A1)"})
    sheet = read_sample_sheet(b'Sample ID,Notes\n=1+1,=2*3\n', "sheet.csv")
    out = build_workbook(export, tmp_path / "text.xlsx", samples=sheet,
                         run_info={"Operator": "=cmd|'/c calc'!A1", "Notes": "+1"}, file_name="=evil.csv")
    wb = openpyxl.load_workbook(out)
    injected = {"=1+1", '=HYPERLINK("http://example.invalid","x")', "=2*3", "=cmd|'/c calc'!A1", "=evil.csv"}
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if c.value in injected]
    assert not [(c.parent.title, c.coordinate) for c in cells if c.data_type != "s"]
    where = {(c.parent.title, c.value) for c in cells}
    assert {("Raw", "=1+1"), ("Raw", '=HYPERLINK("http://example.invalid","x")'), ("Metadata", "=2*3"),
            ("Metadata", "=evil.csv"), ("Setup", "=cmd|'/c calc'!A1")} <= where, where
    positions, _ = process(parse_export(export), Config(), samples=sheet)
    lines = list(csv.reader(io.StringIO(results_csv(positions, Config(), sheet))))
    cells = {c for line in lines for c in line}
    assert "'=1+1" in cells and "=1+1" not in cells and "'@SUM(A1)" in cells and "'=2*3" in cells


def test_r06_configuration_text_stays_text():
    cfg = Config(standards=[("=1+1", "Anchor", -10.0, 1.0, None, None, "=2*3", None, None, None, None)]
                 + list(DEFAULT_STANDARDS[1:]))
    wb = openpyxl.load_workbook(io.BytesIO(config_to_xlsx(cfg)))
    sd = wb["Standards"]
    assert sd["A5"].value == "=1+1" and sd["A5"].data_type == "s"
    assert sd["G5"].value == "=2*3" and sd["G5"].data_type == "s"
    assert read_lab_config(config_to_xlsx(cfg))[0].standards[0][0] == "=1+1"


# ---------------------------------------------------------------- R-07 and R-12, configuration values
def test_r07_switches_in_any_case_and_positions_as_text():
    cfg, notes = read_config_json({"linearity_correction": "off", "drift_model": "BACKGROUND",
                                   "exclude_drift_n": "40 47", "exclude_drift_c": "40; 47", "colour": 1})
    assert cfg.linearity_correction == "Off" and cfg.drift_model == "Background"
    assert cfg.exclude_drift_n == [40, 47] and cfg.exclude_drift_c == [40, 47]
    assert notes == [f"colour is not a setting ea_irms {VERSION} reads, so it was left out"]
    assert config_from_dict({"blank_correction": "on"}).blank_correction == "On"
    for bad, word in (({"drift_model": "Linear"}, "drift_model is 'Linear'"),
                      ({"exclude_drift_n": "40, x"}, "exclude_drift_n is '40, x'"),
                      ({"exclude_elemental": [0]}, "exclude_elemental")):
        with pytest.raises(LabConfigError, match=word):
            read_config_json(bad)
    with pytest.raises(ValueError, match="blank_correction is 'Of'"):
        Config(blank_correction="Of")


def test_r12_standards_percentages_and_1sigma_are_checked():
    rows = [list(r) for r in DEFAULT_STANDARDS]
    rows[0][4], rows[1][7] = 320, -0.1
    with pytest.raises(LabConfigError) as err:
        read_config_json({"standards": rows})
    assert "standards row 1, item 5 (USGS 64) is 320; it needs a percentage from 0 to 100" in str(err.value)
    assert "standards row 2, item 8 (USGS 65) is -0.1; a 1σ is a number of at least 0" in str(err.value)


@pytest.mark.parametrize("text,field", [
    ('{"significance_level": NaN}', "significance_level"),
    ('{"significance_level": Infinity}', "significance_level"),
    ('{"coverage_k": -5}', "coverage_k"),
    ('{"standards": "nope"}', "standards"),
    ('{"min_area_n": "ten"}', "min_area_n"),
    ('{"carryover_mv": true}', "carryover_mv"),
    ('[1, 2]', "one JSON object"),
])
def test_r12_bad_configuration_values_are_a_422_that_names_them(text, field):
    r = post("/api/process", config=text)
    assert r.status_code == 422, r.text
    assert field in r.json()["detail"]


# ---------------------------------------------------------------- R-08, values that belong to one run
def test_r08_scope_to_run():
    cfg = Config(exclude_drift_c=[59], exclude_elemental=[11, 58], first_drift_position=5,
                 u_sample_n=0.181, u_sample_n_source="drift monitors", u_sample_c=0.05,
                 run_file="09222026_PanamaSids3_fullexport_.xls")
    same, notes = scope_to_run(cfg, "09222026_panamasids3_FULLEXPORT_.csv")      # any case, any extension
    assert same is cfg and notes == []
    other, notes = scope_to_run(cfg, "12192025_Panama_Coral7_.xls")
    assert (other.first_drift_position, other.exclude_drift_c, other.exclude_elemental) == (None, [], [])
    assert other.u_sample_n is None and other.u_sample_n_source == MONITORS     # the source stays, the value goes
    assert other.u_sample_c == 0.05                                              # the lab's own value stays
    assert other.standards == cfg.standards and other.coverage_k == cfg.coverage_k
    assert cfg.exclude_drift_c == [59]                                           # the original is untouched
    assert notes == ["the first drift position (5), the position left out of the carbon drift fit (59), the "
                     "positions left out of the elemental correction (11, 58), and the d15N repeatability from "
                     "the drift monitors (0.181 permil) were chosen for 09222026_PanamaSids3_fullexport_.xls, so "
                     "they were not applied to 12192025_Panama_Coral7_.xls"]
    unnamed, notes = scope_to_run(Config(exclude_drift_n=[56]), "Coral5.xls")
    assert unnamed.exclude_drift_n == [56]
    assert notes == ["no run is named for the position left out of the nitrogen drift fit (56), so it was "
                     "applied to Coral5.xls; check that it belongs to this run"]
    assert scope_to_run(Config(run_file="A.xls"), "B.xls")[1] == []              # nothing for one run only
    assert run_key("C:\\data\\Run1.XLS") == run_key("/data/run1.csv") == "run1"
    assert run_key("2026.09.15 run") == run_key("2026.09.15 run.xls") != run_key("2026.09")
    assert run_key("batch.1") != run_key("batch.2")
    assert run_key("Cafe\u0301_run.xls") == run_key("Caf\u00e9_run.CSV")
    shared, notes = scope_to_run(Config(exclude_drift_n=[56]), "Coral5.xls", shared=True)
    assert shared.exclude_drift_n == []
    assert notes == ["no run is named for the position left out of the nitrogen drift fit (56), and the "
                     "configuration was given to several runs at once, so it was applied to none of them"]


def test_r08_the_run_name_and_sources_round_trip():
    cfg = Config(exclude_drift_c=[59], u_sample_n=0.181, u_sample_n_source=MONITORS, run_file="Sids3.xls")
    assert config_from_dict(json.loads(json.dumps(config_to_dict(cfg)))) == cfg
    back, notes = read_lab_config(config_to_xlsx(cfg))
    assert back == cfg and notes == []
    old = openpyxl.load_workbook(io.BytesIO(config_to_xlsx(cfg)))       # a template written before 1.4.10
    st = old["Settings"]
    for r in range(st.max_row, 0, -1):
        if st.cell(r, 4).value in ("run_file", "u_sample_n_source", "u_sample_c_source"):
            st.delete_rows(r)
    buf = io.BytesIO()
    old.save(buf)
    back, notes = read_lab_config(buf.getvalue())
    assert back.run_file is None and back.u_sample_n_source is None and back.exclude_drift_c == [59]
    assert any("run_file" in t for t in notes)
    for bad in ({"run_file": True}, {"run_file": "x" * 300}, {"u_sample_n_source": "anchors"}):
        with pytest.raises(LabConfigError):
            read_config_json(bad)


def test_r08_process_applies_them_to_their_own_run_only():
    cfg = config_to_dict(Config(exclude_drift_n=[40], exclude_drift_c=[40, 47], run_file=RUN.name))
    own = post("/api/process", config=json.dumps(cfg)).json()
    assert own["config"]["exclude_drift_c"] == [40, 47] and own["config_notes"] == []
    other = post("/api/process", name="20260916_raw.csv", config=json.dumps(cfg)).json()
    assert other["config"]["exclude_drift_n"] == [] and other["config"]["exclude_drift_c"] == []
    assert other["config"]["run_file"] == RUN.name
    assert other["config_notes"] == ["the position left out of the nitrogen drift fit (40) and the positions left "
                                     "out of the carbon drift fit (40, 47) were chosen for 20260915_raw.csv, so "
                                     "they were not applied to 20260916_raw.csv"]
    by_row = {p["row"]: p for p in other["positions"]}
    plain = {p["row"]: p for p in post("/api/process", name="20260916_raw.csv").json()["positions"]}
    assert all(by_row[r]["d15n_final"] == plain[r]["d15n_final"] for r in plain)     # as with no exclusions
    unnamed = dict(cfg, run_file=None)
    legacy = post("/api/process", name="20260916_raw.csv", config=json.dumps(unnamed)).json()
    assert legacy["config"]["exclude_drift_c"] == [40, 47]
    assert legacy["config_notes"] == ["no run is named for the position left out of the nitrogen drift fit (40) "
                                      "and the positions left out of the carbon drift fit (40, 47), so they were "
                                      "applied to 20260916_raw.csv; check that they belong to this run"]


def test_r08_downloads_say_what_was_not_applied(tmp_path):
    cfg = json.dumps(config_to_dict(Config(exclude_drift_c=[47], run_file="other.csv")))
    r = post("/api/results.csv", config=cfg)
    assert r.status_code == 200
    assert json.loads(r.headers["x-isoharness-config-notes"])[0].endswith("not applied to 20260915_raw.csv")
    r = post("/api/workbook", config=cfg, values="false")
    assert r.status_code == 200 and "x-isoharness-config-notes" in r.headers
    setup = openpyxl.load_workbook(io.BytesIO(r.content))["Setup"]
    assert setup["A3"].value.startswith("From the configuration: the position left out of the carbon drift fit "
                                        "(47) was chosen for other.csv")
    assert setup["B17"].value in (None, "")
    plain = post("/api/workbook", values="false")
    assert "x-isoharness-config-notes" not in plain.headers
    assert openpyxl.load_workbook(io.BytesIO(plain.content))["Setup"]["A3"].value is None
    files = [("files", ("20260915_raw.csv", RUN.read_bytes())), ("files", ("copy.csv", RUN.read_bytes()))]
    one = json.dumps(config_to_dict(Config(exclude_drift_c=[47], run_file="copy.csv")))
    b = client.post("/api/batch/results.xlsx", files=files, data={"configs": one})
    assert b.status_code == 200, b.text
    runs = openpyxl.load_workbook(io.BytesIO(b.content))["Runs"]
    head = [c.value for c in runs[1]]
    col = head.index("Configuration notes") + 1
    assert runs.cell(2, col).value.endswith("so it was not applied to 20260915_raw.csv")
    assert runs.cell(3, col).value is None


def test_r08_file_downloads_carry_every_note():
    cfg = json.dumps(dict(config_to_dict(Config(exclude_drift_c=[47], run_file="other.csv")), colour="red"))
    r = post("/api/results.csv", config=cfg)
    notes = json.loads(r.headers["x-isoharness-config-notes"])
    assert notes[0] == f"colour is not a setting ea_irms {VERSION} reads, so it was left out" and len(notes) == 2
    files = [("files", ("20260915_raw.csv", RUN.read_bytes())), ("files", ("copy.csv", RUN.read_bytes()))]
    one = json.dumps(config_to_dict(Config(exclude_drift_c=[47])))          # names no run, for both at once
    for path in ("/api/batch/results.csv", "/api/batch/results.xlsx"):
        b = client.post(path, files=files, data={"configs": one})
        assert b.status_code == 200, b.text
        notes = json.loads(b.headers["x-isoharness-config-notes"])
        assert notes == [f"{f}: no run is named for the position left out of the carbon drift fit (47), and the "
                         "configuration was given to several runs at once, so it was applied to none of them"
                         for f in ("20260915_raw.csv", "copy.csv")]
    both = json.dumps([json.loads(one), json.loads(one)])                     # one each: each run's own
    b = client.post("/api/batch/results.csv", files=files, data={"configs": both})
    assert all("so it was applied to" in n and "check that it belongs" in n
               for n in json.loads(b.headers["x-isoharness-config-notes"]))


def test_r08_the_notes_header_is_always_a_json_list():
    from isoharness.app import NOTES_HEADER_MAX, _notes_header
    many = [f"note {i}: " + "x" * 300 for i in range(40)]
    text = _notes_header(many)["X-Isoharness-Config-Notes"]
    got = json.loads(text)
    assert len(text) <= NOTES_HEADER_MAX and got[-1].startswith("and ") and got[-1].endswith(" more")
    assert len(got) - 1 + int(got[-1].split()[1]) == 40
    assert json.loads(_notes_header(["one"])["X-Isoharness-Config-Notes"]) == ["one"]
    assert _notes_header([]) == {}
    one_long = json.loads(_notes_header(["y" * 5000])["X-Isoharness-Config-Notes"])
    assert one_long == ["y" * 1000 + "..."]


def test_r08_command_line_says_so(tmp_path):
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps(config_to_dict(Config(exclude_drift_n=[40], run_file="another.xls"))))
    out = subprocess.run([sys.executable, "-m", "ea_irms.cli", str(RUN), "--config", str(cfg), "--out",
                          str(tmp_path / "o.xlsx"), "--no-recalc"], capture_output=True, text=True, cwd=ROOT)
    assert out.returncode == 0, out.stderr
    assert "note: the position left out of the nitrogen drift fit (40) was chosen for another.xls" in out.stderr
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"coverage_k": -5}))
    out = subprocess.run([sys.executable, "-m", "ea_irms.cli", str(RUN), "--config", str(bad), "--out",
                          str(tmp_path / "b.xlsx"), "--no-recalc"], capture_output=True, text=True, cwd=ROOT)
    assert out.returncode == 2 and "coverage_k is -5" in out.stderr


# ---------------------------------------------------------------- R-09, below detection
def test_r09_a_peak_under_the_intercept_is_below_detection():
    """PanamaSids2 position 25 had 1.808 Vs of N2 against an intercept of 1.907 Vs, and read %N -0.000978
    with 1σ -0.0000776. Here a sample of the test run is given 0.5 mg and 1.5 Vs, under its 1.985 Vs."""
    P = parse_export(RUN)
    q = next(p for p in P if p.row == 21)
    q.amount, q.area_n = 0.5, 1.5
    positions, fits = process(P, Config())
    p = next(x for x in positions if x.row == 21)
    assert fits["k_n"]["intercept"] > 1.5
    assert p.below_detection_n and not p.below_detection_c
    assert p.pct_n_measured is None and p.pct_n_reported is None and p.u_pct_n is None
    assert p.cn_molar_reported is None and p.u_cn_molar is None
    assert p.pct_c_reported is not None and p.pct_c_reported > 0
    # 1.4.15: the acetanilide's whole-run check, which fires on the test run, puts its flag after it, last
    assert p.flags[-2:] == ["%N below detection", "acetanilide d15N scatter above 1σ"]
    assert all(v is None or v >= 0 for x in positions for v in (x.u_pct_n, x.u_pct_c, x.u_cn_molar))


# ---------------------------------------------------------------- R-11, prefix matches
def test_r11_a_prefix_match_takes_no_replicate_of(tmp_path):
    """Coral1 with CI 2-2-H (Replicate of G4) and CI 2-2-S in the sheet: CI 2-2-H+S matched CI 2-2-H by
    prefix and joined its replicate group. Tony, 2026-10-14: the rule stays, and a prefix match does not
    take Replicate of."""
    export = renamed(tmp_path, {21: "CI 2-2-H", 24: "CI 2-2-H+S", 25: "G4", 31: "CI 2-2-S"})
    sheet = read_sample_sheet(b"Sample ID,Site,Replicate of\nCI 2-2-H,CI,G4\nCI 2-2-S,CI,\nG4,CI,\n", "s.csv")
    positions, fits = process(parse_export(export), Config(), samples=sheet)
    at = {p.row: p for p in positions}
    assert at[24].sample_id == "CI 2-2-H" and at[24].matched_by.startswith("prefix")
    assert at[24].replicate_group is None
    assert at[21].replicate_group is not None and at[21].replicate_group == at[25].replicate_group
    assert at[31].replicate_group is None
    groups = fits["replicates"]["groups"]
    assert [sorted(g["rows"]) for g in groups] == [[21, 25]]
    # an analysis matched by a shorter ID is a replicate only of analyses with its own name
    export = renamed(tmp_path, {21: "CI 2-2-H", 24: "CI 2-2-H+S", 25: "G4", 26: "CI 2-2-H+S"}, "twice.csv")
    positions, fits = process(parse_export(export), Config(), samples=sheet)
    assert sorted(sorted(g["rows"]) for g in fits["replicates"]["groups"]) == [[21, 25], [24, 26]]


# ---------------------------------------------------------------- R-13, every reference pulse
def test_r13_reference_gas_sd_over_every_pulse():
    positions, fits = process(parse_export(RUN), Config())
    allc = [v for p in positions for v in p.ref_pulses_d13c]
    assert len(allc) == 112 and fits["ref_gas_c"]["n"] == 112 and fits["ref_gas_c"]["positions"] == 56
    assert abs(fits["ref_gas_c"]["sd"] - statistics.stdev(allc)) < 1e-15
    assert _xl(fits["ref_gas_c"]["sd"], 3) == "0.082"
    first = [p.ref_pulse_d13c for p in positions if p.ref_pulse_d13c is not None]
    assert _xl(statistics.stdev(first), 3) == "0.065"          # the first pulse alone, as before 1.4.10
    assert fits["ref_gas_n"]["n"] == 56


def test_r13_a_second_pulse_off_marks_the_position(tmp_path):
    rows = rows_of()
    h = rows[0]
    c = {k: h.index(k) for k in ("Row", "Ampl  44", "Is Ref _", "d 13C/12C", "Peak Nr")}
    recs = [r for r in rows[1:] if r and r[c["Row"]].strip() == "30" and r[c["Peak Nr"]].strip()
            and len(r) > c["Ampl  44"] and r[c["Ampl  44"]].strip()]
    first_ref = next(i for i, r in enumerate(recs) if r[c["Is Ref _"]].strip() == "1")
    later = [r for r in recs[first_ref + 1:] if r[c["Is Ref _"]].strip() != "1"]
    assert len(later) == 2
    later[1][c["d 13C/12C"]] = repr(float(later[1][c["d 13C/12C"]]) + 1.0)
    positions, _ = process(parse_export(write_rows(rows, tmp_path / "pulse.csv")), Config())
    p = next(x for x in positions if x.row == 30)
    assert any(f.startswith("CO2 reference pulse off by +") for f in p.flags), p.flags


# ---------------------------------------------------------------- R-16, LibreOffice's queue
def test_r16_workbooks_wait_a_limited_time(monkeypatch):
    from isoharness import app as appmod
    monkeypatch.setenv("ISOHARNESS_WORKBOOK_WAIT_S", "0.2")
    monkeypatch.setenv("ISOHARNESS_WORKBOOK_QUEUE", "3")
    monkeypatch.setattr(appmod, "bake_values", lambda path: False)
    assert appmod._libreoffice.acquire(timeout=5)
    try:
        t0 = time.monotonic()
        r = post("/api/workbook")
        assert r.status_code == 503 and r.headers["retry-after"] == "60"
        assert "computing other workbooks" in r.json()["detail"]
        assert time.monotonic() - t0 < 10
        monkeypatch.setenv("ISOHARNESS_WORKBOOK_QUEUE", "0")
        assert post("/api/workbook").status_code == 503
    finally:
        appmod._libreoffice.release()
    monkeypatch.setenv("ISOHARNESS_WORKBOOK_QUEUE", "3")
    assert post("/api/workbook").status_code == 200


def test_r16_the_default_wait_ends_inside_hostings_60_s(monkeypatch):
    from isoharness.app import workbook_queue, workbook_wait_s
    monkeypatch.delenv("ISOHARNESS_WORKBOOK_WAIT_S", raising=False)
    monkeypatch.delenv("ISOHARNESS_WORKBOOK_QUEUE", raising=False)
    assert workbook_wait_s() == 45.0 and workbook_queue() == 3


@pytest.mark.skipif(os.name != "posix", reason="process groups")
def test_r16_a_stuck_libreoffice_is_stopped(tmp_path, monkeypatch):
    fake = tmp_path / "bin" / "soffice"
    fake.parent.mkdir()
    pid_file = tmp_path / "child.pid"
    fake.write_text(f"#!/bin/sh\nsleep 60 &\necho $! > {pid_file}\nsleep 60\n")
    fake.chmod(0o755)
    monkeypatch.setenv("PATH", f"{fake.parent}{os.pathsep}{os.environ['PATH']}")
    book = tmp_path / "w.xlsx"
    book.write_bytes(b"x")
    t0 = time.monotonic()
    assert bake_values(book, timeout=1.0) is False
    assert time.monotonic() - t0 < 15
    pid = int(pid_file.read_text())
    for _ in range(50):
        status = Path(f"/proc/{pid}/status")
        if not status.exists() or "\nState:\tZ" in status.read_text():
            break
        time.sleep(0.1)
    else:
        pytest.fail("the soffice child was left running")


# ---------------------------------------------------------------- R-20, the carbon blank
def test_r20_the_carbon_blank_correction_runs(tmp_path):
    """No run in the repository has a carbon blank, so the correction had never run on any data. The
    constructed variant gives each blank capsule a sample CO2 peak (illustrative values, port_cases.py)."""
    from port_cases import carbon_blank_variant
    export = carbon_blank_variant(RUN, tmp_path / "carbon_blank.csv", Config().standards)
    positions, fits = process(parse_export(export), Config(blank_correction="On"))
    assert fits["blank_c"]["in_force"] and fits["blank_c"]["n"] >= 3
    shifted = [p for p in positions if p.role == "Sample" and p.d13c is not None and p.d13c_blank is not None]
    assert shifted and any(abs(p.d13c_blank - p.d13c) > 1e-6 for p in shifted)
    off, _ = process(parse_export(export), Config(blank_correction="Off"))
    assert any(abs(a.d13c_final - b.d13c_final) > 1e-6 for a, b in zip(positions, off)
               if a.d13c_final is not None and b.d13c_final is not None)


@pytest.mark.skipif(not NODE, reason="needs node for the preview port")
def test_r20_the_port_agrees_on_the_carbon_blank(tmp_path):
    out = tmp_path / "port"
    gen = subprocess.run([sys.executable, str(ROOT / "scripts" / "port_cases.py"), str(RUN), "--out", str(out),
                          "--with-carbon-blank", "--random", "40", "--per-variant", "3"],
                         capture_output=True, text=True, cwd=ROOT)
    assert gen.returncode == 0, gen.stderr
    run = subprocess.run([NODE, str(ROOT / "port" / "test_port.js"), str(out / "port_input.json"),
                          str(out / "cases.json")], capture_output=True, text=True, cwd=ROOT)
    assert run.returncode == 0, run.stdout[-2000:] + run.stderr[-2000:]
    assert "value mismatches 0, flag-list mismatches 0, fit mismatches 0" in run.stdout, run.stdout[-500:]


# ---------------------------------------------------------------- R-21, a half at the last place
@pytest.mark.parametrize("x,places,sign,want", [
    (0.0005, 3, False, "0.001"), (-0.0005, 3, False, "-0.001"), (0.125, 2, True, "+0.13"),
    (-0.125, 2, True, "-0.13"), (-0.0001, 2, True, "-0.00"), (0.0, 2, True, "+0.00"), (-0.0, 2, True, "+0.00"),
    (2.675, 2, False, "2.68"), (1e-20, 2, False, "0.00"),
])
def test_r21_rounding_as_excel_rounds(x, places, sign, want):
    assert _xl(x, places, sign) == want


@pytest.mark.skipif(not NODE, reason="needs node for the preview port")
def test_r21_the_port_rounds_the_same():
    cases = [(0.0005, 3, False), (-0.0005, 3, False), (0.125, 2, True), (-0.125, 2, True), (-0.0001, 2, True),
             (2.675, 2, False), (0.35, 1, True), (123.4565, 3, False)]
    script = ("const E = require(process.argv[1]); const C = JSON.parse(process.argv[2]);"
              "console.log(JSON.stringify(C.map(([x, p, s]) => E.xl(x, p, s))));")
    out = subprocess.run([NODE, "-e", script, str(ROOT / "port" / "engine.js"), json.dumps(cases)],
                         capture_output=True, text=True)
    assert out.returncode == 0, out.stderr
    assert json.loads(out.stdout) == [_xl(*c) for c in cases]


# ---------------------------------------------------------------- the workbook against the engine
CASES = {
    "renumbered": (lambda d: renumbered(d), Config()),
    "carbon_blank": (None, Config(blank_correction="On")),
    "below_detection": (None, Config()),
}


@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    if not SOFFICE:
        pytest.skip("needs LibreOffice")
    from port_cases import carbon_blank_variant
    folder = tmp_path_factory.mktemp("review")
    exports = {"renumbered": renumbered(folder),
               "carbon_blank": carbon_blank_variant(RUN, folder / "carbon_blank.csv", Config().standards)}
    rows = rows_of()
    h = rows[0]
    c = {k: h.index(k) for k in ("Row", "Amount", "Area 28", "Ampl  28", "Is Ref _", "Peak Nr")}
    for r in rows[1:]:
        if r and r[c["Row"]].strip() == "21":
            r[c["Amount"]] = "0.5"
    recs = [r for r in rows[1:] if r and r[c["Row"]].strip() == "21" and r[c["Peak Nr"]].strip()
            and len(r) > c["Ampl  28"] and r[c["Ampl  28"]].strip()]
    recs[-1][c["Area 28"]] = "1.5"                          # the sample N2 peak, under the intercept
    exports["below_detection"] = write_rows(rows, folder / "below_detection.csv")
    files = []
    for name, export in exports.items():
        out = folder / f"{name}.xlsx"
        build_workbook(export, out, CASES[name][1])
        files.append(out)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + [str(f) for f in files], check=True, capture_output=True, timeout=600)
    return exports, folder / "calc"


WORK = {"d15n_blank": "W", "d13c_blank": "AB", "d15n_final": "AA", "d13c_final": "AF", "pct_n_measured": "AI",
        "pct_c_measured": "AJ", "cn_molar": "AK", "u_d15n": "CL", "u_d13c": "CN"}


@pytest.mark.parametrize("name", list(CASES))
def test_workbook_agrees_with_the_engine(computed, name):
    exports, calc = computed
    positions, fits = process(parse_export(exports[name]), CASES[name][1])
    wb = openpyxl.load_workbook(calc / f"{name}.xlsx", data_only=True)
    wk, rs = wb["Work"], wb["Results"]
    for attr, col in WORK.items():
        for p in positions:
            mine, theirs = getattr(p, attr, None), wk[f"{col}{5 + p.row}"].value
            theirs = theirs if isinstance(theirs, (int, float)) else None
            assert (mine is None) == (theirs is None), f"{name} row {p.row} {attr}: {mine} vs {theirs}"
            if mine is not None:
                assert abs(mine - theirs) < 1e-9 * max(1.0, abs(mine)), f"{name} row {p.row} {attr}"
    head = {rs.cell(4, j).value: j for j in range(1, rs.max_column + 1)}
    for p in positions:
        text = rs.cell(4 + p.row, head["Flags"]).value or ""
        assert p.flags == [t.strip() for t in text.split(";") if t.strip()], f"{name} row {p.row}"
        assert (rs.cell(4 + p.row, head["Rerun"]).value or None) == (p.rerun or None), f"{name} row {p.row}"
    if name == "carbon_blank":
        assert fits["blank_c"]["in_force"]
    if name == "below_detection":
        flags = rs.cell(4 + 21, head["Flags"]).value
        # 1.4.15: the acetanilide's whole-run check, which fires on the test run, writes its flag after it, last
        assert flags.rstrip("; ").endswith("%N below detection; acetanilide d15N scatter above 1σ")
