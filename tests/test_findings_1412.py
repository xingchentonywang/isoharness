"""ea_irms 1.4.12: findings W-02, W-03, W-09 and W-10 of 2026-10-14.

Seven runs of another project, put through isoharness.org on 2026-10-14, showed four things this version
changes. Their exports stay out of this repository, so each test writes a copy of the 2026-09-15 run that
shows the case:

W-09  Two unflagged CO2 peaks before the designated pulse. The engine took the one next to the pulse and the
      workbook the first; both now take the one next to it, and Flags says so.
W-10  An anchor material analyzed once added nothing for its mean to the 1σ. It now takes the repeatability
      in use (Setup B30 or B31, else the pooled SD of the materials analyzed more than once).
W-03  In a filtered export, a no-gas CO2 peak before the first designated pulse, where the amplitude
      yardstick is held flat, missed the 2% tolerance. There the tolerance is now 5% (Setup B41), and never
      narrower than the tolerance inside the span.
W-02  Nothing judged the calibration line. Its slope against 1 and the anchors' residual SD are now judged
      at Note and Caution limits (Setup B73 to B76), and either is a problem of the whole run.

A review of the change found three more things, tested at the end: a configuration written before 1.4.12
with a tolerance above 0.05 got a narrower one outside the span; a limit below 0.0001 read 5e-05 in the
engine's text and 0.00005 in the workbook's; and a repeatability of 0 reached Setup B30 as an empty cell. It
also found the workbook's Flags in another order than the engine's, so the flags are compared as lists.

The workbook tests have LibreOffice compute the formulas and skip without it; the port tests need node.
"""
import csv
import json
import math
import shutil
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from ea_irms.labconfig import KEYS, LabConfigError, config_from_xlsx, config_to_xlsx, read_config_json  # noqa: E402
from ea_irms.parse import parse_export                                                          # noqa: E402
from ea_irms.pipeline import DEFAULT_STANDARDS, Config, _num_text, config_to_dict, process       # noqa: E402
from ea_irms.workbook import build_workbook                                                      # noqa: E402
from port_cases import filtered_variant                                                          # noqa: E402

RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NODE = shutil.which("node")
W09_ROW = 20                       # a sample of the test run with an N2 and a CO2 peak
W09_FLAG = "2 CO2 peaks before the reference pulse, the last one used"
LEFT_OUT = "USGS 64 left out"


def _rows(path=RUN):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], rows[1:]


def _write(dst, head, body):
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        w.writerows(body)
    return dst


def two_co2_variant(dst):
    """The test run with a small unflagged CO2 peak (illustrative values) put before the sample CO2 peak of
    position 20; that peak and those after it are numbered one up."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    at = [r for r in body if r[i["Row"]] == str(W09_ROW)]
    sample = next(r for r in at if r[i["Ampl  44"]].strip() and r[i["Is Ref _"]].strip() != "1")
    k = int(float(sample[i["Peak Nr"]]))
    extra = list(sample)
    extra[i["Ampl  44"]], extra[i["Area 44"]], extra[i["d 13C/12C"]] = "40", "2.331", "-58.620"
    out = []
    for r in body:
        if r[i["Row"]] == str(W09_ROW) and r[i["Peak Nr"]].strip() and int(float(r[i["Peak Nr"]])) >= k:
            if r is sample:
                out.append(extra)
            r = list(r)
            r[i["Peak Nr"]] = str(int(float(r[i["Peak Nr"]])) + 1)
        out.append(r)
    return _write(dst, head, out)


def single_anchor_variant(dst):
    """The test run with the second and third USGS 64 analyses renamed, so that a standards row with the role
    Exclude can leave them out and one USGS 64 analysis anchors the line."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    rows64 = []
    for r in body:
        if r[i["Identifier 1"]].strip() == "USGS 64" and r[i["Row"]] not in rows64:
            rows64.append(r[i["Row"]])
    assert len(rows64) == 3
    out = []
    for r in body:
        if r[i["Row"]] in rows64[1:]:
            r = list(r)
            r[i["Identifier 1"]] = LEFT_OUT
        out.append(r)
    return _write(dst, head, out)


SINGLE = Config(standards=list(DEFAULT_STANDARDS) + [
    (LEFT_OUT, "Exclude", None, None, None, None, "left out for the test", None, None, None, None)])


def once_variant(dst):
    """The test run with each anchor material analyzed once: the second and third USGS 64 and USGS 65
    analyses renamed as in single_anchor_variant, and the acetanilide of position 8 made a third anchor, so
    that no material has replicates of its own."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    seen = {}
    for r in body:
        nm = r[i["Identifier 1"]].strip()
        if nm in ("USGS 64", "USGS 65") and r[i["Row"]] not in seen.setdefault(nm, []):
            seen[nm].append(r[i["Row"]])
    assert [len(v) for v in seen.values()] == [3, 3]
    later = {row for rows_ in seen.values() for row in sorted(rows_, key=int)[1:]}
    out = []
    for r in body:
        if r[i["Row"]] in later or r[i["Row"]] == "8":
            r = list(r)
            r[i["Identifier 1"]] = LEFT_OUT if r[i["Row"]] in later else THIRD
        out.append(r)
    return _write(dst, head, out)


THIRD = "Third anchor"
# illustrative accepted values, near what position 8 reads
ONCE = replace(SINGLE, standards=list(SINGLE.standards) + [
    (THIRD, "Anchor", -26.53, -4.99, 71.09, 10.36, "acetanilide as a third anchor, for the test", None, None, None, None)])


def pulse_variant(dst):
    """The test run as a filtered export (port_cases.filtered_variant: designated CO2 pulses at the five
    positions without sample gas, 1, 4, 5, 53 and 56, each with an undesignated pulse ahead of it), with the
    designated pulses of 1, 4 and 5 taken out, so that the yardstick holds the amplitude at 53 flat before it.
    The undesignated pulses at 1 and 4 are set 3% below it and the one at 5 6% below; the one at 56, inside
    the designated pulses' span, 3% below the designated pulse there. Illustrative amplitudes, areas scaled."""
    src = filtered_variant(RUN, dst.with_name("filtered_base.csv"))
    head, body = _rows(src)
    i = {k: n for n, k in enumerate(head)}
    flagged = {r[i["Row"]]: float(r[i["Ampl  44"]]) for r in body
               if r[i["Ampl  44"]].strip() and r[i["Is Ref _"]].strip() == "1"}
    assert sorted(flagged, key=int) == ["1", "4", "5", "53", "56"]
    target = {"1": flagged["53"] * 0.97, "4": flagged["53"] * 0.97, "5": flagged["53"] * 0.94,
              "56": flagged["56"] * 0.97}
    out = []
    for r in body:
        row, ref = r[i["Row"]], r[i["Is Ref _"]].strip() == "1"
        if ref and row in ("1", "4", "5"):
            continue
        if not ref and row in target and r[i["Ampl  44"]].strip():
            r = list(r)
            scale = target[row] / float(r[i["Ampl  44"]])
            r[i["Ampl  44"]] = f"{target[row]:.1f}"
            r[i["Area 44"]] = f"{float(r[i['Area 44']]) * scale:.3f}"
        out.append(r)
    return _write(dst, head, out)


TIGHT = Config(cal_slope_note=0.01, cal_slope_caution=0.02, cal_sd_note=0.1, cal_sd_caution=None)
SMALL = Config(cal_slope_note=0.00005, cal_slope_caution=None, cal_sd_note=0.0001, cal_sd_caution=None)


# ------------------------------------------------------------------------------------------ W-09
def test_w09_the_engine_takes_the_peak_next_to_the_pulse_and_says_so(tmp_path):
    pa, fa = process(parse_export(two_co2_variant(tmp_path / "two_co2.csv")), Config())
    pb, fb = process(parse_export(RUN), Config())
    for a, b in zip(pa, pb):
        assert a.row == b.row
        for k in ("area_c", "ampl_c", "d13c", "d13c_final", "u_d13c", "d15n_final", "pct_c_reported"):
            assert getattr(a, k) == getattr(b, k), (a.row, k)
        if a.row == W09_ROW:
            assert a.c_before == 2
            assert [f for f in a.flags if f != W09_FLAG] == b.flags and W09_FLAG in a.flags
        else:
            assert a.c_before == 0 and a.flags == b.flags
    assert fa["rerun"]["run_level_text"] == fb["rerun"]["run_level_text"]


@pytest.mark.skipif(not NODE, reason="needs node for the preview port")
@pytest.mark.parametrize("variant", ["two_co2", "single_anchor"])
def test_the_port_agrees_on_the_variants(tmp_path, variant):
    if variant == "two_co2":
        export, cfg = two_co2_variant(tmp_path / "two_co2.csv"), Config()
    else:
        export, cfg = single_anchor_variant(tmp_path / "single.csv"), SINGLE
    (tmp_path / "cfg.json").write_text(json.dumps(config_to_dict(cfg)))
    out = tmp_path / "port"
    gen = subprocess.run([sys.executable, str(ROOT / "scripts" / "port_cases.py"), str(export), "--out", str(out),
                          "--config", str(tmp_path / "cfg.json"), "--random", "40", "--per-variant", "2"],
                         capture_output=True, text=True, cwd=ROOT)
    assert gen.returncode == 0, gen.stderr
    run = subprocess.run([NODE, str(ROOT / "port" / "test_port.js"), str(out / "port_input.json"),
                          str(out / "cases.json")], capture_output=True, text=True, cwd=ROOT)
    assert run.returncode == 0, run.stdout[-2000:] + run.stderr[-2000:]
    assert "value mismatches 0, flag-list mismatches 0, fit mismatches 0" in run.stdout, run.stdout[-500:]
    block = json.loads((out / "port_input.json").read_text())
    if variant == "two_co2":
        assert [q["cpre"] for q in block["pos"] if q["cpre"]] == [2]
        cases = json.loads((out / "cases.json").read_text())
        assert all(W09_FLAG in next(r for r in c["rows"] if r[0] == W09_ROW)[-1] for c in cases)


# ------------------------------------------------------------------------------------------ W-10
def _rederived_u(positions, fits, el, single_mean):
    """1σ of each position by the Kragten weights of ea_irms, written out here, with a material analyzed once
    taking ``single_mean`` as the uncertainty of its mean."""
    f, b = fits[f"cal_{el}"], fits[f"unc_{el}"]
    xattr, aattr, uattr = (("d15n_drift", "acc15", "u_acc15") if el == "n" else ("d13c_drift", "acc13", "u_acc13"))
    m, mx, sxx, n = f["slope"], f["mean_x"], f["sxx"], f["n"]
    groups = {}
    for p in positions:
        if p.role == "Anchor" and getattr(p, xattr) is not None and getattr(p, aattr) is not None:
            g = groups.setdefault(p.std_index, {"x": [], "acc": getattr(p, aattr), "u_acc": getattr(p, uattr) or 0.0})
            g["x"].append(getattr(p, xattr))
    for g in groups.values():
        k = len(g["x"])
        mean = sum(g["x"]) / k
        g["u_mean"] = (math.sqrt(sum((x - mean) ** 2 for x in g["x"]) / (k - 1)) / math.sqrt(k) if k > 1
                       else single_mean)
        g["s1"] = sum(x - mx for x in g["x"])
        g["s2"] = sum((g["acc"] - (m * x + f["intercept"])) - m * (x - mx) for x in g["x"])
    out = {}
    for p in positions:
        x = getattr(p, xattr)
        if x is None:
            continue
        total = (b["u_sample"] * m) ** 2
        for g in groups.values():
            k = len(g["x"])
            total += (g["u_acc"] * (k / n + g["s1"] * (x - mx) / sxx)) ** 2
            total += (g["u_mean"] * ((x - mx) * g["s2"] / sxx - k * m / n)) ** 2
        out[p.row] = math.sqrt(total)
    return out


def test_w10_a_material_analyzed_once_takes_the_repeatability_in_use(tmp_path):
    export = single_anchor_variant(tmp_path / "single.csv")
    # 1.4.15: without the terms of the corrections in force and the drift monitors' scatter, 1σ is the
    # calibration's alone, as in 1.4.14, which _rederived_u writes out
    positions, fits = process(parse_export(export), replace(SINGLE, kragten_corrections="Off",
                                                            monitor_repeatability="Off"))
    full, full_fits = process(parse_export(export), SINGLE)
    by_full = {p.row: p for p in full}
    for el, attr in (("n", "u_d15n"), ("c", "u_d13c")):
        b = fits[f"unc_{el}"]
        once = [m for m in b["materials"] if m["n"] == 1]
        assert len(once) == 1 and once[0]["sd"] is None
        assert once[0]["u_mean"] == b["u_sample"] == b["pooled_sd"] and once[0]["u_mean_source"] == "repeatability in use"
        assert all(m["u_mean_source"] == "own replicates" for m in b["materials"] if m["n"] > 1)
        new = _rederived_u(positions, fits, el, b["u_sample"])
        old = _rederived_u(positions, fits, el, 0.0)
        samples = [p for p in positions if p.role == "Sample" and getattr(p, attr) is not None]
        assert len(samples) >= 20
        assert not full_fits[f"unc_{el}"]["monitors"]["used"]        # Auto: the monitors' SD is not in use here
        for p in samples:
            assert abs(getattr(p, attr) - new[p.row]) <= 1e-12 * new[p.row], (el, p.row)
            assert new[p.row] > old[p.row]
            # with the defaults the terms of the corrections in force add to it in quadrature
            terms = [t for t in getattr(by_full[p.row], f"u_terms_{el}").values() if t is not None]
            assert terms and getattr(by_full[p.row], attr) == pytest.approx(
                math.sqrt(new[p.row] ** 2 + sum(t * t for t in terms)), rel=1e-12), (el, p.row)
    # with a repeatability given, the material's mean takes it too
    _, f2 = process(parse_export(export), replace(SINGLE, u_sample_n=0.181))
    assert [m["u_mean"] for m in f2["unc_n"]["materials"] if m["n"] == 1] == [0.181]


def test_w10_runs_with_replicated_anchors_are_unchanged():
    positions, fits = process(parse_export(RUN), Config())
    for el in ("n", "c"):
        assert all(m["n"] == 3 and m["u_mean_source"] == "own replicates" for m in fits[f"unc_{el}"]["materials"])


# ------------------------------------------------------------------------------------------ W-03
def test_w03_a_wider_tolerance_outside_the_designated_pulses(tmp_path):
    export = pulse_variant(tmp_path / "pulses.csv")
    _, fits = process(parse_export(export), Config())
    chk = fits["ref_pulse_check"]
    assert chk["active"] and chk["span"] == [53, 56] and chk["ampl_tol_outside"] == 0.05
    assert chk["rejected"] == [1, 4, 53]                # 5 is 6% off; 56, inside the span, is 3% off of 2%
    _, before = process(parse_export(export), Config(ref_pulse_ampl_tol_outside=None))
    assert before["ref_pulse_check"]["rejected"] == [53]   # as before 1.4.12: 2% everywhere
    _, wide = process(parse_export(export), Config(ref_pulse_ampl_tol=0.04))
    assert wide["ref_pulse_check"]["rejected"] == [1, 4, 53, 56]


def test_w03_and_w02_settings_in_the_lab_configuration(tmp_path):
    for key in ("ref_pulse_ampl_tol_outside", "cal_slope_note", "cal_slope_caution", "cal_sd_note", "cal_sd_caution"):
        assert key in KEYS and key in config_to_dict(Config())
    cfg = replace(Config(), ref_pulse_ampl_tol_outside=None, cal_slope_note=0.03, cal_sd_caution=None)
    back = config_from_xlsx(config_to_xlsx(cfg))
    assert config_to_dict(back) == config_to_dict(cfg)
    with pytest.raises(LabConfigError):
        read_config_json({"cal_sd_note": -0.1})
    with pytest.raises(LabConfigError):
        read_config_json({"ref_pulse_ampl_tol_outside": "wide"})


# ------------------------------------------------------------------------------------------ W-02
def test_w02_the_calibration_line_at_two_levels():
    # 1.4.15: the acetanilide's whole-run check, which fires on the test run and is a problem of its whole run, is
    # off, so the problems listed are the calibration line's alone
    _, f = process(parse_export(RUN), Config(ace_scatter_check="Off"))
    assert {el: f["cal_check"][el]["level"] for el in "nc"} == {"n": "OK", "c": "OK"}
    assert f["rerun"]["run_level_text"] == "none found"
    _, f = process(parse_export(RUN), replace(TIGHT, ace_scatter_check="Off"))
    n, c = f["cal_check"]["n"], f["cal_check"]["c"]
    assert (n["level"], n["slope_level"], n["sd_level"]) == ("Caution", "Caution", "Note")
    assert (c["level"], c["slope_level"], c["sd_level"]) == ("Note", "Note", None)
    assert f["rerun"]["run_level"] == [
        "d15N calibration slope 1.0266, more than 0.02 from 1",
        "d15N anchor residual SD 0.173 permil, above 0.1",
        "d13C calibration slope 1.0142, more than 0.01 from 1"]
    _, off = process(parse_export(RUN), Config(cal_slope_note=None, cal_slope_caution=None, cal_sd_note=None,
                                               cal_sd_caution=None, ace_scatter_check="Off"))
    assert off["cal_check"]["n"]["level"] == "OK" and off["rerun"]["run_level_text"] == "none found"


def test_w02_nothing_calibrated_is_not_judged(tmp_path):
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    out = []
    for r in body:                       # every USGS 65 renamed, so only USGS 64 anchors the line
        if r[i["Identifier 1"]].strip() == "USGS 65":
            r = list(r)
            r[i["Identifier 1"]] = "USGS 64"
        out.append(r)
    _, f = process(parse_export(_write(tmp_path / "one_material.csv", head, out)), Config())
    assert f["cal_check"]["n"]["level"] is None and f["cal_check"]["c"]["level"] is None
    assert not any("calibration slope" in t or "residual SD" in t for t in f["rerun"]["run_level"])


# ------------------------------------------------------------------------------------------ workbooks
@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    """The workbooks of the variants, built by ea_irms and computed by LibreOffice in one call."""
    if not SOFFICE:
        pytest.skip("needs LibreOffice to compute the formulas")
    folder = tmp_path_factory.mktemp("findings")
    cases = {"two_co2": (two_co2_variant(folder / "two_co2.csv"), Config()),
             "single": (single_anchor_variant(folder / "single.csv"), SINGLE),
             "pulses": (pulse_variant(folder / "pulses.csv"), Config()),
             "pulses_before": (folder / "pulses.csv", Config(ref_pulse_ampl_tol_outside=None)),
             "tight": (RUN, TIGHT),
             # the review of 1.4.12
             "pulses_tol08": (folder / "pulses.csv", Config(ref_pulse_ampl_tol=0.08)),
             "single_u0": (folder / "single.csv", replace(SINGLE, u_sample_n=0.0)),
             "small_limits": (RUN, SMALL),
             "order": (RUN, Config(exclude_elemental=[18], min_blank_multiple=1000.0))}
    books = []
    for name, (export, cfg) in cases.items():
        build_workbook(export, folder / f"{name}.xlsx", cfg)
        books.append(str(folder / f"{name}.xlsx"))
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + books, check=True, capture_output=True, timeout=1200)
    out = {}
    for name, (export, cfg) in cases.items():
        wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
        out[name] = (wb, *process(parse_export(export), cfg))
    return out


def _flags(wb, row):
    """Results Flags of a position, in the workbook's order."""
    rs = wb["Results"]
    col = [c for c in range(1, rs.max_column + 1) if rs.cell(4, c).value == "Flags"][0]
    return [t.strip() for t in (rs.cell(4 + row, col).value or "").split(";") if t.strip()]


def _near(a, b):
    if a in (None, "") or b is None:
        return a in (None, "") and b is None
    return abs(a - b) <= 1e-9 * max(1.0, abs(b))


def test_w09_the_workbook_takes_the_same_peak(computed):
    wb, positions, fits = computed["two_co2"]
    wk = wb["Work"]
    for p in positions:
        for col, attr in (("P", "area_c"), ("Q", "ampl_c"), ("R", "d13c"), ("AF", "d13c_final")):
            assert _near(wk[f"{col}{5 + p.row}"].value, getattr(p, attr)), (p.row, attr)
        assert _flags(wb, p.row) == p.flags, p.row
    assert W09_FLAG in _flags(wb, W09_ROW)


def test_w10_the_workbook_gives_the_same_1sigma(computed):
    wb, positions, fits = computed["single"]
    wk = wb["Work"]
    n = 0
    for p in positions:
        for col, attr in (("CL", "u_d15n"), ("CN", "u_d13c")):
            assert _near(wk[f"{col}{5 + p.row}"].value, getattr(p, attr)), (p.row, attr)
            n += getattr(p, attr) is not None
    assert n > 60


def test_w03_the_workbook_rejects_the_same_pulses(computed):
    for name in ("pulses", "pulses_before"):
        wb, positions, fits = computed[name]
        rejected = {p.row for p in positions if "CO2 peak rejected as a reference pulse" in p.flags}
        assert rejected == set(fits["ref_pulse_check"]["rejected"])
        for p in positions:
            assert _flags(wb, p.row) == p.flags, (name, p.row)
            assert _near(wb["Work"][f"P{5 + p.row}"].value, p.area_c), (name, p.row)
        b41 = wb["Setup"]["B41"].value
        assert b41 == 0.05 if name == "pulses" else b41 in (None, "")
    assert computed["pulses"][2]["ref_pulse_check"]["rejected"] == [1, 4, 53]


def test_w02_the_workbook_judges_the_line_as_the_engine_does(computed):
    wb, positions, fits = computed["tight"]
    ft, qc, st = wb["Fits"], wb["QC"], wb["Setup"]
    assert [st[f"B{r}"].value for r in (73, 74, 75)] == [0.01, 0.02, 0.1] and st["B76"].value in (None, "")
    assert ft["B363"].value == fits["cal_check"]["n"]["level"] == "Caution"
    assert ft["B366"].value == fits["cal_check"]["c"]["level"] == "Note"
    assert ft["B356"].value == fits["rerun"]["run_level_text"]
    assert qc["D23"].value.startswith("CHECK, Caution") and qc["D25"].value.startswith("CHECK, Note")
    for name in ("two_co2", "single", "pulses"):          # the defaults: the test run's line is OK
        wb, positions, fits = computed[name]
        assert wb["Fits"]["B356"].value == fits["rerun"]["run_level_text"]


# ------------------------------------------------------------------------------------------ the review of 1.4.12
def test_w03_outside_the_span_the_tolerance_is_never_narrower(tmp_path):
    export = pulse_variant(tmp_path / "pulses.csv")
    # a configuration written before 1.4.12 has no outside tolerance and gets the default, 0.05
    _, f = process(parse_export(export), Config(ref_pulse_ampl_tol=0.08))
    chk = f["ref_pulse_check"]
    assert chk["ampl_tol_outside"] == 0.08 and chk["rejected"] == [1, 4, 5, 53, 56]   # 5 is 6% off, inside 8%
    _, everywhere = process(parse_export(export), Config(ref_pulse_ampl_tol=0.08, ref_pulse_ampl_tol_outside=None))
    assert everywhere["ref_pulse_check"]["rejected"] == chk["rejected"]
    _, off = process(parse_export(export), Config(ref_pulse_ampl_tol=None))
    assert not off["ref_pulse_check"]["active"] and off["ref_pulse_check"]["rejected"] == []


def test_w02_limits_read_as_the_spreadsheet_writes_them():
    for x, text in ((0.05, "0.05"), (0.025, "0.025"), (1.0, "1"), (0, "0"), (0.1 + 0.2, "0.3"), (0.00005, "0.00005"),
                    (2.5e-5, "0.000025"), (1e-9, "0.000000001"), (1e15, "1000000000000000"), (2 / 3, "0.666666666666667")):
        assert _num_text(x) == text, x
    # 1.4.15: the acetanilide's whole-run check, a problem of the test run's whole run, off (see the test above)
    _, f = process(parse_export(RUN), replace(SMALL, ace_scatter_check="Off"))
    assert f["rerun"]["run_level"] == [
        "d15N calibration slope 1.0266, more than 0.00005 from 1",
        "d15N anchor residual SD 0.173 permil, above 0.0001",
        "d13C calibration slope 1.0142, more than 0.00005 from 1",
        "d13C anchor residual SD 0.021 permil, above 0.0001"]


def test_w10_no_source_named_without_a_repeatability(tmp_path):
    export = once_variant(tmp_path / "once.csv")
    _, f = process(parse_export(export), ONCE)
    for el in ("n", "c"):
        b = f[f"unc_{el}"]
        assert f[f"cal_{el}"]["n"] == 3 and b["pooled_sd"] is None and b["u_sample"] is None
        assert [(m["n"], m["u_mean"], m["u_mean_source"]) for m in b["materials"]] == [(1, 0.0, None)] * 3
    _, g = process(parse_export(export), replace(ONCE, u_sample_n=0.2))
    assert {m["u_mean_source"] for m in g["unc_n"]["materials"]} == {"repeatability in use"}


def test_review_the_workbook_agrees(computed):
    # W-03: a tolerance above 0.05 holds outside the span too
    wb, positions, fits = computed["pulses_tol08"]
    assert (wb["Setup"]["B39"].value, wb["Setup"]["B41"].value) == (0.08, 0.05)
    assert fits["ref_pulse_check"]["rejected"] == [1, 4, 5, 53, 56]
    assert wb["Fits"]["B241"].value == 5
    for p in positions:
        assert _flags(wb, p.row) == p.flags, p.row
        assert _near(wb["Work"][f"P{5 + p.row}"].value, p.area_c), p.row
    # a repeatability of 0 is a value, and reaches 1σ as the engine uses it
    wb, positions, fits = computed["single_u0"]
    assert wb["Setup"]["B30"].value == 0 and fits["unc_n"]["u_sample"] == 0.0
    n = 0
    for p in positions:
        for col, attr in (("CL", "u_d15n"), ("CN", "u_d13c")):
            assert _near(wb["Work"][f"{col}{5 + p.row}"].value, getattr(p, attr)), (p.row, attr)
            n += getattr(p, attr) is not None
    assert n > 60
    # limits below 0.0001 read the same in the engine's text and the workbook's
    wb, positions, fits = computed["small_limits"]
    assert wb["Fits"]["B356"].value == fits["rerun"]["run_level_text"]
    assert "more than 0.00005 from 1" in fits["rerun"]["run_level_text"]
    # Flags in the engine's order: the N flags before the C flags, the elemental exclusion before the blank
    wb, positions, fits = computed["order"]
    for p in positions:
        assert _flags(wb, p.row) == p.flags, p.row
    by_row = {p.row: p.flags for p in positions}
    assert by_row[1][:2] == ["N below linearity range", "no C peak"]
    exc = by_row[18].index("excluded from elemental correction")
    assert exc < min(k for k, t in enumerate(by_row[18]) if "x the blank" in t)
