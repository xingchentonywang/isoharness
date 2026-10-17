"""Findings R-23 and R-24 of the 2026-10-14 review, on copies of the 2026-09-15 test run.

R-23  Under the Background drift model a position with no reference background took its position number as
      its x, so positions and millivolts were mixed in one correction. It is now left without a drift
      correction, its drift-corrected value equal to its linearity-corrected one; while the correction is in
      force Flags says "N not drift-corrected: no reference background" (and the same for C), and a sample is
      marked for a check. A drift monitor without a reference background is not in the fit and has no
      studentized residual. The Row model is unchanged. The copy empties BGD 28 of the designated N2 peak at
      one sample and one drift monitor (scripts/port_cases.py, no_background_variant).
R-24  Every comparison of a value with a limit allows the slack of a spreadsheet comparison, in the engine and
      in the preview port. Two copies put a value exactly on a limit where the plain comparison the engine made
      before fell one unit in the last place on the wrong side: three blank capsules whose d15N values (4.002,
      6.002, and 8.002 permil, illustrative) have an SD of 2 permil, the limit of the blank correction on Auto,
      and a sample whose N peak area (6.7212, illustrative) is 3.6 times the blank mean of 1.867, with the
      blank multiple set to 3.6.

The workbook tests have LibreOffice compute the formulas and skip without it; the port tests need node.
"""
import csv
import json
import shutil
import subprocess
import sys
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from ea_irms.parse import parse_export                                   # noqa: E402
from ea_irms.pipeline import Config, _above, _below, config_to_dict, process  # noqa: E402
from ea_irms.workbook import build_workbook                              # noqa: E402
from isoharness.serialize import clean, position_record                  # noqa: E402
from port_cases import no_background_rows, no_background_variant         # noqa: E402
from reference_check import check_values                                 # noqa: E402

RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NODE = shutil.which("node")
NO_BGD = {"n": "N not drift-corrected: no reference background",
          "c": "C not drift-corrected: no reference background"}
BACKGROUND = Config(drift_model="Background", drift_correction="On")
BLANK_D15N = {1: "4.002", 4: "6.002", 56: "8.002"}       # position: d15N of its N peak (illustrative)
MULTIPLE, ON_LIMIT_ROW, ON_LIMIT_AREA = 3.6, 44, "6.7212"   # 3.6 x the blank mean area, (1.676 + 2.058) / 2
MULTIPLE_CONFIG = Config(min_blank_multiple=MULTIPLE)


def _rows(path=RUN):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], rows[1:]


def _write(dst, head, body):
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        w.writerows(body)
    return dst


def _sample_n_peak(i, body, row):
    """The record of a position's sample N2 peak: its last mass-28 peak, unflagged (ea_irms.parse)."""
    peaks = [r for r in body if r[i["Row"]] == str(row) and r[i["Peak Nr"]].strip() and r[i["Ampl  28"]].strip()]
    last = max(peaks, key=lambda r: float(r[i["Peak Nr"]]))
    assert last[i["Is Ref _"]].strip() != "1", row
    return last


def blank_sd_variant(dst):
    """The test run with position 1, an empty carousel position (Blank Hole), named as a blank capsule, so that
    three blank capsules carry an N peak (1, 4, and 56), and their d15N set to BLANK_D15N, whose SD is 2."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    for r in body:
        if r[i["Row"]] == "1":
            r[i["Identifier 1"]] = "Blank Capsule"
    for row, d in BLANK_D15N.items():
        _sample_n_peak(i, body, row)[i["d 15N/14N"]] = d
    return _write(dst, head, body)


def blank_multiple_variant(dst):
    """The test run with the N peak area of sample position ON_LIMIT_ROW set to ON_LIMIT_AREA, MULTIPLE times
    the mean N area of the run's blank capsules."""
    head, body = _rows()
    i = {k: n for n, k in enumerate(head)}
    _sample_n_peak(i, body, ON_LIMIT_ROW)[i["Area 28"]] = ON_LIMIT_AREA
    return _write(dst, head, body)


def _by_row(positions):
    return {p.row: p for p in positions}


# ------------------------------------------------------------------------------------------ R-23, engine
@pytest.fixture(scope="module")
def no_bgd(tmp_path_factory):
    export = no_background_variant(RUN, tmp_path_factory.mktemp("no_bgd") / "no_bgd.csv")
    rows = no_background_rows(export)
    assert len(rows) == 2
    return export, rows


def test_r23_the_variant_blanks_one_sample_and_one_drift_monitor(no_bgd):
    export, rows = no_bgd
    roles = {p.row: p.role for p in process(parse_export(export), Config())[0]}
    assert sorted(roles[r] for r in rows) == ["Drift", "Sample"]
    before = _by_row(parse_export(RUN))
    for p in parse_export(export):
        assert (p.ref_bgd_n is None) == (p.row in rows)
        assert p.ref_bgd_n is None or p.ref_bgd_n == before[p.row].ref_bgd_n


def test_r23_a_position_without_a_reference_background_is_not_drift_corrected(no_bgd):
    export, rows = no_bgd
    positions, F = process(parse_export(export), BACKGROUND)
    for el, lin, drift in (("n", "d15n_lin", "d15n_drift"), ("c", "d13c_lin", "d13c_drift")):
        f = F[f"drift_{el}"]
        assert f["model"] == "Background" and f["in_force"] and f["applied"] != 0
        assert f["no_background"] == rows
        for p in positions:
            if getattr(p, lin) is None:
                continue
            if p.row in rows:
                assert getattr(p, drift) == getattr(p, lin), (el, p.row)
                assert NO_BGD[el] in p.flags
            else:
                want = getattr(p, lin) - f["applied"] * (p.ref_bgd_n - f["pivot"])
                assert getattr(p, drift) == want, (el, p.row)
                assert NO_BGD[el] not in p.flags
    by = _by_row(positions)
    sample = next(r for r in rows if by[r].role == "Sample")
    assert by[sample].rerun == "Check"                    # its other flags are no reason for a mark
    assert [t for t in by[sample].rerun_why.split("; ") if t in NO_BGD.values()] == list(NO_BGD.values())
    monitor = next(r for r in rows if by[r].role == "Drift")
    assert [f for f in by[monitor].flags if f in NO_BGD.values()] == list(NO_BGD.values())
    assert not any("outlier" in f for f in by[monitor].flags)
    assert F["rerun"]["counts"]["Check"] == process(parse_export(RUN), BACKGROUND)[1]["rerun"]["counts"]["Check"] + 1


def test_r23_a_monitor_without_a_reference_background_is_left_out_of_the_fit(no_bgd):
    export, rows = no_bgd
    positions, F = process(parse_export(export), BACKGROUND)
    by = _by_row(positions)
    monitor = next(r for r in rows if by[r].role == "Drift")
    mons = [p for p in positions if p.role == "Drift"]
    for el in ("n", "c"):
        assert F[f"drift_{el}"]["n"] == len(mons) - 1
        assert F[f"drift_{el}"]["pivot"] == pytest.approx(
            sum(p.ref_bgd_n for p in mons if p.row != monitor) / (len(mons) - 1), rel=1e-15)
        assert getattr(by[monitor], f"t_drift_{el}") is None
        assert all(getattr(p, f"t_drift_{el}") is not None for p in mons if p.row != monitor)
        assert monitor not in F[f"drift_{el}"]["outliers"]


def test_r23_the_row_model_and_a_correction_not_in_force_are_unchanged(no_bgd):
    export, rows = no_bgd
    plain = _by_row(process(parse_export(RUN), Config(drift_correction="On"))[0])
    row_model = process(parse_export(export), Config(drift_correction="On"))
    for p in row_model[0]:
        for attr in ("d15n_drift", "d13c_drift", "d15n_final", "d13c_final", "u_d15n", "u_d13c", "t_drift_n"):
            assert getattr(p, attr) == getattr(plain[p.row], attr), (p.row, attr)
        assert p.flags == plain[p.row].flags and p.rerun == plain[p.row].rerun
    assert row_model[1]["drift_n"]["no_background"] == []
    positions, F = process(parse_export(export), replace(BACKGROUND, drift_correction="Off"))
    assert not F["drift_n"]["in_force"] and F["drift_n"]["no_background"] == []
    for p in positions:
        assert p.d15n_drift == p.d15n_lin and p.d13c_drift == p.d13c_lin
        assert not set(NO_BGD.values()) & set(p.flags)


# ------------------------------------------------------------------------------------------ R-24, engine
def test_r24_three_blanks_whose_sd_is_two_permil_support_the_blank_correction(tmp_path):
    positions, F = process(parse_export(blank_sd_variant(tmp_path / "blank_sd.csv")), Config())
    b = F["blank_n"]
    assert b["n"] == 3 and sorted(p.d15n for p in positions if p.role == "Blank" and p.d15n is not None) == [
        float(v) for v in BLANK_D15N.values()]
    # on the limit: 2 in decimal arithmetic, one unit in the last place above 2 in the float SD, which the
    # plain comparison of the engine before this change read as "more than 2", so Auto left the blank out
    assert b["sd"] > 2.0 and not _above(b["sd"], 2.0)
    assert b["in_force"]
    assert any(p.d15n_blank != p.d15n for p in positions if p.d15n is not None and p.role != "Blank")


def test_r24_a_peak_on_the_blank_multiple_is_not_under_it(tmp_path):
    positions, F = process(parse_export(blank_multiple_variant(tmp_path / "multiple.csv")), MULTIPLE_CONFIG)
    p = _by_row(positions)[ON_LIMIT_ROW]
    ab = F["blank_n"]["area"]
    assert p.area_n == float(Decimal(ON_LIMIT_AREA)) and Decimal(ON_LIMIT_AREA) == (
        Decimal("1.676") + Decimal("2.058")) / 2 * Decimal(str(MULTIPLE))
    # one unit in the last place under MULTIPLE x the blank in floats, which the plain comparison flagged
    assert p.area_n < ab * MULTIPLE and not _below(p.area_n, ab * MULTIPLE)
    assert not any(f.startswith("N peak only") for f in p.flags), p.flags
    under = [q for q in positions if q.area_n is not None and q.role != "Blank" and q.area_n < ab * MULTIPLE * 0.999]
    assert under and all(any(f.startswith("N peak only") for f in q.flags) for q in under)


# ------------------------------------------------------------------------------------------ the workbook
@pytest.fixture(scope="module")
def computed(tmp_path_factory, no_bgd):
    """The workbooks of the cases, built by ea_irms and computed by LibreOffice in one call."""
    if not SOFFICE:
        pytest.skip("needs LibreOffice to compute the formulas")
    folder = tmp_path_factory.mktemp("background_limits")
    export, _ = no_bgd
    cases = {"background_on": (export, BACKGROUND),
             "background_auto": (export, replace(BACKGROUND, drift_correction="Auto")),
             "model_auto": (export, replace(BACKGROUND, drift_model="Auto")),
             "background_off": (export, replace(BACKGROUND, drift_correction="Off")),
             "row": (export, Config(drift_correction="On")),
             "blank_sd": (blank_sd_variant(folder / "blank_sd.csv"), Config()),
             "blank_multiple": (blank_multiple_variant(folder / "multiple.csv"), MULTIPLE_CONFIG)}
    books = []
    for name, (path, cfg) in cases.items():
        build_workbook(path, folder / f"{name}.xlsx", cfg)
        books.append(str(folder / f"{name}.xlsx"))
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + books, check=True, capture_output=True, timeout=1200)
    out = {}
    for name, (path, cfg) in cases.items():
        wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
        positions, fits = process(parse_export(path), cfg)
        body = json.loads(json.dumps({"positions": [position_record(p) for p in positions], "fits": clean(fits)}))
        out[name] = (wb, positions, fits, body)
    return out


@pytest.mark.parametrize("name", ["background_on", "background_auto", "model_auto", "background_off", "row",
                                  "blank_sd", "blank_multiple"])
def test_the_workbook_agrees_with_the_engine(computed, name):
    wb, positions, fits, body = computed[name]
    n, bad, errors = check_values(body, wb)        # every value, flag list, mark, and fit (scripts/reference_check)
    assert n > 3000 and not bad and not errors, (bad[:8], errors[:8])


def _flags(wb, row):
    rs = wb["Results"]
    return [t.strip() for t in (rs.cell(4 + row, 17).value or "").split(";") if t.strip()]


def test_r23_the_workbook_leaves_the_same_positions_uncorrected(computed, no_bgd):
    _, rows = no_bgd
    wb, positions, fits, _ = computed["background_on"]
    wk = wb["Work"]
    assert wb["Fits"]["B38"].value == "Background" and wb["Fits"]["B115"].value == "Yes"
    for r in rows:
        assert wk[f"BI{5 + r}"].value in (None, "")
        assert wk[f"Y{5 + r}"].value in (None, "") and wk[f"AD{5 + r}"].value in (None, "")
        assert wk[f"Z{5 + r}"].value == wk[f"X{5 + r}"].value and wk[f"AE{5 + r}"].value == wk[f"AC{5 + r}"].value
        assert [f for f in _flags(wb, r) if f in NO_BGD.values()] == list(NO_BGD.values())
    by = _by_row(positions)
    monitor = next(r for r in rows if by[r].role == "Drift")
    assert wk[f"CB{5 + monitor}"].value in (None, "") and wk[f"CC{5 + monitor}"].value in (None, "")
    sample = next(r for r in rows if by[r].role == "Sample")
    assert wb["Results"].cell(4 + sample, 18).value == "Check"
    assert wb["Fits"]["B107"].value == fits["drift_n"]["n"] == 12
    # the switch off and the Row model: no position flagged
    for name in ("background_off", "row"):
        wb_, positions_, _, _ = computed[name]
        assert not any(set(NO_BGD.values()) & set(_flags(wb_, p.row)) for p in positions_), name


def test_r24_the_workbook_treats_the_values_on_the_limits_as_the_engine_does(computed):
    wb, positions, fits, _ = computed["blank_sd"]
    assert fits["blank_n"]["in_force"] and wb["Fits"]["C5"].value == "Yes"
    wb, positions, fits, _ = computed["blank_multiple"]
    assert not any(f.startswith("N peak only") for f in _flags(wb, ON_LIMIT_ROW))
    assert _flags(wb, ON_LIMIT_ROW) == _by_row(positions)[ON_LIMIT_ROW].flags


# ------------------------------------------------------------------------------------------ the preview port
@pytest.mark.skipif(not NODE, reason="needs node for the preview port")
@pytest.mark.parametrize("variant", ["no_background", "blank_sd", "blank_multiple"])
def test_the_port_agrees(tmp_path, variant):
    args, cfg = [], Config()
    if variant == "no_background":
        export, args = RUN, ["--no-ref-background"]
    elif variant == "blank_sd":
        export = blank_sd_variant(tmp_path / "blank_sd.csv")
    else:
        export, cfg = blank_multiple_variant(tmp_path / "multiple.csv"), MULTIPLE_CONFIG
    (tmp_path / "cfg.json").write_text(json.dumps(config_to_dict(cfg)))
    out = tmp_path / "port"
    gen = subprocess.run([sys.executable, str(ROOT / "scripts" / "port_cases.py"), str(export), "--out", str(out),
                          "--config", str(tmp_path / "cfg.json"), "--random", "60", "--per-variant", "3"] + args,
                         capture_output=True, text=True, cwd=ROOT)
    assert gen.returncode == 0, gen.stderr
    run = subprocess.run([NODE, str(ROOT / "port" / "test_port.js"), str(out / "port_input.json"),
                          str(out / "cases.json")], capture_output=True, text=True, cwd=ROOT)
    assert run.returncode == 0, run.stdout[-2000:] + run.stderr[-2000:]
    assert "value mismatches 0, flag-list mismatches 0, fit mismatches 0" in run.stdout, run.stdout[-500:]
    cases = json.loads((out / "cases.json").read_text())
    if variant == "no_background":
        flagged = [c for c in cases if c["fits"]["driftN"][0] and c["fits"]["driftN"][1] == "Background"]
        assert flagged and all(NO_BGD["n"] in next(r for r in c["rows"] if r[0] == row)[-1]
                               for c in flagged for row in no_background_rows(out / "20260915_raw_no_ref_background.csv"))
    elif variant == "blank_sd":
        assert any(c["fits"]["blankN"] for c in cases if c["case"]["blank"] == "Auto")
