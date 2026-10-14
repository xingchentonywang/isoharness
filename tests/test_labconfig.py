"""ea_irms 1.4.6: the lab configuration as an Excel file."""
import io
import json
from dataclasses import fields
from pathlib import Path

import openpyxl
import pytest

from ea_irms.cli import main
from ea_irms.labconfig import (KEYS, SETTINGS, LabConfigError, config_from_xlsx, config_to_xlsx,
                               read_lab_config)
from ea_irms.pipeline import DEFAULT_STANDARDS, MAX_STANDARDS, Config, config_from_dict, config_to_dict

RUN = Path(__file__).resolve().parent / "data" / "20260915_raw.csv"


def _custom():
    standards = [list(r) for r in DEFAULT_STANDARDS]
    standards[0][9], standards[0][10] = 0.10, 0.06
    standards.append(["Ace Check", "Check", -29.5, 1.2, 71.1, 10.36, "acetanilide check", 0.05, 0.05, None, None])
    return config_from_dict({
        "standards": standards, "blank_correction": "On", "linearity_correction": "Off",
        "linearity_correction_c": "Auto", "drift_correction": "On", "drift_model": "Background",
        "first_drift_position": 12, "exclude_drift_n": [18, 55], "exclude_drift_c": [], "exclude_elemental": [3],
        "significance_level": 0.1, "min_area_n": 9.5, "max_area_c": 400, "min_blank_multiple": None,
        "carryover_mv": None, "u_sample_n": 0.181, "u_sample_c": None, "coverage_k": 3.0, "balance_u_mg": 0.002,
        "elemental_correction": "Off", "elemental_threshold_pct": 4.0, "ref_pulse_ampl_tol": None,
        "ref_pulse_d13c_window": 1.0, "ref_pulse_off_sd": None, "ref_gas_sd_limit": 0.08,
        "rerun_replicates": "Off"})


def _edit(blob, **cells):
    """Change cells of the template: _edit(blob, Settings={'B6': 'x'}, Standards={...})."""
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    for sheet, values in cells.items():
        for ref, value in values.items():
            wb[sheet][ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(blob, key):
    st = openpyxl.load_workbook(io.BytesIO(blob))["Settings"]
    return next(r for r in range(1, st.max_row + 1) if st.cell(r, 4).value == key)


def test_every_setting_has_one_row():
    assert sorted(KEYS) == sorted(f.name for f in fields(Config) if f.name != "standards")
    assert len(KEYS) == len(set(KEYS))
    assert all(kind in ("switch", "switch_or_empty", "onoff", "model", "number", "number_or_empty",
                        "whole_or_empty", "positions", "source", "text_or_empty") for key, _, kind, _ in SETTINGS if key)


@pytest.mark.parametrize("config", [Config(), _custom()], ids=["defaults", "custom"])
def test_round_trip(config):
    back, notes = read_lab_config(config_to_xlsx(config))
    assert config_to_dict(back) == config_to_dict(config)
    assert notes == []


def test_template_is_filled_and_marked():
    wb = openpyxl.load_workbook(io.BytesIO(config_to_xlsx()))
    assert wb.sheetnames == ["README", "Settings", "Standards"]
    st, sd = wb["Settings"], wb["Standards"]
    assert [st.cell(4, j).value for j in range(1, 5)] == ["Setting", "Value", "What it means", "Key"]
    assert st.cell(_row(config_to_xlsx(), "coverage_k"), 2).value == 2.0
    assert sd["A4"].value == "Match text in Identifier 1" and sd["H4"].value == "1σ(d13C)"
    assert [sd.cell(5 + i, 1).value for i in range(len(DEFAULT_STANDARDS))] == [r[0] for r in DEFAULT_STANDARDS]
    assert sd.cell(5 + len(DEFAULT_STANDARDS), 1).value is None
    lists = {str(dv.sqref): dv.formula1 for dv in st.data_validations.dataValidation}
    blank = _row(config_to_xlsx(), "blank_correction")
    assert any(f"B{blank}" in ref.split() and f == '"On,Off,Auto"' for ref, f in lists.items())


def test_problems_are_named_by_cell():
    blob = config_to_xlsx()
    bad = _edit(blob, Settings={f"B{_row(blob, 'blank_correction')}": "maybe",
                                f"B{_row(blob, 'significance_level')}": None,
                                f"B{_row(blob, 'exclude_drift_n')}": "59a, 61",
                                f"B{_row(blob, 'first_drift_position')}": 12.5},
                Standards={"B5": "Anchr", "C6": "abc"})
    with pytest.raises(LabConfigError) as err:
        read_lab_config(bad)
    text = str(err.value)
    for part in ("(blank_correction) is 'maybe'", "(significance_level) is empty", "(exclude_drift_n) is '59a, 61'",
                 "(first_drift_position) is '12.5'", "Standards B5 (USGS 64) is 'Anchr'", "Standards C6"):
        assert part in text, part


def test_empty_means_off_where_a_setting_can_be_off():
    blob = config_to_xlsx()
    cfg = config_from_xlsx(_edit(blob, Settings={f"B{_row(blob, k)}": None for k in
                                                 ("ref_pulse_off_sd", "carryover_mv", "linearity_correction_c",
                                                  "balance_u_mg", "exclude_elemental")}))
    assert cfg.ref_pulse_off_sd is None and cfg.carryover_mv is None and cfg.linearity_correction_c is None
    assert cfg.balance_u_mg is None and cfg.exclude_elemental == []


def test_lower_case_and_text_numbers_are_read():
    blob = config_to_xlsx()
    cfg = config_from_xlsx(_edit(blob, Settings={f"B{_row(blob, 'drift_model')}": "background",
                                                 f"B{_row(blob, 'coverage_k')}": "2.5",
                                                 f"B{_row(blob, 'exclude_drift_c')}": 47}))
    assert cfg.drift_model == "Background" and cfg.coverage_k == 2.5 and cfg.exclude_drift_c == [47]


def test_older_files_and_unknown_keys():
    blob = config_to_xlsx(_custom())
    r = _row(blob, "ref_gas_sd_limit")
    cfg, notes = read_lab_config(_edit(blob, Settings={f"D{r}": None, f"D{_row(blob, 'coverage_k')}": "coverage"}))
    assert cfg.ref_gas_sd_limit == Config().ref_gas_sd_limit and cfg.coverage_k == Config().coverage_k
    assert any("'coverage'" in n for n in notes)
    assert any("package defaults were used" in n and "ref_gas_sd_limit" in n for n in notes)


def test_too_many_standards_and_wrong_files():
    names = {f"A{5 + i}": f"Std {i}" for i in range(MAX_STANDARDS + 1)}
    roles = {f"B{5 + i}": "Sample" for i in range(MAX_STANDARDS + 1)}
    with pytest.raises(LabConfigError, match="at most"):
        read_lab_config(_edit(config_to_xlsx(), Standards={**names, **roles}))
    with pytest.raises(LabConfigError, match="could not be opened"):
        read_lab_config(b"not a workbook")
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(LabConfigError, match="Settings sheet and a Standards sheet"):
        read_lab_config(buf.getvalue())


def test_command_line_writes_and_reads_the_template(tmp_path):
    template = tmp_path / "lab.xlsx"
    assert main(["--write-config", str(template)]) == 0
    assert config_to_dict(config_from_xlsx(template)) == config_to_dict(Config())
    as_json = tmp_path / "lab.json"
    assert main(["--write-config", str(as_json)]) == 0
    assert json.loads(as_json.read_text()) == config_to_dict(Config())
    out = tmp_path / "run.xlsx"
    assert main([str(RUN), "-o", str(out), "--config", str(template), "--no-recalc"]) == 0
    assert openpyxl.load_workbook(out)["Setup"]["B32"].value == 2.0
