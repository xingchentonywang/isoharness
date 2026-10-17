"""The engine must agree with the validated Excel template, cell for cell.

The template was checked against an independent implementation and recalculated
by LibreOffice with no formula errors. Treating it as the reference means any
drift between the two code paths shows up here rather than in someone's data.
"""
import csv
import os
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ea_irms.parse import parse_export            # noqa: E402
from ea_irms.pipeline import Config, process      # noqa: E402

# The recalculated template (with cached values) is the reference. Point
# EA_TEMPLATE at it, or place it in tests/data/. The export it was built from is
# its own Raw sheet, so no second file is needed.
TEMPLATE = Path(os.environ.get(
    "EA_TEMPLATE",
    Path(__file__).resolve().parent / "data" / "EA-IRMS_Correction_Template_v1.xlsx"))
TOL = 1e-9

if not TEMPLATE.exists():
    pytest.skip(f"reference template not found at {TEMPLATE}; set EA_TEMPLATE",
                allow_module_level=True)


@pytest.fixture(scope="module")
def engine(tmp_path_factory):
    raw = openpyxl.load_workbook(TEMPLATE)["Raw"]
    export = tmp_path_factory.mktemp("export") / "raw.csv"
    with open(export, "w", newline="") as fh:
        w = csv.writer(fh)
        for record in raw.iter_rows(max_col=16, values_only=True):
            if record[0] is not None:
                w.writerow(["" if v is None else v for v in record])
    # Run the engine with the switches the template's Setup sheet carries, so the
    # comparison holds whichever defaults the template was built with.
    setup = openpyxl.load_workbook(TEMPLATE)["Setup"]
    switches = {}
    for row in setup.iter_rows(min_col=1, max_col=2, values_only=True):
        label, value = row
        if label in ("Blank correction", "Linearity correction", "Drift correction", "Drift model"):
            switches[label] = value
    return process(parse_export(export), Config(
        blank_correction=switches.get("Blank correction", "On"),
        linearity_correction=switches.get("Linearity correction", "On"),
        drift_correction=switches.get("Drift correction", "On"),
        drift_model=switches.get("Drift model", "Row")))


@pytest.fixture(scope="module")
def sheet():
    wb = openpyxl.load_workbook(TEMPLATE, data_only=True)
    return wb["Work"], wb["Fits"]


def close(a, b, tol=TOL):
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= tol * max(1.0, abs(b)) + tol


def test_positions_parsed(engine):
    positions, _ = engine
    assert len(positions) == 56
    assert [p.row for p in positions if not p.has_n] == [5]
    assert sorted(p.row for p in positions if not p.has_c) == [1, 4, 5, 53, 56]


@pytest.mark.parametrize("cell,path", [
    ("B6", ("blank_n", "area")), ("B7", ("blank_n", "d")),
    ("B17", ("lin_n", "fit", "slope")), ("B25", ("lin_c", "fit", "slope")),
    ("B32", ("drift_n", "row", "slope")), ("B34", ("drift_n", "row", "sd")),
    ("B37", ("drift_n", "background", "sd")),
    ("B56", ("cal_n", "slope")), ("B57", ("cal_n", "intercept")),
    ("B63", ("cal_c", "slope")), ("B64", ("cal_c", "intercept")),
    ("B70", ("k_n", "slope")), ("B71", ("k_n", "intercept")),
    ("B76", ("k_c", "slope")), ("B77", ("k_c", "intercept")),
    ("B83", ("settling", "floor")), ("B85", ("settling", "excess")),
    ("B86", ("settling", "settled_from")),
])
def test_fit_matches_workbook(engine, sheet, cell, path):
    _, fits = engine
    _, ft = sheet
    value = fits
    for key in path:
        value = value[key]
    assert close(value, ft[cell].value), f"{cell} {'.'.join(path)}"


def test_blank_in_force_matches(engine, sheet):
    _, fits = engine
    _, ft = sheet
    assert fits["blank_n"]["in_force"] is (ft["C5"].value == "Yes")


def test_drift_model_matches(engine, sheet):
    _, fits = engine
    _, ft = sheet
    assert fits["drift_n"]["model"] == ft["B38"].value


@pytest.mark.parametrize("attr,col", [
    ("d15n_final", "AA"), ("d13c_final", "AF"),
    ("pct_n_measured", "AI"), ("pct_c_measured", "AJ"),
    ("cn_molar", "AK"), ("cn_mass", "AL"),
    ("ref_bgd_n", "BI"), ("settling_fraction", "BJ"),
])
def test_per_position_matches_workbook(engine, sheet, attr, col):
    positions, _ = engine
    wk, _ = sheet
    worst = 0.0
    for p in positions:
        mine = getattr(p, attr)
        theirs = wk[f"{col}{5 + p.row}"].value
        if not isinstance(theirs, (int, float)):
            theirs = None
        assert (mine is None) == (theirs is None), f"row {p.row} {attr}"
        if mine is not None:
            worst = max(worst, abs(mine - theirs))
    assert worst < 1e-9, f"{attr} worst difference {worst:.3e}"


def test_flags_match_workbook(engine):
    """Flags are text, so compare the ones that carry a decision."""
    positions, _ = engine
    wb = openpyxl.load_workbook(TEMPLATE, data_only=True)
    rs = wb["Results"]
    for p in positions:
        text = (rs[f"K{4 + p.row}"].value or "")
        for flag in ("no N peak", "no C peak", "ran before the background settled"):
            assert (flag in p.flags) == (flag in text), f"row {p.row}: {flag}"
