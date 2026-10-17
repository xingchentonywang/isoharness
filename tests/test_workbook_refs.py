"""The rows and columns the workbook's formulas read are the ones they were written for (finding R-27 of the
2026-10-14 review).

The workbook writer places most rows of Fits, Setup, and QC by number, and formulas on every sheet read them
by number, so a small edit could move a row away from the formulas that read it without any other test
noticing. tests/data/workbook_refs.json, written by scripts/workbook_refs.py, holds the label of every row of
Fits, Setup, and QC and the header of every Work column that a formula reads. These tests rebuild the
2026-09-15 test run's workbook (the defaults and a sample sheet) and fail, naming a formula, when a row or
column it reads carries another label than the file's, is empty, or is not in the file, and when the file
holds a row or column that no formula reads. After a deliberate change to the layout, regenerate the file with
python scripts/workbook_refs.py and read its diff.

The check compares each row number a formula reads with the label recorded for that number, so it catches a
row that moved while formulas kept reading its old number; a formula changed to read another recorded row is
outside what it can see. The last tests move rows and a column, reword a label, and point a formula at an
empty row and at an unrecorded one in a copy of the workbook, to show that each is caught and named; they
compare with the record of the unchanged workbook, so they hold whether or not the file is up to date.
LibreOffice is not needed.
"""
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from workbook_refs import FILE, build, collect, compare, labels, stale  # noqa: E402


@pytest.fixture(scope="module")
def book(tmp_path_factory):
    return build(tmp_path_factory.mktemp("workbook_refs"))


@pytest.fixture(scope="module")
def recorded():
    return json.loads(FILE.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def found(book):
    return collect(book)


def test_every_row_and_column_a_formula_reads_carries_its_recorded_label(found, recorded):
    problems = compare(found, recorded)
    assert not problems, "\n".join(problems[:30])


def test_the_file_holds_only_what_the_formulas_read_in_sheet_order(found, recorded):
    extra = stale(found, recorded)
    assert not extra, "\n".join(extra[:30])
    assert list(recorded["rows"]) == ["Fits", "QC", "Setup"]
    for sheet, held in recorded["rows"].items():
        assert list(held) == sorted(held, key=int), sheet
        assert all(isinstance(v, str) and v for v in held.values()), sheet
    from openpyxl.utils import column_index_from_string
    cols = list(recorded["work_columns"])
    assert cols == sorted(cols, key=column_index_from_string)
    # the tables' data rows are named by the label above them (see scripts/workbook_refs.py)
    assert recorded["rows"]["Setup"]["45"] == "Match text in Identifier 1 (+1)"
    assert recorded["rows"]["Fits"]["276"] == "Material (+16)"
    assert recorded["rows"]["Fits"]["40"] == "pivot applied"
    assert recorded["work_columns"]["Z"] == "d15N drift-corr"


@pytest.fixture(scope="module")
def fresh(found):
    """The record of the unchanged workbook, as the script would write it."""
    return labels(found)


def _problems(book, fresh, change):
    wb = load_workbook(book)
    change(wb)
    return compare(collect(wb), fresh)


def test_the_record_of_the_unchanged_workbook_agrees_with_itself(found, fresh):
    assert compare(found, fresh) == [] and stale(found, fresh) == []


def test_rows_moved_under_their_formulas_are_caught(book, fresh):
    # a row put into Fits above row 39: every number below it now holds the row before, while the formulas
    # still read the old numbers, as when a writer's row moves and a hard-coded reference does not
    problems = _problems(book, fresh, lambda wb: wb["Fits"].insert_rows(39))
    assert any(p.startswith("Work!Z6 reads Fits!$B$40") and "Fits row 40 is 'slope applied' where "
               "workbook_refs.json has 'pivot applied'" in p for p in problems), problems[:5]
    assert any("Fits row 39 is empty" in p for p in problems), problems[:5]


def test_a_reworded_label_is_caught(book, fresh):
    def reword(wb):
        wb["Setup"]["A27"] = "Outlier limit"
    problems = _problems(book, fresh, reword)
    assert len(problems) == 1, problems
    assert " reads Setup!$B$27 (and " in problems[0], problems
    assert problems[0].endswith(": Setup row 27 is 'Outlier limit' where workbook_refs.json has "
                                "'Drift monitor outlier, studentized residual'"), problems


def test_a_formula_reading_an_empty_or_an_unrecorded_row_is_caught(book, fresh):
    def point(wb):
        ws = wb["Work"]
        ws["Z6"] = ws["Z6"].value.replace("Fits!$B$40", "Fits!$B$41")          # between two sections
        ws["AE6"] = ws["AE6"].value.replace("Fits!$B$52", "Fits!$B$18")        # r-squared, which nothing reads
    problems = _problems(book, fresh, point)
    assert "Work!Z6 reads Fits!$B$41: Fits row 41 is empty" in problems, problems
    assert any(p.startswith("Work!AE6 reads Fits!$B$18: Fits row 18 ('r-squared') is not in workbook_refs.json")
               for p in problems), problems


def test_a_work_column_moved_under_its_formulas_is_caught(book, fresh):
    problems = _problems(book, fresh, lambda wb: wb["Work"].insert_cols(26))   # before column Z
    assert any("Work column Z has no header in row 5" in p for p in problems), problems[:5]
    assert any("Work column AA is headed 'd15N drift-corr' where workbook_refs.json has 'd15N calibrated'" in p
               for p in problems), problems[:5]
