"""ea_irms 1.4.4, 1.4.7 and 1.4.14: rerun marks, in the engine and in the workbook.

A sample is marked for what a rerun can fix: a peak below the linearity range or
under the blank multiple, or the reference pulse at its position off; and for a
check when it ran before the background settled, lacks one element's peak, or,
from 1.4.14 (Tony, 2026-10-16), may carry over from an oversized sample before it,
which marked a rerun of N and C until then. From 1.4.7 (Tony, 2026-10-13) a peak above
the linearity range is marked only when its amplitude reaches the instrument's
largest (50,000 mV by default), and a final d15N or d13C whose 1 sigma is above
0.5 permil marks that element. The engine's marks are checked against those rules
on the 2026-09-15 test run and on copies of it with masses, shifted reference
pulses, and an oversized sample with a background rise after it, and the workbook
against the engine under ten configurations, after LibreOffice has computed the
formulas (skipped without it).
"""
import csv
import io
import json
import math
import shutil
import statistics
import subprocess
from pathlib import Path

import openpyxl
import pytest

from ea_irms.parse import parse_export
from ea_irms.pipeline import ROBUST_SD, Config, _mass_window, config_from_dict, config_to_dict, process
from ea_irms.workbook import RERUN_FIRST_ROW, build_workbook
from isoharness.serialize import results_csv

DATA = Path(__file__).resolve().parent / "data"
RUN = DATA / "20260915_raw.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def variant(folder, name, masses=False, shifts=()):
    """A copy of the test run, with masses for its samples and reference pulses moved."""
    rows = list(csv.DictReader(open(RUN, newline="")))
    for r in rows:
        if masses and r["Identifier 1"].strip() == "Random":
            r["Amount"] = f"{0.150 + 0.025 * (int(r['Row']) % 17):.3f}"
        for row, peak, col, delta in shifts:
            if r["Row"] == str(row) and r["Peak Nr"] == str(peak):
                r[col] = f"{float(r[col]) + delta:.3f}"
    out = Path(folder) / f"{name}.csv"
    with open(out, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return out


def check_config(offset):
    data = json.loads((DATA / "config_check_standard.json").read_text())
    data["check_offset_permil"] = offset
    return config_from_dict(data)


# 1.4.14: position 36's N2 peak made larger than the linearity range (Area 28 from 22.876 to 200 Vs), and the
# background under the N2 peak of the sample after it, position 37, raised 9.0 mV (24.5 to 33.5)
CARRY = ((36, 3, "Area 28", 177.124), (37, 3, "BGD 28", 9.0))


def test_marks_on_the_test_run():
    # 1.4.15: the acetanilide's whole-run check is off here, so each mark is the one the sample's own flags give.
    # On, the default, it fires on this run and adds a check to every other sample (test_engine_1415.py).
    positions, F = process(parse_export(RUN), Config(ace_scatter_check="Off"))
    marks = {p.row: p.rerun for p in positions if p.rerun}
    # 1.4.15: the size series sets the floor. Its smallest standard (8.889 Vs) reads 1.02 permil off the line
    # through the three larger ones, so the floor is the next one, 12.058 Vs, and position 21 (10.11 Vs of N2),
    # a check before, is marked for a rerun of N
    assert marks == {20: "Check", 21: "N", 23: "C", 24: "Check", 25: "Check",
                     40: "N and C", 41: "N and C", 43: "N and C", 44: "N and C",
                     45: "N and C", 47: "N and C", 53: "N"}
    assert all(p.rerun is None for p in positions if p.role != "Sample")
    p21, p53 = (next(p for p in positions if p.row == r) for r in (21, 53))
    assert p21.rerun_why == "N below series floor; ran before the background settled"
    assert p53.rerun_why == "N below linearity range; no C peak; N peak only 0.7x the blank"
    assert all(p.rerun_mass is None for p in positions)          # its samples carry no mass
    assert F["rerun"]["counts"] == {"N": 2, "C": 1, "N and C": 6, "Check": 3}
    assert F["rerun"]["run_level_text"] == "none found"
    # 1.4.7: peaks above the linearity range, all under 50,000 mV, are flagged and not marked
    p23, p29, p33 = (next(p for p in positions if p.row == r) for r in (23, 29, 33))
    assert "N above linearity range" in p23.flags and p23.rerun_why == ("C below linearity range; "
                                                                        "ran before the background settled")
    assert "C above linearity range" in p29.flags and p29.rerun is None
    assert {"N above linearity range", "C above linearity range"} <= set(p33.flags) and p33.rerun is None
    # δ outside the anchor range is a flag, never a mark on its own
    outside = [p for p in positions if p.role == "Sample" and p.flags == ["d15N outside anchor range"]]
    assert outside and all(p.rerun is None for p in outside)


def test_a_reference_pulse_off_marks_its_element(tmp_path):
    shifted = variant(tmp_path, "shift", shifts=((20, 1, "d 15N/14N", 1.0), (24, 6, "d 13C/12C", -1.0)))
    positions, F = process(parse_export(shifted), Config())
    ys = [p.ref_pulse_d15n for p in positions if p.ref_pulse_d15n is not None]
    med = statistics.median(ys)
    rsd = ROBUST_SD * statistics.median(abs(y - med) for y in ys)
    assert math.isclose(F["pulse_off"]["n"]["median"], med)
    assert math.isclose(F["pulse_off"]["n"]["limit"], 4.0 * rsd)
    p20, p24 = (next(p for p in positions if p.row == r) for r in (20, 24))
    assert p20.rerun == "N" and p20.rerun_why.startswith("ran before the background settled; N2 reference pulse off by +")
    assert p24.rerun == "C" and "CO2 reference pulse off by -" in p24.rerun_why
    quiet, _ = process(parse_export(shifted), config_from_dict({"ref_pulse_off_sd": None}))
    assert not any("reference pulse off" in f for p in quiet for f in p.flags)
    assert next(p for p in quiet if p.row == 20).rerun == "Check"


def test_possible_carryover_marks_a_check(tmp_path):
    """1.4.14 (Tony, 2026-10-16): possible carryover marks a check, with no suggested mass, where before it
    marked a rerun of N and C. An element with a reason of its own is still marked, with the carryover in
    the Why."""
    carried = variant(tmp_path, "carry", masses=True, shifts=CARRY)
    # 1.4.15: the acetanilide's whole-run check, which fires on this run and would mark 37 for a check by itself, is
    # off, so the check is the carryover's alone
    positions, F = process(parse_export(carried), Config(ace_scatter_check="Off"))
    by = {p.row: p for p in positions}
    assert F["carryover"]["flagged"] == [37] and by[36].oversized
    assert by[37].role == "Sample" and by[37].carry_rise == pytest.approx(9.0)
    assert (by[37].rerun, by[37].rerun_why, by[37].rerun_mass) == (
        "Check", "possible carryover from an oversized sample", None)
    assert by[38].role == "Drift" and by[38].rerun is None          # exposed, with no rise
    # CO2 at position 37 reads 6,395 mV, so a limit of 6,000 mV marks C, and only C
    positions, _ = process(parse_export(carried), config_from_dict({"max_amplitude_mv": 6000,
                                                                    "ace_scatter_check": "Off"}))
    p37 = next(p for p in positions if p.row == 37)
    assert p37.rerun == "C" and p37.rerun_mass not in (None, "none")
    assert p37.rerun_why == "possible carryover from an oversized sample; CO2 peak at or above 6000 mV"


def test_texts_put_carryover_among_the_checks(tmp_path):
    wb = openpyxl.load_workbook(build_workbook(RUN, tmp_path / "texts.xlsx", Config()))
    readme = " ".join(str(c.value) for c in wb["README"]["B"] if c.value)
    assert "when it follows an oversized sample and its background rose (Setup B28)" not in readme
    assert ("It is marked for a check when it ran before the background settled, when it follows an oversized "
            "sample and its background rose by more than Setup B28 (possible carryover, a check from 1.4.14") in readme
    qc = wb["QC"]["A167"].value
    assert qc.index("possible carryover, a check from 1.4.14") > qc.index("It is marked for a check")
    assert "follows an oversized sample with a background rise, or" not in qc
    assert "from 1.4.14 possible carryover marks a check, not a rerun" in wb["Rerun"]["A2"].value


def test_large_peaks_are_marked_only_at_the_instrument_limit():
    old, _ = process(parse_export(RUN), config_from_dict({"max_amplitude_mv": None}))
    assert {p.row: p.rerun for p in old if p.row in (23, 29, 33)} == {23: "N and C", 29: "C", 33: "N and C"}
    # CO2 at position 29 reads 10,125 mV and at 33 12,392 mV; N2 at 29 reads 3,000 mV, inside the range. The
    # acetanilide's whole-run check (1.4.15), whose flag would lead the Why, is off.
    positions, F = process(parse_export(RUN), config_from_dict({"max_amplitude_mv": 10000,
                                                                "ace_scatter_check": "Off"}))
    by = {p.row: p for p in positions}
    assert by[29].rerun == "C" and by[29].rerun_why == "CO2 peak at or above 10000 mV"
    assert by[33].rerun == "C" and by[33].rerun_why == "CO2 peak at or above 10000 mV"
    at, _ = process(parse_export(RUN), config_from_dict({"max_amplitude_mv": 3000}))
    p29 = next(p for p in at if p.row == 29)
    assert p29.ampl_n == 3000.0 and "N2 peak at or above 3000 mV" in p29.rerun_why      # at the limit counts
    assert p29.rerun == "N and C" and "N above linearity range" not in p29.rerun_why
    assert all(p.rerun is None or "linearity range" not in p.rerun_why
               or "below linearity range" in p.rerun_why for p in at)


def test_a_1sigma_above_the_limit_marks_its_element():
    positions, _ = process(parse_export(RUN), Config())
    samples = [p for p in positions if p.role == "Sample"]
    assert max(p.u_d15n for p in samples if p.u_d15n is not None) < 0.5        # 0.2998 at most, position 25
    assert not any("1σ(" in (p.rerun_why or "") for p in samples)
    tight, _ = process(parse_export(RUN), config_from_dict({"rerun_max_1sigma": 0.25}))
    for p in (q for q in tight if q.role == "Sample"):
        over_n = p.u_d15n is not None and p.u_d15n > 0.25
        assert (f"1σ(d15N) {p.u_d15n:.2f} above 0.25" in (p.rerun_why or "")) == over_n, p.row
        if over_n:
            assert p.rerun in ("N", "N and C")
    assert sum(1 for p in tight if p.rerun_why and "1σ(d15N)" in p.rerun_why) > 0
    # both rules off, as before 1.4.7, and the acetanilide's whole-run check of 1.4.15 too
    off, _ = process(parse_export(RUN), config_from_dict({"rerun_max_1sigma": None, "max_amplitude_mv": None,
                                                          "ace_scatter_check": "Off"}))
    old = {p.row: (p.rerun, p.rerun_why) for p in off}
    assert old[29] == ("C", "C above linearity range")


def test_mass_window_by_hand():
    class P:
        area_n, amount = 20.0, 0.5
    assert _mass_window(P, "n", 10.0, 50.0) == (0.25, 1.25)
    assert _mass_window(P, "n", 60.0, 50.0) is None              # no mass reaches the range
    # 1.4.10: the yield above the blank, (20 - 4) / 0.5 = 32 per mg, sets the masses
    assert _mass_window(P, "n", 10.0, 50.0, 4.0) == ((10.0 - 4.0) / 32.0, (50.0 - 4.0) / 32.0)
    assert _mass_window(P, "n", 10.0, 50.0, 20.0) is None        # a peak no larger than the blank
    P.amount = None
    assert _mass_window(P, "n", 10.0, 50.0) is None


def test_suggested_mass_reads_the_yield_above_the_blank():
    """Finding R-10 of the 2026-10-14 review: PanamaSids2 position 25, 1.808 Vs of N2 from 18.642 mg over a
    blank of 1.7268 Vs, needs about 1,728 mg to reach 9.253 Vs, where the area per mg without the blank
    suggested 95.4 mg; and a mass above every capsule of the run says so."""
    class P:
        area_n, amount = 1.808, 18.642
    lo, hi = _mass_window(P, "n", 9.253, 20.0, 1.7268)
    assert abs(lo - (9.253 - 1.7268) / ((1.808 - 1.7268) / 18.642)) < 1e-9 and 1727 < lo < 1729


def test_suggested_mass_from_the_linearity_range(tmp_path):
    weighed = variant(tmp_path, "mass", masses=True)
    # the rule before 1.4.7, under which a peak above the linearity range is marked
    positions, F = process(parse_export(weighed), config_from_dict({"max_amplitude_mv": None}))
    L, mult = F["size_limits"], Config().min_blank_multiple

    def window(p, el):
        ab = F[f"blank_{el}"]["area"]                  # 1.4.10: the yield above the blank
        per_mg = (getattr(p, f"area_{el}") - ab) / p.amount
        lo = max(L[f"min_{el}"], mult * ab)
        return (lo - ab) / per_mg, (L[f"max_{el}"] - ab) / per_mg

    p29 = next(p for p in positions if p.row == 29)                 # C above the range
    (ln, hn), (lc, hc) = window(p29, "n"), window(p29, "c")
    assert p29.rerun == "C" and p29.rerun_mass == f"{max(ln, lc):.3f} to {min(hn, hc):.3f}"
    p23 = next(p for p in positions if p.row == 23)                 # no one mass suits both
    (ln, hn), (lc, hc) = window(p23, "n"), window(p23, "c")
    assert max(ln, lc) > min(hn, hc)
    assert p23.rerun_mass == f"N {ln:.3f} to {hn:.3f}; C {lc:.3f} to {hc:.3f} (two capsules)"
    assert all(p.rerun_mass is None for p in positions if p.rerun == "Check")


def test_problems_of_the_whole_run():
    _, F = process(parse_export(RUN), check_config(0.05))
    assert F["rerun"]["run_level"][0] == "check standard off by more than 0.05 permil at positions 2, 3, 19"
    _, F = process(parse_export(RUN), config_from_dict({"coverage_k": 0.5, "ref_gas_sd_limit": 0.05}))
    items = F["rerun"]["run_level"]
    assert any(t.startswith("anchors outside 0.5σ, %N: ") for t in items)
    assert any(t.startswith("N2 reference pulses SD ") and t.endswith("at or above 0.05") for t in items)
    # 1.4.15: with the defaults the test run has one problem of the whole run, the acetanilide's d15N scatter
    # (test_engine_1415.py); with that check off it has none
    _, F = process(parse_export(RUN), config_from_dict({"ref_gas_sd_limit": None}))
    assert [t.split(",")[0] for t in F["rerun"]["run_level"]] == ["acetanilide d15N"]
    _, F = process(parse_export(RUN), config_from_dict({"ref_gas_sd_limit": None, "ace_scatter_check": "Off"}))
    assert F["rerun"]["run_level_text"] == "none found"


def test_results_table_carries_the_mark_after_flags():
    positions, _ = process(parse_export(RUN), Config())
    text = results_csv(positions, Config())
    header = next(csv.reader(io.StringIO(text)))
    i = header.index("Flags")
    assert header[i + 1:i + 4] == ["Rerun", "Why", "Suggested mass (mg)"]
    rows = {r["Row"]: r for r in csv.DictReader(io.StringIO(text))}
    assert rows["53"]["Rerun"] == "N" and rows["53"]["Why"].startswith("N below linearity range")
    assert rows["11"]["Rerun"] == "" and rows["11"]["Suggested mass (mg)"] == ""


def test_configuration_round_trip_keeps_the_new_settings():
    assert Config().ref_pulse_off_sd == 4.0 and Config().ref_gas_sd_limit == 0.1
    assert Config().max_amplitude_mv == 50000.0 and Config().rerun_max_1sigma == 0.5
    cfg = config_from_dict({"ref_pulse_off_sd": 3.0, "ref_gas_sd_limit": None,
                            "max_amplitude_mv": None, "rerun_max_1sigma": 0.3})
    back = config_from_dict(config_to_dict(cfg))
    assert back.ref_pulse_off_sd == 3.0 and back.ref_gas_sd_limit is None
    assert back.max_amplitude_mv is None and back.rerun_max_1sigma == 0.3


CASES = ("default", "masses", "pulses", "off", "run_level", "check", "limits", "before_147",
         "carryover", "carryover_c")


@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    if not SOFFICE:
        pytest.skip("needs LibreOffice")
    folder = tmp_path_factory.mktemp("rerun")
    pulses = ((20, 1, "d 15N/14N", 1.0), (24, 6, "d 13C/12C", -1.0), (30, 1, "d 15N/14N", -0.8))
    cases = {
        "default": (RUN, Config()),
        "masses": (variant(folder, "masses", masses=True), Config()),
        "pulses": (variant(folder, "pulses", masses=True, shifts=pulses), Config()),
        "off": (variant(folder, "off", masses=True, shifts=pulses[:1]),
                config_from_dict({"ref_pulse_off_sd": None, "min_blank_multiple": None, "ref_gas_sd_limit": None})),
        "run_level": (variant(folder, "run_level", masses=True),
                      config_from_dict({"coverage_k": 0.5, "ref_gas_sd_limit": 0.05, "ref_pulse_off_sd": 2.0})),
        "check": (RUN, check_config(0.05)),
        # 1.4.7: a low amplitude limit and a tight 1σ, so both rules mark samples; then both switched off
        "limits": (variant(folder, "limits", masses=True),
                   config_from_dict({"max_amplitude_mv": 3000, "rerun_max_1sigma": 0.25})),
        "before_147": (variant(folder, "before_147", masses=True),
                       config_from_dict({"max_amplitude_mv": None, "rerun_max_1sigma": None})),
        # 1.4.14: possible carryover alone marks a check, and with a CO2 peak at the limit a rerun of C
        "carryover": (variant(folder, "carryover", masses=True, shifts=CARRY), Config()),
        "carryover_c": (variant(folder, "carryover_c", masses=True, shifts=CARRY),
                        config_from_dict({"max_amplitude_mv": 6000})),
    }
    files = []
    for name, (export, cfg) in cases.items():
        out = folder / f"{name}.xlsx"
        build_workbook(export, out, cfg)
        files.append(out)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(folder / "calc")]
                   + [str(f) for f in files], check=True, capture_output=True, timeout=900)
    return folder, cases


def _close(mine, theirs):
    theirs = theirs if isinstance(theirs, (int, float)) and not isinstance(theirs, bool) else None
    if mine is None:
        return theirs is None
    return theirs is not None and abs(mine - theirs) <= 1e-9 * max(1.0, abs(mine))


@pytest.mark.parametrize("name", CASES)
def test_workbook_agrees_with_engine(computed, name):
    folder, cases = computed
    export, cfg = cases[name]
    positions, F = process(parse_export(export), cfg)
    wb = openpyxl.load_workbook(folder / "calc" / f"{name}.xlsx", data_only=True)
    rs, ft, qc, rr = wb["Results"], wb["Fits"], wb["QC"], wb["Rerun"]
    head = {str(rs.cell(4, c).value).split(",")[0]: c for c in range(1, rs.max_column + 1)}
    assert [rs.cell(4, c).value for c in range(17, 21)] == ["Flags", "Rerun", "Why", "Suggested mass (mg)"]
    for p in positions:
        i = 4 + p.row
        flags = [t.strip() for t in str(rs.cell(i, head["Flags"]).value or "").split(";") if t.strip()]
        assert flags == p.flags, (name, p.row, flags, p.flags)
        for h, attr in (("Rerun", "rerun"), ("Why", "rerun_why"), ("Suggested mass (mg)", "rerun_mass")):
            assert (getattr(p, attr) or "") == (rs.cell(i, head[h]).value or ""), (name, p.row, h)
        for h, attr in (("1σ(%N)", "u_pct_n"), ("1σ(d15N)", "u_d15n"), ("%N as measured", "pct_n_measured")):
            assert _close(getattr(p, attr), rs.cell(i, head[h]).value), (name, p.row, h)
    marked = [p for p in positions if p.rerun]
    top = RERUN_FIRST_ROW                  # 1.4.15: below the corrections fitted but not applied (row 8 before)
    assert [rr.cell(top + k, 1).value for k in range(len(marked))] == [p.row for p in marked]
    assert rr.cell(top + len(marked), 1).value in (None, "")
    for k, p in enumerate(marked):
        assert [rr.cell(top + k, j).value or "" for j in (5, 6, 7)] == [p.rerun, p.rerun_why, p.rerun_mass or ""]
    for base, el in ((339, "n"), (343, "c")):
        po = F["pulse_off"][el]
        for off, key in ((0, "n"), (1, "median"), (2, "robust_sd"), (3, "limit")):
            mine, theirs = po[key], ft.cell(base + off, 2).value
            assert (theirs in (None, "")) if mine is None else _close(mine, theirs), (name, key)
    for row, key, el in ((347, "lo_area", "n"), (348, "hi_area", "n"), (349, "lo_area", "c"), (350, "hi_area", "c")):
        assert _close(F["rerun"][key][el], ft.cell(row, 2).value), (name, row)
    c = F["rerun"]["counts"]
    assert (qc["B168"].value, qc["B169"].value) == (c["N"] + c["C"] + c["N and C"], c["Check"])
    assert ft["B356"].value == rr["A5"].value == F["rerun"]["run_level_text"]
    errors = [x.coordinate for s in (rs, ft, qc, rr) for row in s.iter_rows() for x in row
              if isinstance(x.value, str) and x.value.startswith("#") and x.value[1:4].isupper()]
    assert not errors, errors[:5]


def test_rows_are_shaded_by_their_mark(computed):
    folder, _ = computed
    wb = openpyxl.load_workbook(folder / "default.xlsx")
    for sheet, col in (("Results", "$R5"), ("Rerun", f"$E{RERUN_FIRST_ROW}")):
        rules = [r for rng in wb[sheet].conditional_formatting for r in rng.rules]
        assert {r.formula[0] for r in rules} == {f'{col}="Check"', f'AND({col}<>"",{col}<>"Check")'}
