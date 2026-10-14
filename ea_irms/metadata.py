"""Sample metadata: the sample sheet and the run information (ea_irms 1.4.8).

A sample sheet is a table with one row per sample: an ID column and any others, such as
project, site, species, depth, or notes. Each Sample position of a run is matched to a row
by its identifier (Identifier 1): equal to the row's ID, or beginning with it followed by a
character that is neither a letter nor a digit, so that CI 3-2 matches CI 3-2-S and CI 3-2 H
but not CI 3-20. When several IDs fit, the longest wins. Case, spaces at either end, and runs
of spaces are ignored. When no ID fits as written, every space is taken out of the name and
the IDs and the comparison is made again, so that PD1-H matches PD 1 (Tony, 2026-10-13); such
a match says so. Standards, blanks, and the other roles of the standards table are not
matched.

Analyses of one sample in one run are replicates: Sample positions matched to the sheet that
have the same identifier, or the same value in the sheet's Replicate of column, which names
the sample an analysis repeats. Matching by a shorter ID does not make two analyses
replicates, so CI 3-2-H and CI 3-2-S, fractions of one colony, stay apart.

The run information is four short texts typed for each run: Operator, Instrument, Project,
and Notes. Nothing here changes a value; ea_irms.pipeline compares the replicates.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import math
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

RUN_FIELDS = ("Operator", "Instrument", "Project", "Notes")
REPLICATE_COLUMN = "Replicate of"
# The template's columns (Tony, 2026-10-13), and Replicate of for the replicate rule.
SAMPLE_COLUMNS = ("Sample ID", "Project", "Site", "Species or material", "Fraction", "Collection date",
                  "Depth", "Notes", REPLICATE_COLUMN)
SHEET_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".txt")
MAX_ROWS = 20000
MAX_COLUMNS = 60
MAX_TEXT = 1000          # characters kept of one cell
MAX_RUN_TEXT = 500       # characters kept of one run-information field
HEADER_SEARCH = 30       # rows searched for the header
EMPTY_RUN_STOP = 2000    # consecutive empty rows that end the reading of a sheet
MATCH_HEADER = ("Sample ID (sheet)", "Matched by")
NOT_MATCHED = "not in sheet"
SPACES_IGNORED = ", spaces ignored"      # after exact or prefix, for a match found without spaces
# Column names the results tables already use (and Run, which the combined table puts first), compared
# without regard to case; a sheet column of one of these names, or one headed 1σ(...) or 1sigma(...), is
# written with " (sheet)" after it.
RESULTS_NAMES = ("Run", "Row", "Identifier", "Type", "Mass (mg)", "%C", "%N", "C:N molar", "C:N mass",
                 "d13C (permil)", "d15N (permil)", "Flags", "Rerun", "Why", "Suggested mass (mg)",
                 "Elemental drift factor N", "Elemental drift factor C", "%C as measured", "%N as measured")

# f(n) of ISO 5725-6, section 5.2: the 0.95 quantile of the range of n results from one normal
# distribution, in units of its standard deviation, for n = 2 to 20. Computed here by integrating
# P(W <= w) = n * integral of phi(z) [Phi(z + w) - Phi(z)]^(n - 1) dz and solving for 0.95 (equal to
# scipy.stats.studentized_range.ppf(0.95, n, inf) to 1e-6, checked in tests/test_replicates.py).
# For two results it is 1.96 x sqrt(2) = 2.7718, the factor ISO 5725-6 rounds to 2.8 in r = 2.8 s_r.
CRITICAL_RANGE = {2: 2.771808, 3: 3.314493, 4: 3.633160, 5: 3.857656, 6: 4.030092, 7: 4.169554,
                  8: 4.286309, 9: 4.386509, 10: 4.474124, 11: 4.551864, 12: 4.621655, 13: 4.684920,
                  14: 4.742732, 15: 4.795924, 16: 4.845154, 17: 4.890951, 18: 4.933745, 19: 4.973892,
                  20: 5.011689}


class SampleSheetError(ValueError):
    """The file is not a sample sheet ea_irms can read, or names a column it lacks."""


def _range_cdf(w: float, n: int, lo: float = -9.0, hi: float = 9.0, m: int = 6000) -> float:
    """P(range of n standard normal values <= w), by Simpson's rule."""
    h = (hi - lo) / m
    total = 0.0
    r2 = math.sqrt(2.0)
    for i in range(m + 1):
        z = lo + i * h
        phi = math.exp(-0.5 * z * z) / math.sqrt(2.0 * math.pi)
        d = 0.5 * (math.erf((z + w) / r2) - math.erf(z / r2))
        total += phi * d ** (n - 1) * (1 if i in (0, m) else 4 if i % 2 else 2)
    return n * total * h / 3.0


@lru_cache(maxsize=None)
def _range_quantile(n: int, p: float = 0.95) -> float:
    a, b = 0.0, 12.0
    for _ in range(60):
        c = 0.5 * (a + b)
        if _range_cdf(c, n) < p:
            a = c
        else:
            b = c
    return round(0.5 * (a + b), 6)


def critical_range_factor(n: int) -> float | None:
    """f(n) for n results (ISO 5725-6): the table for 2 to 20, computed beyond; None below 2."""
    if n < 2:
        return None
    return CRITICAL_RANGE.get(n) or _range_quantile(n)


# ------------------------------------------------------------------ text
def cell_text(value) -> str:
    """A cell as text: whole numbers without a decimal point, dates as ISO dates."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else f"{value:.15g}"
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    return str(value)


def norm(value) -> str:
    """Text as the matching compares it: case folded, spaces at either end dropped, runs of spaces as one."""
    return " ".join(cell_text(value).split()).casefold()


def excel_trim(value) -> str:
    """Text as Excel's TRIM leaves it: spaces at either end dropped and runs of spaces as one."""
    return re.sub(" +", " ", cell_text(value).strip(" "))


# ------------------------------------------------------------------ the sheet
@dataclass
class SampleSheet:
    """A sample sheet as read: every column (the ID column among them) and one row per sample."""
    name: str
    columns: list
    rows: list
    id_column: str
    sheet_name: str = ""
    header_row: int = 1
    notes: list = field(default_factory=list)

    def __post_init__(self):
        self.id_index = self.columns.index(self.id_column)
        folded = [norm(c) for c in self.columns]
        self.replicate_index = folded.index(norm(REPLICATE_COLUMN)) if norm(REPLICATE_COLUMN) in folded else None
        self.lookup: dict[str, int] = {}
        self.duplicates: dict[str, int] = {}
        self.lookup_nospace: dict[str, int] = {}     # the IDs without any space, for the second pass
        for i, row in enumerate(self.rows):
            key = norm(row[self.id_index])
            if key in self.lookup:
                self.duplicates[key] = self.duplicates.get(key, 1) + 1
            else:
                self.lookup[key] = i
            self.lookup_nospace.setdefault(key.replace(" ", ""), i)

    def sample_id(self, i: int) -> str:
        return cell_text(self.rows[i][self.id_index]).strip()

    def value(self, i: int, j: int):
        return self.rows[i][j]

    def match(self, identifier) -> tuple[int | None, str | None]:
        """The row an identifier matches and how: exact or prefix, as written, and when nothing fits
        that way, with every space taken out of both ("exact, spaces ignored" or "prefix, spaces
        ignored"); (None, None) when none does."""
        t = norm(identifier)
        if not t:
            return None, None
        for table, text, tag in ((self.lookup, t, ""), (self.lookup_nospace, t.replace(" ", ""), SPACES_IGNORED)):
            if text in table:
                return table[text], "exact" + tag
            for k in range(len(text) - 1, 0, -1):
                if text[k - 1].isalnum() and text[k].isalnum():
                    continue          # inside a word or a number: CI 3-2 does not match CI 3-20
                hit = table.get(text[:k])
                if hit is not None:
                    return hit, "prefix" + tag
        return None, None

    def describe(self) -> dict:
        """The sheet as the page shows it."""
        return {"name": self.name, "sheet": self.sheet_name, "header_row": self.header_row,
                "columns": list(self.columns), "id_column": self.id_column, "rows": len(self.rows),
                "replicate_column": None if self.replicate_index is None else self.columns[self.replicate_index],
                "duplicates": sorted(self.duplicates), "notes": list(self.notes)}


def _clean_value(value, notes_count: dict):
    """A cell's value as kept: text trimmed, empty as None, numbers and dates as they are."""
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return None
        if len(v) > MAX_TEXT:
            notes_count["cut"] = notes_count.get("cut", 0) + 1
            v = v[:MAX_TEXT]
        return v
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, _dt.datetime):
        return value.date().isoformat() if value.time() == _dt.time(0, 0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    return value


_PLAIN_NUMBER = re.compile(r"-?(0|[1-9][0-9]*)(\.[0-9]+)?")


def _number_from_text(v):
    """Text that is a plain decimal number becomes the number; text such as 001 or 3-2 stays text."""
    if isinstance(v, str) and _PLAIN_NUMBER.fullmatch(v):
        return float(v) if "." in v else int(v)
    return v


def _xlsx_grid(data: bytes):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise SampleSheetError(f"it could not be opened as an Excel workbook ({type(exc).__name__})") from exc
    try:
        names = wb.sheetnames
        pick = next((n for n in names if n.strip().casefold() == "samples"), names[0] if names else None)
        if pick is None:
            raise SampleSheetError("the workbook has no sheets")
        ws = wb[pick]
        grid, empty = [], 0
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
                empty += 1
                if empty >= EMPTY_RUN_STOP:
                    break
                grid.append([])
                continue
            empty = 0
            grid.append(vals)
            if len(grid) > MAX_ROWS + HEADER_SEARCH + EMPTY_RUN_STOP:
                break
        return pick, grid
    finally:
        wb.close()


def _xls_grid(data: bytes):
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise SampleSheetError(f"it could not be opened as an Excel 97-2003 workbook ({type(exc).__name__})") from exc
    names = book.sheet_names()
    pick = next((n for n in names if n.strip().casefold() == "samples"), names[0] if names else None)
    if pick is None:
        raise SampleSheetError("the workbook has no sheets")
    sh = book.sheet_by_name(pick)
    grid = []
    for r in range(min(sh.nrows, MAX_ROWS + HEADER_SEARCH + EMPTY_RUN_STOP)):
        row = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            v = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                except Exception:
                    pass
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                v = None
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                v = int(v)
            row.append(v)
        grid.append(row)
    return pick, grid


def _text_grid(data: bytes, suffix: str):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:8192]
    if suffix == ".tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delim = "\t" if sample.count("\t") > sample.count(",") else ","
    grid = []
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        grid.append([_number_from_text(v.strip()) if isinstance(v, str) else v for v in row])
        if len(grid) > MAX_ROWS + HEADER_SEARCH + EMPTY_RUN_STOP:
            break
    return "", grid


def _guess_id_column(columns: list) -> str:
    folded = [norm(c) for c in columns]
    for want in ("sample id", "sample_id", "sampleid", "sample-id", "id"):
        if want in folded:
            return columns[folded.index(want)]
    for c, f in zip(columns, folded):
        if {"id", "identifier"} & set(re.split(r"[^0-9a-z]+", f)):
            return c
    for c, f in zip(columns, folded):
        if "sample" in f or "name" in f:
            return c
    return columns[0]


def read_sample_sheet(data: bytes, name: str, id_column: str | None = None) -> SampleSheet:
    """Read a sample sheet (.xlsx, .xlsm, .xls, .csv, .tsv, or .txt) from its bytes.

    The sheet is the one named Samples, or the first. The header is the first row, among the
    first 30, with as many filled cells as the fullest of them up to two, so a title line above
    it is passed over. The ID column is ``id_column`` when given (matched without regard to case),
    else a column named Sample ID, else one whose name has the word ID, else one named for a
    sample or a name, else the first. Rows without an ID are left out, and the notes say so.
    """
    name = Path(name or "sheet").name
    suffix = Path(name).suffix.lower()
    if suffix not in SHEET_SUFFIXES:
        raise SampleSheetError(f"{name} is not a sample sheet ea_irms reads; use one of {', '.join(SHEET_SUFFIXES)}")
    if not data:
        raise SampleSheetError(f"{name} is empty")
    if suffix in (".xlsx", ".xlsm"):
        sheet_name, grid = _xlsx_grid(data)
    elif suffix == ".xls":
        sheet_name, grid = _xls_grid(data)
    else:
        sheet_name, grid = _text_grid(data, suffix)

    def filled(row):
        return sum(1 for v in row if v is not None and not (isinstance(v, str) and not v.strip()))

    head = [filled(r) for r in grid[:HEADER_SEARCH]]
    if not head or max(head) == 0:
        raise SampleSheetError(f"{name} holds no table")
    need = min(2, max(head))
    h = next(i for i, k in enumerate(head) if k >= need)
    body = [r for r in grid[h + 1:]]
    width = max([len(grid[h])] + [len(r) for r in body])
    # the last column that holds anything, in the header or below it
    last = 0
    for r in [grid[h]] + body:
        for j in range(len(r) - 1, -1, -1):
            v = r[j]
            if v is not None and not (isinstance(v, str) and not v.strip()):
                last = max(last, j + 1)
                break
    width = min(width, last)
    notes, counts = [], {}
    if width > MAX_COLUMNS:
        notes.append(f"only the first {MAX_COLUMNS} of its {width} columns were read")
        width = MAX_COLUMNS
    columns, seen = [], {}
    for j in range(width):
        raw = grid[h][j] if j < len(grid[h]) else None
        label = " ".join(cell_text(raw).split()) or f"Column {_letters(j + 1)}"
        key = label.casefold()
        seen[key] = seen.get(key, 0) + 1
        columns.append(label if seen[key] == 1 else f"{label} ({seen[key]})")
    if id_column is not None and str(id_column).strip():
        want = norm(id_column)
        hit = next((c for c in columns if norm(c) == want), None)
        if hit is None:
            raise SampleSheetError(f"{name} has no column named {str(id_column).strip()!r}; its columns are "
                                   + ", ".join(columns))
        idc = hit
    else:
        idc = _guess_id_column(columns)
    ix = columns.index(idc)
    rows, no_id = [], 0
    for r in body:
        vals = [_clean_value(r[j] if j < len(r) else None, counts) for j in range(width)]
        if all(v is None for v in vals):
            continue
        if vals[ix] is None or not cell_text(vals[ix]).strip():
            no_id += 1
            continue
        rows.append(vals)
        if len(rows) > MAX_ROWS:
            raise SampleSheetError(f"{name} has more than {MAX_ROWS} sample rows, the most this service reads")
    if not rows:
        raise SampleSheetError(f"{name} has no row with a value in its ID column, {idc}")
    if h:
        notes.append(f"the header was read from row {h + 1}")
    if no_id:
        notes.append(f"{no_id} row{'s' if no_id != 1 else ''} without a value in {idc} {'were' if no_id != 1 else 'was'} left out")
    if counts.get("cut"):
        notes.append(f"{counts['cut']} cell{'s' if counts['cut'] != 1 else ''} longer than {MAX_TEXT} characters "
                     f"{'were' if counts['cut'] != 1 else 'was'} cut to {MAX_TEXT}")
    sheet = SampleSheet(name=name, columns=columns, rows=rows, id_column=idc, sheet_name=sheet_name,
                        header_row=h + 1, notes=notes)
    for key, k in sorted(sheet.duplicates.items()):
        sheet.notes.append(f"the ID {sheet.sample_id(sheet.lookup[key])} is in {k} rows; the first is used")
    return sheet


def _letters(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# ------------------------------------------------------------------ matching and replicates
def match_positions(positions, sheet: SampleSheet | None) -> list[dict]:
    """Match each Sample position to the sheet and find the replicates; returns the replicate groups.

    Sets on every position ``sheet_index`` (the row of the sheet, or None), ``sample_id``,
    ``matched_by`` (exact, prefix, or None), and ``replicate_group`` (a number from 1 in order
    of the group's first position, or None). Roles must be assigned first. A group is two or more
    Sample positions matched to the sheet with the same identifier or the same Replicate of.
    """
    keyed: dict[str, list] = {}
    for p in sorted(positions, key=lambda q: q.row):
        p.sheet_index = p.sample_id = p.matched_by = None
        p.replicate_group = None
        if sheet is None or p.role != "Sample":
            continue
        i, how = sheet.match(p.identifier)
        if i is None:
            continue
        p.sheet_index, p.matched_by, p.sample_id = i, how, sheet.sample_id(i)
        rep = sheet.value(i, sheet.replicate_index) if sheet.replicate_index is not None else None
        key = norm(rep) if norm(rep) else norm(p.identifier)
        if key:
            keyed.setdefault(key, []).append((p, cell_text(rep).strip() if norm(rep) else " ".join(p.identifier.split())))
    groups = []
    for key, members in keyed.items():
        if len(members) < 2:
            continue
        g = len(groups) + 1
        for p, _ in members:
            p.replicate_group = g
        by = "name"
        if sheet.replicate_index is not None and any(norm(sheet.value(p.sheet_index, sheet.replicate_index))
                                                     for p, _ in members):
            by = REPLICATE_COLUMN
        groups.append({"group": g, "key": key, "name": members[0][1], "rows": [p.row for p, _ in members],
                       "by": by})
    return groups


def match_summary(positions, sheet: SampleSheet | None) -> dict:
    """How the run's samples matched: counts, and the samples the sheet does not hold."""
    samples = [p for p in positions if p.role == "Sample"]
    if sheet is None:
        return {"sheet": None, "samples": len(samples), "matched": 0, "exact": 0, "prefix": 0, "spaces": 0,
                "unmatched": []}
    how = [getattr(p, "matched_by", None) or "" for p in samples]
    exact = sum(1 for h in how if h.startswith("exact"))
    prefix = sum(1 for h in how if h.startswith("prefix"))
    return {"sheet": sheet.name, "samples": len(samples), "matched": exact + prefix, "exact": exact,
            "prefix": prefix, "spaces": sum(1 for h in how if h.endswith(SPACES_IGNORED)),
            "unmatched": [{"row": p.row, "identifier": p.identifier} for p in samples
                          if getattr(p, "sheet_index", None) is None]}


def summary_text(m: dict) -> str:
    """The match counts in one line, as the workbook and the combined results write them."""
    if not m.get("sheet"):
        return "no sample sheet"
    parts = []
    if m["exact"]:
        parts.append(f"{m['exact']} by the exact name")
    if m["prefix"]:
        parts.append(f"{m['prefix']} by an ID the name begins with")
    text = f"{m['matched']} of {m['samples']} sample{'s' if m['samples'] != 1 else ''} matched"
    if parts and m.get("spaces"):
        parts[-1] += f"; {m['spaces']} of the {m['matched']} only with spaces ignored"
    return text + (f" ({', '.join(parts)})" if parts else "")


# ------------------------------------------------------------------ results columns
def output_columns(sheet: SampleSheet | None, position_lists) -> list[tuple[int, str]]:
    """The sheet's columns a results table writes after Sample ID (sheet) and Matched by: each
    column other than the ID that holds a value for at least one matched sample, in the sheet's
    order, as (column index, header). A name the table already uses gets " (sheet)" after it."""
    if sheet is None:
        return []
    used = set()
    for positions in position_lists:
        for p in positions:
            if getattr(p, "sheet_index", None) is None:
                continue
            row = sheet.rows[p.sheet_index]
            for j, v in enumerate(row):
                if j != sheet.id_index and v is not None and cell_text(v).strip():
                    used.add(j)
    taken = {t.casefold() for t in RESULTS_NAMES + MATCH_HEADER}
    out = []
    for j, c in enumerate(sheet.columns):
        if j in used:
            clash = c.casefold() in taken or c.casefold().startswith(("1σ(", "1sigma("))
            out.append((j, f"{c} (sheet)" if clash else c))
    return out


def output_header(sheet: SampleSheet | None, columns) -> list[str]:
    return [] if sheet is None else list(MATCH_HEADER) + [h for _, h in columns]


def output_values(p, sheet: SampleSheet | None, columns) -> list:
    """One position's metadata cells: its sheet ID, how it matched, and the sheet's columns."""
    if sheet is None:
        return []
    if p.role != "Sample":
        return [None, None] + [None] * len(columns)
    i = getattr(p, "sheet_index", None)
    if i is None:
        return [None, NOT_MATCHED] + [None] * len(columns)
    row = sheet.rows[i]
    return [sheet.sample_id(i), p.matched_by] + [row[j] for j, _ in columns]


# ------------------------------------------------------------------ run information
def clean_run_info(data) -> dict:
    """The run information as kept: the four fields, trimmed, empty ones left out.

    Accepts the field names in any case; raises ValueError for anything else."""
    if data is None:
        return {}
    if not isinstance(data, dict):
        raise ValueError("the run information must be a JSON object with Operator, Instrument, Project, and Notes")
    names = {f.casefold(): f for f in RUN_FIELDS}
    out, unknown = {}, []
    for k, v in data.items():
        f = names.get(str(k).strip().casefold())
        if f is None:
            unknown.append(str(k))
            continue
        if v is None:
            continue
        if not isinstance(v, (str, int, float)) or isinstance(v, bool):
            raise ValueError(f"the run information's {f} must be text")
        t = " ".join(str(v).split()) if f != "Notes" else str(v).strip()
        if t:
            out[f] = t[:MAX_RUN_TEXT]
    if unknown:
        raise ValueError("the run information has fields ea_irms does not read (" + ", ".join(unknown)
                         + "); it reads Operator, Instrument, Project, and Notes")
    return {f: out[f] for f in RUN_FIELDS if f in out}


# ------------------------------------------------------------------ the template
def sample_names(name_lists) -> list[str]:
    """The distinct sample names of one or more runs, in order of first appearance."""
    seen, out = set(), []
    for names in name_lists:
        for n in names:
            t = " ".join(cell_text(n).split())
            if t and norm(t) not in seen:
                seen.add(norm(t))
                out.append(t)
    return out


def template_xlsx(names) -> bytes:
    """A sample sheet to fill in: the Samples sheet with one row per name, and How to fill."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from .workbook import VERSION
    arial = "Arial"
    f_title, f_h1 = Font(name=arial, size=14, bold=True), Font(name=arial, size=11, bold=True)
    f_hdr, f_sm = Font(name=arial, size=10, bold=True, color="FFFFFF"), Font(name=arial, size=10)
    f_in, f_note = Font(name=arial, size=10, color="0000FF"), Font(name=arial, size=9, italic=True, color="595959")
    fill_hdr, fill_in = PatternFill("solid", fgColor="1F4E79"), PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"
    widths = {"Sample ID": 24, "Project": 16, "Site": 14, "Species or material": 26, "Fraction": 12,
              "Collection date": 15, "Depth": 10, "Notes": 36, REPLICATE_COLUMN: 20}
    for j, h in enumerate(SAMPLE_COLUMNS, start=1):
        c = ws.cell(1, j, h)
        c.font, c.fill = f_hdr, fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[_letters(j)].width = widths.get(h, 14)
    names = list(names)
    for i, n in enumerate(names, start=2):
        c = ws.cell(i, 1, n)
        c.font = f_sm
        c.number_format = "@"
        for j in range(2, len(SAMPLE_COLUMNS) + 1):
            c = ws.cell(i, j)
            c.font, c.fill = f_in, fill_in
            if SAMPLE_COLUMNS[j - 1] == "Collection date":
                c.number_format = "yyyy-mm-dd"
            elif SAMPLE_COLUMNS[j - 1] in (REPLICATE_COLUMN, "Fraction"):
                c.number_format = "@"
    ws.freeze_panes = "B2"
    hw = wb.create_sheet("How to fill")
    hw.sheet_view.showGridLines = False
    hw.column_dimensions["A"].width = 3
    hw.column_dimensions["B"].width = 100
    lines = [
        ("Sample sheet", f_title),
        (f"ea_irms {VERSION}. For isoharness.org and the ea_irms workbook.", f_note),
        ("", f_sm),
        ("How to fill it", f_h1),
        ("1. On Samples, each row is one sample, and Sample ID holds the sample's name as the sequence wrote it "
         f"(Identifier 1). This file came with the {len(names)} sample name{'s' if len(names) != 1 else ''} of "
         "the runs on the page; add rows for other samples if you like.", f_sm),
        ("2. Fill in the yellow columns that apply and leave the rest empty. A column left empty for every "
         "sample is not written into the results. Add columns of your own after Replicate of, with a name in "
         "row 1; they are written into the results too.", f_sm),
        ("3. Choose this file as the sample sheet on the page. It stays in your browser and goes with each "
         "download; the server reads it for that request and keeps nothing.", f_sm),
        ("", f_sm),
        ("How samples are matched", f_h1),
        ("A sample matches a row when its name equals the Sample ID, or begins with it followed by a space, a "
         "hyphen, or another character that is neither a letter nor a digit. So CI 3-2 matches CI 3-2-S and "
         "CI 3-2 H but not CI 3-20, and when several IDs fit, the longest wins. Upper and lower case, spaces at either "
         "end, and double spaces count as the same. When no ID fits as written, the spaces are taken out of the "
         "name and the IDs and they are compared again, so PD1-H matches PD 1, and the results say so. One row can "
         "describe several analyses, such as the fractions of one colony. Standards and blanks are not matched.",
         f_sm),
        ("", f_sm),
        ("Replicates", f_h1),
        ("Within one run, samples that match the sheet and share a name, or share a Replicate of, are compared "
         "as replicates. Write in Replicate of the name of the analysis a row repeats, such as "
         "CI 3-2-H in the row of CI 3-2-H rep, or one name for all of them. When their range is larger than the "
         "critical range of ISO 5725-6, f(n) times the repeatability of one analysis (2.77 for two analyses, "
         "which the standard rounds to 2.8), each of them is marked for a rerun of that element. Analyses that "
         "only share a shorter ID, such as CI 3-2-H and CI 3-2-S, are not replicates.", f_sm),
        ("", f_sm),
        ("Your own file", f_h1),
        ("The page also reads your lab's own sheet, .xlsx or .csv, with one row per sample and a header row. "
         "Choose its ID column on the page if it is not the one found.", f_sm),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = hw.cell(i, 2, text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
