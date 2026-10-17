"""The rows of Fits, Setup, and QC and the columns of Work that the workbook's formulas read, with the label of
each, for tests/test_workbook_refs.py (finding R-27 of the 2026-10-14 review).

The workbook writer (ea_irms.workbook) puts most rows of Fits, Setup, and QC at fixed numbers, and formulas on
every sheet read them by number, so a row moved without every formula that reads it would break the workbook
with no test noticing. This script builds the 2026-09-15 test run's workbook with the package defaults and a
one-row sample sheet, so that the Replicates and Metadata sheets and the sample columns of Results carry
formulas too, and reads every formula on every sheet and the data ranges of every chart. It collects each
reference to a row of Fits, Setup, or QC, absolute or relative, from another sheet or from the same one, a
range counting every row it spans, and each reference to a column of Work. It writes
tests/data/workbook_refs.json: for each row referenced, the text of its column A (the row's label), and for
each Work column referenced, its header in row 5, in the order of the sheet. LibreOffice is not needed, since
only the formulas are read.

The label of a row is the text in its column A when that is a constant, neither a formula nor a yellow input
cell. A referenced row without such a label but with something in it, such as a data row of a table (the
standards table on Setup, the uncertainty tables on Fits), is recorded by the nearest label above it in column
A and how far below that label it lies, written "Material (+3)" for the third row under the label Material,
so that a table moved by a row is caught too. A referenced row with nothing in it is an error, and so is a
referenced Work column without a header; the script names the formula and writes nothing.

    python scripts/workbook_refs.py            writes tests/data/workbook_refs.json
    python scripts/workbook_refs.py --check    compares a fresh workbook with the file, as the test does

Regenerating the file is a deliberate act, after a change to the layout of Fits, Setup, QC, or Work or to the
wording of a label: read the diff, check that every formula that read a row or column that moved now reads it
where it went, and commit the file with the change. A file regenerated only to make the test pass defeats it.
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                                      # noqa: E402
from openpyxl.formula.tokenizer import Token, Tokenizer                 # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from openpyxl.utils.cell import range_boundaries                        # noqa: E402

RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
FILE = ROOT / "tests" / "data" / "workbook_refs.json"
ROW_SHEETS = ("Fits", "QC", "Setup")
WORK = "Work"
WORK_HEADER_ROW = 5
INPUT_FILL = "FFF2CC"                  # the yellow of the cells a student types in (ea_irms.workbook, FILL_IN)
# One row naming every Sample position of the test run (all are named Random), so that the run has a sample
# sheet, a replicate group, and the sample columns on Results
SAMPLE_SHEET = "Sample ID,Project,Site\nRandom,Test run,Boston College\n"
ABOUT = ("The rows of Fits, Setup, and QC and the columns of Work that the formulas of the workbook read, with the "
         "label of each row (column A; a row with none is named by the nearest label above it, as 'Material (+3)') "
         "and the header of each column (row 5), from the 2026-09-15 test run with the defaults and a sample sheet. "
         "tests/test_workbook_refs.py fails when a formula reads a row or column whose label is not the one here. "
         "Written by python scripts/workbook_refs.py: regenerate it only after a deliberate change to the layout, "
         "and read the diff before committing it.")


def build(folder) -> Path:
    """The test run's workbook with the package defaults and the one-row sample sheet, formulas only."""
    from ea_irms.metadata import read_sample_sheet
    from ea_irms.pipeline import Config
    from ea_irms.workbook import build_workbook
    out = Path(folder) / "workbook_refs.xlsx"
    build_workbook(RUN, out, Config(), run_date="2026-09-15",
                   samples=read_sample_sheet(SAMPLE_SHEET.encode(), "samples.csv"))
    return out


def _is_input(cell) -> bool:
    rgb = getattr(getattr(cell.fill, "fgColor", None), "rgb", None)
    return isinstance(rgb, str) and rgb.upper().endswith(INPUT_FILL)


def _label(ws, row):
    """Column A's text at ``row`` when it is a label: a constant, neither a formula nor an input cell."""
    c = ws.cell(row, 1)
    v = c.value
    if isinstance(v, str) and v.strip() and not v.startswith("=") and not _is_input(c):
        return v
    return None


def _has_content(ws, row) -> bool:
    return any(c.value not in (None, "") or _is_input(c) for c in ws[row])


def row_label(ws, row):
    """The label a referenced row is recorded by, or None for an empty row (see the module docstring)."""
    own = _label(ws, row)
    if own is not None:
        return own
    if not _has_content(ws, row):
        return None
    for k in range(1, row):
        above = _label(ws, row - k)
        if above is not None:
            return f"{above} (+{k})"
    return None


def _refs(formula):
    """The references in one formula, as (sheet or None, the reference as written, its range text)."""
    for tok in Tokenizer(formula).items:
        if tok.type != Token.OPERAND or tok.subtype != Token.RANGE:
            continue
        sheet, _, rng = tok.value.rpartition("!")
        yield (sheet.strip("'") or None), tok.value, rng


def collect(book):
    """Every reference to a row of Fits, Setup, or QC and to a column of Work in ``book`` (a path, or a workbook
    opened with its formulas). Returns {"rows": {(sheet, row): [(cell, ref), ...]}, "columns": {column letter:
    [(cell, ref), ...]}, "book": the workbook}, each reference named by the cell whose formula holds it."""
    wb = load_workbook(book) if isinstance(book, (str, Path)) else book
    rows, columns = {}, {}

    def note(where, own_sheet, ref_sheet, ref, rng):
        sheet = ref_sheet or own_sheet
        if sheet not in ROW_SHEETS and sheet != WORK:
            return
        try:
            c0, r0, c1, r1 = range_boundaries(rng.replace("$", ""))
        except ValueError:
            raise ValueError(f"{where} holds {ref!r}, which is not a cell reference") from None
        if sheet in ROW_SHEETS:
            if r0 is None:
                raise ValueError(f"{where} reads whole columns of {sheet} ({ref}), which no row label can check")
            for r in range(r0, r1 + 1):
                rows.setdefault((sheet, r), []).append((where, ref))
        else:
            if c0 is None:
                raise ValueError(f"{where} reads whole rows of Work ({ref}), which no column header can check")
            for c in range(c0, c1 + 1):
                columns.setdefault(get_column_letter(c), []).append((where, ref))

    for ws in wb.worksheets:
        for line in ws.iter_rows():
            for cell in line:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    for ref_sheet, ref, rng in _refs(v):
                        note(f"{ws.title}!{cell.coordinate}", ws.title, ref_sheet, ref, rng)
        for k, chart in enumerate(getattr(ws, "_charts", []), start=1):
            for s in chart.series:
                for part in (getattr(s, "xVal", None), getattr(s, "yVal", None), getattr(s, "val", None)):
                    f = getattr(getattr(part, "numRef", None), "f", None)
                    if f:
                        for ref_sheet, ref, rng in _refs("=" + f):
                            note(f"{ws.title} chart {k}", ws.title, ref_sheet, ref, rng)
    return {"rows": rows, "columns": columns, "book": wb}


def labels(found) -> dict:
    """The file's content for the references ``found`` (collect): each row's label and each column's header.
    Raises ValueError naming the formula when a referenced row is empty or a Work column has no header."""
    wb = found["book"]
    out = {"about": ABOUT, "rows": {s: {} for s in ROW_SHEETS}, "work_columns": {}}
    bad = []
    for (sheet, row), where in sorted(found["rows"].items(), key=lambda kv: (kv[0][0], kv[0][1])):
        lab = row_label(wb[sheet], row)
        if lab is None:
            bad.append(f"{where[0][0]} reads {where[0][1]}, and {sheet} row {row} is empty")
        out["rows"][sheet][str(row)] = lab
    for col, where in sorted(found["columns"].items(), key=lambda kv: column_index_from_string(kv[0])):
        head = wb[WORK].cell(WORK_HEADER_ROW, column_index_from_string(col)).value
        if head in (None, ""):
            bad.append(f"{where[0][0]} reads {where[0][1]}, and Work column {col} has no header in row 5")
        out["work_columns"][col] = head
    if bad:
        raise ValueError("; ".join(bad))
    return out


def _who(where) -> str:
    first = f"{where[0][0]} reads {where[0][1]}"
    return first + (f" (and {len(where) - 1} more formula{'s' if len(where) > 2 else ''})" if len(where) > 1 else "")


def compare(found, recorded: dict) -> list[str]:
    """What differs between the references ``found`` in a workbook and the ``recorded`` file, one message each,
    naming a formula that holds the reference; empty when they agree."""
    wb = found["book"]
    problems = []
    rec_rows = recorded.get("rows", {})
    for (sheet, row), where in sorted(found["rows"].items(), key=lambda kv: (kv[0][0], kv[0][1])):
        now = row_label(wb[sheet], row)
        if now is None:
            problems.append(f"{_who(where)}: {sheet} row {row} is empty")
            continue
        was = rec_rows.get(sheet, {}).get(str(row))
        if was is None:
            problems.append(f"{_who(where)}: {sheet} row {row} ({now!r}) is not in {FILE.name}; after a deliberate "
                            "change to the layout, regenerate it with python scripts/workbook_refs.py")
        elif was != now:
            problems.append(f"{_who(where)}: {sheet} row {row} is {now!r} where {FILE.name} has {was!r}")
    rec_cols = recorded.get("work_columns", {})
    for col, where in sorted(found["columns"].items(), key=lambda kv: column_index_from_string(kv[0])):
        now = wb[WORK].cell(WORK_HEADER_ROW, column_index_from_string(col)).value
        was = rec_cols.get(col)
        if now in (None, ""):
            problems.append(f"{_who(where)}: Work column {col} has no header in row 5")
        elif was is None:
            problems.append(f"{_who(where)}: Work column {col} ({now!r}) is not in {FILE.name}; after a deliberate "
                            "change to the layout, regenerate it with python scripts/workbook_refs.py")
        elif was != now:
            problems.append(f"{_who(where)}: Work column {col} is headed {now!r} where {FILE.name} has {was!r}")
    return problems


def stale(found, recorded: dict) -> list[str]:
    """Entries of the ``recorded`` file that no formula reads any longer."""
    out = [f"{FILE.name} holds {sheet} row {row}, which no formula reads"
           for sheet, held in recorded.get("rows", {}).items() for row in held
           if (sheet, int(row)) not in found["rows"]]
    out += [f"{FILE.name} holds Work column {col}, which no formula reads"
            for col in recorded.get("work_columns", {}) if col not in found["columns"]]
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--check", action="store_true", help="compare a fresh workbook with the file; write nothing")
    args = ap.parse_args(argv)
    with tempfile.TemporaryDirectory() as tmp:
        found = collect(build(tmp))
        if args.check:
            recorded = json.loads(FILE.read_text(encoding="utf-8"))
            problems = compare(found, recorded) + stale(found, recorded)
            for p in problems:
                print(p)
            print(f"{len(found['rows'])} rows and {len(found['columns'])} Work columns read by formulas; "
                  f"{len(problems)} problem{'' if len(problems) == 1 else 's'}")
            return 1 if problems else 0
        try:
            data = labels(found)
        except ValueError as exc:
            print(f"not written: {exc}")
            return 1
    FILE.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    n_rows = sum(len(v) for v in data["rows"].values())
    print(f"{FILE.relative_to(ROOT)}: {n_rows} rows of Fits, QC, and Setup, and {len(data['work_columns'])} "
          "Work columns")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
