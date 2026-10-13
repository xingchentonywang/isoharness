"""Check a running isoharness service end to end on one export.

    python scripts/smoke.py http://localhost:8000 tests/data/20260915_raw.csv

Waits for /api/health, then asks for the corrected values and for the workbook,
and checks that the workbook came back with its formulas computed and that its
Results sheet holds the values /api/process returned (within 1e-9 relative).
Needs only the standard library and openpyxl. Exits 1 on any failure.
"""
from __future__ import annotations

import io
import json
import sys
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

import openpyxl

# Results columns of ea_irms 1.4.3: each value followed by its U, the u columns after Flags.
RESULTS = {5: "pct_c_reported", 6: "U_pct_c", 7: "pct_n_reported", 8: "U_pct_n", 9: "cn_molar_reported",
           10: "U_cn_molar", 11: "cn_mass_reported", 12: "U_cn_mass", 13: "d13c_final", 14: "U_d13c",
           15: "d15n_final", 16: "U_d15n", 18: "u_pct_c", 19: "u_pct_n", 20: "u_cn_molar", 21: "u_cn_mass",
           22: "u_d13c", 23: "u_d15n"}
FLAGS = 17


def post(url, path: Path, fields=None, timeout=300):
    boundary = uuid.uuid4().hex
    parts = []
    for k, v in (fields or {}).items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
                 f"Content-Type: application/octet-stream\r\n\r\n".encode() + path.read_bytes() + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(url, data=b"".join(parts), method="POST",
                                 headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    t0 = time.perf_counter()
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read(), dict(r.headers), time.perf_counter() - t0


def wait_healthy(base, seconds=180):
    end = time.time() + seconds
    while time.time() < end:
        try:
            with urllib.request.urlopen(f"{base}/api/health", timeout=5) as r:
                if json.loads(r.read()) == {"status": "ok"}:
                    return
        except (urllib.error.URLError, ConnectionError, TimeoutError, json.JSONDecodeError):
            pass
        time.sleep(2)
    raise SystemExit(f"{base} did not become healthy within {seconds} s")


def main(base, export):
    base = base.rstrip("/")
    export = Path(export)
    wait_healthy(base)
    with urllib.request.urlopen(f"{base}/api/version", timeout=10) as r:
        print("version", json.loads(r.read()))

    body, _, t_process = post(f"{base}/api/process", export)
    data = json.loads(body)
    positions = data["positions"]
    print(f"process: {len(positions)} positions in {t_process:.2f} s")

    blob, headers, t_book = post(f"{base}/api/workbook", export)
    computed = headers.get("X-Isoharness-Values") or headers.get("x-isoharness-values")
    print(f"workbook: {len(blob)} bytes in {t_book:.2f} s, values {computed}")
    problems = []
    if computed != "computed":
        problems.append(f"workbook values are {computed!r}, not computed")
    rs = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)["Results"]
    compared = 0
    for p in positions:
        for col, key in RESULTS.items():
            mine, theirs = p.get(key), rs.cell(4 + p["row"], col).value
            theirs = theirs if isinstance(theirs, (int, float)) and not isinstance(theirs, bool) else None
            if (mine is None) != (theirs is None):
                problems.append((p["row"], key, mine, theirs))
            elif mine is not None:
                compared += 1
                if abs(mine - theirs) > 1e-9 * max(1.0, abs(mine)):
                    problems.append((p["row"], key, mine, theirs))
        flags = {t.strip() for t in (rs.cell(4 + p["row"], FLAGS).value or "").split(";") if t.strip()}
        if flags != set(p["flags"]):
            problems.append((p["row"], "flags", sorted(flags ^ set(p["flags"]))))
    print(f"workbook against /api/process: {compared} values and {len(positions)} flag lists compared, "
          f"{len(problems)} problems")
    for item in problems[:10]:
        print("  ", item)
    return 1 if problems or compared == 0 else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    raise SystemExit(main(sys.argv[1], sys.argv[2]))
