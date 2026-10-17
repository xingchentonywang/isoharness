"""The engine check before a release: on every reference run, the API's values
must equal its workbook, and the workbook must equal the one delivered.

    python scripts/reference_check.py MANIFEST.json [--url http://localhost:8000]

The reference runs are lab data and stay out of this public repository. The
manifest names them, with paths relative to the manifest's own folder:

    {"runs": [{"name": "PanamaSids2",
               "export": "runs/PanamaSids2_09232026_RAW.xls",
               "config": "configs/cfg_PanamaSids2.json",          (optional)
               "delivered": "delivered/09232026_PanamaSids2.xlsx"} (optional)]}

For each run it asks the API (at --url, or in this process when --url is left
out) for the corrected values and for the workbook, with the run name and date
read from the delivered workbook's Setup sheet, and checks three things:

1. Values. Every position value, flag list and fit result /api/process returns
   equals the workbook /api/workbook returns, after LibreOffice has computed it
   (1e-9 relative), with no formula error and no stray row. From ea_irms 1.4.15
   that includes the floor the size series sets and each step of its walk, the
   blank below detection, the drift monitors' SD against the repeatability in
   use, each correction's Kragten term at every position and its largest, the
   acetanilide's whole-run check, the corrections fitted but not applied, and
   the Rerun sheet's problems of the whole run, unapplied list and samples.
2. Engine. /api/process returns what ea_irms returns when called directly here,
   value for value (only when ea_irms imports).
3. Delivered. The workbook equals the delivered one in every cell: formulas and
   text exactly, computed numbers to 1e-9 relative. A workbook delivered by an
   earlier ea_irms version differs by design where that version changed the
   layout (1.4.3 moved the Results columns and added the uncertainty of %N, %C
   and C:N; 1.4.4 added the rerun mark, its reason and a suggested mass after
   Flags; 1.4.5 put each value's 1 sigma beside it and dropped the u columns;
   1.4.15 put the corrections fitted but not applied on the Rerun sheet and its
   samples from row 16), so compare against workbooks built by the same version.

Exits 1 on any mismatch.
"""
from __future__ import annotations

import argparse
import io
import json
import math
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

TOL = 1e-9
# Work sheet columns by letter, and the 1.4.2 columns found by their row-5 header
WORK = {"d15n_final": "AA", "d13c_final": "AF", "pct_n_measured": "AI", "pct_c_measured": "AJ",
        "cn_molar": "AK", "cn_mass": "AL", "ref_bgd_n": "BI", "settling_fraction": "BJ",
        "ref_pulse_d15n": "BT", "ref_pulse_d13c": "BW", "t_drift_n": "CB", "t_drift_c": "CC",
        "u_d15n": "CL", "U_d15n": "CM", "u_d13c": "CN", "U_d13c": "CO",
        "area_n": "L", "d15n": "O", "area_c": "P", "d13c": "R", "bgd_n": "N",
        "d15n_blank": "W", "d15n_lin": "X", "d15n_drift": "Z", "d13c_blank": "AB", "d13c_lin": "AC",
        "d13c_drift": "AE", "ug_n": "AG", "ug_c": "AH"}
WORK_BY_HEADER = {"elem_bias_n": "elem bias N", "elem_bias_c": "elem bias C",
                  "elem_factor_n": "elem factor N", "elem_factor_c": "elem factor C",
                  "pct_n_reported": "%N corrected", "pct_c_reported": "%C corrected",
                  "cn_molar_reported": "C:N molar corrected", "cn_mass_reported": "C:N mass corrected",
                  "pulse_ampl_c": "pulse ampl", "ampl_c": "Ampl 44"}
# Results columns of ea_irms 1.4.5: each value followed by its 1 sigma, Flags in Q, the rerun mark in R to T,
# then the elemental drift factors and %C and %N as measured in U to X.
RESULTS = {"pct_c_reported": 5, "u_pct_c": 6, "pct_n_reported": 7, "u_pct_n": 8, "cn_molar_reported": 9,
           "u_cn_molar": 10, "cn_mass_reported": 11, "u_cn_mass": 12, "elem_factor_n": 21, "elem_factor_c": 22,
           "pct_c_measured": 23, "pct_n_measured": 24}
FLAGS_COL = "Q"
RERUN_TEXT = {18: "rerun", 19: "rerun_why", 20: "rerun_mass"}
FITS = {"B20": ("lin_n", "applied"), "B95": ("lin_n", "fit", "p"), "B96": ("lin_n", "in_force"),
        "B28": ("lin_c", "applied"), "B102": ("lin_c", "fit", "p"), "B103": ("lin_c", "in_force"),
        "B39": ("drift_n", "applied"), "B114": ("drift_n", "fit", "p"), "B115": ("drift_n", "in_force"),
        "B51": ("drift_c", "applied"), "B127": ("drift_c", "fit", "p"), "B128": ("drift_c", "in_force"),
        "B56": ("cal_n", "slope"), "B57": ("cal_n", "intercept"), "B63": ("cal_c", "slope"),
        "B64": ("cal_c", "intercept"), "B70": ("k_n", "slope"), "B71": ("k_n", "intercept"),
        "B76": ("k_c", "slope"), "B77": ("k_c", "intercept"),
        "B135": ("unc_n", "pooled_sd"), "B155": ("unc_c", "pooled_sd"),
        "B138": ("cal_n", "acc_min"), "B159": ("cal_c", "acc_max"),
        "B201": ("ref_gas_n", "n"), "B205": ("ref_gas_n", "fit", "slope"), "B206": ("ref_gas_n", "fit", "p"),
        "B208": ("ref_gas_c", "n"), "B213": ("ref_gas_c", "fit", "p"),
        "B217": ("checks", "n"), "B97": ("size_limits", "max_n"),
        "B85": ("settling", "excess"), "B86": ("settling", "settled_from"),
        "B229": ("carryover", "oversized"), "B230": ("carryover", "followed"),
        "B234": ("settling", "level"), "B235": ("settling", "excess_level"),
        "B237": ("ref_pulse_check", "n_designated"), "B238": ("ref_pulse_check", "mean_ampl"),
        "B239": ("ref_pulse_check", "mean_d13c"), "B240": ("ref_pulse_check", "active"),
        "B244": ("elem_n", "n"), "B245": ("elem_n", "max_bias_pct"), "B246": ("elem_n", "in_force"),
        "B247": ("elem_c", "n"), "B248": ("elem_c", "max_bias_pct"), "B249": ("elem_c", "in_force"),
        # ea_irms 1.4.13: the anchors read as one material (finding W-13), and the elemental drift correction
        # against the anchors (empty where the correction is not in force)
        "B370": ("anchor_material", "limits", "ratio"), "B371": ("anchor_material", "limits", "spread"),
        "B389": ("anchor_material", "level"),
        "B394": ("elem_check", "n", "n"), "B395": ("elem_check", "n", "rms_measured"),
        "B396": ("elem_check", "n", "rms_reported"), "B397": ("elem_check", "n", "added"),
        "B398": ("elem_check", "n", "level"),
        "B399": ("elem_check", "c", "n"), "B400": ("elem_check", "c", "rms_measured"),
        "B401": ("elem_check", "c", "rms_reported"), "B402": ("elem_check", "c", "added"),
        "B403": ("elem_check", "c", "level")}
# ea_irms 1.4.13: the anchors of each row of the standards table, one Fits row each from 373 (16 rows, as
# ea_irms.workbook.NSTD), compared row by row in check_values because a row holds a material only where the
# run has anchors of it; a row without one reads 0 anchors and no level
ANCHOR_TOP, ANCHOR_ROWS = 373, 16
ANCHOR_COLS = {"B": "n", "C": "n_ratio", "D": "ratio_min", "E": "ratio_max", "F": "ratio_x", "G": "d15n_min",
               "H": "d15n_max", "I": "d15n_spread", "J": "d13c_min", "K": "d13c_max", "L": "d13c_spread",
               "M": "level"}
# ea_irms 1.4.15: the size limits the flags use, the floor the size series sets and where it comes from (None in
# the engine is an empty cell), the blank below detection, the drift monitors after every correction against the
# repeatability in use, and the acetanilide's whole-run check
FITS.update({"B81": ("size_limits", "min_n"), "B82": ("size_limits", "min_c"),
             "B104": ("size_limits", "max_c"),
             "B407": ("blank_n", "capsules"), "B408": ("blank_n", "below_detection"),
             "B409": ("blank_c", "below_detection"),
             "B413": ("floor_n", "smallest"), "B415": ("floor_n", "floor"), "B416": ("size_limits", "min_n_source"),
             "B419": ("floor_c", "smallest"), "B421": ("floor_c", "floor"), "B422": ("size_limits", "min_c_source"),
             "B484": ("rerun", "unapplied_text")})
SIGMA_TOP = {"n": 425, "c": 439}         # + 0 to 6: the monitors' n, SD, repeatability, df, p, larger, used
ACE_TOP = {"n": 454, "c": 463}           # + 0 to 6: n, n_series, SD, basis, df, p, over
for _el in ("n", "c"):
    FITS.update({f"B{SIGMA_TOP[_el] + k}": (f"unc_{_el}", "monitors", key)
                 for k, key in enumerate(("n", "sd", "repeatability", "df", "p", "larger", "used"))})
    FITS.update({f"B{ACE_TOP[_el] + k}": (f"ace_{_el}", key)
                 for k, key in enumerate(("n", "n_series", "sd", "basis", "df", "p", "over"))})
# the row of each correction in the 1σ table (Fits, from SIGMA_TOP + 10): term in 1σ (C), fitted but not applied
# (D), the largest term (I), and the samples moved, the smallest and the largest move (J to L); the Work column
# of each correction's shift, headed "kb shift N" and so on
CORRECTIONS = (("blank", "kb"), ("linearity", "kl"), ("drift", "kd"))
FLOOR_N_TOP = 488                        # the walk up the nitrogen series; the carbon one follows it
FLOOR_COLS = {"B": "row", "C": "area", "G": "n_line", "N": "line_sd", "O": "limit", "P": "resid"}
UNAPPLIED_TOP, UNAPPLIED_COUNT = 477, 483


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def near(a, b):
    return abs(a - b) <= TOL * max(1.0, abs(a))


def dig(node, path):
    for k in path:
        node = node.get(k) if isinstance(node, dict) else None
        if node is None:
            return None
    return node


class Client:
    """POST to a running service, or to the app in this process."""

    def __init__(self, url):
        if url:
            import httpx
            self.http = httpx.Client(base_url=url.rstrip("/"), timeout=600)
        else:
            from fastapi.testclient import TestClient
            from isoharness.app import app
            self.http = TestClient(app)

    def post(self, path, export: Path, **form):
        r = self.http.post(path, files={"file": (export.name, export.read_bytes())},
                           data={k: v for k, v in form.items() if v is not None})
        if r.status_code != 200:
            raise RuntimeError(f"{path} returned {r.status_code}: {r.text[:300]}")
        return r


def check_values(body, wb):
    """1. The API's values against the workbook it returned."""
    positions, fits = body["positions"], json.loads(json.dumps(body["fits"]))
    for el in ("elem_n", "elem_c"):
        mb = fits[el]["max_bias"]
        fits[el]["max_bias_pct"] = None if mb is None else mb * 100
    wk, ft, rs = wb["Work"], wb["Fits"], wb["Results"]
    head = {str(wk.cell(5, c).value): get_column_letter(c) for c in range(1, wk.max_column + 1)}
    cols = dict(WORK)
    cols.update({attr: head[name] for attr, name in WORK_BY_HEADER.items()})
    feeds = {c for c in range(1, wk.max_column + 1) if str(wk.cell(5, c).value or "").startswith("ch ")}
    bad, n = [], 0
    errors = [(s.title, c.coordinate, c.value) for s in wb for row in s.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper()
              and not (s.title == "Work" and c.column in feeds and c.value == "#N/A")
              and not (s.title == "Raw" and str(s.cell(1, c.column).value or "").startswith("chart ")
                       and c.value == "#N/A")]   # 1.4.10: the pulses chart feeds, NA() by design
    for p in positions:
        for attr, col in cols.items():
            mine, theirs = p.get(attr), num(wk[f"{col}{5 + p['row']}"].value)
            n += 1
            if (mine is None) != (theirs is None) or (mine is not None and not near(mine, theirs)):
                bad.append(("Work", attr, p["row"], mine, theirs))
        for attr, col in RESULTS.items():
            mine, theirs = p.get(attr), num(rs.cell(4 + p["row"], col).value)
            n += 1
            if (mine is None) != (theirs is None) or (mine is not None and not near(mine, theirs)):
                bad.append(("Results", attr, p["row"], mine, theirs))
        text = rs[f"{FLAGS_COL}{4 + p['row']}"].value or ""
        n += 1
        if [t.strip() for t in text.split(";") if t.strip()] != list(p["flags"]):     # in order (1.4.12)
            bad.append(("flags", p["row"], text, p["flags"]))
        # the rerun mark, its reason and the suggested mass (text, R to T), from 1.4.10 here too
        for col, attr in RERUN_TEXT.items():
            n += 1
            mine, theirs = p.get(attr) or "", rs.cell(4 + p["row"], col).value or ""
            if mine != theirs:
                bad.append(("Results", attr, p["row"], mine, theirs))
    for cell, path in FITS.items():
        mine, theirs = dig(fits, path), ft[cell].value
        n += 1
        if isinstance(mine, bool):
            mine = "Yes" if mine else "No"
        if isinstance(mine, float):
            if num(theirs) is None or not near(mine, theirs):
                bad.append(("Fits", cell, path, mine, theirs))
        elif mine != theirs:
            bad.append(("Fits", cell, path, mine, theirs))
    materials = {m["index"]: m for m in (fits.get("anchor_material") or {}).get("materials", [])}
    for k in range(ANCHOR_ROWS):
        m = materials.get(k)
        for col, key in ANCHOR_COLS.items():
            cell = f"{col}{ANCHOR_TOP + k}"
            mine, theirs = (m[key] if m else 0 if key in ("n", "n_ratio") else None), ft[cell].value
            n += 1
            if isinstance(mine, float):
                if num(theirs) is None or not near(mine, theirs):
                    bad.append(("Fits", cell, ("anchor_material", k, key), mine, theirs))
            elif mine != theirs:
                bad.append(("Fits", cell, ("anchor_material", k, key), mine, theirs))
    n += 1
    if ft["B356"].value != fits["rerun"]["run_level_text"]:
        bad.append(("Fits", "B356", "problems of the whole run", fits["rerun"]["run_level_text"], ft["B356"].value))
    last = max([125] + [5 + p["row"] for p in positions])
    extra = [r for r in range(6, last + 1) if wk[f"B{r}"].value not in (None, "")
             and r - 5 not in {p["row"] for p in positions}]
    if extra:
        bad.append(("extra Work rows", extra[:5]))
    m, more = check_1415(positions, fits, wb, head)          # ea_irms 1.4.15
    return n + m, bad + more, errors


def _same(mine, theirs):
    """The engine's value against a computed cell: numbers to TOL, booleans as Yes or No, None as empty."""
    if isinstance(mine, bool):
        mine = "Yes" if mine else "No"
    if mine is None or mine == "":
        return theirs in (None, "")
    if isinstance(mine, (int, float)):
        return num(theirs) is not None and near(mine, theirs)
    return mine == theirs


def check_1415(positions, fits, wb, head):
    """ea_irms 1.4.15: what check_values does not reach through FITS, against the engine's fits (the whole of F,
    as /api/process returns it) and its positions."""
    from ea_irms.workbook import RERUN_FIRST_ROW, RERUN_UNAPPLIED_ROW
    wk, ft, rr = wb["Work"], wb["Fits"], wb["Rerun"]
    n, bad = 0, []

    def cmp(what, mine, theirs):
        nonlocal n
        n += 1
        if not _same(mine, theirs):
            bad.append(what + (mine, theirs))
    # each correction's Kragten term at every position: the size of its shift (Work, "kb shift N" and so on)
    for p in positions:
        for el in ("n", "c"):
            for key, ch in CORRECTIONS:
                if key in (p.get(f"u_terms_{el}") or {}):
                    v = num(wk[f"{head[f'{ch} shift {el.upper()}']}{5 + p['row']}"].value)
                    cmp(("Work", f"{ch} shift {el.upper()}", p["row"]), p[f"u_terms_{el}"][key],
                        None if v is None else abs(v))
    # the 1σ table: a term in force, its largest; a correction fitted but not applied, how far it moves the samples
    unapplied = {(u["el"], u["correction"]): u for u in fits.get("unapplied", [])}
    for el in ("n", "c"):
        corr = fits[f"unc_{el}"]["corrections"]
        for k, (key, _) in enumerate(CORRECTIONS):
            r = SIGMA_TOP[el] + 10 + k
            u = unapplied.get((el, key))
            cmp(("Fits", f"C{r}", el, key, "term in 1σ"), key in corr, ft[f"C{r}"].value)
            cmp(("Fits", f"I{r}", el, key, "largest term"), corr.get(key), ft[f"I{r}"].value)
            cmp(("Fits", f"D{r}", el, key, "fitted, not applied"), u is not None, ft[f"D{r}"].value)
            for col, attr in (("J", "n"), ("K", "lo"), ("L", "hi")):
                cmp(("Fits", f"{col}{r}", el, key, attr), None if u is None else u[attr], ft[f"{col}{r}"].value)
    # the walk up the size series, step by step, and no step after the last
    tops = {"n": FLOOR_N_TOP}
    for r in range(FLOOR_N_TOP, ft.max_row + 1):
        if str(ft.cell(r, 1).value or "").startswith("Size series floor, carbon"):
            tops["c"] = r + 2
            break
    for el in ("n", "c"):
        fl = fits[f"floor_{el}"]
        top = tops.get(el)
        if top is None:
            bad.append(("Fits", "no table of the carbon series floor"))
            continue
        for k, step in enumerate(fl["steps"]):
            for col, key in FLOOR_COLS.items():
                cmp(("Fits", f"{col}{top + k}", el, "step", key), step[key], ft[f"{col}{top + k}"].value)
            cmp(("Fits", f"Q{top + k}", el, "tested"), 1, ft[f"Q{top + k}"].value)
            cmp(("Fits", f"R{top + k}", el, "moved"), 1 if step["moved"] else 0, ft[f"R{top + k}"].value)
        cmp(("Fits", f"Q{top + len(fl['steps'])}", el, "not tested"), 0, ft[f"Q{top + len(fl['steps'])}"].value)
    # the corrections fitted but not applied, on Fits in the engine's order and on the Rerun sheet, one a line
    items = fits["rerun"]["unapplied"]
    cmp(("Fits", f"B{UNAPPLIED_COUNT}", "unapplied"), len(items), ft[f"B{UNAPPLIED_COUNT}"].value)
    order = [(el, key) for el in ("n", "c") for key, _ in CORRECTIONS]
    texts = iter(items)
    for k, pair in enumerate(order):
        cmp(("Fits", f"B{UNAPPLIED_TOP + k}", "unapplied", pair),
            next(texts) if pair in unapplied else None, ft[f"B{UNAPPLIED_TOP + k}"].value)
    lines = [rr.cell(RERUN_UNAPPLIED_ROW + k, 1).value for k in range(6)]
    want = (list(items) + [None] * 6)[:6] if items else ["none"] + [None] * 5
    for k, (mine, theirs) in enumerate(zip(want, lines)):
        cmp(("Rerun", f"A{RERUN_UNAPPLIED_ROW + k}", "unapplied"), mine, theirs)
    # the Rerun sheet: the problems of the whole run, and every sample marked, in the order they ran
    cmp(("Rerun", "A5", "problems of the whole run"), fits["rerun"]["run_level_text"], rr["A5"].value)
    marked = [p for p in positions if p.get("rerun")]
    for k, p in enumerate(marked):
        row = RERUN_FIRST_ROW + k
        for col, attr in ((1, "row"), (5, "rerun"), (6, "rerun_why"), (7, "rerun_mass")):
            cmp(("Rerun", f"row {row}", attr), p.get(attr), rr.cell(row, col).value)
    cmp(("Rerun", f"row {RERUN_FIRST_ROW + len(marked)}", "after the last mark"), None,
        rr.cell(RERUN_FIRST_ROW + len(marked), 1).value)
    return n, bad


def check_engine(body, export: Path, config_text):
    """2. The API's positions against ea_irms called directly."""
    from ea_irms.parse import parse_export
    from ea_irms.pipeline import Config, config_from_dict, process, scope_to_run
    from isoharness.serialize import position_record
    cfg = config_from_dict(json.loads(config_text)) if config_text else Config()
    cfg, _ = scope_to_run(cfg, export.name)        # as the service applies it (ea_irms 1.4.10)
    positions, _ = process(parse_export(export), cfg)
    want = [json.loads(json.dumps(position_record(p))) for p in positions]
    got = body["positions"]
    n, bad = 0, []
    if len(want) != len(got):
        return 1, [("position count", len(got), len(want))]
    for g, w in zip(got, want):
        for k in w:
            n += 1
            a, b = g.get(k), w[k]
            if a != b and not (isinstance(a, float) and isinstance(b, float) and math.isnan(a) and math.isnan(b)):
                bad.append((w["row"], k, a, b))
    return n, bad


def check_delivered(blob, delivered: Path):
    """3. Every cell of the API's workbook against the delivered workbook."""
    n, bad = 0, []
    for data_only in (False, True):
        mine = openpyxl.load_workbook(io.BytesIO(blob), data_only=data_only)
        theirs = openpyxl.load_workbook(delivered, data_only=data_only)
        if mine.sheetnames != theirs.sheetnames:
            return 1, [("sheets", mine.sheetnames, theirs.sheetnames)]
        for name in mine.sheetnames:
            a, b = mine[name], theirs[name]
            rows, cols = max(a.max_row, b.max_row), max(a.max_column, b.max_column)
            for r in range(1, rows + 1):
                for c in range(1, cols + 1):
                    x, y = a.cell(r, c).value, b.cell(r, c).value
                    n += 1
                    if data_only and num(x) is not None and num(y) is not None:
                        if not near(x, y):
                            bad.append((name, a.cell(r, c).coordinate, "value", x, y))
                    elif x != y:
                        bad.append((name, a.cell(r, c).coordinate, "value" if data_only else "formula", x, y))
    return n, bad


def main(argv=None):
    ap = argparse.ArgumentParser(description="engine check over the reference runs")
    ap.add_argument("manifest")
    ap.add_argument("--url", help="a running service; default: the app in this process")
    args = ap.parse_args(argv)
    manifest = Path(args.manifest).resolve()
    runs = json.loads(manifest.read_text())["runs"]
    client = Client(args.url)
    failed = 0
    totals = [0, 0, 0]
    for run in runs:
        export = manifest.parent / run["export"]
        config_text = (manifest.parent / run["config"]).read_text() if run.get("config") else None
        delivered = manifest.parent / run["delivered"] if run.get("delivered") else None
        name = date = None
        if delivered:
            setup = openpyxl.load_workbook(delivered)["Setup"]
            name, date = setup["B5"].value, setup["B6"].value
        body = client.post("/api/process", export, config=config_text).json()
        r = client.post("/api/workbook", export, config=config_text, run_name=name, run_date=date)
        blob = r.content
        computed = r.headers.get("x-isoharness-values")
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
        n1, bad1, errors = check_values(body, wb)
        try:
            n2, bad2 = check_engine(body, export, config_text)
        except ImportError:
            n2, bad2 = 0, []
        n3, bad3 = check_delivered(blob, delivered) if delivered else (0, [])
        totals = [totals[0] + n1, totals[1] + n2, totals[2] + n3]
        ok = computed == "computed" and not (bad1 or bad2 or bad3 or errors)
        failed += not ok
        print(f"{run['name']}: {len(body['positions'])} positions, values {computed}; "
              f"values {n1} compared, {len(bad1)} mismatches, {len(errors)} formula errors; "
              f"engine {n2} compared, {len(bad2)} mismatches; "
              f"delivered {n3} cells compared, {len(bad3)} differ" + ("" if delivered else " (none given)"))
        for item in (bad1 + bad2 + bad3 + errors)[:8]:
            print("   ", item)
    print(f"{len(runs)} runs: values {totals[0]}, engine {totals[1]}, delivered cells {totals[2]} compared; "
          + ("ALL OK" if not failed else f"{failed} run(s) FAILED"))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
