"""End-to-end check of the page in Chromium, against a service started for it.

    python scripts/e2e_web.py [--shots DIR] [--keep DIR] [--url URL]

Starts ``python -m uvicorn isoharness.app:app`` on a free port from the repository
root (LibreOffice on the PATH computes the workbook), or with --url uses a service
already running there (the container image, say), opens the page in Chromium
(Python Playwright, with PLAYWRIGHT_BROWSERS_PATH set) and:

0. opens the home page, checks that it lists the three instruments with two of them
   planned, and follows its link to the EA-IRMS page;
1. uploads two files that are not exports and checks that the page shows the
   service's message for each, word for word;
2. uploads tests/data/20260915_raw.csv and checks the upload screen (the twelve
   columns found, the positions, the page's agreement with the engine);
3. downloads the Excel template of the lab configuration and checks that it reads
   as the package defaults; loads a lab configuration after the upload, as JSON and
   then as the Excel file the service writes for it, checks that the run is processed
   again with it and the decisions start from it, then goes back to the defaults;
4. visits all ten step screens, checking the instrument screen's charts and its
   settling line, each correction step's explanation with its worked example and the
   checks that step shows (isoharness 0.4.0; the blank's mass balance against the value
   after the step, the button that hides every explanation and the browser keeping that
   choice), the charts of each sample's change against kσ, and the quality-check
   summary's links back to every step, and changing decisions on the way: the blank switch,
   a drift monitor left out of the d15N fit, the drift model, a monitor left out of
   the elemental correction, and the d15N repeatability (the drift monitors' SD); and (finding R-29) that each
   figure's drawing is an image named for its number, what it plots against what, its points, and the selected
   one, that the figures' lines and gray marks reach 3:1 against the page, that the sequence's bars are at least
   24 px wide, that the current step, the selected position and the chosen filter say so, and that each button
   back from the quality checks names its check;
5. checks that every value the results screen shows (mass, d15N, 1σ, d13C, 1σ, %N,
   %C, C:N, and the flags) equals, to the shown precision, the results table the
   service returns for the same decisions, built here from the processed
   configuration without the page;
6. downloads the workbook, the results table and the configuration, checks their
   names (from Content-Disposition), the configuration's fields (the Excel file read
   back by the service), the table, and the workbook's computed Results sheet, and
   posts the downloaded configuration back to /api/results.csv to check that it
   reproduces the page's values;
7. adds two more runs at once (isoharness 0.2.0), copies of the test run with each
   sample named by its position, the second with each sample's N2 peak 0.5 permil
   higher; checks the batch list, that the first run kept its decisions when shown
   again, the samples measured in more than one run, and the combined results and
   table, each run's rows against the service's table for that run;
8. types run information (isoharness 0.3.0) and gives it to every run, downloads the
   sample sheet the page fills with the batch's sample names, chooses a sample sheet,
   checks how the samples matched and that choosing another ID column matches them
   again, and checks the sheet's columns and the run information in the combined
   results and in a run's workbook; then removes a run;
9. saves a screenshot of every screen at 1440 by 1000 (full page) and of the upload
   and results screens at 400 px wide;
10. loads the page afresh, checks the landing page at 400 px and 1440 px wide, follows
   the brand back to the home page and checks it at 400 px, and reads the example run
   through its button;
11. (finding W-11, with W-12) adds two copies of the test run at once, each sample named by its
   position, as S 20 in the first and S20 in the second; downloads the first run's configuration,
   switches its blank correction On, and chooses it for the second run alone; checks the batch list, that
   the second run's results follow that configuration (the service's table for it, and its blank switch)
   and the first run's do not, that its results table, configuration, and workbook are its own, that the
   combined results table equals each run's own table, that the combined Excel file names both
   configurations and pairs the names that differ only in their spaces (Name matched), as the page does;
   that a configuration naming another run gives the second run a note of its own (R-08); goes back to the
   batch's configuration; and checks the sequence at 400 px wide.

Any console error, page error, failed request, or HTTP error the test did not
cause on purpose fails the run. Prints one result line; exits 1 on failure.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import uuid
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "tests" / "data" / "20260915_raw.csv"
SCREENS = ["inst", "setup", "blank", "lin", "drift", "cal", "elem", "unc", "qc", "results"]
TITLES = {"inst": "Instrument performance", "setup": "Standards and roles", "blank": "Blank", "lin": "Linearity", "drift": "Drift",
          "cal": "Calibration", "elem": "Elemental content", "unc": "Uncertainty",
          "qc": "Quality checks", "results": "Results"}
# The step names the Back and Continue buttons use, where they differ from the headings.
BUTTONS = dict(TITLES, inst="Instrument", setup="Setup")
# Results screen column -> results table column, and the decimals the page shows.
SHOWN = {"amt": ("Mass (mg)", 3), "dN": ("d15N (permil)", 2), "UN": ("1sigma(d15N)", 2),
         "dC": ("d13C (permil)", 2), "UC": ("1sigma(d13C)", 2), "pN": ("%N", 2), "pC": ("%C", 2),
         "cn": ("C:N molar", 1)}
# Workbook Results sheet column -> results table column (as scripts/smoke.py reads it).
BOOK = {5: "%C", 6: "1sigma(%C)", 7: "%N", 8: "1sigma(%N)", 9: "C:N molar", 10: "1sigma(C:N molar)",
        11: "C:N mass", 12: "1sigma(C:N mass)", 13: "d13C (permil)", 14: "1sigma(d13C)", 15: "d15N (permil)",
        16: "1sigma(d15N)", 21: "Elemental drift factor N", 22: "Elemental drift factor C", 23: "%C as measured",
        24: "%N as measured"}
# The steps that explain their correction, and the checks each step shows, in the order the Quality
# checks step lists them after the instrument's five (isoharness 0.4.0).
EXPLAINED = ("blank", "lin", "drift", "cal", "elem", "unc")
MOVED = {"lin": ["Effect of the linearity choice", "Sample size against the size series"],
         "drift": ["Effect of the drift choice", "Drift monitors agree"],
         "cal": ["Calibration line", "Anchors bracket the samples", "Check standard"],    # the first from 1.4.12
         "elem": ["Elemental content"]}
# Two checks the page shows when the engine sends their data (the anchors' material, and the elemental
# correction judged at the anchors); without it the page shows neither, so the steps' lists above hold either way.
ENGINE_CHECKS = {"cal": ("Anchors read as one material each", "C/N peak-area ratio"),
                 "elem": ("Elemental correction at the anchors", "the elemental gate (")}
# The Wong colors of the data marks, which keep their hue whatever their contrast (finding R-29), and white.
WONG = {"#0072B2", "#D55E00", "#E69F00", "#56B4E9", "#CC79A7", "#009E73", "#F0E442", "#FFFFFF"}
# The decisions the test makes, and the drift monitors it leaves out.
EXCL_N, EXCL_E = 18, 55


class Failed(Exception):
    pass


def check(cond, what):
    if not cond:
        raise Failed(what)


def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def post(base, path, file: Path, fields=None, name=None):
    """POST a multipart form; returns (status, headers, body bytes)."""
    boundary = uuid.uuid4().hex
    parts = []
    for k, v in (fields or {}).items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name or file.name}"\r\n'
                 "Content-Type: application/octet-stream\r\n\r\n".encode() + file.read_bytes() + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(base + path, data=b"".join(parts), method="POST",
                                 headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    try:
        with urllib.request.urlopen(req, timeout=300) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def get(base, path):
    """GET; returns (status, headers, body bytes)."""
    try:
        with urllib.request.urlopen(base + path, timeout=60) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def post_json(base, path, obj):
    """POST a JSON body; returns (status, headers, body bytes)."""
    req = urllib.request.Request(base + path, data=json.dumps(obj).encode(), method="POST",
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def table_diff(a_text, b_text, rel=1e-12):
    """Cells of two results tables that differ, text exactly and numbers by more than rel, and the
    largest relative difference between numbers. An Excel file keeps 16 significant digits, so a
    configuration read back from one may carry a value computed to more digits a few 1e-16 apart."""
    a, b = list(csv.reader(io.StringIO(a_text))), list(csv.reader(io.StringIO(b_text)))
    if len(a) != len(b):
        return [f"{len(a)} rows against {len(b)}"], 0.0
    out, worst = [], 0.0
    for i, (ra, rb) in enumerate(zip(a, b)):
        if len(ra) != len(rb):
            out.append(f"row {i}: {len(ra)} cells against {len(rb)}")
            continue
        for j, (x, y) in enumerate(zip(ra, rb)):
            if x == y:
                continue
            try:
                fx, fy = float(x), float(y)
            except ValueError:
                out.append(f"row {i}, column {j}: {x!r} against {y!r}")
                continue
            d = abs(fx - fy) / max(abs(fx), 1e-300)
            worst = max(worst, d)
            if abs(fx - fy) > rel * max(1.0, abs(fx)):
                out.append(f"row {i}, column {j}: {x} against {y}")
    return out, worst


def rows_of(text):
    return {int(r["Row"]): r for r in csv.DictReader(io.StringIO(text))}


def variant(path: Path, samples: set, shift=0.0, name="S{}") -> bytes:
    """The test run with each sample (its positions in samples) named S<position>, or by the pattern name, and
    its N2 peak's d15N moved by shift, as tests/test_batch.py builds it: the sample peak is the position's last
    mass-28 peak."""
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8"))))
    h = rows[0]
    col = {k: h.index(k) for k in ("Row", "Identifier 1", "Ampl  28", "Is Ref _", "d 15N/14N")}
    last28 = {r[col["Row"]].strip(): i for i, r in enumerate(rows[1:], start=1)
              if len(r) > col["Ampl  28"] and r[col["Ampl  28"]].strip()}
    for i, r in enumerate(rows[1:], start=1):
        row = r[col["Row"]].strip() if r else ""
        if row in samples:
            r[col["Identifier 1"]] = name.format(row)
            if shift and last28.get(row) == i and r[col["Is Ref _"]].strip() != "1":
                r[col["d 15N/14N"]] = repr(float(r[col["d 15N/14N"]]) + shift)
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def sheet_diff(rows, text, rel=1e-15):
    """Rows of an Excel sheet (values) against a results table's rows (CSV text, header left out):
    the cells that differ, and the largest relative difference between numbers."""
    table = list(csv.reader(io.StringIO(text)))[1:]
    if len(rows) != len(table):
        return [f"{len(rows)} rows against {len(table)}"], 0.0
    out, worst = [], 0.0
    for i, (ra, rb) in enumerate(zip(rows, table)):
        for j, (a, b) in enumerate(zip(ra, rb)):
            if b == "":
                if a not in (None, ""):
                    out.append(f"row {i}, column {j}: {a!r} against empty")
                continue
            if isinstance(a, str) or a is None:
                if a != b:
                    out.append(f"row {i}, column {j}: {a!r} against {b!r}")
                continue
            d = abs(a - float(b))
            worst = max(worst, d / max(abs(float(b)), 1e-300))
            if d > rel * max(1.0, abs(float(b))):
                out.append(f"row {i}, column {j}: {a!r} against {b}")
    return out, worst


def ftext(flag):
    """The page's wording of an engine flag (logic.js ftext)."""
    t = re.sub(r"(\d)x the blank", r"\1× the blank", flag.replace("d15N", "δ¹⁵N").replace("d13C", "δ¹³C"), count=1)
    return re.sub(r" by -(\d)", r" by −\1", t, count=1)


def js_fixed(value, d):
    """value.toFixed(d) as the page writes it, with its minus sign."""
    q = Decimal(abs(value)).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP)
    s = f"{q:.{d}f}"
    return ("−" if value < 0 and float(s) != 0 else "") + s


def agree(shown, value, d):
    """Whether a shown value equals the table's to the shown precision; and whether exactly as formatted."""
    if value in ("", None):
        return shown == "", shown == ""
    if shown == "":
        return False, False
    v = float(value)
    close = abs(float(shown.replace("−", "-")) - v) <= 0.5 * 10 ** -d + 1e-9
    return close, shown == js_fixed(v, d)


def compare(page_rows, table, label):
    """Every shown value of every sample against a results table; returns counts."""
    n = values = exact = 0
    for r in page_rows:
        row = table.get(int(r["r"]))
        check(row is not None, f"{label}: position {r['r']} is not in the results table")
        check(row["Type"] == "Sample", f"{label}: position {r['r']} is a {row['Type']}, not a sample")
        for key, (col, d) in SHOWN.items():
            ok, same = agree(r[key], row[col], d)
            check(ok, f"{label}: position {r['r']} {col}: page shows {r[key]!r}, table has {row[col]!r}")
            values += 1
            exact += same
        want = [ftext(f) for f in row["Flags"].split("; ") if f]
        check(r["flags"] == want, f"{label}: position {r['r']} flags: page {r['flags']}, table {want}")
        n += 1
    return n, values, exact


def sd(values):
    m = sum(values) / len(values)
    return math.sqrt(sum((v - m) ** 2 for v in values) / (len(values) - 1))


def book_cells(path, table, label="workbook"):
    """A downloaded workbook's computed Results sheet against a results table's rows, at every column of BOOK;
    returns how many values LibreOffice computed."""
    book = openpyxl.load_workbook(path, data_only=True)["Results"]
    cells = 0
    for row, rec in table.items():
        for c, col in BOOK.items():
            mine_v = rec[col]
            theirs = book.cell(4 + row, c).value
            theirs = theirs if isinstance(theirs, (int, float)) and not isinstance(theirs, bool) else None
            check((mine_v == "") == (theirs is None), f"{label} row {row} {col}: {theirs!r} against {mine_v!r}")
            if theirs is not None:
                cells += 1
                check(abs(float(mine_v) - theirs) <= 1e-9 * max(1.0, abs(theirs)),
                      f"{label} row {row} {col}: {theirs!r} against {mine_v!r}")
    return cells


def results_rows(page):
    """The values the results screen shows for each sample, as the cells read."""
    return page.eval_on_selector_all(
        'section[aria-label="Samples"] [role="table"] > [role="row"]:not(:first-child)',
        """els => els.map(b => { const c = [...b.children].map(x => x.textContent);
             return { r: c[0], id: c[1], amt: c[2], aN: c[3], dN: c[4], UN: c[5], dC: c[6], UC: c[7],
                      pN: c[8], pC: c[9], cn: c[10],
                      flags: [...b.children[11].children].map(x => x.getAttribute('title')) }; })""")


def contrast(color, ground="#FFFFFF"):
    """The WCAG 2 contrast ratio of two colors written #RRGGBB."""
    def lum(h):
        def lin(c):
            c /= 255
            return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
        h = h.lstrip("#")
        r, g, b = (lin(int(h[i:i + 2], 16)) for i in (0, 2, 4))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    hi, lo = sorted((lum(color), lum(ground)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def figures(page, scr):
    """Finding R-29 on the screen shown: each figure's drawing is an image whose name begins with the figure's
    number and title and says what it plots against what and how many marks; every line and gray mark (the
    axes, gray dots, error bars or rug ticks, lines, points, and rings; the bands and the grid aside) reaches
    3:1 against the white page unless it is a Wong color of the data. Returns the names, by figure number."""
    figs = page.eval_on_selector_all("figure.dc-figure", """fs => fs.map(f => {
        const h = f.querySelector('h2, h3'), svg = f.querySelector('svg[role=img]');
        return { head: h ? h.textContent.trim() : '', alt: svg ? svg.getAttribute('aria-label') : null,
                 marks: svg ? [...svg.querySelectorAll(':scope > path')].map((p, i) => [i, p.getAttribute('stroke'),
                   p.getAttribute('fill'), p.getAttribute('d') || '']) : [] }; })""")
    names = {}
    for f in figs:
        num, _, title = f["head"].partition(" ")[2].partition(" ")
        if f["alt"] is None:
            continue            # a figure with nothing to draw says why in its text
        check(f["alt"].startswith(f"Fig. {num}, {title}: ") and " against " in f["alt"]
              and re.search(r"\b\d+ [a-z]", f["alt"].split(" against ", 1)[1]), f"{scr}: Fig. {num} is named {f['alt']!r}")
        names[num] = f["alt"]
        for i, stroke, fill, d in f["marks"]:
            if i < 3 or not d:
                continue
            for what, c in (("stroke", stroke), ("fill", fill)):
                if c and c != "none" and c.upper() not in WONG:
                    check(contrast(c) >= 3.0, f"{scr}: Fig. {num}, mark {i}: its {what} {c} is {contrast(c):.2f}:1 against white")
    return names


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--shots", default=str(ROOT / "build" / "e2e_web" / "shots"), help="folder for the screenshots")
    ap.add_argument("--keep", help="folder to keep the downloads in (default: a temporary folder)")
    ap.add_argument("--url", help="a running service to check instead of starting one")
    args = ap.parse_args(argv)
    shots = Path(args.shots)
    shots.mkdir(parents=True, exist_ok=True)
    work = Path(args.keep) if args.keep else Path(tempfile.mkdtemp(prefix="e2e_web-"))
    work.mkdir(parents=True, exist_ok=True)

    log = open(work / "server.log", "w")
    if args.url:
        base, server = args.url.rstrip("/"), None
    else:
        port = free_port()
        base = f"http://127.0.0.1:{port}"
        server = subprocess.Popen([sys.executable, "-m", "uvicorn", "isoharness.app:app", "--host", "127.0.0.1",
                                   "--port", str(port)], cwd=ROOT, stdout=log, stderr=subprocess.STDOUT)
    noise = []          # console errors, page errors, failed requests and HTTP errors
    expected = []      # the ones this test causes on purpose
    state = {"expect": None}
    stats = {}
    try:
        for _ in range(150):
            try:
                with urllib.request.urlopen(base + "/api/health", timeout=2) as r:
                    if json.loads(r.read()) == {"status": "ok"}:
                        break
            except (urllib.error.URLError, ConnectionError, TimeoutError, json.JSONDecodeError):
                time.sleep(0.2)
        else:
            raise Failed("the service did not start; see " + str(work / "server.log"))
        with urllib.request.urlopen(base + "/api/version", timeout=10) as r:
            version = json.loads(r.read())
        check(version["libreoffice"], "LibreOffice is not on the PATH, so the workbook would carry formulas only")
        ver = version["ea_irms"]

        def note(kind, text, url=""):
            want = state["expect"]
            if want and ((kind == "console" and re.search(r"status of (415|422)", text))
                         or (kind == "http" and url.endswith("/api/process") and text in ("415", "422"))):
                expected.append((kind, text))
                return
            noise.append(f"{kind}: {text} {url}".strip())

        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            context = browser.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            page = context.new_page()
            page.on("console", lambda m: note("console", m.text) if m.type == "error" else None)
            page.on("pageerror", lambda e: note("pageerror", str(e)))
            page.on("requestfailed", lambda r: note("requestfailed", r.failure or "", r.url))
            page.on("response", lambda r: note("http", str(r.status), r.url) if r.status >= 400 else None)

            def shot(name, full=True):
                page.wait_for_timeout(250)
                page.screenshot(path=str(shots / name), full_page=full)

            def no_prototype():
                check("Prototype" not in page.inner_text("body"), "the page says Prototype")

            def heading(scr):
                expect(page.get_by_role("heading", level=1, name=TITLES[scr], exact=True)).to_be_visible()

            # ---- the home page, which lists the instruments
            page.goto(base + "/")
            expect(page.get_by_role("heading", level=1, name="Stable isotope data reduction")).to_be_visible()
            names = page.locator("main li h3").all_text_contents()
            check(names == ["EA-IRMS", "GasBench N₂O", "GasBench carbonate"], f"the home page lists {names}")
            status = page.locator("main li .status").all_text_contents()
            check(status == ["Available", "Planned", "Planned"], f"the home page marks them {status}")
            shot("1440-home.png")
            page.get_by_role("link", name="Open EA-IRMS").click()
            page.wait_for_url(re.compile(r"/ea-irms$"))

            # ---- the landing page
            expect(page.get_by_text("Drop exports here, or choose files")).to_be_visible()
            check("deletes each file before it sends its reply" in page.inner_text("body"),
                  "the upload screen does not say that the file is deleted")
            no_prototype()
            shot("1440-01-upload-empty.png")

            # ---- files that are not exports: the service's message, as written
            pdf = work / "notes.pdf"
            pdf.write_bytes(b"%PDF-1.4\n% not an export\n")
            bad_csv = work / "not_an_export.csv"
            bad_csv.write_text("sample,mass\nA,0.4\n")
            state["expect"] = True
            for bad in (pdf, bad_csv):
                status, _, body = post(base, "/api/process", bad)
                detail = json.loads(body)["detail"]
                check(status in (415, 422), f"{bad.name}: the service answered {status}")
                page.set_input_files("#dc-export", str(bad))
                alert = page.locator("[role=alert]")
                expect(alert).to_contain_text(f"{bad.name} was not processed")
                expect(page.locator('[data-dc="error-detail"]')).to_have_text(detail)
            state["expect"] = None
            check(len([e for e in expected if e[0] == "http"]) == 2, f"expected two refused uploads, saw {expected}")
            shot("1440-02-upload-refused.png")

            # ---- the test run
            page.set_input_files("#dc-export", str(RUN))
            expect(page.get_by_text("Read without errors")).to_be_visible(timeout=30000)
            check(page.locator("[role=alert]").count() == 0, "the refusal is still shown after a good upload")
            cols = page.locator('[aria-label="Column mapping"] > [role=row]')
            check(cols.count() == 13, f"the column table has {cols.count() - 1} rows, not 12")
            found = page.locator('[aria-label="Column mapping"] [role=cell]:has-text("Found")').count()
            check(found == 12, f"{found} of 12 columns marked found")
            body_text = page.inner_text("body")
            check("56 positions" in body_text, "the upload screen does not count 56 positions")
            check("match the server’s at all 56 positions" in body_text,
                  "the page's own check against the engine did not pass")
            check("Run date 15 Sep 2026, read from the file name" in body_text, "the run date was not read from the file name")
            shot("1440-03-upload.png")

            # ---- the Excel template of the lab configuration (ea_irms 1.4.6), as the package defaults
            with page.expect_download(timeout=30000) as dl:
                page.get_by_role("button", name="Download the template").click()
            template = work / dl.value.suggested_filename
            dl.value.save_as(str(template))
            check(template.name == f"lab_configuration_ea_irms_v{ver}.xlsx", f"template named {template.name}")
            status, _, body = post(base, "/api/config/read", template)
            check(status == 200, f"/api/config/read answered {status} for the template")
            status, _, dflt = get(base, "/api/defaults")
            check(json.loads(body)["config"] == json.loads(dflt)["config"], "the template is not the package defaults")

            # ---- a lab configuration loaded after the upload is processed on the server again,
            # first as JSON, then as the Excel file the service writes for the same settings
            lab = work / "lab-config.json"
            lab.write_text(json.dumps({"blank_correction": "On", "coverage_k": 3.0}))
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-config", str(lab))
            expect(page.locator(".dc-toast").get_by_text("The run was processed again with lab-config.json")).to_be_visible(timeout=30000)
            status, _, body = post_json(base, "/api/config.xlsx", {"config": json.loads(lab.read_text()), "name": "lab"})
            check(status == 200, f"/api/config.xlsx answered {status}")
            lab = work / "lab-config.xlsx"
            lab.write_bytes(body)
            # (k = 3 on the uncertainty step can only come from the server, through the port block's settings)
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-config", str(lab))
            expect(page.locator(".dc-toast").get_by_text("The run was processed again with lab-config.xlsx")).to_be_visible(timeout=30000)
            page.get_by_role("button", name="Start the corrections").click()
            heading("inst")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Setup").click()
            heading("setup")
            expect(page.get_by_text("From lab-config.xlsx.")).to_be_visible()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Blank").click()
            heading("blank")
            expect(page.get_by_role("group", name="Blank correction").get_by_role("button", name="On")).to_have_attribute("aria-pressed", "true")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Uncertainty").click()
            heading("unc")
            expect(page.get_by_text("by more than their uncertainty, it uses 3σ")).to_be_visible()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Results").click()
            heading("results")
            expect(page.get_by_text("with 1σ, %N, %C, and C:N")).to_be_visible()
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.get_by_role("button", name="Use the package defaults").click()
            expect(page.locator(".dc-toast").get_by_text("processed again with the package defaults")).to_be_visible(timeout=30000)

            # isoharness 0.4.1 (R-17): with a run on the page, leaving it asks first
            check(page.evaluate("() => { const e = new Event('beforeunload', { cancelable: true }); "
                                "window.dispatchEvent(e); return e.defaultPrevented; }"),
                  "leaving the page with a run on it does not ask first")

            # ---- the steps, with decisions on the way
            page.get_by_role("button", name="Start the corrections").click()
            for scr in SCREENS:
                heading(scr)
                no_prototype()
                if scr == "inst":
                    for title in ("N2 reference pulses", "CO2 reference pulses", "Reference peak size", "Background at m/z 28"):
                        expect(page.locator("figure", has_text=title)).to_have_count(1)
                    expect(page.get_by_text("26 positions ran before the background settled at position 27")).to_be_visible()
                    expect(page.get_by_text("under the 0.1‰ limit the FIRMS guide")).to_have_count(2)
                    check(page.get_by_role("button", name=re.compile(r"Step 1, Instrument")).count() == 0,
                          "the instrument screen links to itself")
                if scr == "setup":
                    expect(page.get_by_text("Package defaults, ea_irms " + ver + ".")).to_be_visible()
                how = page.locator('section[aria-label="How this step works"]')
                if scr in EXPLAINED:
                    # isoharness 0.4.0: each correction explains itself, shown by default, with a worked example
                    expect(how).to_be_visible()
                    expect(how.get_by_role("button", name="Hide explanations")).to_be_visible()
                    expect(how).to_contain_text("Worked example, ")
                    rows = how.locator('[aria-label="Worked example"] > [role=row]')
                    check(rows.count() >= 5, f"{scr}: the worked example has {rows.count() - 1} rows")
                else:
                    check(how.count() == 0, f"{scr}: an explanation on a screen that corrects nothing")
                titles = page.locator('section[aria-label="Checks"] strong').all_text_contents()
                extra = {t for t, _ in ENGINE_CHECKS.values()}
                if scr in MOVED:
                    check([t for t in titles if t not in extra] == MOVED[scr], f"{scr}: its checks are {titles}, not {MOVED[scr]}")
                    if scr in ENGINE_CHECKS and ENGINE_CHECKS[scr][0] in titles:
                        box = page.locator('section[aria-label="Checks"] > div', has_text=ENGINE_CHECKS[scr][0])
                        expect(box).to_contain_text(ENGINE_CHECKS[scr][1])
                # finding R-29: the figures' names and marks, the current step, and the selected position
                figures(page, scr)
                rail = page.locator("nav[aria-label='Correction steps'] button[aria-current=step]")
                expect(rail).to_have_count(1)
                expect(rail).to_contain_text(BUTTONS[scr])
                if scr not in ("results",):
                    bars = page.locator(".dc-ribbon button")
                    widths = bars.evaluate_all("bs => bs.map(b => b.getBoundingClientRect().width)")
                    check(len(widths) == 56 and min(widths) >= 24 - 1e-6, f"{scr}: the sequence's bars are {min(widths):.1f} px at the least")
                    expect(page.locator('.dc-ribbon button[aria-current="true"]')).to_have_count(1)
                if scr in ("blank", "lin", "drift"):
                    for el in ("δ¹⁵N", "δ¹³C"):
                        expect(page.locator("figure", has_text="Change in " + el)).to_have_count(1)
                if scr == "blank":
                    group = page.get_by_role("group", name="Blank correction")
                    expect(group.get_by_role("button", name="Auto")).to_have_attribute("aria-pressed", "true")
                    group.get_by_role("button", name="On").click()
                    expect(group.get_by_role("button", name="On")).to_have_attribute("aria-pressed", "true")
                    expect(page.get_by_text("δ¹⁵N: subtracted")).to_be_visible()
                    # with the blank subtracted, the mass balance is the value after this step
                    cells = lambda label: how.locator("[role=row]", has_text=label).first.locator("[role=cell]").all_text_contents()
                    mb, after = cells("Mass balance"), cells("After this step")
                    check(mb[0] != "" and mb[0] == after[0], f"blank: the mass balance reads {mb}, the step {after}")
                    stats["mass balance"] = mb[0]
                if scr == "lin":
                    # an expert's choice to hide the explanations holds for every step and is kept in the browser
                    how.get_by_role("button", name="Hide explanations").click()
                    expect(how.get_by_role("button", name="Show explanations")).to_be_visible()
                    check(how.locator('[aria-label="Worked example"]').count() == 0, "the explanation did not hide")
                    check(page.evaluate("localStorage.getItem('isoharness.explanations')") == "hidden",
                          "the choice to hide the explanations was not kept")
                    page.locator("nav[aria-label='Correction steps'] button", has_text="Drift").click()
                    heading("drift")
                    expect(page.get_by_role("button", name="Show explanations")).to_be_visible()
                    page.get_by_role("button", name="Show explanations").click()
                    check(page.evaluate("localStorage.getItem('isoharness.explanations')") is None,
                          "showing the explanations again left the choice to hide them")
                    page.locator("nav[aria-label='Correction steps'] button", has_text="Linearity").click()
                    heading("lin")
                    expect(how.get_by_role("button", name="Hide explanations")).to_be_visible()
                if scr == "drift":
                    chart = page.locator("figure", has_text="δ¹⁵N of the drift monitors")
                    # isoharness 0.4.1 (R-03): a figure's points are one tab stop, each at least a dot wide
                    check(chart.locator('button.dc-hit[tabindex="0"]').count() == 1,
                          "the drift monitors' figure has more than one tab stop")
                    small = chart.locator("button.dc-hit").evaluate_all(
                        "bs => bs.filter(b => { const r = b.getBoundingClientRect(); return r.width < 4 || r.height < 4; }).length")
                    check(small == 0, f"{small} points of the drift monitors' figure are under 4 px")
                    point = chart.get_by_role("button", name=re.compile(rf", position {EXCL_N}: "))
                    point.click()
                    expect(chart.get_by_role("button", name=re.compile(rf", position {EXCL_N}: .*Left out"))).to_have_count(1)
                    expect(chart).to_contain_text("12 of 13 in the fit")
                    # R-29: the figure's name gives what it plots, its points, and the selected one, which is current
                    alt = figures(page, scr)["1"]
                    check(alt.startswith("Fig. 1, δ¹⁵N of the drift monitors: δ¹⁵N after linearity (‰) against position in the run, 13 points; ")
                          and re.search(rf"the selected position is [^,]+, position {EXCL_N}: .*left out$", alt), f"the drift figure is named {alt!r}")
                    expect(chart.locator('button.dc-hit[aria-current="true"]')).to_have_count(1)
                    expect(page.locator('.dc-ribbon button[aria-current="true"]')).to_have_attribute("aria-label", re.compile(rf"^Position {EXCL_N}, "))
                    page.get_by_role("group", name="Model").get_by_role("button", name="Background").click()
                    expect(chart).to_contain_text("against background")
                if scr == "elem":
                    chart = page.locator("figure", has_text="%N of the drift monitors against certified")
                    chart.get_by_role("button", name=re.compile(rf", position {EXCL_E}: ")).click()
                    expect(chart).to_contain_text("12 of 13 monitors used")
                if scr == "qc":
                    # every check of the run, each with a button back to the step it belongs to, which reads
                    # Step 1, Instrument and is named for its check (finding R-29), so no two share a name
                    back = page.get_by_role("button", name=re.compile(r"^Go to Step 1, Instrument, for "))
                    expect(back).to_have_count(5)
                    labels = back.evaluate_all("bs => bs.map(b => [b.getAttribute('aria-label'), b.textContent.trim()])")
                    check(len({a for a, _ in labels}) == 5 and {t for _, t in labels} == {"Step 1, Instrument"}
                          and [a.split(", for ", 1)[1] for a, _ in labels] == titles[:5],
                          f"the buttons back to the instrument read {labels}")
                    for step, names in MOVED.items():
                        label = f"Go to Step {SCREENS.index(step) + 1}, {BUTTONS[step]}, for "
                        for t in names:
                            expect(page.get_by_role("button", name=label + t, exact=True)).to_have_count(1)
                    check([t for t in titles[5:] if t not in extra] == [t for step in MOVED for t in MOVED[step]],
                          f"quality checks listed as {titles}")
                    for label, step in (("Go to Step 1, Instrument, for Carryover", "inst"),
                                        ("Go to Step 4, Linearity, for Effect of the linearity choice", "lin")):
                        page.get_by_role("button", name=label, exact=True).click()
                        heading(step)
                        page.locator("nav[aria-label='Correction steps'] button", has_text="Quality checks").click()
                        heading("qc")
                if scr == "results":
                    filters = page.get_by_role("group", name="Filter the samples").get_by_role("button")
                    pressed = filters.evaluate_all("bs => bs.map(b => [b.textContent.trim(), b.getAttribute('aria-pressed')])")
                    check([p for _, p in pressed] == ["true", "false", "false", "false"] and pressed[0][0].startswith("All "),
                          f"the results filters read {pressed}")
                if scr == "unc":
                    group = page.get_by_role("group", name="δ¹⁵N repeatability")
                    group.get_by_role("button", name="Drift monitors").click()
                    expect(group.get_by_role("button", name="Drift monitors")).to_have_attribute("aria-pressed", "true")
                    expect(page.get_by_text("SD of the drift monitors in the fit")).to_be_visible()
                shot(f"1440-{SCREENS.index(scr) + 4:02d}-{scr}.png")
                if scr != "results":
                    page.get_by_role("button", name="Continue to " + BUTTONS[SCREENS[SCREENS.index(scr) + 1]]).first.click()

            # ---- the results screen against the service, decisions built here without the page
            page_rows = results_rows(page)
            check(len(page_rows) >= 5, f"the results screen lists {len(page_rows)} samples")
            status, _, body = post(base, "/api/process", RUN)
            processed = json.loads(body)["config"]
            mine = dict(processed, blank_correction="On", exclude_drift_n=[EXCL_N], drift_model="Background",
                        exclude_elemental=[EXCL_E])
            status, _, body = post(base, "/api/process", RUN, {"config": json.dumps(mine)})
            check(status == 200, f"/api/process answered {status} for the test's own configuration")
            positions = json.loads(body)["positions"]
            mons = [p["d15n_drift"] for p in positions
                    if p["role"] == "Drift" and p["row"] != EXCL_N and p["d15n_drift"] is not None]
            mine["u_sample_n"] = sd(mons)
            status, _, body = post(base, "/api/results.csv", RUN, {"config": json.dumps(mine)})
            check(status == 200, f"/api/results.csv answered {status}")
            independent = rows_of(body.decode())
            n, values, exact = compare(page_rows, independent, "results screen against the service")
            stats["compared"] = (n, values, exact)
            status, _, body = post(base, "/api/results.csv", RUN)
            defaults = rows_of(body.decode())
            moved = sum(1 for r in page_rows if not agree(r["dN"], defaults[int(r["r"])]["d15N (permil)"], 2)[0])
            check(moved > 0, "the decisions changed no shown d15N")
            stats["moved"] = moved

            # ---- downloads
            files = page.locator('section[aria-label="Files and methods"]')
            got = {}
            for label, timeout in (("Workbook", 180000), ("Results table", 60000), ("Configuration", 30000)):
                with page.expect_download(timeout=timeout) as dl:
                    files.get_by_role("button", name=re.compile("^" + label)).click()
                d = dl.value
                path = work / d.suggested_filename
                d.save_as(str(path))
                got[label] = path
                expect(files.get_by_role("button", name=re.compile("^" + label))).to_be_enabled(timeout=timeout)
            check(got["Workbook"].name == f"20260915_raw_ea_irms_v{ver}.xlsx", f"workbook named {got['Workbook'].name}")
            check(got["Results table"].name == f"20260915_raw_results_ea_irms_v{ver}.csv", f"table named {got['Results table'].name}")
            check(got["Configuration"].name == "20260915_raw_config.xlsx", f"configuration named {got['Configuration'].name}")
            # the message bar (isoharness 0.4.1: .dc-toast; the live region beside it is always there)
            check(page.locator(".dc-toast").count() == 0, "a download left a message: " + (
                page.locator(".dc-toast").first.inner_text() if page.locator(".dc-toast").count() else ""))

            # the configuration, read back by the service: the decisions in their fields, every other field
            # as processed
            status, _, body = post(base, "/api/config/read", got["Configuration"])
            check(status == 200, f"/api/config/read answered {status} for the downloaded configuration")
            check(json.loads(body)["notes"] == [], f"reading the configuration left notes: {json.loads(body)['notes']}")
            cfg = json.loads(body)["config"]
            want = dict(processed, blank_correction="On", linearity_correction="Auto", linearity_correction_c="Auto",
                        drift_correction="Auto", drift_model="Background", exclude_drift_n=[EXCL_N], exclude_drift_c=[],
                        elemental_correction="Auto", exclude_elemental=[EXCL_E], u_sample_c=None,
                        # ea_irms 1.4.10 (R-08): the run these were chosen for, and the repeatability's source
                        run_file="20260915_raw.csv", u_sample_n_source="Drift monitors", u_sample_c_source=None)
            check(set(cfg) == set(want), f"configuration fields differ: {sorted(set(cfg) ^ set(want))}")
            for k, v in want.items():
                if k != "u_sample_n":
                    check(cfg[k] == v, f"configuration {k} is {cfg[k]!r}, not {v!r}")
            check(abs(cfg["u_sample_n"] - mine["u_sample_n"]) <= 1e-12 * mine["u_sample_n"],
                  f"u_sample_n is {cfg['u_sample_n']}, the monitors' SD is {mine['u_sample_n']}")

            # the downloaded configuration posted back reproduces the page and the downloaded table
            status, _, body = post(base, "/api/results.csv", RUN, {"config": json.dumps(cfg)})
            check(status == 200, f"/api/results.csv answered {status} for the downloaded configuration")
            diff, worst = table_diff(body.decode(), got["Results table"].read_text())
            check(not diff, "the downloaded table differs from the one its configuration gives: " + "; ".join(diff[:3]))
            stats["round trip"] = worst
            round_trip = rows_of(body.decode())
            compare(page_rows, round_trip, "results screen against the downloaded configuration")
            compare(page_rows, rows_of(got["Results table"].read_text()), "results screen against the downloaded table")

            # the workbook: computed by LibreOffice, and holding the same values
            cells = book_cells(got["Workbook"], round_trip)
            check(cells > 100, f"only {cells} computed values in the workbook")
            stats["book"] = cells

            # the name comes from Content-Disposition: a stand-in answer names the file differently
            page.route("**/api/results.csv", lambda route: route.fulfill(
                status=200, content_type="text/csv", body="Row\n",
                headers={"Content-Disposition": 'attachment; filename="e2e-named-by-the-service.csv"'}))
            with page.expect_download(timeout=60000) as dl:
                files.get_by_role("button", name=re.compile("^Results table")).click()
            check(dl.value.suggested_filename == "e2e-named-by-the-service.csv",
                  f"the page named the table {dl.value.suggested_filename}, not as the service did")
            page.unroute("**/api/results.csv")
            expect(files.get_by_role("button", name=re.compile("^Results table"))).to_be_enabled(timeout=60000)

            # ---- several runs at once (isoharness 0.2.0)
            before = page.locator('section[aria-label="Samples"]').inner_text()
            run_a, run_b = work / "20260916_run_a.csv", work / "20260917_run_b.csv"
            status, _, body = post(base, "/api/process", RUN)
            samples = {str(q["row"]) for q in json.loads(body)["positions"] if q["role"] == "Sample"}
            run_a.write_bytes(variant(RUN, samples))
            run_b.write_bytes(variant(RUN, samples, 0.5))
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            page.set_input_files("#dc-export", [str(run_a), str(run_b)])
            expect(page.get_by_text("Batch of 3 runs")).to_be_visible(timeout=60000)
            expect(page.get_by_role("heading", level=1, name="Correct 3 runs")).to_be_visible()
            listed = page.locator('section[aria-label="Runs in this batch"] li')
            check(listed.count() == 3, f"the batch lists {listed.count()} runs, not 3")
            dated = (("20260915_raw.csv", "15 Sep 2026"), ("20260916_run_a.csv", "16 Sep 2026"),
                     ("20260917_run_b.csv", "17 Sep 2026"))
            for k, (name, date) in enumerate(dated):
                text = listed.nth(k).inner_text()
                check(name in text and "run of " + date in text and "Read without errors" in text,
                      f"batch row {k + 1} reads {text!r}")
            check(listed.nth(2).get_attribute("aria-current") == "true", "the last run added is not the one shown")
            shot("1440-batch-upload.png")
            page.get_by_role("button", name="Show 20260915_raw.csv").click()
            page.get_by_role("button", name="Start the corrections").click()
            heading("results")
            expect(page.get_by_text("Run 1 of 3")).to_be_visible()
            check(page.locator('section[aria-label="Samples"]').inner_text() == before,
                  "the first run's results changed when two runs were added and it was shown again")
            reps = page.locator('[aria-label="Samples measured in more than one run"] [role=row]')
            status, _, body = post(base, "/api/process", run_a)
            slope = json.loads(body)["fits"]["cal_n"]["slope"]
            check(reps.count() - 1 == 26, f"{reps.count() - 1} samples measured in more than one run, not 26")
            ranges = {reps.nth(i).locator("[role=cell]").nth(5).inner_text() for i in range(1, reps.count())}
            want = js_fixed(0.5 * slope, 2)
            check(ranges == {want}, f"the d15N ranges read {ranges}, not {want}")
            shot("1440-batch-results.png")
            # the combined results: each run's rows equal the service's table for that run, run 1 with its decisions
            files = page.locator('section[aria-label="Files and methods"]')
            combined = {}
            for label, key in (("Combined results, all 3 runs", "xlsx"), ("Combined results table", "csv")):
                with page.expect_download(timeout=120000) as dl:
                    files.get_by_role("button", name=re.compile("^" + label)).click()
                d = dl.value
                combined[key] = work / d.suggested_filename
                d.save_as(str(combined[key]))
                expect(files.get_by_role("button", name=re.compile("^" + label))).to_be_enabled(timeout=120000)
            for key in ("xlsx", "csv"):
                check(combined[key].name == f"combined_3_runs_results_ea_irms_v{ver}.{key}",
                      f"the combined {key} file is named {combined[key].name}")
            book = openpyxl.load_workbook(combined["xlsx"])
            check(book.sheetnames == ["README", "Results", "Runs", "Repeats"], f"combined sheets {book.sheetnames}")
            rows = list(book["Results"].iter_rows(min_row=2, values_only=True))
            check(len(rows) == 168, f"the combined Results sheet has {len(rows)} rows, not 168")
            lines = combined["csv"].read_text().splitlines()
            check(len(lines) == 169 and lines[0].startswith("Run,Row,Identifier"),
                  f"the combined table has {len(lines)} lines, the first {lines[0][:40]!r}")
            tables = [got["Results table"].read_text()]
            for f in (run_a, run_b):
                status, _, body = post(base, "/api/results.csv", f)
                check(status == 200, f"/api/results.csv answered {status} for {f.name}")
                tables.append(body.decode())
            worst = 0.0
            for k, (label, table) in enumerate(zip(("20260915_raw", "20260916_run_a", "20260917_run_b"), tables)):
                part = rows[56 * k:56 * (k + 1)]
                check({r[0] for r in part} == {label},
                      f"combined rows {56 * k + 1} to {56 * (k + 1)} are not all {label}")
                diff, w = sheet_diff([r[1:] for r in part], table)
                check(not diff, f"combined results, {label}: " + "; ".join(diff[:3]))
                worst = max(worst, w)
                own = [ln.split(",", 1)[1] for ln in lines[1 + 56 * k:1 + 56 * (k + 1)]]
                check(own == table.splitlines()[1:], f"combined table, {label}: not the run's own table")
            runs = list(book["Runs"].iter_rows(min_row=2, values_only=True))
            check([r[2] for r in runs] == ["2026-09-15", "2026-09-16", "2026-09-17"],
                  f"run dates {[r[2] for r in runs]}")
            status, _, body = post(base, "/api/process", RUN, {"config": json.dumps(mine)})
            fits = json.loads(body)["fits"]
            n_on, c_on = fits["blank_n"]["in_force"], fits["blank_c"]["in_force"]
            blank = "N and C" if n_on and c_on else "N only" if n_on else "C only" if c_on else "no"
            check(runs[0][7] == blank and blank != "no",
                  f"the first run's blank reads {runs[0][7]!r} on the Runs sheet, not {blank!r}")
            rep = list(book["Repeats"].iter_rows(values_only=True))
            ix = {h: i for i, h in enumerate(rep[0])}
            off = [r for r in rep[1:] if r[ix["d15N minus mean"]] is not None
                   and abs(abs(r[ix["d15N minus mean"]]) - 0.25 * slope) > 1e-9]
            check(len(rep) - 1 == 52 and not off, f"Repeats: {len(rep) - 1} rows, {len(off)} off by more than 1e-9")
            stats["batch"] = (len(rows), worst, len(rep) - 1)

            # ---- run information and a sample sheet (isoharness 0.3.0)
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            for field, value in (("operator", "XT Wang"), ("instrument", "Delta V Plus"), ("project", "e2e")):
                page.fill(f"#dc-info-{field}", value)
            page.get_by_role("button", name="Use for every run").click()
            expect(page.locator(".dc-toast").get_by_text("Every run of the batch now has the run information of 20260915_raw.csv.")).to_be_visible()
            page.fill("#dc-info-notes", "first run")
            with page.expect_download(timeout=30000) as dl:
                page.get_by_role("button", name="Download a sheet with these sample names").click()
            tpl = work / dl.value.suggested_filename
            dl.value.save_as(str(tpl))
            check(tpl.name == f"sample_sheet_ea_irms_v{ver}.xlsx", f"the sample sheet is named {tpl.name}")
            ids = [c.value for c in openpyxl.load_workbook(tpl)["Samples"]["A"]][1:]
            want_ids = ["Random"] + [f"S{r}" for r in sorted(int(x) for x in samples)]
            check(ids == want_ids, f"the sample sheet lists {len(ids)} names, {ids[:3]}, not {len(want_ids)}")
            sheet = work / "samples.csv"
            sheet.write_text("Sample ID,Project,Depth,Replicate of,Alt ID\nS21,P1,10,,X21\nS24,P1,12,S21,X24\n"
                             "S2,P2,,,X2\n")
            with page.expect_request(lambda r: r.url.endswith("/api/samplesheet/read") and r.method == "POST"):
                page.set_input_files("#dc-sheet", str(sheet))
            summary = page.locator('[data-dc="sheet-summary"]')
            # S21 and S24 of the two copies match; S2 matches none of S20 to S29, whose next character is a digit.
            # Findings R-30 and R-34: the run in which no group of replicates formed is named.
            expect(summary).to_have_text("4 of 78 samples matched (4 by the exact name); 74 not in the sheet. "
                                         "2 groups of replicates within a run will be compared; none formed in "
                                         "20260915_raw.csv.", timeout=30000)
            matched = page.locator('[aria-label="How each sample matched the sheet"] [role=row]')
            check(matched.count() == 79, f"the match table has {matched.count() - 1} rows, not 78")
            shot("1440-sample-sheet.png")
            page.select_option("#dc-sheet-id", "Alt ID")
            expect(summary).to_have_text("0 of 78 samples matched; 78 not in the sheet. "
                                         "No group of replicates formed in any run.", timeout=30000)
            page.select_option("#dc-sheet-id", "Sample ID")
            expect(summary).to_contain_text("4 of 78 samples matched", timeout=30000)
            page.get_by_role("button", name="Start the corrections").click()
            heading("results")
            files = page.locator('section[aria-label="Files and methods"]')
            meta = {}
            for label, key in (("Combined results, all 3 runs", "xlsx"), ("Combined results table", "csv")):
                with page.expect_download(timeout=120000) as dl:
                    files.get_by_role("button", name=re.compile("^" + label)).click()
                meta[key] = work / ("with_sheet_" + dl.value.suggested_filename)
                dl.value.save_as(str(meta[key]))
                expect(files.get_by_role("button", name=re.compile("^" + label))).to_be_enabled(timeout=120000)
            mlines = meta["csv"].read_text().splitlines()
            check(mlines[0] == lines[0] + ",Sample ID (sheet),Matched by,Project,Depth,Replicate of,Alt ID",
                  f"the combined table with the sheet is headed {mlines[0][-80:]!r}")
            mrows, orig = list(csv.reader(mlines[1:])), list(csv.reader(lines[1:]))
            check(len(mrows) == 168 and all(m[:25] == o for m, o in zip(mrows, orig)),
                  "the sheet changed the combined table's own 25 columns (no replicates here are over the limit)")
            got_meta = {(m[0], m[1]): m[25:] for m in mrows if m[25:] != ["", "", "", "", "", ""]}
            want_meta = {}
            for run in ("20260916_run_a", "20260917_run_b"):
                want_meta[(run, "21")] = ["S21", "exact", "P1", "10", "", "X21"]
                want_meta[(run, "24")] = ["S24", "exact", "P1", "12", "S21", "X24"]
                for r in samples - {"21", "24"}:
                    want_meta[(run, r)] = ["", "not in sheet", "", "", "", ""]
            for r in samples:
                want_meta[("20260915_raw", r)] = ["", "not in sheet", "", "", "", ""]
            off_meta = sorted(k for k in set(got_meta) | set(want_meta) if got_meta.get(k) != want_meta.get(k))
            check(not off_meta, f"the sheet's cells in the combined table differ at {off_meta[:4]}")
            runs_sheet = list(openpyxl.load_workbook(meta["xlsx"])["Runs"].iter_rows(values_only=True))
            check(list(runs_sheet[0][-5:]) == ["Operator", "Instrument", "Project", "Notes", "Samples matched"],
                  f"the Runs sheet ends {runs_sheet[0][-5:]}")
            check([list(r[-5:]) for r in runs_sheet[1:]] == [
                ["XT Wang", "Delta V Plus", "e2e", "first run", "0 of 26 samples matched"],
                ["XT Wang", "Delta V Plus", "e2e", None, "2 of 26 samples matched (2 by the exact name)"],
                ["XT Wang", "Delta V Plus", "e2e", None, "2 of 26 samples matched (2 by the exact name)"]],
                f"the Runs sheet reads {[r[-5:] for r in runs_sheet[1:]]}")
            with page.expect_download(timeout=180000) as dl:
                page.get_by_role("button", name="Workbook of 20260916_run_a.csv").click()
            run_book = work / ("with_sheet_" + dl.value.suggested_filename)
            dl.value.save_as(str(run_book))
            expect(page.get_by_role("button", name="Workbook of 20260916_run_a.csv")).to_be_enabled(timeout=180000)
            wbm = openpyxl.load_workbook(run_book, data_only=True)
            md, rsm = wbm["Metadata"], wbm["Results"]
            check([md.cell(r, 2).value for r in range(5, 12)] == [
                "20260916_run_a", "20260916_run_a.csv", "2026-09-16", "XT Wang", "Delta V Plus", "e2e", None],
                f"the run's Metadata sheet reads {[md.cell(r, 2).value for r in range(5, 12)]}")
            check([rsm.cell(4 + 24, c).value for c in range(25, 31)] == ["S24", "exact", "P1", 12, "S21", "X24"],
                  f"the run's Results sheet carries {[rsm.cell(4 + 24, c).value for c in range(25, 31)]} at position 24")
            rpm = wbm["Replicates"]
            check([rpm.cell(9, c).value for c in (1, 3, 10, 14, 15)] == [1, 21, None, None, "21, 24"]
                  and rpm.cell(10, 3).value == 24, "the run's Replicates sheet does not hold S21 and S24 as one group")
            stats["sheet"] = (len(want_meta), len(runs_sheet) - 1)
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            page.get_by_role("button", name="Remove 20260917_run_b.csv").click()
            expect(page.get_by_text("Batch of 2 runs")).to_be_visible()
            page.get_by_role("button", name="Start the corrections").click()
            heading("results")

            # ---- ea_irms 1.4.10 (R-08): the downloaded configuration chosen for the batch. Its positions left
            # out and its repeatability from the drift monitors belong to 20260915_raw.csv; the other run keeps
            # none of them, takes the SD of its own monitors, and says so.
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-config", str(got["Configuration"]))
            expect(page.locator(".dc-toast").get_by_text(
                "The 2 runs were processed again with 20260915_raw_config.xlsx")).to_be_visible(timeout=60000)
            expect(page.get_by_text(re.compile(
                rf"the position left out of the nitrogen drift fit \({EXCL_N}\), the position left out of the elemental "
                rf"correction \({EXCL_E}\), and the d15N repeatability from the drift monitors \(0\.\d{{3}} permil\) were "
                r"chosen for 20260915_raw\.csv and are used for that run only"))).to_be_visible()
            listed = page.locator('section[aria-label="Runs in this batch"] li')
            first_row, second_row = listed.nth(0).inner_text(), listed.nth(1).inner_text()
            check("Read without errors" in first_row and "note" not in first_row, f"batch row 1 reads {first_row!r}")
            check("Read without errors; 1 note on the configuration" in second_row, f"batch row 2 reads {second_row!r}")
            page.get_by_role("button", name="Show 20260916_run_a.csv").click()
            expect(page.get_by_text(re.compile(
                rf"In the configuration: the position left out of the nitrogen drift fit \({EXCL_N}\), the position "
                rf"left out of the elemental correction \({EXCL_E}\), and the d15N repeatability from the drift "
                r"monitors \(0\.\d{3} permil\) were chosen for 20260915_raw\.csv, so they were not applied to "
                r"20260916_run_a\.csv"))).to_be_visible()
            page.get_by_role("button", name="Start the corrections").click()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Drift").click()
            heading("drift")
            chart = page.locator("figure", has_text="δ¹⁵N of the drift monitors")
            expect(chart).to_contain_text("13 of 13 in the fit")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Uncertainty").click()
            heading("unc")
            group = page.get_by_role("group", name="δ¹⁵N repeatability")
            expect(group.get_by_role("button", name="Drift monitors")).to_have_attribute("aria-pressed", "true")
            check(group.get_by_role("button", name="Configuration").count() == 0,
                  "another run's monitor SD is offered as the configuration's value")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Results").click()
            heading("results")
            files = page.locator('section[aria-label="Files and methods"]')
            with page.expect_download(timeout=30000) as dl:
                files.get_by_role("button", name=re.compile("^Configuration")).click()
            other = work / ("run_a_" + dl.value.suggested_filename)
            dl.value.save_as(str(other))
            status, _, body = post(base, "/api/config/read", other)
            mine_a = json.loads(body)["config"]
            check(mine_a["run_file"] == "20260916_run_a.csv" and mine_a["exclude_drift_n"] == []
                  and mine_a["exclude_elemental"] == [] and mine_a["u_sample_n_source"] == "Drift monitors",
                  f"the second run's configuration reads {[mine_a[k] for k in ('run_file', 'exclude_drift_n', 'exclude_elemental', 'u_sample_n_source')]}")
            status, _, body = post(base, "/api/process", run_a, {"config": json.dumps(dict(mine_a, u_sample_n=None))})
            mons_a = [p["d15n_drift"] for p in json.loads(body)["positions"]
                      if p["role"] == "Drift" and p["d15n_drift"] is not None]
            check(abs(mine_a["u_sample_n"] - sd(mons_a)) <= 1e-12 * sd(mons_a),
                  f"the second run's repeatability is {mine_a['u_sample_n']}, its own monitors' SD {sd(mons_a)}")
            page.get_by_role("navigation", name="Runs").get_by_role("button", name=re.compile("20260915_raw.csv")).click()
            expect(page.get_by_text("Run 1 of 2")).to_be_visible()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Drift").click()
            heading("drift")
            chart = page.locator("figure", has_text="δ¹⁵N of the drift monitors")
            expect(chart).to_contain_text("12 of 13 in the fit")
            stats["scoped"] = len(mons_a)
            # a configuration that names no run, chosen for the two runs at once, leaves its positions out of both
            nameless = work / "nameless.json"
            nameless.write_text(json.dumps({"exclude_drift_n": [EXCL_N]}))
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-config", str(nameless))
            expect(page.locator(".dc-toast").get_by_text("The 2 runs were processed again with nameless.json")).to_be_visible(timeout=60000)
            expect(page.get_by_text(f"no run is named for the position left out of the nitrogen drift fit ({EXCL_N}), so "
                                    "it goes with none of the runs, since the configuration was given to several at "
                                    "once")).to_be_visible()
            listed = page.locator('section[aria-label="Runs in this batch"] li')
            for k in range(2):
                text = listed.nth(k).inner_text()
                check("1 note on the configuration" in text, f"batch row {k + 1} reads {text!r}")
            expect(page.get_by_text(f"In the configuration: no run is named for the position left out of the "
                                    f"nitrogen drift fit ({EXCL_N}), and the configuration was given to several runs "
                                    "at once, so it was applied to none of them")).to_be_visible()
            page.get_by_role("button", name="Start the corrections").click()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Drift").click()
            heading("drift")
            expect(page.locator("figure", has_text="δ¹⁵N of the drift monitors")).to_contain_text("13 of 13 in the fit")

            # ---- 400 px wide
            page.set_viewport_size({"width": 400, "height": 1000})
            shot("400-results.png")
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            expect(page.get_by_text("Read without errors").first).to_be_visible()
            shot("400-upload.png")
            wide = page.evaluate("document.documentElement.scrollWidth")
            check(wide <= 400, f"the upload screen is {wide} px wide at a 400 px viewport")

            # ---- the landing page, fresh, at both widths; the home page at 400 px; the example run
            page.goto(base + "/ea-irms")
            expect(page.get_by_role("heading", level=1, name="EA-IRMS data reduction", exact=True)).to_be_visible()
            shot("400-landing.png")
            wide = page.evaluate("document.documentElement.scrollWidth")
            check(wide <= 400, f"the landing page is {wide} px wide at a 400 px viewport")
            page.get_by_role("link", name="isoharness, all instruments").click()
            expect(page.get_by_role("heading", level=1, name="Stable isotope data reduction")).to_be_visible()
            shot("400-home.png")
            wide = page.evaluate("document.documentElement.scrollWidth")
            check(wide <= 400, f"the home page is {wide} px wide at a 400 px viewport")
            page.set_viewport_size({"width": 1440, "height": 1000})
            page.goto(base + "/ea-irms")
            shot("1440-00-landing.png")
            with page.expect_request(lambda r: r.url.endswith("/api/example") and r.method == "GET"):
                page.get_by_role("button", name="Try the 15 Sep 2026 test run").click()
            expect(page.get_by_text("Read without errors")).to_be_visible(timeout=30000)
            check("56 positions" in page.inner_text("body"), "the example run does not count 56 positions")
            check(page.get_by_role("heading", level=1, name="Correct a run").count() == 1,
                  "the upload screen has no heading once a run is read")

            # ---- finding W-11: a configuration for one run of a batch. With W-12: the two copies name each sample
            # S 20 and S20, which the samples measured in more than one run pair with the spaces taken out.
            page.goto(base + "/ea-irms")
            copy_a, copy_b = work / "20260918_copy_a.csv", work / "20260919_copy_b.csv"
            copy_a.write_bytes(variant(RUN, samples, name="S {}"))
            copy_b.write_bytes(variant(RUN, samples, name="S{}"))
            page.set_input_files("#dc-export", [str(copy_a), str(copy_b)])
            expect(page.get_by_text("Batch of 2 runs")).to_be_visible(timeout=60000)
            listed = page.locator('section[aria-label="Runs in this batch"] li')
            for k in range(2):
                expect(listed.nth(k)).to_contain_text("Configuration: the batch’s, the package defaults")
            batch_back = page.get_by_role("button", name=re.compile(r"^Use the batch’s configuration for "))
            expect(batch_back).to_have_count(0)
            # the first run's configuration, downloaded from the page, with one switch changed: the blank On
            page.get_by_role("button", name=f"Show {copy_a.name}").click()
            page.get_by_role("button", name="Start the corrections").click()
            heading("inst")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Results").click()
            heading("results")
            files = page.locator('section[aria-label="Files and methods"]')
            with page.expect_download(timeout=30000) as dl:
                files.get_by_role("button", name=re.compile("^Configuration")).click()
            cfg_a = work / ("w11_" + dl.value.suggested_filename)
            dl.value.save_as(str(cfg_a))
            expect(files.get_by_role("button", name=re.compile("^Configuration"))).to_be_enabled(timeout=30000)
            status, _, body = post(base, "/api/config/read", cfg_a)
            check(status == 200, f"/api/config/read answered {status} for {cfg_a.name}")
            from_a = json.loads(body)["config"]
            check(from_a["blank_correction"] == "Auto" and from_a["run_file"] == copy_a.name,
                  f"the first run's configuration reads {from_a['blank_correction']!r} for {from_a['run_file']!r}")
            own = dict(from_a, blank_correction="On")
            own_file = work / "copy_b_blank_on.json"
            own_file.write_text(json.dumps(own))
            # chosen for the second run alone, on its row of the batch
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-run-config-2", str(own_file))
            expect(page.locator(".dc-toast").get_by_text(f"{copy_b.name} was processed again with its own configuration, "
                                                         f"{own_file.name}; its decisions now start from it.")).to_be_visible(timeout=60000)
            expect(listed.nth(0)).to_contain_text("Configuration: the batch’s, the package defaults")
            expect(listed.nth(1)).to_contain_text(f"Configuration: its own, {own_file.name}")
            expect(page.get_by_role("button", name=f"Use the batch’s configuration for {copy_b.name}", exact=True)).to_have_count(1)
            check("note" not in listed.nth(1).inner_text(), f"batch row 2 reads {listed.nth(1).inner_text()!r}")
            shot("1440-run-config.png")
            status, _, body = post(base, "/api/results.csv", copy_b, {"config": json.dumps(own)})
            own_text = body.decode()
            own_table = rows_of(own_text)
            status, _, body = post(base, "/api/results.csv", copy_a)
            table_a = rows_of(body.decode())
            status, _, body = post(base, "/api/results.csv", copy_b)
            plain_b = rows_of(body.decode())
            # the first run is as the package defaults give it; its own table is the one the combined table holds
            page.get_by_role("button", name="Start the corrections").click()
            heading("results")
            expect(page.get_by_text("Run 1 of 2")).to_be_visible()
            compare(results_rows(page), table_a, "the first run, with the batch's configuration")
            with page.expect_download(timeout=60000) as dl:
                files.get_by_role("button", name=re.compile("^Results table")).click()
            table_own_a = work / ("w11_" + dl.value.suggested_filename)
            dl.value.save_as(str(table_own_a))
            expect(files.get_by_role("button", name=re.compile("^Results table"))).to_be_enabled(timeout=60000)
            # the second run follows its own configuration: the service's table for it, and its blank switch
            page.get_by_role("navigation", name="Runs").get_by_role("button", name=re.compile(copy_b.name)).click()
            expect(page.get_by_text("Run 2 of 2")).to_be_visible()
            rows_b = results_rows(page)
            n_b, _, _ = compare(rows_b, own_table, "the second run, with its own configuration")
            moved_b = sum(1 for r in rows_b if not agree(r["dN"], plain_b[int(r["r"])]["d15N (permil)"], 2)[0])
            check(moved_b > 0, "the second run's own configuration moved no shown d15N")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Blank").click()
            heading("blank")
            expect(page.get_by_role("group", name="Blank correction").get_by_role("button", name="On")).to_have_attribute("aria-pressed", "true")
            page.locator("nav[aria-label='Correction steps'] button", has_text="Setup").click()
            heading("setup")
            expect(page.get_by_text(f"From {own_file.name}.")).to_be_visible()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Results").click()
            heading("results")
            # its workbook, results table, and configuration are its own
            got_b = {}
            for label, timeout in (("Workbook", 180000), ("Results table", 60000), ("Configuration", 30000)):
                with page.expect_download(timeout=timeout) as dl:
                    files.get_by_role("button", name=re.compile("^" + label)).click()
                got_b[label] = work / ("w11_" + dl.value.suggested_filename)
                dl.value.save_as(str(got_b[label]))
                expect(files.get_by_role("button", name=re.compile("^" + label))).to_be_enabled(timeout=timeout)
            diff, _ = table_diff(got_b["Results table"].read_text(), own_text)
            check(not diff, "the second run's results table is not its configuration's: " + "; ".join(diff[:3]))
            status, _, body = post(base, "/api/config/read", got_b["Configuration"])
            cfg_b = json.loads(body)["config"]
            check(cfg_b["blank_correction"] == "On" and cfg_b["run_file"] == copy_b.name,
                  f"the second run's configuration reads {cfg_b['blank_correction']!r} for {cfg_b['run_file']!r}")
            cells_b = book_cells(got_b["Workbook"], own_table, "the second run's workbook")
            check(cells_b > 100, f"only {cells_b} computed values in the second run's workbook")
            # the combined results: each run's rows are its own table, and the Excel file names both configurations
            comb = {}
            for label, key in (("Combined results, all 2 runs", "xlsx"), ("Combined results table", "csv")):
                with page.expect_download(timeout=120000) as dl:
                    files.get_by_role("button", name=re.compile("^" + label)).click()
                comb[key] = work / ("w11_" + dl.value.suggested_filename)
                dl.value.save_as(str(comb[key]))
                expect(files.get_by_role("button", name=re.compile("^" + label))).to_be_enabled(timeout=120000)
            clines = comb["csv"].read_text().splitlines()
            check(len(clines) == 113, f"the combined table has {len(clines)} lines, not 113")
            for k, (stem, own_t) in enumerate(((copy_a.stem, table_own_a.read_text()), (copy_b.stem, got_b["Results table"].read_text()))):
                check(clines[1 + 56 * k:1 + 56 * (k + 1)] == [stem + "," + ln for ln in own_t.splitlines()[1:]],
                      f"the combined table's rows of {stem} are not that run's own table")
            cb = openpyxl.load_workbook(comb["xlsx"])
            readme = [c.value for c in cb["README"]["B"] if c.value]
            want = (f"Configuration: the package defaults for run 1; {own_file.name} for run 2. Each run's choices were written "
                    "into its configuration; the Runs sheet shows which corrections were in force, and in its Configuration "
                    "column which configuration each run used.")
            check(want in readme, f"the combined README names the configurations as {[t for t in readme if t.startswith('Configuration')]}")
            # isoharness 0.5.0: config_names gives each run's configuration its own cell on Runs
            runs_rows = list(cb["Runs"].iter_rows(values_only=True))
            ccol = runs_rows[0].index("Configuration") if "Configuration" in runs_rows[0] else None
            check(ccol is not None and [r[ccol] for r in runs_rows[1:3]] == ["the package defaults", own_file.name],
                  f"the Runs sheet names the configurations as {[r[ccol] for r in runs_rows[1:3]] if ccol is not None else None}")
            runs_c = list(cb["Runs"].iter_rows(values_only=True))
            col = runs_c[0].index("Blank subtracted")
            check([r[col] for r in runs_c[1:]] == ["no", "N only"], f"the Runs sheet's blanks read {[r[col] for r in runs_c[1:]]}")
            # W-12: S 20 of the first run and S20 of the second are one sample, matched with the spaces taken out
            rep = list(cb["Repeats"].iter_rows(values_only=True))
            ix = {h: i for i, h in enumerate(rep[0])}
            check("Name matched" in ix and len(rep) - 1 == 52, f"the Repeats sheet has {len(rep) - 1} rows, headed {rep[0]}")
            for r in rep[1:]:
                pos, first = r[ix["Row"]], r[ix["Run"]] == copy_a.stem
                check(r[ix["Name"]] == f"S {pos}" and r[ix["Identifier"]] == (f"S {pos}" if first else f"S{pos}")
                      and r[ix["Name matched"]] == ("as written" if first else "spaces ignored") and r[ix["Runs"]] == 2,
                      f"a Repeats row reads {r}")
            reps = page.locator('[aria-label="Samples measured in more than one run"] [role=row]')
            check(reps.count() - 1 == 26, f"{reps.count() - 1} samples measured in more than one run, not 26")
            for i in range(1, reps.count()):
                cell = reps.nth(i).locator("[role=cell]").first.inner_text().split("\n")
                check(len(cell) == 2 and cell[1] == "with " + cell[0].replace(" ", "") + ", spaces ignored",
                      f"the page pairs {cell}")
            expect(page.locator('section[aria-label="All runs"]')).to_contain_text("with the spaces taken out (all 26)")
            # R-08 per run: a configuration naming another run, chosen for the second run, gives it a note of its own
            page.get_by_role("navigation", name="Main").get_by_role("button", name="EA-IRMS").click()
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-run-config-2", str(got["Configuration"]))
            expect(page.locator(".dc-toast").get_by_text(f"its own configuration, {got['Configuration'].name}")).to_be_visible(timeout=60000)
            expect(listed.nth(1)).to_contain_text("Read without errors; 1 note on the configuration")
            check("note" not in listed.nth(0).inner_text(), f"batch row 1 reads {listed.nth(0).inner_text()!r}")
            # a configuration for the batch processes the run that takes it, and the other keeps its own
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.set_input_files("#dc-config", str(work / "lab-config.json"))
            expect(page.locator(".dc-toast").get_by_text(
                f"The run that takes the batch’s configuration was processed again with lab-config.json; the decisions now "
                f"start from it. {copy_b.name} keeps its own configuration.")).to_be_visible(timeout=60000)
            expect(listed.nth(0)).to_contain_text("Configuration: the batch’s, lab-config.json")
            expect(listed.nth(1)).to_contain_text(f"Configuration: its own, {got['Configuration'].name}")
            # and the second run goes back to the batch's
            with page.expect_request(lambda r: r.url.endswith("/api/process") and r.method == "POST"):
                page.get_by_role("button", name=f"Use the batch’s configuration for {copy_b.name}", exact=True).click()
            expect(page.locator(".dc-toast").get_by_text(f"{copy_b.name} was processed again with the batch’s configuration, "
                                                         "lab-config.json; its decisions now start from it.")).to_be_visible(timeout=60000)
            expect(listed.nth(1)).to_contain_text("Configuration: the batch’s, lab-config.json")
            expect(batch_back).to_have_count(0)
            check("note" not in listed.nth(1).inner_text(), f"batch row 2 reads {listed.nth(1).inner_text()!r}")
            # R-29 at 400 px: the sequence keeps 24 px bars and scrolls within its box, not the page
            page.set_viewport_size({"width": 400, "height": 1000})
            page.get_by_role("button", name="Start the corrections").click()
            page.locator("nav[aria-label='Correction steps'] button", has_text="Drift").click()
            heading("drift")
            widths = page.locator(".dc-ribbon button").evaluate_all("bs => bs.map(b => b.getBoundingClientRect().width)")
            box = page.locator(".dc-ribbon").evaluate("t => { const s = t.closest('.dc-scroll'); return [s.scrollWidth, s.clientWidth]; }")
            check(min(widths) >= 24 - 1e-6 and box[0] > box[1], f"at 400 px the bars are {min(widths):.1f} px and the strip {box}")
            wide = page.evaluate("document.documentElement.scrollWidth")
            check(wide <= 400, f"the drift screen is {wide} px wide at a 400 px viewport")
            shot("400-drift.png")
            page.set_viewport_size({"width": 1440, "height": 1000})
            stats["own"] = (n_b, moved_b, cells_b, len(rep) - 1)

            context.close()
            browser.close()
        check(not noise, "unexpected console errors or failed requests:\n  " + "\n  ".join(noise))
    except (Failed, PlaywrightError) as exc:
        print(f"e2e_web: FAILED: {exc}")
        if noise:
            print("  also seen:\n    " + "\n    ".join(noise))
        return 1
    finally:
        if server is not None:
            server.terminate()
            try:
                server.wait(timeout=20)
            except subprocess.TimeoutExpired:
                server.kill()
        log.close()
    n, values, exact = stats["compared"]
    print(f"e2e_web: passed. {len(SCREENS) + 1} screens and the home page; {n} samples on the results screen equal the service's table for the "
          f"same decisions ({values} values and {n} flag lists, {exact} formatted identically; {stats['moved']} "
          f"d15N values moved by the decisions); workbook, table and configuration downloaded, {stats['book']} "
          f"workbook values checked, configuration round trip within {stats['round trip']:.1e} relative; "
          f"2 refused uploads shown as written; a batch of 3 runs, {stats['batch'][0]} combined rows equal to each "
          f"run's own table within {stats['batch'][1]:.1e} relative, {stats['batch'][2]} repeat analyses; "
          f"a sample sheet matched and carried into {stats['sheet'][0]} rows of the combined table, run information "
          f"on {stats['sheet'][1]} runs and in a run's workbook; one run's configuration kept to that run in a "
          f"batch, the other taking its own {stats['scoped']} monitors; a run with a configuration of its own "
          f"({stats['own'][0]} samples equal to the service's table for it, {stats['own'][1]} d15N values moved by it, "
          f"{stats['own'][2]} values of its workbook), the combined table equal to each run's own, and "
          f"{stats['own'][3]} repeat analyses paired with the spaces taken out; "
          f"0 console errors, 0 failed requests; screenshots in {shots}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
